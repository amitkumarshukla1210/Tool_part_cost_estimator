"""Country comparison and landed cost in India (roadmap stage 11), for the part and its tool.

For every country the part is costed as if part AND tool were made there (the same engines as the part panel: RM,
tool cost, part cost, each with that country's approved rates). The ex-works price is the part cost without the local
logistics; imports to India then add

    freight (sea / air, per kg) + insurance (% of ex-works + freight) = CIF (assessable value)
    + Basic Customs Duty (HS code and origin, preferential FTA rate when you tick it) + SWS (% of BCD)
    + IGST (% of CIF + BCD + SWS; credit shown separately when claimable) + clearance + inland transport (per kg)

Every rate is an approved version in the rate library (categories freight, duty, import, process_unit); the HS codes
are master data (Mould & die database → HS codes). Nothing is calculated that is not a rate times a quantity.
"""

from __future__ import annotations

from . import costing_setup, master_data, mould_db, piece_cost, rates, register_setup

HOME = "IN"


def _hs(scope: str, part_type: str) -> str:
    row = next((r for r in mould_db.tables()["hs_codes"] if r["scope"] == scope and r["part_type"] == part_type), None)
    if not row:
        raise ValueError(f"Mould & die database → HS codes has no row for {scope} / {part_type}.")
    return row["hs_code"]


def _rate(used: dict, cat: str, item: str, country: str, as_of: str, optional: bool = False):
    try:
        r = rates.lookup(cat, item, country, as_of)
    except rates.RateMissing:
        if optional:
            return None
        raise
    used[r["id"]] = r
    return r["value"]


def import_cost(exw: float, weight_kg: float, origin: str, hs: str, as_of: str, mode: str, use_fta: bool, used: dict) -> dict:
    """Landed cost in India of one item (INR) with every step."""
    inland = _rate(used, "process_unit", "Transport (per kg)", HOME, as_of) * weight_kg
    if origin == HOME:
        return {"exw": exw, "freight": 0.0, "insurance": 0.0, "cif": exw, "bcd_pct": 0.0, "bcd": 0.0, "sws": 0.0, "igst_pct": 0.0, "igst": 0.0,
                "clearance": 0.0, "inland": inland, "landed_with_igst": exw + inland, "igst_credit": 0.0, "landed": exw + inland,
                "duty_basis": "domestic: no import duty", "hs": hs}
    freight = _rate(used, "freight", f"{mode}_per_kg", origin, as_of) * weight_kg
    insurance = _rate(used, "import", "insurance_pct", "", as_of) / 100 * (exw + freight)
    cif = exw + freight + insurance
    fta = _rate(used, "duty", f"bcd_pct_fta:{hs}", origin, as_of, optional=True) if use_fta else None
    bcd_pct = fta if fta is not None else _rate(used, "duty", f"bcd_pct:{hs}", origin, as_of)
    bcd = cif * bcd_pct / 100
    sws = bcd * _rate(used, "duty", "sws_pct", HOME, as_of) / 100
    igst_pct = _rate(used, "duty", f"igst_pct:{hs}", HOME, as_of)
    igst = (cif + bcd + sws) * igst_pct / 100
    clearance = cif * _rate(used, "import", "clearance_pct", HOME, as_of) / 100
    claim = bool(master_data.parameters("all").get("igst_claimable"))
    gross = cif + bcd + sws + igst + clearance + inland
    return {"exw": exw, "freight": freight, "insurance": insurance, "cif": cif, "bcd_pct": bcd_pct, "bcd": bcd, "sws": sws,
            "igst_pct": igst_pct, "igst": igst, "clearance": clearance, "inland": inland, "landed_with_igst": gross,
            "igst_credit": igst if claim else 0.0, "landed": gross - (igst if claim else 0.0), "hs": hs,
            "duty_basis": ("FTA preferential rate" if fta is not None else "standard rate")}


def compare(part_id: int, countries: list[str] | None = None, mode: str = "sea", use_fta: bool = True) -> dict:
    from . import register
    d = register.part_detail(part_id, with_rm=False)
    if register.tool_eligibility(d):
        raise ValueError(register.tool_eligibility(d))
    part_type = register_setup.processes()[d["process"]]["part_type"]
    setup = costing_setup.tables()
    cmap = {c["code"]: c for c in setup["countries"]}
    wanted = [c.upper() for c in (countries or list(cmap))]
    unknown = [c for c in wanted if c not in cmap]
    if unknown:
        raise ValueError(f"Country '{unknown[0]}' is not in Master data → Countries.")
    pct = float(master_data.parameters("all").get("fx_scenario_pct", 5))
    rows = []
    for c in wanted:
        used: dict = {}
        row = {"country": c, "name": cmap[c]["name"], "currency": cmap[c]["currency"], "error": ""}
        try:
            pc = piece_cost.run(part_id, c)
            as_of = pc["as_of"]
            n = pc["nodes"]
            k = pc["value_inr"] / pc["value"]                                   # INR per unit of the country's currency
            usd = rates.lookup("forex", "USD", "", as_of)["value"]
            used |= {r["id"]: r for r in pc["rates_used"]}
            comp = lambda *names: sum(n[f"calc:{x}"]["value"] for x in names) * k
            exw_local = pc["value"] - n["calc:logistics"]["value"]
            exw = exw_local * k
            parts = {"rm": n["in:rm_cost"]["value"] * k, "process": comp("process_cost", "secondary_cost"),
                     "tool": comp("amortisation"), "packing": comp("packing")}
            parts["other"] = exw - sum(parts.values())
            wkg = n["in:total_weight_kg"]["value"]
            tl = pc["tool_summary"]
            land = import_cost(exw, wkg, c, _hs("part", part_type), as_of, mode, use_fta, used)
            tool_land = import_cost(tl["value_inr"], tl["weight_kg"], c, _hs("tool", part_type), as_of, mode, use_fta, used)
            scen = {}
            for sign in (-1, 1):
                u: dict = {}
                scen[f"{sign * pct:+g}%"] = import_cost(exw * (1 + sign * pct / 100), wkg, c, _hs("part", part_type), as_of, mode, use_fta, u)["landed"] \
                    if cmap[c]["currency"] != "INR" else land["landed"]
            row |= {"as_of": as_of, "exw_local": exw_local, "exw_inr": exw, "exw_usd": exw / usd, "components_inr": parts, "weight_kg": wkg,
                    "landed": land, "tool": {"cost_local": tl["value"], "cost_inr": tl["value_inr"], "cost_usd": tl["value_inr"] / usd,
                                             "weight_kg": tl["weight_kg"], "landed": tool_land},
                    "forex": {"usd_inr": usd, "fx_inr_per_unit": k, "scenario": scen | {"0%": land["landed"]}},
                    "piece_machine": pc["machine"], "is_example": pc["is_example"] or any(r.get("is_example") for r in used.values()),
                    "rates_used": [{"id": r["id"], "category": r["category"], "item": r["item"], "country": r["country"], "month": r["month"],
                                    "value": r["value"], "unit": r["unit"], "is_example": r["is_example"]} for r in sorted(used.values(), key=lambda r: r["id"])]}
        except (piece_cost.PieceError, rates.RateError, ValueError) as exc:
            row["error"] = str(exc)
        rows.append(row)
    ok = [r for r in rows if not r["error"]]
    cheapest = min(ok, key=lambda r: r["landed"]["landed"])["country"] if ok else None
    cheapest_tool = min(ok, key=lambda r: r["tool"]["landed"]["landed"])["country"] if ok else None
    return {"part_id": part_id, "part_code": d["part_code"], "mode": mode, "use_fta": use_fta, "fx_pct": pct, "rows": rows,
            "cheapest": cheapest, "cheapest_tool": cheapest_tool, "home": HOME}


def _flat(res: dict) -> list[list]:
    out = []
    for r in res["rows"]:
        if r["error"]:
            out.append([r["country"], r["name"], r["error"]] + [""] * 17)
            continue
        c, l, t = r["components_inr"], r["landed"], r["tool"]["landed"]
        out.append([r["country"], r["name"], "", r["currency"], r["exw_local"], r["exw_usd"], r["exw_inr"], c["rm"], c["process"], c["tool"],
                    c["packing"], c["other"], l["freight"], l["insurance"], l["bcd"], l["sws"], l["igst"], l["clearance"], l["inland"],
                    l["landed_with_igst"], l["igst_credit"], l["landed"], r["tool"]["cost_inr"], t["landed"]])
    return out


HEAD = ["Country", "Name", "Error", "Currency", "Ex-works (local)", "Ex-works (USD)", "Ex-works (INR)", "RM", "Process", "Tool amortisation",
        "Packing", "Overhead, rejection, profit, other", "Freight", "Insurance", "BCD", "SWS", "IGST", "Clearance", "Inland transport",
        "Landed incl. IGST", "IGST credit", "Landed (net of credit)", "Tool ex-works (INR)", "Tool landed (INR)"]


def to_excel(res: dict) -> bytes:
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from .table_set import HEADER_FILL
    wb = Workbook()
    sheets = [("Comparison (INR)", HEAD, _flat(res)),
              ("Forex scenario", ["Country", "Currency"] + list(next((r["forex"]["scenario"] for r in res["rows"] if not r["error"]), {})),
               [[r["country"], r["currency"]] + list(r["forex"]["scenario"].values()) for r in res["rows"] if not r["error"]]),
              ("Duty basis", ["Country", "Part HS", "Part BCD %", "Basis", "IGST %", "Tool HS", "Tool BCD %"],
               [[r["country"], r["landed"]["hs"], r["landed"]["bcd_pct"], r["landed"]["duty_basis"], r["landed"]["igst_pct"],
                 r["tool"]["landed"]["hs"], r["tool"]["landed"]["bcd_pct"]] for r in res["rows"] if not r["error"]]),
              ("Rates used", ["Rate #", "Country", "Category", "Item", "Month", "Value", "Unit", "Example"],
               [[x["id"], r["country"], x["category"], x["item"], x["month"], x["value"], x["unit"], "EXAMPLE" if x["is_example"] else ""]
                for r in res["rows"] if not r["error"] for x in r["rates_used"]])]
    for i, (name, head, rows) in enumerate(sheets):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = name
        ws.append(head)
        for c in ws[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), HEADER_FILL
        for r in rows:
            ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def chart_png(res: dict) -> bytes:
    import io
    import matplotlib
    matplotlib.use("Agg")
    from matplotlib import pyplot as plt
    ok = [r for r in res["rows"] if not r["error"]]
    keys = [("RM", lambda r: r["components_inr"]["rm"]), ("Process", lambda r: r["components_inr"]["process"]),
            ("Tool", lambda r: r["components_inr"]["tool"]), ("Packing", lambda r: r["components_inr"]["packing"]),
            ("Other", lambda r: r["components_inr"]["other"]), ("Freight + insurance", lambda r: r["landed"]["freight"] + r["landed"]["insurance"]),
            ("Duty (BCD + SWS + IGST net)", lambda r: r["landed"]["bcd"] + r["landed"]["sws"] + r["landed"]["igst"] - r["landed"]["igst_credit"]),
            ("Clearance + inland", lambda r: r["landed"]["clearance"] + r["landed"]["inland"])]
    colours = ["#1f6feb", "#2da44e", "#bf8700", "#8250df", "#8b949e", "#0969da", "#cf222e", "#6e7781"]
    fig, ax = plt.subplots(figsize=(8, 4), dpi=110)
    bottom = [0.0] * len(ok)
    for (label, f), col in zip(keys, colours):
        vals = [f(r) for r in ok]
        ax.bar([r["country"] for r in ok], vals, bottom=bottom, label=label, color=col)
        bottom = [b + v for b, v in zip(bottom, vals)]
    for i, r in enumerate(ok):
        if r["country"] == res["cheapest"]:
            ax.annotate("cheapest", (i, bottom[i]), ha="center", va="bottom", fontsize=8, color="#2da44e")
    ax.set_ylabel("INR per part, landed in India")
    ax.legend(fontsize=7, ncol=2)
    buf = io.BytesIO()
    fig.savefig(buf, format="png", bbox_inches="tight")
    plt.close(fig)
    return buf.getvalue()


def to_pdf(res: dict, part: dict) -> bytes:
    import io
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer
    from .pdf_report import _example_banner, _table, _theme
    theme, ss = _theme(), getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=ss["Title"], textColor=theme["ink"], fontSize=18, alignment=0, spaceAfter=2)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], textColor=theme["ink"], fontSize=11, spaceBefore=8, spaceAfter=4)
    note = ParagraphStyle("n", parent=ss["Normal"], textColor=theme["grey"], fontSize=7.5)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=14 * mm, rightMargin=14 * mm, topMargin=12 * mm, bottomMargin=12 * mm,
                            title=f"Design to Cost — Country comparison {part['part_code']}", author="Design to Cost")
    ok = [r for r in res["rows"] if not r["error"]]
    f0 = lambda v: f"{v:,.2f}"
    story = [Paragraph("COUNTRY COMPARISON AND LANDED COST IN INDIA", ParagraphStyle("k", parent=ss["Normal"], textColor=theme["brand"], fontSize=8)),
             Paragraph(f"{part['part_code']} · {part['name']}", h1),
             Paragraph(f"{'Sea' if res['mode'] == 'sea' else 'Air'} freight · {'FTA rates used where available' if res['use_fta'] else 'standard duty rates'} · "
                       f"all amounts INR per part unless stated · cheapest landed: {res['cheapest']}", note)]
    if any(r["is_example"] for r in ok):
        story += [Spacer(1, 2 * mm), _example_banner(ss, "Rates and duties are EXAMPLE values.")]
    story += [Spacer(1, 3 * mm), Image(io.BytesIO(chart_png(res)), 190 * mm, 95 * mm), Paragraph("Landed cost per part", h2),
              _table([["Country", "Ex-works", "USD", "Freight+ins.", "BCD+SWS", "IGST", "Clear.+inland", "IGST credit", "Landed", "Tool landed"]] +
                     [[r["country"], f0(r["exw_inr"]), f0(r["exw_usd"]), f0(r["landed"]["freight"] + r["landed"]["insurance"]),
                       f0(r["landed"]["bcd"] + r["landed"]["sws"]), f0(r["landed"]["igst"]), f0(r["landed"]["clearance"] + r["landed"]["inland"]),
                       f0(r["landed"]["igst_credit"]), f0(r["landed"]["landed"]), f"{r['tool']['landed']['landed']:,.0f}"] for r in ok], theme),
              Paragraph(f"Forex ±{res['fx_pct']:g} %", h2),
              _table([["Country"] + list(ok[0]["forex"]["scenario"])] + [[r["country"]] + [f0(v) for v in r["forex"]["scenario"].values()] for r in ok], theme)
              if ok else Spacer(1, 1)]
    for r in res["rows"]:
        if r["error"]:
            story.append(Paragraph(f"<font color='#b91c1c'>■</font> {r['country']}: {r['error']}", ss["Normal"]))
    doc.build(story)
    return buf.getvalue()
