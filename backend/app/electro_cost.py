"""Harness and PCB costing (roadmap stage 15): from a reviewed BOM extract to a traced cost per harness / per board.

    harness   wires: length × price per metre, from the Part library by part number, else copper content × (copper rate +
              conversion rate per kg of copper), both approved rates of the country (copper-index linked);
              connectors, terminals, seals, sleeves, tapes, clips, ...: quantity × Part library price (by part number);
              processing: drivers (wires, crimps, connectors, ...) × seconds ÷ 3600 × (MHR + operators × labour rate)
    PCB       components: quantity × Part library price; bare board: area × rate per cm² of its layer count (rate library),
              plus panel waste; assembly: SMT / THT placements, boards, area × seconds × (MHR + labour)
    then      the Cost formulas of scope harness / pcb (rejection, overhead, profit, packing)

No price is guessed: a part that is not in the library is listed as not priced and counts as 0 until you add it (the result
says so). Distributor price APIs are not used. Every number is a rate, a library row, a process row or a formula in master data.
"""

from __future__ import annotations

from . import bom_extract, costing_setup, master_data, piece_cost, rate_engine, rates, register_setup
from .trace import Trace

COPPER_G_CM3 = 8.96
DRIVERS = {
    "wires": "Wires (Σ quantity of wire lines)", "wire_m": "Wire length (m)", "crimps": "Crimps = terminal quantity, else 2 per wire",
    "connectors": "Connectors", "seals": "Seals and plugs", "sleeves": "Sleeves / tubes", "tapes": "Tapes", "clips": "Clips and ties", "harness": "Harness (1)",
    "smt_placements": "SMT placements", "tht_placements": "THT placements", "boards": "Boards (1)", "board_cm2": "Board area (cm²)",
}


class ElectroError(ValueError):
    pass


def _library() -> dict:
    out = {}
    for r in register_setup.tables()["part_library"]:
        out.setdefault(r["part_number"].strip().upper(), r)
    return out


def _fx(cur: str, as_of: str) -> float:
    return 1.0 if cur == "INR" else rates.lookup("forex", cur, "", as_of)["value"]


def _price_in(row: dict, cur: str, as_of: str) -> float:
    return row["unit_price"] * _fx(row["currency"], as_of) / _fx(cur, as_of)


def drivers(e: dict) -> dict:
    L = e["lines"]
    q = lambda *cats: sum((l["qty"] or 0) for l in L if l["category"] in cats)
    wires = q("wire")
    d = {"wires": wires, "wire_m": sum((l["length_mm"] or 0) for l in L if l["category"] == "wire") / 1000,
         "crimps": q("terminal") or 2 * wires, "connectors": q("connector"), "seals": q("seal"), "sleeves": q("sleeve"), "tapes": q("tape"),
         "clips": q("clip", "grommet"), "harness": 1.0}
    smt = sum((l["qty"] or 0) for l in L if l["category"] == "component" and l["mount"] == "SMT")
    tht = sum((l["qty"] or 0) for l in L if l["category"] == "component" and l["mount"] == "THT")
    board = next((l for l in L if l["category"] == "pcb"), None)
    note = (board or {}).get("notes") or ""
    import re
    m = re.match(r"\s*([\d.]+) × ([\d.]+) mm", note)
    d |= {"smt_placements": smt, "tht_placements": tht, "boards": 1.0, "board_cm2": float(m.group(1)) * float(m.group(2)) / 100 if m else 0.0}
    return d


def _unit_costs(e, lib, cur, as_of, country, params, missing):
    """(wire cost, parts/components cost, lines) in the country's currency; unpriced lines are listed in `missing`."""
    wire_total, parts_total, lines = 0.0, 0.0, []
    cu_rate = conv = None
    for l in e["lines"]:
        cat = l["category"]
        if cat == "pcb":
            continue
        key = (l["part_number"] or "").strip().upper()
        hit = lib.get(key)
        cost, basis = None, ""
        if cat == "wire":
            m = (l["length_mm"] or 0) / 1000
            if hit and hit["unit"] == "m":
                cost, basis = m * _price_in(hit, cur, as_of), f"Part library {hit['part_number']} per m"
            elif l["cross_section_mm2"]:
                try:
                    cu = rates.lookup("material", "Copper", country, as_of)
                    cv = rates.lookup("harness", "wire_conversion_per_kg_copper", country, as_of)
                except rates.RateMissing as exc:
                    missing.append(f"{l['ref'] or l['description']}: {exc}")
                else:
                    kg_m = l["cross_section_mm2"] * COPPER_G_CM3 / 1000
                    cost = m * kg_m * (cu["value"] * _fx(cu["currency"], as_of) + cv["value"] * _fx(cv["currency"], as_of)) / _fx(cur, as_of)
                    basis = f"{kg_m * 1000:.1f} g copper/m × (copper {cu['value']:g} + conversion {cv['value']:g}) {cu['currency']}/kg"
            else:
                missing.append(f"Wire {l['ref'] or l['description']}: no cross-section and no library price per metre")
        elif hit:
            cost, basis = (l["qty"] or 0) * _price_in(hit, cur, as_of), f"Part library {hit['part_number']}"
        else:
            missing.append(f"{l['category']} {l['part_number'] or l['description'] or l['ref']}: not in the Part library")
        if cost is not None:
            if cat == "wire":
                wire_total += cost
            else:
                parts_total += cost
        lines.append({"id": l.get("id"), "category": cat, "ref": l["ref"], "part_number": l["part_number"], "description": l["description"],
                      "qty": l["qty"], "cost": cost, "basis": basis or "not priced", "low_confidence": l["confidence"] < bom_extract.LOW})
    return wire_total, parts_total, lines


def run(extract_id: int, country: str | None = None, as_of: str | None = None) -> dict:
    e = bom_extract.get(extract_id)
    kind = e["kind"]
    country = (country or register_setup.settings()["default_country"]).upper()
    as_of = as_of or rates.current_month()
    setup = costing_setup.tables()
    cmap = {c["code"]: c for c in setup["countries"]}
    if country not in cmap:
        raise ElectroError(f"Country '{country}' is not in Master data → Countries.")
    cur = cmap[country]["currency"]
    params = master_data.parameters("all")
    lib, missing = _library(), []
    wire, parts, lines = _unit_costs(e, lib, cur, as_of, country, params, missing)
    t = Trace(cur)
    src = lambda note: {"type": "register", "note": note}
    d = drivers(e)
    procs = []
    machines = {m["machine_id"]: m for m in setup["machines"]}
    names, proc_nodes = {}, []
    for i, p in enumerate(register_setup.electro_processes(kind)):
        units = d.get(p["driver"])
        if units is None:
            raise ElectroError(f"Process '{p['code']}': driver '{p['driver']}' is not one of {', '.join(DRIVERS)}.")
        if not units:
            continue
        pre = f"p{i}:"
        try:
            lab = piece_cost.graft(t, rate_engine.labour_rate(p["labour_grade"], country, as_of), pre + "lab:")
            parts_ = {"units": f"{pre}units", "sec": f"{pre}sec", "ops": f"{pre}ops", "lab": lab}
            t.add_input(parts_["units"], f"{p['code']}: {DRIVERS.get(p['driver'], p['driver'])}", units, "", src("From the BOM extract"))
            t.add_input(parts_["sec"], f"{p['code']}: seconds per unit", p["seconds"], "s", {"type": "master", "sheet": "Harness and PCB processes", "row": f"{p['scope']} / {p['code']}", "is_example": p["is_example"]})
            t.add_input(parts_["ops"], f"{p['code']}: operators", p["operators"], "", {"type": "master", "sheet": "Harness and PCB processes", "row": f"{p['scope']} / {p['code']}", "is_example": p["is_example"]})
            if p["machine_id"]:
                if p["machine_id"] not in machines:
                    raise ElectroError(f"Process '{p['code']}': machine '{p['machine_id']}' is not in Machines.")
                parts_["mhr"] = piece_cost.graft(t, rate_engine.mhr(p["machine_id"], country, as_of), pre + "mhr:")
                f = "units * sec / 3600 * (mhr + ops * lab)"
            else:
                f = "units * sec / 3600 * ops * lab"
            t.compute(f"proc:{p['code']}", f"{p['code']} ({p['description'] or p['driver']})", f, cur, money=True, names=parts_,
                      description="Driver units × seconds ÷ 3600 × (machine hour rate + operators × labour rate)")
        except (rate_engine.EngineError, rates.RateError) as exc:
            missing.append(f"Process {p['code']}: {exc}")
            continue
        proc_nodes.append(f"proc:{p['code']}")
    if proc_nodes:
        t.compute("process_total", "Processing cost", " + ".join(f"n{i}" for i in range(len(proc_nodes))), cur, money=True,
                  names={f"n{i}": n for i, n in enumerate(proc_nodes)}, description="Sum of the processes")
    else:
        t.add_input("process_total", "Processing cost", 0.0, cur, src("No process applies"), money=True)
    board = None
    if kind == "harness":
        t.add_input("in:wire_cost", "Wires", wire, cur, src(f"{len([l for l in lines if l['category'] == 'wire'])} wire lines priced from the Part library or copper rate"), money=True)
        t.add_input("in:parts_cost", "Connectors, terminals and other parts", parts, cur, src("Part library"), money=True)
        inputs = {"wire_cost": "in:wire_cost", "parts_cost": "in:parts_cost", "process_cost": "process_total"}
    else:
        b = next((l for l in e["lines"] if l["category"] == "pcb"), None)
        layers = (e["board"] or {}).get("layers")
        board_cost = 0.0
        if d["board_cm2"] and layers:
            try:
                r = rates.lookup("pcb_board", f"{int(layers)}L", country, as_of)
                board_cost = d["board_cm2"] * r["value"] * _fx(r["currency"], as_of) / _fx(cur, as_of) * (1 + params.get("pcb_panel_waste_pct", 0) / 100)
                board = {"cm2": d["board_cm2"], "layers": int(layers), "rate": r["value"], "unit": r["unit"]}
            except rates.RateMissing as exc:
                missing.append(f"Bare board: {exc}")
        else:
            missing.append("Bare board: layers and size are needed (type them when you read the BOM)")
        t.add_input("in:components_cost", "Components", parts, cur, src("Part library"), money=True)
        t.add_input("in:board_cost", "Bare board", board_cost, cur, src("Rate library pcb_board × area × (1 + panel waste)"), money=True)
        inputs = {"components_cost": "in:components_cost", "board_cost": "in:board_cost", "process_cost": "process_total"}
    for inp in (r for r in setup["formula_inputs"] if r["scope"] == kind):
        var = inp["variable"]
        if inp["source"] == "piece":
            if var not in inputs:
                raise ElectroError(f"Formula input '{var}' of scope {kind} is not a value of the {kind} engine ({', '.join(inputs)}).")
            continue
        nid = f"in:{var}"
        if inp["source"] == "parameter":
            v = master_data.parameters("all").get(inp["item"])
            if not isinstance(v, (int, float)):
                raise ElectroError(f"{inp['label']}: Master data → Parameters has no number '{inp['item']}'.")
            t.add_input(nid, inp["label"], v, inp["unit"], {"type": "master", "sheet": "Parameters", "row": f"all / {inp['item']}"})
        else:
            try:
                rate = rates.lookup(inp["category"], inp["item"], country, as_of)
            except rates.RateMissing as exc:
                raise ElectroError(str(exc)) from None
            t.add_rate(nid, inp["label"], rate)
        inputs[var] = nid
    try:
        result, _, order = rate_engine.evaluate_scope(t, setup, kind, dict(inputs), cur, as_of)
    except (rate_engine.EngineError, rates.RateMissing) as exc:
        raise ElectroError(str(exc)) from None
    out = t.to_dict(result)
    out |= {"kind": kind, "extract_id": extract_id, "name": e["name"], "country": country, "currency": cur, "as_of": as_of, "value": t.value(result),
            "value_inr": t.value("result:inr"), "priced": not missing, "missing": missing, "lines": lines, "drivers": d, "board": board,
            "low_confidence_lines": e["summary"]["low_confidence"], "reviewed": e["status"] == "reviewed",
            "steps": [{"name": r["name"], "node": f"calc:{r['name']}", "is_result": bool(r["is_result"])} for r in order],
            "is_example": any(x.get("is_example") for x in out["rates_used"]) or any(
                (n.get("source") or {}).get("is_example") for n in t.nodes.values()) or any(l.get("is_example") for l in []),
            "library_examples": any(lib[k]["is_example"] for k in lib if k in {(l["part_number"] or "").strip().upper() for l in e["lines"]})}
    return out


def to_excel(result: dict) -> bytes:
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from .table_set import HEADER_FILL
    from .trace import source_text
    n = result["nodes"]
    wb = Workbook()
    sheets = [("Summary", ["Item", "Value"], [["Extract", result["name"]], ["Kind", result["kind"]], ["Country", result["country"]], ["Costing month", result["as_of"]],
                                                [f"Cost ({result['currency']})", result["value"]], ["Cost (INR)", result["value_inr"]],
                                                ["All parts priced", "yes" if result["priced"] else "NO: see Not priced"],
                                                ["BOM reviewed", "yes" if result["reviewed"] else "no"], ["Lines to check in the BOM", result["low_confidence_lines"]],
                                                ["Example values", "yes: not for negotiation" if result["is_example"] or result["library_examples"] else "no"]]),
              ("Steps", ["Step", "Formula", "Value", "Unit"], [[n[s["node"]]["label"], n[s["node"]].get("formula", ""), n[s["node"]]["value"], n[s["node"]]["unit"]] for s in result["steps"]]),
              ("Lines", ["Category", "Ref", "Part number", "Description", "Qty", "Cost", "Basis"],
               [[l["category"], l["ref"], l["part_number"], l["description"], l["qty"], l["cost"], l["basis"]] for l in result["lines"]]),
              ("Processes", ["Process", "Value", "Unit", "Formula"], [[n[k]["label"], n[k]["value"], n[k]["unit"], n[k].get("formula", "")] for k in n if k.startswith("proc:")]),
              ("Inputs", ["Input", "Value", "Unit", "Source"], [[n[k]["label"], n[k]["value"], n[k]["unit"], source_text(n[k])] for k in n if k.startswith("in:")]),
              ("Not priced", ["Why"], [[m] for m in result["missing"]]),
              ("Rates used", ["Rate #", "Category", "Item", "Country", "Month", "Value", "Unit", "Example"],
               [[x["id"], x["category"], x["item"], x["country"], x["month"], x["value"], x["unit"], "EXAMPLE" if x.get("is_example") else ""] for x in result["rates_used"]])]
    for i, (name, head, rows) in enumerate(sheets):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = name
        ws.append(head)
        for c in ws[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), HEADER_FILL
        for r in rows:
            ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_pdf(result: dict) -> bytes:
    import io
    from datetime import date
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    from .pdf_report import _example_banner, _table, _theme
    from .trace import source_text
    theme, ss = _theme(), getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=ss["Title"], textColor=theme["ink"], fontSize=18, alignment=0, spaceAfter=2)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], textColor=theme["ink"], fontSize=11, spaceBefore=8, spaceAfter=4)
    note = ParagraphStyle("n", parent=ss["Normal"], textColor=theme["grey"], fontSize=7.5)
    cell = ParagraphStyle("c", parent=ss["Normal"], fontSize=7.5, leading=9)
    wrap = lambda x: Paragraph(str(x if x is not None else "").replace("&", "&amp;").replace("<", "&lt;"), cell)
    f = lambda v: "—" if v is None else f"{v:,.4f}".rstrip("0").rstrip(".")
    n = result["nodes"]
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=14 * mm, rightMargin=14 * mm, topMargin=12 * mm, bottomMargin=12 * mm,
                            title=f"Design to Cost — {result['kind']} cost {result['name']}", author="Design to Cost")
    story = [Paragraph(f"{result['kind'].upper()} COST SHEET", ParagraphStyle("k", parent=ss["Normal"], textColor=theme["brand"], fontSize=8)),
             Paragraph(result["name"], h1),
             Paragraph(f"Country {result['country']} · month {result['as_of']} · {date.today().isoformat()} · BOM {'reviewed' if result['reviewed'] else 'NOT reviewed'}"
                       f" · {result['low_confidence_lines']} BOM line(s) to check", note)]
    if result["is_example"] or result["library_examples"]:
        story += [Spacer(1, 2 * mm), _example_banner(ss)]
    if not result["priced"]:
        story += [Spacer(1, 2 * mm), Paragraph("<font color='#b91c1c'><b>NOT FULLY PRICED</b></font>: the lines below count as 0 until they are in the Part library.", ss["Normal"])]
        story += [Paragraph(f"■ {m}", ss["Normal"]) for m in result["missing"][:25]]
    story += [Paragraph("Result", h2), _table([["", result["currency"], "INR"], ["Cost per " + ("harness" if result["kind"] == "harness" else "board"), f(result["value"]), f(result["value_inr"])]], theme, [80 * mm, 40 * mm, 40 * mm]),
              Paragraph("Calculation steps", h2),
              _table([["Step", "Formula", "Value", "Unit"]] + [[wrap(n[s["node"]]["label"]), wrap(n[s["node"]].get("formula", "")), f(n[s["node"]]["value"]), n[s["node"]]["unit"]] for s in result["steps"]], theme,
                     [60 * mm, 130 * mm, 35 * mm, 25 * mm], total_rows=(len(result["steps"]),)),
              Paragraph("BOM lines priced", h2),
              _table([["Category", "Part", "Description", "Qty", "Cost", "Basis"]] + [[l["category"], wrap(l["part_number"] or l["ref"]), wrap(l["description"]), f(l["qty"]), f(l["cost"]), wrap(l["basis"])]
                                                                                    for l in result["lines"]], theme, [22 * mm, 45 * mm, 70 * mm, 18 * mm, 28 * mm, 70 * mm]),
              Paragraph("Processes", h2),
              _table([["Process", "Cost", "Formula"]] + [[wrap(n[k]["label"]), f(n[k]["value"]), wrap(n[k].get("formula", ""))] for k in n if k.startswith("proc:")], theme, [100 * mm, 35 * mm, 100 * mm])]
    doc.build(story)
    return buf.getvalue()
