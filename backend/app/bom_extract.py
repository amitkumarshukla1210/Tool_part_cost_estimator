"""Harness and PCB BOM extraction (roadmap stage 14, "Page A"): read a BOM from a file, review it, edit it, export it.

    harness   KBL / VEC XML (VDA 4964 / 4968)       wires, connectors, terminals, seals, sleeves, tapes, clips, ...
              (a wire line's length_mm is the TOTAL length of its qty wires)
              text-based PDF drawing with a BOM table (pdfplumber); scanned PDFs / images by OCR and DXF / DWG drawings: external.py
    PCB       BOM CSV / Excel (KiCad, Altium, generic), pick-and-place CPL CSV (counts, board extents), typed bare-board data

Every line carries a CONFIDENCE (0-1) and the reasons it is lower; lines under the review threshold LOW (0.8)
are marked for the engineer in the review table. Editing a line sets its confidence to 1.
Nothing here prices anything: costing is stage 15.
"""

from __future__ import annotations

import csv
import io
import re
import xml.etree.ElementTree as ET
from collections import defaultdict

LOW = 0.8
HARNESS_CATEGORIES = ("wire", "connector", "terminal", "seal", "sleeve", "tape", "clip", "grommet", "label", "other")
PCB_CATEGORIES = ("component", "pcb")
FIELDS = ("category", "ref", "part_number", "maker", "description", "qty", "unit", "length_mm", "wire_type", "cross_section_mm2",
          "colour", "package", "mount", "notes")


class ExtractError(ValueError):
    pass


def line(category: str, source: str, confidence: float = 1.0, flags: list | None = None, **kw) -> dict:
    out = {k: None for k in FIELDS} | {"unit": "pcs", "qty": 1.0} | kw
    out |= {"category": category, "source": source, "confidence": round(confidence, 2), "flags": flags or [], "edited": False}
    return out


def _num(v):
    if v is None:
        return None
    s = str(v).strip().replace(",", ".")
    m = re.search(r"-?\d+(?:\.\d+)?", s)
    return float(m.group()) if m else None


# ---------- harness: KBL / VEC XML ----------

KEYWORDS = (("clip", ("clip", "fixing", "clamp", "cable_tie", "cable tie", "cabletie", "tie wrap")), ("tape", ("tape",)),
            ("sleeve", ("protection", "tube", "sleeve", "conduit", "corrugated", "shrink")), ("grommet", ("grommet",)),
            ("label", ("label", "marker")), ("seal", ("seal", "plug", "cavityplug")), ("terminal", ("terminal", "contact")),
            ("connector", ("connector", "housing")), ("wire", ("wire", "cable", "core")))        # first match wins: specific before general


def _local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _classify(name: str) -> str | None:
    n = name.lower()
    for cat, keys in KEYWORDS:
        if any(k in n for k in keys):
            return cat
    return None


def _text(el, *names) -> str:
    """First descendant text whose local name (lower, no _ or -) is one of names."""
    want = {n.replace("_", "").lower() for n in names}
    for d in el.iter():
        if _local(d.tag).replace("_", "").replace("-", "").lower() in want and (d.text or "").strip():
            return d.text.strip()
    return ""


def _value(el, *names):
    """Number from <name><Value_component>12</Value_component></name> or <name>12</name>."""
    want = {n.replace("_", "").lower() for n in names}
    for d in el.iter():
        if _local(d.tag).replace("_", "").lower() in want:
            v = _text(d, "Value_component", "ValueComponent") or (d.text or "")
            if _num(v) is not None:
                return _num(v)
    return None


def parse_xml(content: bytes) -> list[dict]:
    try:
        root = ET.fromstring(content)
    except ET.ParseError as exc:
        raise ExtractError(f"The XML file could not be read: {exc}") from None
    rootname = _local(root.tag).lower()
    fmt = "VEC" if "vec" in rootname else "KBL"
    parts: dict = {}
    for el in root.iter():
        i = el.attrib.get("id") or el.attrib.get("xml:id")
        if i and (i not in parts or not _text(parts[i], "Part_number", "PartNumber")):
            parts[i] = el
    groups: dict = defaultdict(lambda: {"qty": 0.0, "length": 0.0, "n": 0, "el": None})
    for el in root.iter():
        name = _local(el.tag)
        if not re.search(r"occurrence|usage|element$|instance", name, re.I) or re.search(r"cavity$|contact_point|segment|node", name, re.I):
            continue
        cat = _classify(name)
        if not cat:
            continue
        ref_id = el.attrib.get("Part") or el.attrib.get("part") or ""
        spec = parts.get(ref_id, el)
        pn = _text(spec, "Part_number", "PartNumber", "Identification", "Id") if spec is not el else _text(el, "Part_number", "PartNumber")
        maker = _text(spec, "Company_name", "CompanyName", "Manufacturer", "Supplier")
        desc = _text(spec, "Description", "Part_description", "Name")
        ident = _text(el, "Special_wire_id", "Wire_number", "Identification", "Designator", "Wire_name") or el.attrib.get("id", "")
        key = (cat, pn or desc or name, maker, _text(spec, "Colour", "Cover_colour", "Colour_name", "CoreColor") if cat == "wire" else "")
        g = groups[key]
        g["qty"] += 1
        g["n"] += 1
        g["el"] = g["el"] if g["el"] is not None else (el, spec, name)
        g.setdefault("ids", []).append(ident)
        ln = _value(el, "Length_value", "Length", "Cut_length", "LengthValue")
        if ln:
            g["length"] += ln
    if not groups:
        raise ExtractError(f"No wires, connectors or other parts were found in this {fmt} file. Is it a harness (KBL / VEC) export?")
    out = []
    for (cat, pn, maker, colour), g in groups.items():
        el, spec, name = g["el"]
        flags = []
        conf = 1.0
        if not _text(spec, "Part_number", "PartNumber", "Identification", "Id"):
            flags.append("no part number in the file")
            conf -= 0.25
        if cat in ("connector", "terminal", "seal") and not maker:
            flags.append("no manufacturer in the file")
            conf -= 0.1
        kw = dict(ref=", ".join(i for i in g["ids"][:12] if i) + (" …" if len(g["ids"]) > 12 else ""), part_number=_text(spec, "Part_number", "PartNumber"),
                  maker=maker, description=_text(spec, "Description", "Part_description", "Name") or pn)
        if cat == "wire":
            xs = _value(spec, "Cross_section_area", "CrossSectionArea", "Nominal_cross_section")
            if xs is None:
                flags.append("no cross-section in the file")
                conf -= 0.2
            if not g["length"]:
                flags.append("no length in the file")
                conf -= 0.25
            kw |= {"wire_type": _text(spec, "Wire_type", "Type", "Insulation_material") or None, "cross_section_mm2": xs, "colour": colour or None,
                   "qty": g["qty"], "unit": "wires", "length_mm": g["length"] or None}
        else:
            kw["qty"] = g["qty"]
        out.append(line(cat, f"{fmt} {name} ×{int(g['n'])}", conf, flags, **kw))
    return out


# ---------- harness: text PDF with a BOM table ----------

HEADER_ALIASES = {
    "ref": ("item", "pos", "position", "no", "ref", "designator", "wire no", "id"), "part_number": ("part number", "part no", "p/n", "pn", "mpn", "article", "art. no", "number"),
    "description": ("description", "designation", "name", "text", "bezeichnung"), "qty": ("qty", "quantity", "pcs", "amount", "count", "menge"),
    "length_mm": ("length", "len", "länge", "cut length"), "colour": ("colour", "color", "farbe"),
    "cross_section_mm2": ("cross section", "cross-section", "section", "size", "gauge", "mm2", "mm²", "csa"),
    "maker": ("maker", "manufacturer", "supplier", "mfr", "vendor"), "wire_type": ("type", "grade", "wire type", "spec"),
}


def _map_header(cells: list[str]) -> dict:
    mapping = {}
    for i, c in enumerate(cells):
        h = re.sub(r"\s+", " ", (c or "").strip().lower())
        if not h:
            continue
        for field, aliases in HEADER_ALIASES.items():
            if field not in mapping.values() and any(h == a or h.startswith(a + " ") or h.endswith(" " + a) or a in h.split() for a in aliases):
                mapping[i] = field
                break
    return mapping


def _kind_from_text(text: str) -> str:
    t = text.lower()
    for cat, keys in KEYWORDS:
        if any(k.replace("_", " ") in t for k in keys):
            return cat
    return "other"


def table_to_lines(table: list, source: str, conf_table: list | None = None) -> list[dict] | None:
    """Lines of one table (rows of cell texts) whose header row names a quantity or part number column; None if there is no such header.
    conf_table (same shape, 0-1): confidence of each cell from OCR; a line is no more sure than its least sure cell."""
    start = mapping = None
    for ri, row in enumerate(table):
        mp = _map_header([c or "" for c in row])
        if len(mp) >= 3 and ("qty" in mp.values() or "part_number" in mp.values()):
            mapping, start = mp, ri + 1
            break
    if mapping is None:
        return None
    out = []
    for ri, row in enumerate(table[start:], start=start + 1):
        cells = {f: (row[i] or "").replace("\n", " ").strip() for i, f in mapping.items() if i < len(row)}
        if not any(cells.values()):
            continue
        flags, conf = [], 1.0
        if conf_table:
            cc = [conf_table[ri - 1][i] for i, f in mapping.items() if i < len(row) and (row[i] or "").strip()]
            if cc and min(cc) < 0.9:
                conf = min(conf, min(cc))
                flags.append(f"OCR: least sure word {min(cc):.0%}")
        qty = _num(cells.get("qty"))
        if qty is None:
            flags.append("quantity not a number")
            conf -= 0.4
        if not cells.get("part_number"):
            flags.append("no part number")
            conf -= 0.2
        length = _num(cells.get("length_mm"))
        desc = cells.get("description") or ""
        cat = _kind_from_text(desc + " " + cells.get("wire_type", ""))
        if cat == "other":
            flags.append("category guessed: other")
            conf -= 0.1
        if cat == "wire" and length is None:
            flags.append("wire without a length")
            conf -= 0.2
        each = length
        if cat == "wire" and length is not None:
            length = length * (qty if qty is not None else 1.0)      # a wire line's length is the TOTAL of its wires
        out.append(line(cat, f"{source} row {ri}", max(conf, 0.0), flags, ref=cells.get("ref"),
                        part_number=cells.get("part_number"), maker=cells.get("maker") or None, description=desc, qty=qty if qty is not None else 1.0,
                        notes=f"{each:g} mm each" if cat == "wire" and each is not None else None, length_mm=length, colour=cells.get("colour") or None,
                        cross_section_mm2=_num(cells.get("cross_section_mm2")), wire_type=cells.get("wire_type") or None, unit="wires" if cat == "wire" else "pcs"))
    return out


def parse_pdf(content: bytes) -> list[dict]:
    import pdfplumber
    out, any_table = [], False
    try:
        pdf = pdfplumber.open(io.BytesIO(content))
    except Exception as exc:
        raise ExtractError(f"The PDF could not be opened: {exc}") from None
    with pdf:
        n_chars = 0
        for pno, page in enumerate(pdf.pages, start=1):
            n_chars += len(page.chars)
            for tno, table in enumerate(page.extract_tables(), start=1):
                if not table or len(table) < 2:
                    continue
                lines = table_to_lines(table, f"PDF page {pno} table {tno}")
                if lines is not None:
                    any_table = True
                    out += lines
        if n_chars == 0:
            from . import external
            return external.ocr_pdf(content)               # a scan: Tesseract, or a clear message when it is not installed
    if not any_table:
        raise ExtractError("No BOM table with a quantity or part number column was found in the PDF. Check the table headers.")
    return out


# ---------- PCB: BOM and CPL ----------

PCB_ALIASES = {
    "ref": ("reference", "references", "designator", "designators", "ref", "refdes", "ref des", "part reference"),
    "description": ("value", "comment", "description", "val"), "package": ("footprint", "package", "pattern", "case"),
    "qty": ("qty", "quantity", "count", "amount"), "part_number": ("mpn", "manufacturer part number", "manufacturer part", "mfr part number", "part number",
                                                              "partnumber", "manufacturer_part_number", "mfr. part #", "lcsc part #"),
    "maker": ("manufacturer", "mfr", "mfr.", "maker"), "mount": ("mount", "type", "smd", "mounting"),
}
CPL_ALIASES = {"ref": ("designator", "ref", "refdes", "reference"), "x": ("mid x", "pos x", "posx", "x", "center-x(mm)", "midx"), "y": ("mid y", "pos y", "posy", "y", "center-y(mm)", "midy"),
               "rot": ("rotation", "rot"), "side": ("layer", "side", "tb"), "package": ("package", "footprint"), "description": ("val", "value", "comment")}
SMT_HINT = re.compile(r"(^|[^a-z])(0201|0402|0603|0805|1206|1210|2010|2512|sot|soic|ssop|tssop|qfn|dfn|qfp|lqfp|tqfp|bga|lga|sod|sma|smb|smc|msop|powerpad|sc70|son|smd|c_|r_|led_smd)",
                      re.I)
THT_HINT = re.compile(r"(through|tht|\bdip|pin_?header|pinheader|\bto-?(92|220|247|3)|axial|radial|screw|terminalblock|jack|barrel|socket|_th\b|pth)", re.I)


def _read_table(content: bytes, name: str) -> list[list[str]]:
    if name.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(content), read_only=True, data_only=True).worksheets[0]
        return [["" if c is None else str(c) for c in row] for row in ws.iter_rows(values_only=True)]
    text = content.decode("utf-8-sig", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return [row for row in csv.reader(io.StringIO(text), dialect)]


def _find_header(rows, aliases) -> tuple[int, dict]:
    for i, row in enumerate(rows[:30]):
        cells = [re.sub(r"\s+", " ", c.strip().lower()) for c in row]
        mp = {}
        for j, c in enumerate(cells):
            for field, al in aliases.items():
                if field not in mp.values() and c in al:
                    mp[j] = field
                    break
        if len(mp) >= 3 and "ref" in mp.values():
            return i, mp
    return -1, {}


def split_designators(s: str) -> list[str]:
    return [x.strip() for x in re.split(r"[,\s;]+", s or "") if x.strip()]


def mount_of(package: str, desc: str, given: str = "") -> tuple[str, float, str]:
    g = (given or "").strip().lower()
    if g in ("smd", "smt", "yes", "true", "1", "top"):
        return "SMT", 1.0, ""
    if g in ("tht", "th", "pth", "no", "false", "0"):
        return "THT", 1.0, ""
    if SMT_HINT.search(package or "") and not THT_HINT.search(package or ""):
        return "SMT", 0.9, ""
    if THT_HINT.search(package or ""):
        return "THT", 0.9, ""
    return "SMT", 0.5, "mount type guessed (unknown package)"


def parse_pcb_bom(content: bytes, name: str) -> list[dict]:
    rows = _read_table(content, name)
    hi, mp = _find_header(rows, PCB_ALIASES)
    if hi < 0:
        raise ExtractError("No BOM header found. A KiCad / Altium / generic BOM needs at least a designator column plus two of value, footprint, quantity, part number.")
    grouped = {}
    out = []
    for ri, row in enumerate(rows[hi + 1:], start=hi + 2):
        cells = {f: row[j].strip() for j, f in mp.items() if j < len(row)}
        refs = split_designators(cells.get("ref", ""))
        if not refs and not any(cells.values()):
            continue
        qty = _num(cells.get("qty")) or float(len(refs) or 1)
        flags, conf = [], 1.0
        if refs and cells.get("qty") and _num(cells["qty"]) != len(refs):
            flags.append(f"quantity {cells['qty']} differs from {len(refs)} designators")
            conf -= 0.3
        if not cells.get("part_number"):
            flags.append("no manufacturer part number")
            conf -= 0.15
        mount, mconf, mflag = mount_of(cells.get("package", ""), cells.get("description", ""), cells.get("mount", ""))
        if mflag:
            flags.append(mflag)
        out.append(line("component", f"{name} row {ri}", min(conf, mconf + 0.1) if mflag else conf, flags, ref=", ".join(refs), part_number=cells.get("part_number") or None,
                        maker=cells.get("maker") or None, description=cells.get("description") or None, package=cells.get("package") or None, mount=mount, qty=qty))
    if not out:
        raise ExtractError("The BOM has a header but no component rows.")
    return out


def parse_cpl(content: bytes, name: str) -> dict:
    """Placement counts per side and board extents (component centres: a lower bound of the board size)."""
    rows = _read_table(content, name)
    hi, mp = _find_header(rows, CPL_ALIASES)
    if hi < 0 or "x" not in mp.values() or "y" not in mp.values():
        raise ExtractError("No pick-and-place header found (designator, X, Y columns).")
    xs, ys, sides, refs = [], [], defaultdict(int), []
    for row in rows[hi + 1:]:
        c = {f: row[j].strip() for j, f in mp.items() if j < len(row)}
        x, y = _num(c.get("x")), _num(c.get("y"))
        if x is None or y is None:
            continue
        xs.append(x)
        ys.append(y)
        sides[(c.get("side") or "top").lower()[:3]] += 1
        refs.append(c.get("ref"))
    if not xs:
        raise ExtractError("The pick-and-place file has no placements.")
    return {"placements": len(xs), "sides": dict(sides), "width_mm": max(xs) - min(xs), "height_mm": max(ys) - min(ys), "refs": refs}


def pcb_board_line(cpl: dict | None, typed: dict) -> dict:
    """The bare board as a BOM line: typed layers / size / thickness win; else extents of the placements (low confidence)."""
    w, h = typed.get("width_mm"), typed.get("height_mm")
    flags, conf = [], 1.0
    if not (w and h):
        if cpl:
            w, h = cpl["width_mm"], cpl["height_mm"]
            flags.append("size from component positions only: the board outline is larger")
            conf = 0.5
        else:
            flags.append("board size missing: type it")
            conf = 0.3
    if not typed.get("layers"):
        flags.append("layer count missing")
        conf = min(conf, 0.4)
    return line("pcb", "typed" if typed else "pick-and-place extents", conf, flags, description=f"Bare PCB {typed.get('layers') or '?'} layers",
                length_mm=w, notes=f"{w:.1f} × {h:.1f} mm, {typed.get('thickness_mm') or '?'} mm thick" if w and h else "", qty=1.0,
                package=f"{typed.get('layers') or ''}L" if typed.get("layers") else None)


def summary(lines: list[dict]) -> dict:
    cats = defaultdict(float)
    for l in lines:
        cats[l["category"]] += l["qty"] or 0
    mount = defaultdict(float)
    for l in lines:
        if l["category"] == "component" and l["mount"]:
            mount[l["mount"]] += l["qty"] or 0
    low = [i for i, l in enumerate(lines) if l["confidence"] < LOW]
    return {"lines": len(lines), "by_category": dict(cats), "by_mount": dict(mount), "low_confidence": len(low),
            "wire_length_m": sum((l["length_mm"] or 0) * (l["qty"] if l["unit"] != "wires" else 1) for l in lines if l["category"] == "wire") / 1000}


def extract(content: bytes, filename: str, kind: str = "auto", typed: dict | None = None, cpl: tuple[bytes, str] | None = None) -> dict:
    """-> {"kind", "format", "lines", "board"}. `typed` is the bare-board data of a PCB; `cpl` an optional (content, name)."""
    n = filename.lower()
    typed = typed or {}
    if n.endswith(".xml"):
        return {"kind": "harness", "format": "KBL/VEC XML", "lines": parse_xml(content)}
    if n.endswith(".pdf"):
        return {"kind": "harness", "format": "PDF table", "lines": parse_pdf(content)}
    if n.endswith((".dxf", ".dwg")):
        from . import external
        return {"kind": "harness", "format": "DXF drawing" if n.endswith(".dxf") else "DWG drawing (converted to DXF)", "lines": external.parse_drawing(content, n)}
    if n.endswith((".png", ".jpg", ".jpeg", ".tif", ".tiff")):
        from . import external
        return {"kind": "harness", "format": "image (OCR)", "lines": external.ocr_image(content)}
    if n.endswith((".csv", ".tsv", ".txt", ".xlsx", ".xlsm")):
        lines = parse_pcb_bom(content, filename)
        cp = parse_cpl(*cpl) if cpl else None
        board = pcb_board_line(cp, typed)
        if cp:
            counts = {"placements": cp["placements"], "sides": cp["sides"]}
            bom_refs = sum(len(split_designators(l["ref"])) for l in lines)
            if bom_refs and bom_refs != cp["placements"]:
                board["flags"].append(f"BOM has {bom_refs} designators but the pick-and-place file {cp['placements']} placements")
                board["confidence"] = min(board["confidence"], 0.6)
            board["notes"] = (board["notes"] + f"; {counts['placements']} placements ({counts['sides']})").strip("; ")
        return {"kind": "pcb", "format": "BOM + CPL" if cpl else "BOM", "lines": lines + [board]}
    raise ExtractError("Unsupported file. Use KBL / VEC .xml, a .pdf (text or scanned), a .dxf / .dwg drawing or a scan image for a harness, and .csv / .xlsx for a PCB BOM (plus an optional pick-and-place CSV).")


# ---------- storage, review and export ----------

def _db():
    from . import projects_db as db
    return db


def create(content: bytes, filename: str, name: str = "", project_id: int | None = None, typed: dict | None = None,
           cpl: tuple[bytes, str] | None = None) -> dict:
    from sqlalchemy import select
    from . import file_store, register, scope
    db = _db()
    res = extract(content, filename, typed=typed, cpl=cpl)
    with db.engine().begin() as conn:
        if project_id and not scope.project_ok(conn, project_id):
            scope.deny("This project is not yours.")
        f = file_store.put(content, filename, "bom", conn)
        eid = conn.execute(db.bom_extracts.insert().values(
            project_id=project_id, name=name or filename, kind=res["kind"], format=res["format"], file_id=f["id"], board=typed or {},
            created_by=register.user(), owner=scope.me())).inserted_primary_key[0]
        conn.execute(db.bom_extract_lines.insert(), [{"extract_id": eid, "position": i, "data": l} for i, l in enumerate(res["lines"])])
        register._log(conn, "bom_extract", eid, "created", {"file": filename, "format": res["format"], "lines": len(res["lines"])})
    return get(eid)


def get(eid: int) -> dict:
    from sqlalchemy import select
    from . import register
    db = _db()
    with db.engine().connect() as conn:
        e = conn.execute(select(db.bom_extracts).where(db.bom_extracts.c.id == eid)).mappings().first()
        if not e:
            raise ExtractError(f"BOM extract {eid} does not exist.")
        rows = conn.execute(select(db.bom_extract_lines).where(db.bom_extract_lines.c.extract_id == eid)
                            .order_by(db.bom_extract_lines.c.position)).mappings().all()
    lines = [r["data"] | {"id": r["id"]} for r in rows]
    return register._plain(e) | {"lines": lines, "summary": summary(lines), "low_threshold": LOW,
                                 "categories": HARNESS_CATEGORIES if e["kind"] == "harness" else PCB_CATEGORIES}


def list_all(project_id: int | None = None) -> list[dict]:
    from sqlalchemy import func, or_, select
    from . import register, scope
    db = _db()
    with db.engine().connect() as conn:
        q = select(db.bom_extracts)
        if project_id:
            q = q.where(db.bom_extracts.c.project_id == project_id)
        if scope.restricted():          # the user's own extracts and those of projects he / she sees
            q = q.where(or_(db.bom_extracts.c.owner == scope.me(), db.bom_extracts.c.project_id.in_(select(db.projects.c.id).where(scope.project_clause()))))
        rows = conn.execute(q.order_by(db.bom_extracts.c.id.desc())).mappings().all()
        counts = dict(conn.execute(select(db.bom_extract_lines.c.extract_id, func.count()).group_by(db.bom_extract_lines.c.extract_id)).all())
    return [register._plain(r) | {"line_count": counts.get(r["id"], 0)} for r in rows]


def _clean_line(data: dict, kind: str, old: dict | None = None) -> dict:
    out = dict(old or line("other" if kind == "harness" else "component", "typed by you"))
    for k in FIELDS:
        if k in data:
            v = data[k]
            if k in ("qty", "length_mm", "cross_section_mm2"):
                v = None if v in (None, "") else _num(v)
                if v is not None and v < 0:
                    raise ExtractError(f"{k.replace('_', ' ')} must not be negative.")
                if k == "qty" and v is None:
                    raise ExtractError("Quantity must be a number.")
            else:
                v = str(v).strip() if v not in (None, "") else None
            out[k] = v
    cats = HARNESS_CATEGORIES if kind == "harness" else PCB_CATEGORIES
    if out["category"] not in cats:
        raise ExtractError(f"Category must be one of {', '.join(cats)}.")
    if out.get("mount") not in (None, "SMT", "THT"):
        raise ExtractError("Mount must be SMT or THT.")
    out |= {"confidence": 1.0, "flags": [], "edited": True}
    return out


def save_line(eid: int, line_id: int | None, data: dict) -> dict:
    from sqlalchemy import func, select
    from . import register
    db = _db()
    with db.engine().begin() as conn:
        e = conn.execute(select(db.bom_extracts).where(db.bom_extracts.c.id == eid)).mappings().first()
        if not e:
            raise ExtractError(f"BOM extract {eid} does not exist.")
        if line_id:
            r = conn.execute(select(db.bom_extract_lines).where((db.bom_extract_lines.c.id == line_id) & (db.bom_extract_lines.c.extract_id == eid))).mappings().first()
            if not r:
                raise ExtractError(f"Line {line_id} does not exist.")
            new = _clean_line(data, e["kind"], r["data"])
            conn.execute(db.bom_extract_lines.update().where(db.bom_extract_lines.c.id == line_id).values(data=new))
            register._log(conn, "bom_extract", eid, "line_edited", {k: {"old": r["data"].get(k), "new": new[k]} for k in FIELDS if r["data"].get(k) != new[k]})
        else:
            pos = (conn.execute(select(func.max(db.bom_extract_lines.c.position)).where(db.bom_extract_lines.c.extract_id == eid)).scalar() or 0) + 1
            new = _clean_line(data, e["kind"])
            conn.execute(db.bom_extract_lines.insert().values(extract_id=eid, position=pos, data=new))
            register._log(conn, "bom_extract", eid, "line_added", {"line": new.get("description") or new.get("part_number")})
        if e["status"] != "draft":
            conn.execute(db.bom_extracts.update().where(db.bom_extracts.c.id == eid).values(status="draft", reviewed_by=""))
    return get(eid)


def delete_line(eid: int, line_id: int) -> dict:
    from . import register
    db = _db()
    with db.engine().begin() as conn:
        n = conn.execute(db.bom_extract_lines.delete().where((db.bom_extract_lines.c.id == line_id) & (db.bom_extract_lines.c.extract_id == eid))).rowcount
        if not n:
            raise ExtractError(f"Line {line_id} does not exist.")
        conn.execute(db.bom_extracts.update().where(db.bom_extracts.c.id == eid).values(status="draft", reviewed_by=""))
        register._log(conn, "bom_extract", eid, "line_removed", {"line_id": line_id})
    return get(eid)


def set_status(eid: int, status: str) -> dict:
    from . import register
    db = _db()
    if status not in ("draft", "reviewed"):
        raise ExtractError("Status must be draft or reviewed.")
    with db.engine().begin() as conn:
        n = conn.execute(db.bom_extracts.update().where(db.bom_extracts.c.id == eid)
                         .values(status=status, reviewed_by=register.user() if status == "reviewed" else "")).rowcount
        if not n:
            raise ExtractError(f"BOM extract {eid} does not exist.")
        register._log(conn, "bom_extract", eid, "status", {"status": status})
    return get(eid)


def delete(eid: int) -> dict:
    db = _db()
    with db.engine().begin() as conn:
        if not conn.execute(db.bom_extracts.delete().where(db.bom_extracts.c.id == eid)).rowcount:
            raise ExtractError(f"BOM extract {eid} does not exist.")
    return {"deleted": eid}


def to_excel(eid: int) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from .table_set import HEADER_FILL
    e = get(eid)
    wb = Workbook()
    ws = wb.active
    ws.title = "BOM"
    head = ["Category", "Ref / designators", "Part number (MPN)", "Maker", "Description", "Qty", "Unit", "Length (mm)", "Wire type", "Cross-section (mm²)",
            "Colour", "Package", "Mount", "Notes", "Confidence", "Why not sure", "Source", "Edited"]
    ws.append(head)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), HEADER_FILL
    for l in e["lines"]:
        ws.append([l["category"], l["ref"], l["part_number"], l["maker"], l["description"], l["qty"], l["unit"], l["length_mm"], l["wire_type"],
                   l["cross_section_mm2"], l["colour"], l["package"], l["mount"], l["notes"], l["confidence"], "; ".join(l["flags"]), l["source"], "yes" if l["edited"] else ""])
        if l["confidence"] < LOW:
            for c in ws[ws.max_row]:
                c.fill = PatternFill("solid", fgColor="FFF2CC")
    s = e["summary"]
    ws2 = wb.create_sheet("Summary")
    for r in [["Extract", e["name"]], ["Kind", e["kind"]], ["Format", e["format"]], ["Status", e["status"]], ["Reviewed by", e["reviewed_by"]],
              ["Lines", s["lines"]], ["Lines to check (confidence < %.1f)" % LOW, s["low_confidence"]], ["Wire length (m)", s["wire_length_m"]]] + \
             [[f"Qty {k}", v] for k, v in s["by_category"].items()] + [[f"Components {k}", v] for k, v in s["by_mount"].items()]:
        ws2.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
