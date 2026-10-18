"""Versioned country rate library with approval.

Every rate (forex, electricity, labour, rent, interest, machine prices, ...)
is one row per version:

    category, item, country   what the rate is ("power", "electricity_industrial", "IN")
    value, unit, currency     the number and what it means
    month                     the month it is valid from, "YYYY-MM"
    source, source_url, fetched_date, comment, is_example
    status                    draft -> reviewed -> approved   (or rejected)
    approved_by, approved_date

Rules:
- Approved rows are never changed or deleted; a new value is a new row.
- Costing uses only approved rows: the latest approved version whose month is
  on or before the costing month.
- Forex rows: category "forex", item = currency code, country "" (global),
  value = INR per 1 unit of that currency.
"""

from __future__ import annotations

import io
import re
from contextlib import closing
from datetime import date

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from . import costing_db
from .table_set import EXAMPLE_FILL, HEADER_FILL

STATUSES = ("draft", "reviewed", "approved", "rejected")
OPEN = ("draft", "reviewed")
BASE_CURRENCY = "INR"
COLUMNS = ("category", "item", "country", "value", "unit", "currency", "month", "source", "source_url",
           "fetched_date", "status", "approved_by", "approved_date", "comment", "is_example", "origin", "created_at")
EDITABLE = ("value", "unit", "currency", "month", "source", "source_url", "fetched_date", "comment", "is_example")
MONTH = re.compile(r"^\d{4}-(0[1-9]|1[0-2])$")
DEFAULT_SETTINGS = {"approver_name": "", "review_threshold_pct": "5"}


class RateError(ValueError):
    pass


class RateMissing(RateError):
    pass


def _connect():
    con = costing_db.connect()
    con.execute(
        "CREATE TABLE IF NOT EXISTS rates (id INTEGER PRIMARY KEY, category TEXT NOT NULL, item TEXT NOT NULL, "
        "country TEXT NOT NULL DEFAULT '', value REAL NOT NULL, unit TEXT, currency TEXT, month TEXT NOT NULL, "
        "source TEXT, source_url TEXT, fetched_date TEXT, status TEXT NOT NULL, approved_by TEXT, "
        "approved_date TEXT, comment TEXT, is_example INTEGER DEFAULT 0, origin TEXT, created_at TEXT)"
    )
    con.execute("CREATE INDEX IF NOT EXISTS rates_key ON rates (category, item, country, month)")
    return con


def current_month() -> str:
    return date.today().strftime("%Y-%m")


def _row(r) -> dict:
    d = dict(r)
    d["is_example"] = bool(d["is_example"])
    v = d["value"]
    d["value"] = int(v) if float(v).is_integer() else v
    for k in COLUMNS:
        if d.get(k) is None and k not in ("value", "is_example"):
            d[k] = ""
    return d


# ---------- settings ----------

def settings() -> dict:
    with closing(costing_db.connect()) as con:
        stored = dict(con.execute("SELECT key, value FROM settings").fetchall())
    return DEFAULT_SETTINGS | stored


def save_settings(values: dict) -> dict:
    clean = {}
    for k, v in values.items():
        if k not in DEFAULT_SETTINGS:
            raise RateError(f"Unknown setting '{k}'.")
        v = str(v).strip()
        if k == "review_threshold_pct":
            try:
                if float(v) < 0:
                    raise ValueError
            except ValueError:
                raise RateError("Review threshold must be a number of 0 or more (%).") from None
        clean[k] = v
    with closing(costing_db.connect()) as con, con:
        con.executemany("INSERT INTO settings (key, value) VALUES (?, ?) "
                        "ON CONFLICT(key) DO UPDATE SET value = excluded.value", clean.items())
    return settings()


# ---------- validation ----------

def _clean(data: dict, where: str = "Rate") -> dict:
    out = {}
    for k in ("category", "item"):
        out[k] = str(data.get(k) or "").strip()
        if not out[k]:
            raise RateError(f"{where}: {k} must not be empty.")
    out["country"] = str(data.get("country") or "").strip().upper()
    v = data.get("value")
    try:
        out["value"] = float(str(v).replace(",", "").strip())
    except (TypeError, ValueError):
        raise RateError(f"{where}: value '{v}' is not a number.") from None
    if out["value"] < 0:
        raise RateError(f"{where}: value must not be negative.")
    out["month"] = str(data.get("month") or current_month()).strip()
    if not MONTH.match(out["month"]):
        raise RateError(f"{where}: month '{out['month']}' must look like 2026-10.")
    out["currency"] = str(data.get("currency") or "").strip().upper()
    if out["currency"] and not re.match(r"^[A-Z]{3}$", out["currency"]):
        raise RateError(f"{where}: currency '{out['currency']}' must be a 3-letter code such as INR or USD.")
    for k in ("unit", "source", "source_url", "fetched_date", "comment"):
        out[k] = str(data.get(k) or "").strip()
    out["is_example"] = bool(data.get("is_example"))
    return out


# ---------- reading ----------

def get(rate_id: int) -> dict:
    with closing(_connect()) as con:
        r = con.execute("SELECT * FROM rates WHERE id = ?", (rate_id,)).fetchone()
    if not r:
        raise RateError(f"Rate {rate_id} does not exist.")
    return _row(r)


def lookup(category: str, item: str, country: str = "", as_of: str | None = None) -> dict:
    """Latest approved version on or before the month; raises RateMissing."""
    as_of = as_of or current_month()
    with closing(_connect()) as con:
        r = con.execute(
            "SELECT * FROM rates WHERE category = ? AND item = ? AND country = ? AND status = 'approved' "
            "AND month <= ? ORDER BY month DESC, approved_date DESC, id DESC LIMIT 1",
            (category, item, country, as_of),
        ).fetchone()
    if not r:
        where = f" in {country}" if country else ""
        raise RateMissing(f"No approved rate '{category} / {item}'{where} for {as_of} or earlier. "
                          f"Add or approve one on the Rates page.")
    return _row(r)


def history(category: str, item: str, country: str = "") -> list[dict]:
    with closing(_connect()) as con:
        rows = con.execute("SELECT * FROM rates WHERE category = ? AND item = ? AND country = ? "
                           "ORDER BY month, id", (category, item, country)).fetchall()
    return [_row(r) for r in rows]


def current(as_of: str | None = None, country: str | None = None, category: str | None = None) -> list[dict]:
    """One row per rate key: its approved value for the month, plus any open draft."""
    as_of = as_of or current_month()
    with closing(_connect()) as con:
        keys = con.execute("SELECT DISTINCT category, item, country FROM rates ORDER BY category, item, country").fetchall()
        out = []
        for k in keys:
            if country is not None and k["country"] not in (country, ""):
                continue
            if category and k["category"] != category:
                continue
            approved = con.execute(
                "SELECT * FROM rates WHERE category = ? AND item = ? AND country = ? AND status = 'approved' "
                "AND month <= ? ORDER BY month DESC, approved_date DESC, id DESC LIMIT 1", tuple(k) + (as_of,)
            ).fetchone()
            pending = con.execute(
                "SELECT COUNT(*) FROM rates WHERE category = ? AND item = ? AND country = ? AND status IN "
                "('draft', 'reviewed')", tuple(k)).fetchone()[0]
            out.append({"category": k["category"], "item": k["item"], "country": k["country"],
                        "approved": _row(approved) if approved else None, "open_drafts": pending})
        return out


def keys_used(trace_rates: list[dict]) -> list[dict]:
    """Compact list of rate versions a calculation used (for reproducible reports)."""
    return [{"id": r["id"], "category": r["category"], "item": r["item"], "country": r["country"],
             "month": r["month"], "value": r["value"]} for r in trace_rates]


# ---------- writing (append-only for approved rows) ----------

def create(data: dict, approve: bool = False, origin: str = "manual") -> dict:
    row = _clean(data)
    now = date.today().isoformat()
    row |= {"status": "draft", "approved_by": "", "approved_date": "", "origin": origin, "created_at": now}
    if approve:
        row |= {"status": "approved", "approved_by": _approver(), "approved_date": now}
    cols = list(row)
    with closing(_connect()) as con, con:
        cur = con.execute(f"INSERT INTO rates ({', '.join(cols)}) VALUES ({', '.join('?' * len(cols))})",
                          [int(row[c]) if isinstance(row[c], bool) else row[c] for c in cols])
        new_id = cur.lastrowid
    return get(new_id)


def update(rate_id: int, data: dict, approve: bool = False) -> dict:
    old = get(rate_id)
    if old["status"] not in OPEN:
        raise RateError(f"Rate {rate_id} is {old['status']} and cannot be changed. Save a new version instead.")
    merged = _clean({**old, **{k: v for k, v in data.items() if k in EDITABLE}})
    with closing(_connect()) as con, con:
        con.execute(f"UPDATE rates SET {', '.join(f'{k} = ?' for k in EDITABLE)} WHERE id = ?",
                    [int(merged[k]) if isinstance(merged[k], bool) else merged[k] for k in EDITABLE] + [rate_id])
    if approve:
        set_status([rate_id], "approved")
    return get(rate_id)


def _approver() -> str:
    name = settings()["approver_name"].strip()
    if not name:
        raise RateError("Set your name for approvals first (Rates → Rate review → Approver name).")
    return name


def set_status(ids: list[int], status: str, comment: str = "") -> list[dict]:
    if status not in ("reviewed", "approved", "rejected"):
        raise RateError(f"Status must be reviewed, approved or rejected, not '{status}'.")
    rows = [get(i) for i in ids]
    for r in rows:
        if r["status"] not in OPEN:
            raise RateError(f"Rate {r['id']} ({r['category']} / {r['item']} {r['country']}) is already {r['status']}.")
    approver = _approver() if status == "approved" else ""
    today = date.today().isoformat()
    with closing(_connect()) as con, con:
        for r in rows:
            note = "\n".join(x for x in (r["comment"], comment.strip()) if x)
            if status == "approved":
                con.execute("UPDATE rates SET status = ?, approved_by = ?, approved_date = ?, comment = ? WHERE id = ?",
                            (status, approver, today, note, r["id"]))
            else:
                con.execute("UPDATE rates SET status = ?, comment = ? WHERE id = ?", (status, note, r["id"]))
    return [get(i) for i in ids]


def review() -> dict:
    """Open drafts next to the approved value they would replace."""
    threshold = float(settings()["review_threshold_pct"] or 0)
    with closing(_connect()) as con:
        drafts = con.execute("SELECT * FROM rates WHERE status IN ('draft', 'reviewed') "
                             "ORDER BY category, item, country, month").fetchall()
    out = []
    for d in map(_row, drafts):
        try:
            prev = lookup(d["category"], d["item"], d["country"], d["month"])
        except RateMissing:
            prev = None
        change = None
        if prev and prev["value"]:
            change = round((d["value"] - prev["value"]) / prev["value"] * 100, 2)
        out.append({"draft": d, "approved": prev, "change_pct": change,
                    "over_threshold": change is not None and abs(change) > threshold})
    return {"threshold_pct": threshold, "items": out}


# ---------- Excel ----------

def to_excel() -> bytes:
    with closing(_connect()) as con:
        rows = [_row(r) for r in con.execute("SELECT * FROM rates ORDER BY category, item, country, month, id")]
    wb = Workbook()
    ws = wb.active
    ws.title = "Rates"
    wb.properties.title = "Design to Cost — Country rate library (full history)"
    wb.properties.creator = "Design to Cost"
    keys = ("id",) + COLUMNS
    ws.append(list(keys))
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
    for r in rows:
        ws.append(["yes" if r[k] is True else "no" if r[k] is False else r[k] for k in keys])
        if r["is_example"]:
            for cell in ws[ws.max_row]:
                cell.fill = EXAMPLE_FILL
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def import_excel(content: bytes) -> list[dict]:
    """Every row of the upload becomes a new DRAFT; nothing existing is changed."""
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise RateError(f"Could not read the uploaded file as Excel (.xlsx): {exc}") from exc
    ws = wb["Rates"] if "Rates" in wb.sheetnames else wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        raise RateError("The Excel file is empty.")
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    for need in ("category", "item", "value"):
        if need not in header:
            raise RateError(f"The Rates sheet needs a '{need}' column.")
    records = []
    for i, row in enumerate(rows[1:], start=2):
        if not any(v not in (None, "") for v in row):
            continue
        data = {h: row[j] for j, h in enumerate(header) if h and j < len(row)}
        data["is_example"] = str(data.get("is_example", "")).lower() in ("1", "true", "yes")
        records.append(_clean(data, f"Rates row {i}"))
    return [create(r, origin="excel import") for r in records]
