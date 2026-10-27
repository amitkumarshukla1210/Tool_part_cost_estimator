"""Process parameters for the best part quality (roadmap stage 34): the setting sheet for the first trial, worked out from the
simulation of the part (stage 6) or the sheet-metal forces (stage 9) and the material.

    window   Mould & die database → Process windows: melt / pour / preform temperature, mould or die temperature, ejection
             temperature and drying of each material, with the range
    rules    Mould & die database → Process parameters: one row per setting (formula, min, max, unit, why). The formulas use the
             simulation result (fill and cooling time, flow ratio, injection pressure, clamp force …), its quality maps (thin,
             thick / sink-risk, low-draft and undercut areas), the material window, Parameters and the numeric Design rules;
             sheet metal: the stage forces, draw ratio, springback. `risk` rows (result above 0) are quality warnings with the remedy.
The simulation itself (its rules and results) is not changed: the setting sheet uses the material window on top of it.
All seeded rows are EXAMPLE values until you replace them with your moulding / casting experience.
"""

from __future__ import annotations

import io

from . import formula, master_data, mould_db, register, register_setup

EXAMPLE = "EXAMPLE - typical value: replace with your own data sheet or shop experience"
WINDOW_FIELDS = ("melt_c", "melt_min_c", "melt_max_c", "mould_c", "mould_min_c", "mould_max_c", "eject_c", "dry_c", "dry_h")
KINDS = ("parameter", "risk")

# material, melt, min, max, mould, min, max, eject, dry °C, dry h
SEED_WINDOWS = [
    ("PP (polypropylene)", 230, 200, 260, 40, 20, 60, 90, 0, 0), ("ABS", 235, 210, 260, 60, 40, 80, 90, 80, 3), ("PC (polycarbonate)", 295, 270, 320, 90, 80, 110, 120, 120, 4),
    ("PC/ABS", 260, 240, 280, 70, 60, 90, 100, 100, 3), ("PA6", 260, 240, 290, 80, 60, 95, 120, 80, 4), ("PA66 GF30", 290, 275, 305, 85, 70, 100, 150, 80, 4),
    ("PBT GF30", 260, 245, 275, 80, 60, 100, 140, 120, 3), ("POM (acetal)", 205, 190, 230, 90, 80, 110, 140, 0, 0), ("PET preform grade", 280, 270, 295, 10, 5, 15, 70, 160, 4),
    ("Aluminium ADC12 / A383", 670, 650, 690, 220, 180, 260, 400, 0, 0), ("Aluminium A380", 670, 650, 690, 220, 180, 260, 400, 0, 0), ("Zinc Zamak 3", 420, 400, 440, 180, 150, 220, 250, 0, 0),
    ("Magnesium AZ91D", 650, 630, 670, 220, 180, 260, 350, 0, 0),
    ("Grey cast iron FG260", 1380, 1340, 1420, 20, 10, 40, 0, 0, 0), ("SG / ductile iron", 1400, 1360, 1440, 20, 10, 40, 0, 0, 0), ("Cast steel", 1580, 1540, 1620, 20, 10, 40, 0, 0, 0),
    ("Aluminium LM6 / A356", 720, 690, 750, 20, 10, 40, 0, 0, 0), ("Bronze", 1150, 1100, 1200, 20, 10, 40, 0, 0, 0),
    ("EPS bead grade S (20 kg/m3)", 115, 108, 125, 40, 30, 50, 0, 0, 0), ("EPS bead grade F, flame retardant (25 kg/m3)", 115, 108, 125, 40, 30, 50, 0, 0, 0),
    ("PET bottle grade (IV 0.80)", 105, 95, 115, 12, 8, 15, 60, 160, 4), ("rPET food grade", 105, 95, 115, 12, 8, 15, 60, 160, 4),
    # materials added to Master data → Materials later (stage 29); typical values, EXAMPLE
    ("PP copolymer (impact)", 230, 200, 260, 40, 20, 60, 90, 0, 0), ("PP talc filled 20 %", 230, 210, 260, 45, 30, 60, 100, 0, 0),
    ("ABS flame retardant V0", 220, 200, 240, 55, 40, 70, 85, 80, 3), ("PA66 GF30 heat stabilised", 290, 275, 305, 85, 70, 100, 150, 80, 4),
    ("Aluminium AlSi10Mg (EN AC-43400)", 680, 660, 700, 220, 180, 260, 400, 0, 0),
]
# part type, name, formula, min, max, unit, label, why, kind
P = "parameter"
SEED_RULES = [
    ("plastic", "thin_share", "thin_cm2 / max(surface_cm2, 1)", "", "", "", "Share of the surface that is thin (short-shot risk)", "from the wall-thickness map of the simulation", P),
    ("plastic", "sink_share", "sink_risk_cm2 / max(surface_cm2, 1)", "", "", "", "Share of the surface that is thick (sink-mark risk)", "from the wall-thickness map of the simulation", P),
    ("plastic", "melt_temp_c", "min(melt_max_c, max(melt_min_c, melt_c + 10 * (flow_ratio > 120) + 10 * (thin_share > 0.02) - 10 * (sink_share > 0.02) * (thin_share <= 0.02)))",
     "melt_min_c", "melt_max_c", "°C", "Melt temperature", "long flow path or thin walls: hotter melt fills; thick sections: cooler melt shrinks less", P),
    ("plastic", "mould_temp_c", "min(mould_max_c, max(mould_min_c, mould_c + 10 * (flow_ratio > 120) + 5 * (thin_share > 0.02) - 5 * (sink_share > 0.02)))",
     "mould_min_c", "mould_max_c", "°C", "Mould temperature", "warmer mould keeps thin, long flow open; cooler mould freezes thick sections sooner", P),
    ("plastic", "fill_time_s", "fill_s", "fill_s * 0.7", "fill_s * 1.3", "s", "Injection (fill) time", "fill time of the simulation for the shot weight", P),
    ("plastic", "shot_volume_cm3", "shot_weight_g / max(density_g_cm3, 0.1)", "", "", "cm³", "Shot volume of the parts", "shot weight / material density", P),
    ("plastic", "dosing_volume_cm3", "shot_volume_cm3 * (1 + runner_pct / 100)", "", "", "cm³", "Dosing volume with runner", "shot volume + runner share (Parameters → runner_pct)", P),
    ("plastic", "injection_rate_cm3_s", "shot_volume_cm3 / fill_s", "shot_volume_cm3 / fill_s * 0.7", "shot_volume_cm3 / fill_s * 1.3", "cm³/s", "Injection rate", "shot volume / fill time", P),
    ("plastic", "injection_pressure_bar", "injection_bar", "injection_bar * 0.8", "injection_bar * 1.2", "bar", "Injection pressure limit", "flow ratio of the simulation (flow length / wall)", P),
    ("plastic", "hold_pressure_bar", "injection_bar * (0.55 + min(0.2, 5 * sink_share))", "injection_bar * 0.4", "injection_bar * 0.85", "bar", "Hold (pack) pressure",
     "more packing where thick sections can sink", P),
    ("plastic", "hold_time_s", "cooling_s * (0.45 + min(0.25, 5 * sink_share))", "cooling_s * 0.3", "cooling_s * 0.8", "s", "Hold (pack) time", "until the gate freezes: longer for thick sections", P),
    ("plastic", "cooling_time_s", "cooling_s", "cooling_s * 0.9", "cooling_s * 1.3", "s", "Cooling time", "thickest wall cooled to the ejection temperature", P),
    ("plastic", "back_pressure_bar", "6", "3", "12", "bar", "Back pressure", "even melt without degrading the material", P),
    ("plastic", "clamp_force_t", "clamp_t", "clamp_t", "clamp_t * 1.3", "t", "Clamp force", "projected area × cavities × cavity pressure with the safety factor", P),
    ("plastic", "ejection_temp_c", "eject_c", "", "", "°C", "Part temperature at ejection", "rigid enough to eject without deforming", P),
    ("plastic", "drying_temp_c", "dry_c", "", "", "°C", "Drying temperature (0 = no drying)", "material window", P),
    ("plastic", "drying_time_h", "dry_h", "", "", "h", "Drying time", "material window", P),
    ("plastic", "cycle_time_s", "cycle_s", "", "", "s", "Cycle time", "fill + cooling + open / eject (+ insert loading)", P),
    ("plastic", "risk_sink", "sink_share > 0.02", "", "", "", "Sink marks likely on {sink_risk_cm2:.1f} cm²: core out the thick sections, raise hold pressure and time, cool the mould there.", "", "risk"),
    ("plastic", "risk_short", "thin_share > 0.02", "", "", "", "Short shot likely on {thin_cm2:.1f} cm² of thin wall: faster injection, hotter melt and mould, wider gate.", "", "risk"),
    ("plastic", "risk_flow", "flow_ratio > 250", "", "", "", "Flow ratio {flow_ratio:.0f} is very long: add a second gate or thicken the wall.", "", "risk"),
    ("plastic", "risk_draft", "draft_low_cm2 > 0", "", "", "", "{draft_low_cm2:.1f} cm² of wall has too little draft: scuffing and high ejection force, add draft.", "", "risk"),
    ("plastic", "risk_undercut", "undercut_cm2 > 0", "", "", "", "Undercuts ({undercut_cm2:.1f} cm²) need slides or lifters: watch for flash and witness lines.", "", "risk"),
    ("die_casting", "thin_share", "thin_cm2 / max(surface_cm2, 1)", "", "", "", "Share of the surface that is thin", "from the wall-thickness map", P),
    ("die_casting", "sink_share", "sink_risk_cm2 / max(surface_cm2, 1)", "", "", "", "Share of the surface that is thick (porosity risk)", "from the wall-thickness map", P),
    ("die_casting", "metal_temp_c", "min(melt_max_c, max(melt_min_c, melt_c + 10 * (thin_share > 0.02) - 10 * (sink_share > 0.02) * (thin_share <= 0.02)))", "melt_min_c", "melt_max_c", "°C",
     "Metal (holding furnace) temperature", "thin walls need hotter metal; thick sections shrink less when cooler", P),
    ("die_casting", "die_temp_c", "min(mould_max_c, max(mould_min_c, mould_c + 10 * (thin_share > 0.02) - 10 * (sink_share > 0.02)))", "mould_min_c", "mould_max_c", "°C", "Die temperature", "balance of fill and solidification", P),
    ("die_casting", "fast_shot_speed_m_s", "min(6, max(1.5, 4.5 - 0.7 * max_wall_mm))", "1.5", "6", "m/s", "Fast-shot plunger speed", "thinner walls need a faster fill", P),
    ("die_casting", "slow_shot_speed_m_s", "0.3", "0.2", "0.5", "m/s", "Slow-shot plunger speed", "pushes the metal to the gate without trapping air", P),
    ("die_casting", "fill_time_s", "fill_s", "fill_s * 0.7", "fill_s * 1.3", "s", "Cavity fill time", "simulation", P),
    ("die_casting", "intensification_pressure_bar", "cavity_pressure_bar * 1.0", "cavity_pressure_bar * 0.8", "cavity_pressure_bar * 1.25", "bar", "Intensification (squeeze) pressure", "cavity pressure of the part type (Parameters)", P),
    ("die_casting", "dwell_time_s", "solidify_s", "solidify_s * 0.8", "solidify_s * 1.3", "s", "Dwell (solidification) time", "thickest wall solidified before opening", P),
    ("die_casting", "spray_time_s", "spray_s", "", "", "s", "Die spray time", "die lubricant and cooling", P),
    ("die_casting", "clamp_force_t", "clamp_t", "clamp_t", "clamp_t * 1.3", "t", "Clamp force", "projected area with overflow × cavity pressure", P),
    ("die_casting", "cycle_time_s", "cycle_s", "", "", "s", "Cycle time", "ladle + fill + solidify + spray + open / eject", P),
    ("die_casting", "risk_porosity", "sink_share > 0.02", "", "", "", "Porosity likely in {sink_risk_cm2:.1f} cm² of thick section: core it out, raise intensification pressure, cool the die there.", "", "risk"),
    ("die_casting", "risk_cold_shut", "thin_share > 0.02", "", "", "", "Cold shuts / misruns likely on {thin_cm2:.1f} cm² of thin wall: faster shot, hotter die, better venting.", "", "risk"),
    ("die_casting", "risk_draft", "draft_low_cm2 > 0", "", "", "", "{draft_low_cm2:.1f} cm² of wall has too little draft: soldering and drag marks, add draft.", "", "risk"),
    ("sheet_metal", "press_tonnage_t", "force_t * press_safety_factor", "force_t * press_safety_factor", "force_t * press_safety_factor * 1.25", "t", "Press tonnage to set / needed",
     "biggest stage force (total of all stations for a progressive die) with the press safety factor", P),
    ("sheet_metal", "blank_holder_force_t", "0.3 * force_t if draw_depth > 0 else 0", "0.2 * force_t if draw_depth > 0 else 0", "0.4 * force_t if draw_depth > 0 else 0", "t", "Blank-holder (cushion) force",
     "about 30 % of the draw force: wrinkles below, tears above", P),
    ("sheet_metal", "strokes_per_min", "60 / max(cycle_time_s, 0.5) * (0.6 if draw_depth > 0 else 1)", "60 / max(cycle_time_s, 0.5) * 0.5", "60 / max(cycle_time_s, 0.5) * 1.2", "1/min",
     "Press speed", "from the planned cycle time (Parameters → cycle_time_s); slower for deep draws", P),
    ("sheet_metal", "die_clearance_mm", "cut_clearance_pct / 100 * t", "cut_clearance_pct / 100 * t * 0.8", "cut_clearance_pct / 100 * t * 1.2", "mm/side", "Die clearance per side", "Design rules → cut_clearance_pct of the thickness", P),
    ("sheet_metal", "springback_comp_deg", "springback_deg", "springback_deg * 0.8", "springback_deg * 1.2", "°", "Overbend to compensate springback", "from the sheet-material rules", P),
    ("sheet_metal", "draw_speed_mm_s", "min(400, max(100, 600 / (1 + draw_ratio)))", "100", "400", "mm/s", "Draw speed at the punch", "deeper draws and higher draw ratios need slower speed", P),
    ("sheet_metal", "risk_draw", "draw_ratio > ldr", "", "", "", "Draw ratio {draw_ratio:.2f} is above the limit {ldr:.2f} of the material: use more draw stages or anneal.", "", "risk"),
    ("sheet_metal", "risk_springback", "springback_deg > max_springback_deg", "", "", "", "Springback {springback_deg:.1f}° is above the allowed {max_springback_deg:.1f}°: overbend or restrike.", "", "risk"),
    ("sheet_metal", "risk_force", "progressive_ok < 1 and progressive > 0", "", "", "", "A progressive die is not recommended for this part by the sheet rules: check the strip layout and feed.", "", "risk"),
    ("foundry", "pour_temp_c", "melt_c", "melt_min_c", "melt_max_c", "°C", "Pouring temperature", "material window (thin sections: upper half)", P),
    ("foundry", "poured_weight_kg", "part_weight_kg * cavities * 1.4", "", "", "kg", "Metal poured per mould", "castings plus about 40 % gating and riser (EXAMPLE share: your yield)", P),
    ("foundry", "pour_time_s", "ceil(2.5 * sqrt(poured_weight_kg) * 10) / 10", "2 * sqrt(poured_weight_kg)", "3.5 * sqrt(poured_weight_kg)", "s", "Pouring time", "weight rule: fast enough to avoid cold laps, slow enough to avoid erosion", P),
    ("foundry", "solidification_min", "3 * (part_modulus_mm / 10) ** 2", "", "", "min", "Solidification time (Chvorinov)", "modulus (volume / surface) squared × mould constant", P),
    ("foundry", "shakeout_min", "ceil(1.5 * 3 * (part_modulus_mm / 10) ** 2)", "", "", "min", "Earliest shake-out after pouring", "casting solid and below the critical temperature", P),
    ("foundry", "mould_temp_c", "mould_c", "mould_min_c", "mould_max_c", "°C", "Sand / mould temperature", "material window", P),
    ("foundry", "risk_shrink", "part_modulus_mm > 8", "", "", "", "Heavy section (modulus {part_modulus_mm:.1f} mm): riser and chills are needed against shrinkage porosity.", "", "risk"),
    ("eps", "steam_pressure_bar", "min(1.3, max(0.7, 0.9 + 0.01 * (wall_mm - 20)))", "0.7", "1.3", "bar", "Steam pressure in the chest", "thicker walls need a little more steam to fuse the beads", P),
    ("eps", "steam_time_s", "8 + 0.5 * wall_mm", "6 + 0.4 * wall_mm", "10 + 0.6 * wall_mm", "s", "Steaming time", "fuse the beads through the wall", P),
    ("eps", "cooling_time_s", "15 + 0.04 * wall_mm ** 2", "10", "30 + 0.06 * wall_mm ** 2", "s", "Cooling (vacuum) time", "until the part holds its shape on ejection", P),
    ("eps", "bead_density_kg_m3", "density_g_cm3 * 1000", "", "", "kg/m³", "Pre-expanded bead density", "the material grade", P),
    ("eps", "risk_thick", "wall_mm > 60", "", "", "", "Wall {wall_mm:.0f} mm is thick: steam may not reach the core, use steam-through cores.", "", "risk"),
    ("blow", "preform_temp_c", "melt_c", "melt_min_c", "melt_max_c", "°C", "Preform temperature at the blow station", "material window", P),
    ("blow", "pre_blow_bar", "8", "5", "12", "bar", "Pre-blow pressure", "forms the bubble before the stretch rod bottoms", P),
    ("blow", "blow_pressure_bar", "min(40, 20 + 0.05 * part_h_mm)", "18", "40", "bar", "High blow pressure", "taller containers need more pressure to take the mould detail", P),
    ("blow", "mould_temp_c", "mould_c", "mould_min_c", "mould_max_c", "°C", "Blow-mould water temperature", "material window", P),
    ("blow", "mould_cooling_s", "1.2 + 4 * wall_mm", "1 + 3 * wall_mm", "2 + 5 * wall_mm", "s", "Cooling time in the mould", "thicker walls hold heat longer", P),
    ("blow", "risk_thin", "wall_mm < 0.3", "", "", "", "Wall {wall_mm:.2f} mm is very thin: top-load and burst strength are at risk.", "", "risk"),
]
SAMPLE = {"plastic": ("thin_cm2", "sink_risk_cm2", "surface_cm2", "draft_low_cm2", "undercut_cm2", "fill_s", "cooling_s", "cycle_s", "flow_ratio",
                      "injection_bar", "cavity_bar", "clamp_t", "shot_weight_g", "wall_mm", "max_wall_mm", "min_wall_mm", "cavities"),
          "die_casting": ("thin_cm2", "sink_risk_cm2", "surface_cm2", "draft_low_cm2", "undercut_cm2", "fill_s", "solidify_s", "spray_s", "cycle_s",
                          "clamp_t", "wall_mm", "max_wall_mm", "min_wall_mm", "cavities", "shot_weight_g"),
          "sheet_metal": ("t", "force_t", "total_force_t", "draw_ratio", "draw_depth", "springback_deg", "progressive", "progressive_ok", "ldr"),
          "foundry": ("part_modulus_mm", "part_weight_kg", "cavities", "wall_mm", "max_wall_mm"),
          "eps": ("wall_mm", "part_h_mm", "max_wall_mm"),
          "blow": ("wall_mm", "part_h_mm", "max_wall_mm")}
WINDOW_SAMPLE = {"melt_c": 230, "melt_min_c": 200, "melt_max_c": 260, "mould_c": 50, "mould_min_c": 30, "mould_max_c": 70, "eject_c": 90,
                 "dry_c": 80, "dry_h": 3, "density_g_cm3": 1.0}


def seed() -> dict:
    """Add the missing Process windows and Process parameters rows (existing rows are never changed)."""
    d = mould_db.tables()
    out = {"windows": 0, "rules": 0}
    have = {r["material"] for r in d["process_windows"]}
    for m, *vals in SEED_WINDOWS:
        if m not in have:
            d["process_windows"].append(dict(zip(("material",) + WINDOW_FIELDS, (m, *vals))) | {"is_example": True, "notes": EXAMPLE})
            out["windows"] += 1
    have = {(r["part_type"], r["name"]) for r in d["process_parameters"]}
    step = {}
    for pt, name, f, lo, hi, unit, label, why, kind in SEED_RULES:
        step[pt] = step.get(pt, 0) + 1
        if (pt, name) not in have:
            d["process_parameters"].append({"part_type": pt, "step": step[pt], "name": name, "formula": f, "min_formula": lo, "max_formula": hi,
                                            "unit": unit, "label": label, "why": why, "kind": kind, "is_example": True, "notes": EXAMPLE})
            out["rules"] += 1
    if any(out.values()):
        mould_db.write(d)
    return out


def window(material: str) -> dict:
    row = next((r for r in mould_db.tables()["process_windows"] if r["material"] == material), None)
    return {k: float(row[k] or 0) for k in WINDOW_FIELDS} if row else {}


def _rules_env(part_type: str) -> dict:
    env = {k: float(v) for k, v in master_data.parameters(part_type).items() if isinstance(v, (int, float))}
    env |= {k: float(r["value"]) for k, r in mould_db.design_rules(part_type).items() if r["value"] is not None}
    return env


def sample_env(part_type: str) -> dict:
    return {k: 1.0 for k in SAMPLE.get(part_type, ())} | WINDOW_SAMPLE | _rules_env(part_type)


def validate(clean: dict) -> None:
    rows = clean.get("process_parameters", [])
    for pt in {r["part_type"] for r in rows}:
        env = sample_env(pt)
        for r in sorted((x for x in rows if x["part_type"] == pt), key=lambda x: x["step"]):
            if not r["name"].isidentifier():
                raise mould_db.MouldDbError(f"Process parameters '{r['name']}' must be one word (letters, digits and _).")
            if r["kind"] not in KINDS:
                raise mould_db.MouldDbError(f"Process parameters ({r['name']}): kind must be parameter or risk.")
            for col in ("formula", "min_formula", "max_formula"):
                if r.get(col):
                    try:
                        v = formula.evaluate(r[col], env)
                        if col == "formula":
                            env[r["name"]] = float(v)
                    except formula.FormulaError as exc:
                        raise mould_db.MouldDbError(f"Process parameters ({r['name']}, {col}): {exc}") from None
                    except (ArithmeticError, ValueError):
                        env[r["name"]] = 1.0


# ---------- the inputs of a part ----------

def _env(part_id: int) -> tuple[str, dict, str, dict]:
    """(part type, rule environment, where the numbers come from, part) of the part's current revision."""
    d = register.part_detail(part_id, with_rm=False)
    pt = (register_setup.processes().get(d["process"]) or {}).get("part_type")
    if pt not in SAMPLE:
        raise register.ToolNotApplicable("Process parameters are worked out for moulded, cast, sheet-metal, EPS and blow-moulded parts.")
    density, _ = register._density(d["material"])
    env = _rules_env(pt) | {"density_g_cm3": float(density or 1.0)}
    win = window(d["material"])
    if pt == "sheet_metal":
        res = register.sheet_tools_result(part_id)
        ops = res.get("operations") or []
        prog = res.get("mode") == "progressive"
        env |= {k: float(v) for k, v in (res.get("features") or {}).items() if isinstance(v, (int, float))}
        env |= {k: float(v) for k, v in (res.get("values") or {}).items() if isinstance(v, (int, float))}
        env |= {k: float(v) for k, v in (res.get("material_values") or {}).items() if isinstance(v, (int, float))}
        env |= {"force_t": float(res.get("total_force_t") if prog else max((o["force_t"] for o in ops), default=0)),
                "total_force_t": float(res.get("total_force_t") or 0), "progressive": 1.0 if prog else 0.0}
        return pt, env | win, "stage forces, draw ratio and springback of the sheet-metal tools and the sheet material", d
    _, rev, res = register._simulate(part_id)
    env |= {k: float(v) for k, v in res["results"].items() if isinstance(v, (int, float))}
    env |= {k: float(v) for k, v in res["maps"].items()}
    g = rev.get("geometry") or {}
    env |= {"part_h_mm": float(max((g.get("bbox_axes_mm") or {"x": 0}).values())),
            "part_weight_kg": float(res["results"].get("shot_weight_g", 0)) / max(float(res["results"].get("cavities", 1) or 1), 1) / 1000,
            "part_modulus_mm": float(g.get("volume_cm3") or 0) * 1000 / max(float(g.get("surface_area_mm2") or 0), 1.0)}
    env |= win                                   # the material window wins over the simulation's generic temperatures
    src = "simulation of the part (wall thickness, flow, cooling, clamp force), its quality maps and the material window"
    return pt, env, src + ("" if win else " (no Process windows row for this material: the simulation's temperatures are used)"), d


def compute(part_id: int) -> dict:
    pt, env, source, d = _env(part_id)
    rows = sorted((r for r in mould_db.tables()["process_parameters"] if r["part_type"] == pt), key=lambda r: r["step"])
    if not rows:
        raise register.ToolNotApplicable(f"Mould & die database → Process parameters has no rows for '{pt}'.")
    for k in ("melt_c", "melt_min_c", "melt_max_c", "mould_c", "mould_min_c", "mould_max_c", "eject_c", "dry_c", "dry_h"):
        env.setdefault(k, float(env.get(k.replace("_min", "").replace("_max", ""), 0)))
    params, risks, derived = [], [], dict(env)
    for r in rows:
        try:
            v = float(formula.evaluate(r["formula"], derived))
            lo = float(formula.evaluate(r["min_formula"], derived)) if r["min_formula"] else None
            hi = float(formula.evaluate(r["max_formula"], derived)) if r["max_formula"] else None
        except (formula.FormulaError, ArithmeticError, ValueError) as exc:
            raise register.RegisterError(f"Process parameter '{r['name']}' ({r['formula']}): {exc}") from None
        derived[r["name"]] = v
        if r["kind"] == "risk":
            if v > 0:
                try:
                    risks.append(r["label"].format(**derived))
                except (KeyError, ValueError, IndexError):
                    risks.append(r["label"])
            continue
        if r["unit"] == "" and not r["min_formula"] and r["name"].endswith("_share"):
            continue                       # helper values that only feed other rows
        params.append({"name": r["name"], "label": r["label"], "value": round(v, 3), "min": None if lo is None else round(lo, 3),
                       "max": None if hi is None else round(hi, 3), "unit": r["unit"], "why": r["why"], "is_example": bool(r["is_example"])})
    return {"part_id": part_id, "part_code": d["part_code"], "part_type": pt, "material": d["material"], "process": d["process"], "source": source,
            "parameters": params, "risks": risks, "window": window(d["material"]), "is_example": any(p["is_example"] for p in params),
            "note": "Starting values for the first trial, worked out from the simulation of the part. Confirm them in the try-out and keep "
                    "what works in Mould & die database → Process parameters / Process windows."}


# ---------- downloads ----------

def xlsx(part_id: int) -> tuple[bytes, str]:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    res = compute(part_id)
    wb = Workbook()
    ws = wb.active
    ws.title = "Process parameters"
    if res["is_example"]:
        ws.append(["Uses EXAMPLE process windows / rules - confirm in the try-out"])
    ws.append([f"{res['part_code']} · {res['material']} · {res['process']}"])
    ws.append([f"Worked out from {res['source']}."])
    ws.append([])
    ws.append(["Setting", "Value", "Min", "Max", "Unit", "Why"])
    for c in ws[ws.max_row]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="152238")
    for p in res["parameters"]:
        ws.append([p["label"], p["value"], p["min"], p["max"], p["unit"], p["why"]])
    if res["risks"]:
        ws.append([])
        ws.append(["Quality warnings"])
        for r in res["risks"]:
            ws.append([r])
    ws.append([])
    ws.append([res["note"]])
    for col, w in zip("ABCDEF", (36, 12, 12, 12, 10, 80)):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), f"{res['part_code']}_process_parameters.xlsx"


def pdf(part_id: int) -> tuple[bytes, str]:
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    from . import drawing_style
    from .drawings import _styles
    from .pdf_report import _table
    res = compute(part_id)
    theme, ss, st = _styles()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm, topMargin=10 * mm, bottomMargin=10 * mm,
                            title=f"Design to Cost - process parameters {res['part_code']}")
    story = [drawing_style.pdf_header(st, "PROCESS PARAMETERS", f"{res['part_code']} · {res['material']} · {res['process']}"),
             Paragraph(f"Worked out from {res['source']}." + (" <b>EXAMPLE process windows / rules: confirm in the try-out.</b>" if res["is_example"] else ""), ss["Normal"]),
             Spacer(1, 3 * mm),
             _table([["Setting", "Value", "Range", "Unit", "Why"]] + [[p["label"], f"{p['value']:g}",
                     "" if p["min"] is None else f"{p['min']:g} – {p['max']:g}", p["unit"], Paragraph(p["why"], ss["Normal"])] for p in res["parameters"]],
                    theme, [70 * mm, 24 * mm, 36 * mm, 18 * mm, 122 * mm])]
    if res["risks"]:
        story += [Spacer(1, 3 * mm), Paragraph("<b>Quality warnings</b>", ss["Normal"])] + [Paragraph(f"■ {r}", ss["Normal"]) for r in res["risks"]]
    story += [Spacer(1, 3 * mm), Paragraph(res["note"], ss["Normal"])]
    doc.build(story)
    return buf.getvalue(), f"{res['part_code']}_process_parameters.pdf"
