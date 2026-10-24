"""Production tool designs of a part (roadmap stage 21): generate, view, correct by STEP upload, check, release.

    generate     3D tool design (moulds: plates, pillars, bushes, inserts cut from the corrected part; sheet metal: every stage
                 or the progressive die) split into components with part numbers, class, steel, features, raw stock,
                 machining operations and costs; one STEP file per component (stored by checksum), bought-out standard parts
                 and assembly steps in the bill of materials
    correct      download a component STEP, edit it in any CAD (radius, thickness, …), upload it: the solid is validated, size /
                 stock / weight / machining are recalculated, every identical part follows, and the whole design is checked
    check        valid solids and interference between components, listed with the volume
    release      freezes the design as a drawing revision (A, B, …) for the purchase orders; a released design is never edited, a new
                 revision is started from it instead
"""

from __future__ import annotations

import base64
import io
import math
import os
import tempfile
from collections import OrderedDict
from datetime import datetime
from pathlib import Path

import cadquery as cq
import numpy as np
from sqlalchemy import func, select, update

from . import (auth, corrections, design_pack, drawing_data, drawing_sheet, design_skills, fea, file_store, foundry_design, geometry, machine_spec, master_data, mould_db, outputs3d, rates, register, register_setup,
               tool_components, tool_design, tool_review)
from . import projects_db as db
from .register import RegisterError, ToolNotApplicable

_SHAPES: OrderedDict = OrderedDict()
SHAPE_CACHE = 160
REVS = "ABCDEFGHJKLMNPRSTUVWXYZ"


def _user() -> str:
    return register.user()


def _safe(s: str) -> str:
    return "".join(ch if ch.isalnum() or ch in "-_" else "_" for ch in s)


def _next_rev(rev: str) -> str:
    i = REVS.index(rev) if rev in REVS else 0
    return REVS[min(i + 1, len(REVS) - 1)]


# ------------------------------------------------------------ files and shapes

def _shape_of(file_row: dict) -> cq.Shape:
    key = file_row["sha256"]
    if key in _SHAPES:
        _SHAPES.move_to_end(key)
        return _SHAPES[key]
    shape = geometry.load_step(str(file_store.path_of(file_row)))
    _SHAPES[key] = shape
    while len(_SHAPES) > SHAPE_CACHE:
        _SHAPES.popitem(last=False)
    return shape


def _file_row(conn, file_id: int) -> dict:
    return dict(conn.execute(select(db.files).where(db.files.c.id == file_id)).mappings().one())


def _store_step(conn, shape: cq.Shape, name: str, color) -> dict:
    return file_store.put(outputs3d.solid_step_bytes(shape), f"{name}.step", "tool_step", conn)


# ------------------------------------------------------------ generate

def _tool_group(part_type: str) -> str:
    if part_type == "sheet_metal":
        return "sheet_metal"
    if part_type in ("plastic", "die_casting", "eps", "blow"):
        return "mould"
    if part_type == "foundry":
        return "foundry"
    raise ToolNotApplicable("Production tool drawings cover moulds, dies (plastic, die casting, EPS, blow), sheet-metal tools and foundry patterns.")


def _signature(shape: cq.Shape, cls: str) -> tuple:
    bb = shape.BoundingBox()
    com = shape.Center()
    # the centre of mass inside the bounding box tells mirror images (left / right halves) apart: they are not the same part
    lean = tuple(round((c - (lo + hi) / 2) / 0.5) * 0.5 for c, lo, hi in ((com.x, bb.xmin, bb.xmax), (com.y, bb.ymin, bb.ymax), (com.z, bb.zmin, bb.zmax)))
    return (cls, tuple(sorted(round(x, 1) for x in (bb.xlen, bb.ylen, bb.zlen))), round(shape.Volume(), -1), lean)


def _prepare(part_id: int, mode: str | None = None) -> dict:
    d = register.part_detail(part_id, with_rm=False)
    proc = register_setup.processes().get(d["process"])
    if not proc:
        raise ToolNotApplicable("Set the part's process first.")
    group = _tool_group(proc["part_type"])
    code = _safe(d["part_code"])
    comps, std, steps, summary, stage_of, fea_vtk = [], [], [], {}, {}, None
    if group == "mould":
        P = design_pack.mould_prepare(part_id, production=True)
        m3d, tool = P["m3d"], P["tool"]
        country, month = tool["country"], tool["as_of"]
        base = m3d["base_mm"]
        env = {"cavities": float(m3d["cavities_total"]), "base_l_mm": base[0], "base_w_mm": base[1], "plate_a_mm": m3d["plate_a_mm"],
               "plate_b_mm": m3d["plate_b_mm"], "pillar_dia_mm": m3d["pillar_dia_mm"], "pillars": 4.0,
               "part_h_mm": tool.get("part_h", 0) or 0, "screw_dia_mm": float(min(30, max(8, 2 * round(min(base) / 60)))),
               "dowel_dia_mm": float(min(30, max(8, 2 * round(min(base) / 60))))}
        det = m3d.get("details") or {"counts": {}, "rules": {}}
        c3 = det["counts"]
        env |= {k: float(x) for k, x in (det.get("rules") or {}).items()}
        env |= {"screws_total": float(c3.get("screws", 0)), "dowels_total": float(c3.get("dowels", 0))}
        std = tool_components.apply_counts(tool_components.bom_standard_lines(tool), {
            "ejector pin": c3.get("ejector_pins"), "return pin": c3.get("return_pins"), "cooling connectors": c3.get("cooling_connectors"),
            "slide mechanism": c3.get("slides") and c3["slides"], "lifter guide": c3.get("lifters"), "hot runner nozzle": c3.get("nozzles")})
        std += tool_components.standard_lines("mould", env)
        steps = tool_components.assembly_steps("mould", env)
        for c in m3d["components"]:
            comps.append(dict(c, prefix="M-", stage=None))
        base_line = next((l for l in tool.get("lines", []) if l.get("is_base")), None)
        summary = {"mould_base": m3d["base"], "base_mm": m3d["base_mm"], "cavities": m3d["cavities_total"], "notes": m3d["note"], "parting_z_mm": m3d["parting_z_mm"],
                   "opening_axis": P["axis"], "tool_cost": round(tool["value"]), "tool_cost_inr": round(tool["value_inr"]),
                   "currency": tool["currency"], "catalogue_base": ({"item": base_line["element"], "price": round(base_line["total"])} if base_line else None),
                   "draft": ({"before_cm2": P["corr"]["before"]["low_draft_cm2"], "after_cm2": P["corr"]["after"]["low_draft_cm2"]} if P["corr"] else None),
                   "example": bool(tool.get("example", {}).get("any"))}
        trial = tool.get("trial") or {}
        mach = trial.get("machine") or {}
        tie = min(mach.get("length_mm") or 0, mach.get("width_mm") or 0)
        b_ins = next((c["shape"] for c in m3d["components"] if c["name"] == "B side insert"), None)
        pressure = master_data.parameters(P["part_type"]).get("cavity_pressure_bar", 400) / 10.0
        fea_res = (fea.solve(b_ins, pressure, "cavity", master_data.parameters(P["part_type"])) if b_ins is not None
                   else {"available": False, "note": "No B side insert to check."})
        det = m3d.get("details") or {"counts": {}, "rules": {}}
        facts = {"fea_ok": (1 if fea_res.get("ok") else 0) if fea_res.get("max_deflection_mm") is not None else -1,
                 "fea_deflection_mm": fea_res.get("max_deflection_mm") or 0, "fea_limit_mm": fea_res.get("deflection_limit_mm") or 0,
                 "pins_placed": det["counts"].get("ejector_pins", 0), "pins_wanted": det["rules"].get("ejector_pins", 0),
                 "cool_placed": det["counts"].get("cooling_channels", 0), "cool_wanted": 2 * det["rules"].get("cool_channels", 0),
                 "slides": det["counts"].get("slides", 0), "lifters": det["counts"].get("lifters", 0),
                 "undercut_unresolved": det["counts"].get("unresolved", 0), "hot_runner": int(bool(det["counts"].get("nozzles"))),
                 "draft_left_cm2": (P["corr"]["after"]["low_draft_cm2"] if P["corr"] else 0), "clamp_t": trial.get("force_t", 0),
                 "machine_t": mach.get("tonnage_t", 0), "tie_bar_fit": int(bool(tie) and min(base) <= 0.95 * tie), "cavities": m3d["cavities_total"]}
        summary["fea"] = {k: v for k, v in fea_res.items() if k != "vtk"}
        spec = d.get("machine") or None
        mach_sum = machine_spec.summary(P["part_type"], spec, machine_spec.recommend(P["part_type"], [det["rules"]]),
                                        {"clamp_t": trial.get("force_t", 0), "base_l_mm": base[0], "base_w_mm": base[1], "mould_h_mm": m3d["mould_h_mm"], "riser_mm": m3d["riser_mm"]})
        summary["machine"] = mach_sum
        facts["machine_fit"] = mach_sum["fit_ok"]
        tie_row = next((r for r in mach_sum["fit"] if r["item"] == "Tie-bar clear distance"), None)
        if tie_row:
            facts["tie_bar_fit"] = int(tie_row["ok"])
        summary["facts"] = facts
        fea_vtk = fea_res.get("vtk")
    elif group == "foundry":
        _, rev, shape = register.part_shape(part_id)
        density, _ = register._density(d["material"])
        if not density:
            raise ToolNotApplicable("Set the part's material (with a density) before designing the pattern.")
        axis = rev["geometry"]["pull_direction"]
        vol = shape.Volume() / 1000.0
        facts_in = {"part_weight_kg": vol * density / 1000.0, "part_volume_cm3": vol, "part_surface_cm2": shape.Area() / 100.0}
        F = foundry_design.build(shape, axis, d["material"], facts_in)
        country = (d["tool_country"] or d["country"] or register_setup.settings()["default_country"]).upper()
        month = d["costing_month"] or rates.current_month()
        for c in F["components"]:
            comps.append(dict(c, prefix="F-", stage=None))
        c3 = F["counts"]
        env = {k: float(x) for k, x in F["rules"].items()} | {"screws_total": float(c3.get("screws", 0)), "dowels_total": float(c3.get("dowels", 0))}
        std = tool_components.standard_lines("foundry", env)
        steps = tool_components.assembly_steps("foundry", env)
        mach_sum = machine_spec.summary("foundry", d.get("machine") or None, machine_spec.recommend("foundry", [F["rules"]]), {})
        facts = {"check_errors": 0, "check_warnings": 0, "pattern_ok": F["pattern_ok"], "gating_ok": F["gating_ok"], "cores": F["cores"], "machine_fit": mach_sum["fit_ok"]}
        summary = {"notes": " ".join(F["notes"]), "facts": facts, "foundry": {k: F[k] for k in ("shrinkage_pct", "allowance_mm", "flask_mm", "cores", "draft")},
                   "rules": F["rules"], "example": True, "opening_axis": axis, "currency": "INR", "machine": mach_sum}
        P = {"axis": axis}
    else:
        P = design_pack.sheet_prepare(part_id, mode)
        design, fp, shape = P["design"], P["fp"], P["shape"]
        country, month = P["country"], design["month"]
        progressive = design["mode"] == "progressive"
        if progressive:
            sol = tool_design.progressive_solids(design, fp, shape, open_state=False, details=True)
            for c in sol:
                comps.append(dict(c, prefix="P-", stage=None))
            env = _stage_env(design["stages"][0]) | _detail_env(design, "progressive")
            std = tool_components.standard_lines("sheet_metal", env)
            steps = tool_components.assembly_steps("sheet_metal", env)
        else:
            agg, x0 = {}, 0.0
            for s in design["stages"]:
                width = s["values"]["die_set_l_mm"]
                # the stage dies stand side by side (one after the other along X) so that they can be viewed and checked together
                for c in tool_design.stage_solids(design, s, fp, shape, open_state=False, details=True):
                    shifted = c["shape"].translate(cq.Vector(x0 + width / 2, 0, 0))
                    comps.append(dict(c, shape=shifted, prefix=f"S{s['n']}-", stage=s["n"]))
                x0 += width + 150
                env = _stage_env(s) | _detail_env(design, s["n"])
                for l in tool_components.standard_lines("sheet_metal", env):
                    key = (l["item"], l["spec"])
                    if key in agg:
                        agg[key]["quantity"] += l["quantity"]
                        agg[key]["amount_inr"] += l["amount_inr"]
                    else:
                        agg[key] = dict(l, item=f"{l['item']}")
                    agg[key].setdefault("stages", []).append(s["n"])
            std = list(agg.values())
            steps = []
            for s in design["stages"]:
                env = _stage_env(s) | _detail_env(design, s["n"])
                for st in tool_components.assembly_steps("sheet_metal", env):
                    steps.append(dict(st, text=f"Stage {s['n']} ({s['name']}): {st['text']}"))
            steps = [dict(st, step=i) for i, st in enumerate(steps, start=1)]
        biggest = max(design["stages"], key=lambda x: x["force_t"])
        die_c = next((c for c in comps if c["name"] == "Die block" and (progressive or c["stage"] == biggest["n"])), None)
        fea_res = {"available": False, "note": "No die block to check."}
        if die_c is not None:
            bbd = die_c["shape"].BoundingBox()
            fea_res = fea.solve(die_c["shape"], (design["total_force_t"] if progressive else biggest["design_force_t"]) * 9806.65 / (bbd.xlen * bbd.ylen),
                                "top", master_data.parameters("sheet_metal"))
        dets = list((design.get("_details") or {}).values())
        sp_placed = sum(d_["counts"].get("springs", 0) for d_ in dets)
        sp_wanted = sum(int(st["values"].get("springs", 0)) for st in (design["stages"][:1] if progressive else design["stages"]))
        facts = {"fea_ok": (1 if fea_res.get("ok") else 0) if fea_res.get("max_deflection_mm") is not None else -1,
                 "fea_deflection_mm": fea_res.get("max_deflection_mm") or 0, "fea_limit_mm": fea_res.get("deflection_limit_mm") or 0,
                 "press_fit": int(all(s["press"].get("fits") for s in design["stages"])), "springs_placed": sp_placed,
                 "springs_wanted": sp_wanted or sp_placed, "form_undercuts": sum("undercut" in n for n in design.get("form_notes", [])),
                 "stages": len(design["stages"])}
        spec = d.get("machine") or None
        dets_rules = [x["rules"] for x in (design.get("_details") or {}).values()]
        mach_sum = machine_spec.summary("sheet_metal", spec, machine_spec.recommend("sheet_metal", dets_rules), {
            "design_force_t": max(s["design_force_t"] for s in design["stages"]),
            "progressive": progressive, "die_set_l_mm": max(s["values"]["die_set_l_mm"] for s in design["stages"]),
            "die_set_w_mm": max(s["values"]["die_set_w_mm"] for s in design["stages"]), "shut_height_mm": max(s["values"]["shut_height_mm"] for s in design["stages"]),
            "shut_height_low_mm": min(s["values"]["shut_height_mm"] for s in design["stages"]), "riser_mm": design.get("riser_mm", 0)})
        facts["machine_fit"] = mach_sum["fit_ok"]
        fea_vtk = fea_res.get("vtk")
        summary = {"fea": {k: v for k, v in fea_res.items() if k != "vtk"}, "facts": facts,
                   "mode": design["mode"], "stages": len(design["stages"]), "total_force_t": design["total_force_t"],
                   "presses": [s["press"].get("name") for s in design["stages"]], "notes": " ".join(design["warnings"] + design.get("form_notes", []) + [n for d_ in (design.get("_details") or {}).values() for n in d_["notes"]]),
                   "corrections": [c for c in P["corr"]["changes"] if c.get("count")], "steel_cost": design["steel_cost"],
                   "currency": design["currency"], "example": design["is_example"], "thickness_mm": design["thickness_mm"], "machine": mach_sum}
    return {"d": d, "group": group, "part_type": proc["part_type"], "code": code, "comps": comps, "std": std, "steps": steps,
            "summary": summary, "country": country, "month": month, "P": P, "vtk": fea_vtk}


def _detail_env(design: dict, key) -> dict:
    """Rule values and the counts of screws, dowels, springs and bolts placed in the 3D die set, for the bill of materials."""
    d = (design.get("_details") or {}).get(key) or {"counts": {}, "rules": {}}
    c = d["counts"]
    return {k: float(x) for k, x in d["rules"].items()} | {"screws_total": float(c.get("screws", 0)), "dowels_total": float(c.get("dowels", 0)),
                                                         "springs_total": float(c.get("springs", 0)), "stripper_bolts_total": float(c.get("stripper_bolts", 0))}


def _stage_env(stage: dict) -> dict:
    env = {k: float(v) for k, v in stage["values"].items()}
    env |= {k: float(v) for k, v in master_data.parameters("sheet_metal").items() if isinstance(v, (int, float))}
    env["force_t"] = stage["force_t"]
    return env


def generate(part_id: int, note: str = "", mode: str | None = None, use_skills: bool = True) -> dict:
    """Build a new draft design version of the part's tool. Returns the design (see get_design). The user's skills (design_skills.py) shape it unless use_skills is False."""
    d0 = register.part_detail(part_id, with_rm=False)
    proc = register_setup.processes().get(d0["process"]) or {}
    token = design_skills.begin(_tool_group(proc.get("part_type", "")) if proc.get("part_type") else "mould", proc.get("part_type", ""), use_skills)
    try:
        pre = _prepare(part_id, mode)
    finally:
        used = design_skills.end(token)
    pre["summary"]["skills"] = used
    if used["applied"]:
        pre["summary"]["notes"] = (pre["summary"].get("notes") or "") + " Your skills used: " + "; ".join(
            f"{a['target'].split('.')[-1]} {a['from']:g} → {a['to']:g}" for a in used["applied"]) + "."
    code, part_type, country, month = pre["code"], pre["part_type"], pre["country"], pre["month"]
    # part numbers: identical parts share one, per prefix
    counters, numbers, infos = {}, {}, {}
    rows = []
    comps = sorted(pre["comps"], key=lambda c: (c["stage"] or 0, {"lower": 0, "upper": 1}.get(c["group"], 2), c["name"]))
    for c in comps:
        cls = c["cls"]
        if cls in tool_components.NON_TOOL:
            kind = cls
        else:
            row = tool_components.class_row(cls, part_type)
            kind = row["kind"] if row else "made"
        c["kind"] = kind
        if kind == "made":
            sig = (c["prefix"],) + _signature(c["shape"], cls)
            if sig not in numbers:
                counters[c["prefix"]] = counters.get(c["prefix"], 0) + 1
                numbers[sig] = f"{code}-{c['prefix']}{counters[c['prefix']]:02d}"
                infos[sig] = tool_components.component_info(c["shape"], cls, part_type, country, month)
            c["part_no"], c["info"] = numbers[sig], infos[sig]
            c["instance"] = sum(1 for r in rows if r["part_no"] == c["part_no"]) + 1
        else:
            tag = {"standard": "STD", "work": "WRK", "part": "PRT", "proposal": "PRP"}[kind]
            counters[tag] = counters.get(tag, 0) + 1
            c["part_no"], c["instance"] = f"{code}-{tag}{counters[tag]:02d}", 1
            bb = c["shape"].BoundingBox()
            c["info"] = {"class": cls, "bbox_mm": [round(bb.xlen, 2), round(bb.ylen, 2), round(bb.zlen, 2)],
                         "bbox_min": [round(bb.xmin, 2), round(bb.ymin, 2), round(bb.zmin, 2)],
                         "volume_cm3": round(c["shape"].Volume() / 1000, 2)}
        rows.append(c)
    with db.engine().begin() as conn:
        version = (conn.execute(select(func.max(db.tool_designs.c.version)).where(db.tool_designs.c.part_id == part_id)).scalar() or 0) + 1
        summary = pre["summary"] | {"part_code": pre["d"]["part_code"], "part_rev": pre["d"]["revisions"][0]["rev"] if pre["d"]["revisions"] else ""}
        did = conn.execute(db.tool_designs.insert().values(
            part_id=part_id, version=version, rev_label="A", status="draft", kind=pre["group"], part_type=part_type, country=country,
            as_of=month, summary=summary, bom={"standard": pre["std"], "assembly": pre["steps"]}, note=note, created_by=_user())).inserted_primary_key[0]
        for c in rows:
            f = _store_step(conn, c["shape"], c["part_no"] + (f"_{c['instance']}" if c["instance"] > 1 else ""), c["color"])
            cid = conn.execute(db.tool_components.insert().values(
                design_id=did, part_no=c["part_no"], instance=c["instance"], name=c["name"], group=c["group"], stage_no=c["stage"],
                cls=c["cls"], kind=c["kind"], color=list(c["color"]), file_id=f["id"], rev="A", info=c["info"], edited=0)).inserted_primary_key[0]
            conn.execute(db.tool_component_revs.insert().values(component_id=cid, rev="A", file_id=f["id"], note="generated", info=c["info"],
                                                                created_by=_user()))
        register._log(conn, "part", part_id, "tool_design_generated", {"version": version, "components": len(rows)})
        tool_review.sync(conn, did, pre["group"], summary.get("facts", {}))
        if pre.get("vtk"):
            f = file_store.put(pre["vtk"], f"{code}_v{version}_fea.vtk", "tool_fea", conn)
            s2 = dict(summary, fea_file_id=f["id"])
            conn.execute(update(db.tool_designs).where(db.tool_designs.c.id == did).values(summary=s2))
    check_design(did)
    return get_design(did)


# ------------------------------------------------------------ read

def _design_row(conn, design_id: int) -> dict:
    r = conn.execute(select(db.tool_designs).where(db.tool_designs.c.id == design_id)).mappings().first()
    if not r:
        raise RegisterError("Tool design not found.")
    return dict(r)


def _comp_rows(conn, design_id: int) -> list[dict]:
    rows = conn.execute(select(db.tool_components).where(db.tool_components.c.design_id == design_id)
                        .order_by(db.tool_components.c.id)).mappings().all()
    return [dict(r) for r in rows]


def list_designs(part_id: int) -> list[dict]:
    with db.engine().connect() as conn:
        rows = conn.execute(select(db.tool_designs).where(db.tool_designs.c.part_id == part_id)
                            .order_by(db.tool_designs.c.version.desc())).mappings().all()
    out = []
    for r in rows:
        r = dict(r)
        r["created_at"] = r["created_at"].isoformat() if r["created_at"] else None
        r["released_at"] = r["released_at"].isoformat() if r["released_at"] else None
        s = r.pop("summary") or {}
        r.pop("bom")
        r["summary"] = {k: s.get(k) for k in ("mould_base", "cavities", "mode", "stages", "total_force_t", "tool_cost_inr", "steel_cost", "checks")}
        out.append(r)
    return out


def get_design(design_id: int) -> dict:
    with db.engine().connect() as conn:
        d = _design_row(conn, design_id)
        comps = _comp_rows(conn, design_id)
    d["created_at"] = d["created_at"].isoformat() if d["created_at"] else None
    d["released_at"] = d["released_at"].isoformat() if d["released_at"] else None
    d["components"] = [{k: c[k] for k in ("id", "part_no", "instance", "name", "group", "stage_no", "cls", "kind", "rev", "edited", "color")}
                       | {"info": c["info"]} for c in comps]
    d["bill"] = bill(design_id)
    with db.engine().connect() as conn:
        d["review"] = tool_review.items(conn, design_id)
        d["quotes"] = [{**dict(q), "created_at": q["created_at"].isoformat() if q["created_at"] else None} for q in
                       conn.execute(select(db.tool_quotes).where(db.tool_quotes.c.design_id == design_id).order_by(db.tool_quotes.c.id)).mappings()]
    return d


def bill(design_id: int) -> dict:
    """Made parts grouped by part number (quantity = instances), bought-out parts, assembly steps, totals."""
    with db.engine().connect() as conn:
        d = _design_row(conn, design_id)
        comps = _comp_rows(conn, design_id)
        over = drawing_data.stored(conn, design_id)
    made = {}
    for c in comps:
        if c["kind"] != "made":
            continue
        m = made.setdefault(c["part_no"], {"part_no": c["part_no"], "name": c["name"].split(":")[0] if c["instance"] > 1 else c["name"],
                                           "cls": c["cls"], "stage_no": c["stage_no"], "rev": c["rev"], "qty": 0,
                                           "info": drawing_data.apply(c["info"], over.get(c["part_no"]), d["country"], d["as_of"]), "ids": []})
        m["qty"] += 1
        m["ids"].append(c["id"])
        m["rev"] = max(m["rev"], c["rev"])
    for m in made.values():
        if m["qty"] > 1:
            m["name"] = m["name"].rstrip("0123456789 ").rstrip() or m["name"]
    std3d = [c for c in comps if c["kind"] == "standard"]
    lines = list(d["bom"]["standard"])
    totals = {"made_parts": len(made), "made_qty": sum(m["qty"] for m in made.values()),
              "steel_kg": round(sum(m["info"]["stock"]["weight_kg"] * m["qty"] for m in made.values()), 1),
              "material_cost": sum((m["info"]["material_cost"] or 0) * m["qty"] for m in made.values()),
              "machining_cost": sum(m["info"]["machining_cost"] * m["qty"] for m in made.values()),
              "standard_cost_inr": sum(l["amount_inr"] for l in lines),
              "assembly_hours": round(sum(a["hours"] for a in d["bom"]["assembly"]), 1)}
    return {"made": sorted(made.values(), key=lambda m: m["part_no"]), "standard": lines, "standard_3d": [c["name"] for c in std3d],
            "assembly": d["bom"]["assembly"], "totals": totals, "currency": d["summary"].get("currency", ""), "country": d["country"]}


# ------------------------------------------------------------ files, mesh, assembly

def component_step(design_id: int, comp_id: int) -> tuple[bytes, str]:
    with db.engine().connect() as conn:
        c = conn.execute(select(db.tool_components).where((db.tool_components.c.id == comp_id) & (db.tool_components.c.design_id == design_id))).mappings().first()
        if not c:
            raise RegisterError("Component not found.")
        f = _file_row(conn, c["file_id"])
    return file_store.path_of(f).read_bytes(), f"{c['part_no']}_rev{c['rev']}" + (f"_{c['instance']}" if c["instance"] > 1 else "") + ".step"


def _live(design_id: int, kinds=("made", "standard")) -> list[dict]:
    with db.engine().connect() as conn:
        comps = _comp_rows(conn, design_id)
        out = []
        for c in comps:
            if c["kind"] in kinds:
                out.append(dict(c, shape=_shape_of(_file_row(conn, c["file_id"]))))
    return out


def assembly_step(design_id: int, include_reference: bool = False) -> bytes:
    kinds = ("made", "standard") + (("work", "part") if include_reference else ())
    items = _live(design_id, kinds)
    return outputs3d.step_bytes([{"name": f"{c['part_no']}_{c['instance']}", "shape": c["shape"], "color": tuple(c["color"] or (0.7, 0.7, 0.7))}
                                 for c in items], f"tool_design_{design_id}")


def mesh_payload(design_id: int) -> dict:
    """Meshes of every component for the 3D viewer (base64 float32 vertices, uint32 triangles) and the features to mark."""
    items = _live(design_id, ("made", "standard", "work", "part", "proposal"))
    allbb = None
    for c in items:
        bb = c["shape"].BoundingBox()
        allbb = bb if allbb is None else allbb.add(bb)
    size = max(allbb.xlen, allbb.ylen, allbb.zlen) if allbb else 100.0
    tol = max(0.15, size / 900)
    out = []
    for c in items:
        v, f = outputs3d.mesh(c["shape"], tol)
        if not len(f):
            continue
        out.append({"id": c["id"], "part_no": c["part_no"], "instance": c["instance"], "name": c["name"], "group": c["group"], "kind": c["kind"],
                    "stage": c["stage_no"], "rev": c["rev"], "edited": bool(c["edited"]),
                    "color": [float(x) for x in (c["color"] or (0.7, 0.7, 0.7))],
                    "v": base64.b64encode(v.astype("<f4").tobytes()).decode(), "f": base64.b64encode(f.astype("<u4").tobytes()).decode(),
                    "info": {k: (c["info"] or {}).get(k) for k in ("bbox_mm", "bbox_min", "volume_cm3", "steel", "hardness", "stock", "class_name")}
                            | {"features": ((c["info"] or {}).get("features") or {})}})
    return {"design_id": design_id, "components": out, "span_mm": round(size * 0.25, 1)}


# ------------------------------------------------------------ check

def check_design(design_id: int) -> dict:
    items = _live(design_id, ("made", "standard"))
    issues = []
    for c in items:
        if not c["shape"].isValid():
            issues.append({"level": "error", "component": c["part_no"], "text": f"{c['name']}: the solid is not valid."})
    for x in tool_components.interference([{"id": c["id"], "name": f"{c['part_no']} {c['name']}", "shape": c["shape"]} for c in items]):
        issues.append({"level": "error" if x["volume_mm3"] > 100 else "warning", "component": x["a_name"].split(" ")[0],
                       "text": f"{x['a_name']} and {x['b_name']} overlap by {x['volume_mm3']:g} mm³.", "volume_mm3": x["volume_mm3"]})
    with db.engine().begin() as conn:
        d = _design_row(conn, design_id)
        s = dict(d["summary"] or {})
        s["checks"] = {"issues": issues, "errors": sum(i["level"] == "error" for i in issues), "warnings": sum(i["level"] == "warning" for i in issues),
                       "at": datetime.utcnow().isoformat(timespec="seconds"), "components_checked": len(items)}
        conn.execute(update(db.tool_designs).where(db.tool_designs.c.id == design_id).values(summary=s))
        tool_review.sync(conn, design_id, "mould" if d["kind"] == "mould" else d["kind"], dict(s.get("facts", {}), check_errors=s["checks"]["errors"],
                                                                                            check_warnings=s["checks"]["warnings"]))
    return s["checks"]


# ------------------------------------------------------------ correct by STEP upload

def _require_draft(conn, design_id: int) -> dict:
    d = _design_row(conn, design_id)
    if d["status"] != "draft":
        raise RegisterError(f"This design is {d['status']}: start a new revision to change it.")
    return d


def upload_component(design_id: int, comp_id: int, content: bytes, filename: str, note: str) -> dict:
    """Replace a component by a corrected STEP file (coordinates as downloaded). Identical parts follow; the design is re-checked."""
    if not note.strip():
        raise RegisterError("Say what you changed (for the drawing revision).")
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp, "upload.step")
        path.write_bytes(content)
        try:
            new = geometry.load_step(str(path))
        except Exception as exc:
            raise RegisterError(f"The STEP file could not be read as a solid: {exc}") from None
    if not new.isValid() or new.Volume() <= 0:
        raise RegisterError("The uploaded solid is not valid (or has no volume). Repair it in your CAD and upload again.")
    with db.engine().begin() as conn:
        d = _require_draft(conn, design_id)
        comp = conn.execute(select(db.tool_components).where((db.tool_components.c.id == comp_id) & (db.tool_components.c.design_id == design_id))).mappings().first()
        if not comp:
            raise RegisterError("Component not found.")
        if comp["kind"] not in ("made", "standard"):
            raise RegisterError("Only tool parts can be replaced.")
        peers = [dict(r) for r in conn.execute(select(db.tool_components).where((db.tool_components.c.design_id == design_id)
                                               & (db.tool_components.c.part_no == comp["part_no"]))).mappings()]
        old_shape = _shape_of(_file_row(conn, comp["file_id"]))
        old_ref = old_shape.BoundingBox()
        old_vol = comp["info"]["volume_cm3"]
        rev = _next_rev(comp["rev"])
        changed = []
        for peer in peers:
            pb = _shape_of(_file_row(conn, peer["file_id"])).BoundingBox()
            shape = new if peer["id"] == comp_id else new.translate(cq.Vector(pb.xmin - old_ref.xmin, pb.ymin - old_ref.ymin, pb.zmin - old_ref.zmin))
            info = (tool_components.component_info(shape, comp["cls"], d["part_type"], d["country"], d["as_of"]) if comp["kind"] == "made"
                    else {**comp["info"], "bbox_mm": [round(x, 2) for x in (shape.BoundingBox().xlen, shape.BoundingBox().ylen, shape.BoundingBox().zlen)],
                          "volume_cm3": round(shape.Volume() / 1000, 2)})
            f = _store_step(conn, shape, f"{peer['part_no']}_rev{rev}" + (f"_{peer['instance']}" if peer["instance"] > 1 else ""), peer["color"])
            conn.execute(update(db.tool_components).where(db.tool_components.c.id == peer["id"]).values(file_id=f["id"], rev=rev, info=info, edited=1))
            conn.execute(db.tool_component_revs.insert().values(component_id=peer["id"], rev=rev, file_id=f["id"], note=note.strip(),
                                                                info=info, created_by=_user()))
            changed.append(peer["id"])
        new_info = info
    checks = check_design(design_id)
    mine = [i for i in checks["issues"] if comp["part_no"] in i["text"]]
    nb = new.BoundingBox()
    try:
        learned = design_skills.learn(d, comp, old_shape, new, (d["summary"].get("base_mm") or [None])[0])
    except Exception:
        learned = []                      # learning must never stop a correction
    return {"component": comp["part_no"], "rev": rev, "applied_to": len(changed), "volume_cm3": {"before": old_vol, "after": new_info["volume_cm3"]},
            "size_mm": {"before": comp["info"]["bbox_mm"], "after": new_info["bbox_mm"]},
            "stock": new_info.get("stock"), "material_cost": new_info.get("material_cost"), "machining_cost": new_info.get("machining_cost"),
            "issues_for_this_part": mine, "checks": {k: checks[k] for k in ("errors", "warnings")}, "learned": learned}


def revert_component(design_id: int, comp_id: int, rev: str) -> dict:
    with db.engine().begin() as conn:
        d = _require_draft(conn, design_id)
        comp = conn.execute(select(db.tool_components).where((db.tool_components.c.id == comp_id) & (db.tool_components.c.design_id == design_id))).mappings().first()
        old = conn.execute(select(db.tool_component_revs).where((db.tool_component_revs.c.component_id == comp_id)
                           & (db.tool_component_revs.c.rev == rev))).mappings().first()
        if not comp or not old:
            raise RegisterError("Revision not found.")
        conn.execute(update(db.tool_components).where(db.tool_components.c.id == comp_id).values(file_id=old["file_id"], info=old["info"],
                     rev=old["rev"], edited=0 if old["rev"] == "A" else 1))
    check_design(design_id)
    return {"component": comp["part_no"], "rev": rev}


def component_revisions(design_id: int, comp_id: int) -> list[dict]:
    with db.engine().connect() as conn:
        rows = conn.execute(select(db.tool_component_revs).where(db.tool_component_revs.c.component_id == comp_id)
                            .order_by(db.tool_component_revs.c.id)).mappings().all()
    return [{"rev": r["rev"], "note": r["note"], "by": r["created_by"], "at": r["created_at"].isoformat() if r["created_at"] else None,
             "size_mm": (r["info"] or {}).get("bbox_mm"), "volume_cm3": (r["info"] or {}).get("volume_cm3")} for r in rows]


# ------------------------------------------------------------ release and revisions

def release(design_id: int, accept_warnings: bool = False) -> dict:
    checks = check_design(design_id)
    with db.engine().begin() as conn:
        d = _require_draft(conn, design_id)
        if checks["errors"]:
            raise RegisterError(f"{checks['errors']} error(s) in the check (solid not valid or parts overlapping): correct them first.")
        if checks["warnings"] and not accept_warnings:
            raise RegisterError(f"{checks['warnings']} warning(s) in the check: look at them, then release with 'accept warnings'.")
        blocked = tool_review.blockers(conn, design_id)
        if blocked:
            raise RegisterError(f"{len(blocked)} review item(s) are not signed off: " + "; ".join(blocked[:3]) + (" …" if len(blocked) > 3 else "")
                                + " Open the Review tab: sign each one, or waive it with a reason.")
        conn.execute(update(db.tool_designs).where((db.tool_designs.c.part_id == d["part_id"]) & (db.tool_designs.c.status == "released"))
                     .values(status="superseded"))
        conn.execute(update(db.tool_designs).where(db.tool_designs.c.id == design_id)
                     .values(status="released", released_by=_user(), released_at=datetime.utcnow()))
        register._log(conn, "part", d["part_id"], "tool_design_released", {"version": d["version"], "rev": d["rev_label"]})
    return get_design(design_id)


def new_revision(design_id: int, note: str = "") -> dict:
    """A new draft design (next drawing revision) starting from the components of this one."""
    with db.engine().begin() as conn:
        d = _design_row(conn, design_id)
        comps = _comp_rows(conn, design_id)
        version = conn.execute(select(func.max(db.tool_designs.c.version)).where(db.tool_designs.c.part_id == d["part_id"])).scalar() + 1
        nid = conn.execute(db.tool_designs.insert().values(
            part_id=d["part_id"], version=version, rev_label=_next_rev(d["rev_label"]), status="draft", kind=d["kind"], part_type=d["part_type"],
            country=d["country"], as_of=d["as_of"], summary=d["summary"], bom=d["bom"], note=note or f"revision of version {d['version']}",
            created_by=_user())).inserted_primary_key[0]
        for c in comps:
            cid = conn.execute(db.tool_components.insert().values(
                design_id=nid, part_no=c["part_no"], instance=c["instance"], name=c["name"], group=c["group"], stage_no=c["stage_no"],
                cls=c["cls"], kind=c["kind"], color=c["color"], file_id=c["file_id"], rev=c["rev"], info=c["info"], edited=0)).inserted_primary_key[0]
            conn.execute(db.tool_component_revs.insert().values(component_id=cid, rev=c["rev"], file_id=c["file_id"], note=f"carried from version {d['version']}",
                                                                info=c["info"], created_by=_user()))
    with db.engine().begin() as conn:
        tool_review.sync(conn, nid, "mould" if d["kind"] == "mould" else d["kind"], (d["summary"] or {}).get("facts", {}))
        for r in conn.execute(select(db.tool_drawing_data).where(db.tool_drawing_data.c.design_id == design_id)).mappings().all():
            conn.execute(db.tool_drawing_data.insert().values(design_id=nid, part_no=r["part_no"], fields=r["fields"], notes=r["notes"],
                                                              edited_by=r["edited_by"], edited_at=r["edited_at"]))
        for r in conn.execute(select(db.tool_drawing_edits).where(db.tool_drawing_edits.c.design_id == design_id).order_by(db.tool_drawing_edits.c.id)).mappings().all():
            conn.execute(db.tool_drawing_edits.insert().values(design_id=nid, part_no=r["part_no"], summary=r["summary"], by=r["by"], at=r["at"]))
    check_design(nid)
    return get_design(nid)


# ------------------------------------------------------------ drawings

def _head(conn, d: dict, part_row) -> dict:
    proj = conn.execute(select(db.projects).where(db.projects.c.id == part_row["project_id"])).mappings().first()
    tool = ("Mould" if d["kind"] == "mould" else "Sheet-metal die") + f" · {d['part_type'].replace('_', ' ')}"
    return {"part_code": part_row["part_code"], "part_name": part_row["name"], "tool": tool, "revision_label": d["rev_label"],
            "status": d["status"], "country": d["country"], "project": (proj["code"] if proj else ""), "date": datetime.utcnow().date().isoformat()}


def _part_row(conn, part_id: int):
    return conn.execute(select(db.parts).where(db.parts.c.id == part_id)).mappings().one()


_DRAWINGS: OrderedDict = OrderedDict()      # finished drawings by (component, revision, file, format, status): the supplier ZIP reuses them
DRAWING_CACHE = 400


def component_drawing(design_id: int, comp_id: int, fmt: str = "pdf") -> tuple[bytes, str]:
    with db.engine().connect() as conn:
        c0 = conn.execute(select(db.tool_components).where((db.tool_components.c.id == comp_id) & (db.tool_components.c.design_id == design_id))).mappings().first()
        d0 = _design_row(conn, design_id) if c0 else None
        dv = _data_version(conn, design_id, c0["part_no"]) if c0 else None
    key = (db.engine().url.render_as_string(), comp_id, c0["rev"], c0["file_id"], fmt, d0["status"], d0["rev_label"], d0["released_at"], dv) if c0 else None
    if key and key in _DRAWINGS:
        _DRAWINGS.move_to_end(key)
        return _DRAWINGS[key]
    out = _component_drawing(design_id, comp_id, fmt)
    if key:
        _DRAWINGS[key] = out
        while len(_DRAWINGS) > DRAWING_CACHE:
            _DRAWINGS.popitem(last=False)
    return out


def _component_drawing(design_id: int, comp_id: int, fmt: str = "pdf") -> tuple[bytes, str]:
    from . import tool_drawings
    with db.engine().connect() as conn:
        d = _design_row(conn, design_id)
        c = conn.execute(select(db.tool_components).where((db.tool_components.c.id == comp_id) & (db.tool_components.c.design_id == design_id))).mappings().first()
        if not c:
            raise RegisterError("Component not found.")
        if c["kind"] != "made":
            raise RegisterError("Only made parts have a manufacturing drawing: bought-out parts are ordered from the catalogue.")
        shape = _shape_of(_file_row(conn, c["file_id"]))
        head = _head(conn, d, _part_row(conn, d["part_id"]))
        qty = conn.execute(select(func.count()).select_from(db.tool_components).where((db.tool_components.c.design_id == design_id)
                                                                                      & (db.tool_components.c.part_no == c["part_no"]))).scalar()
    with db.engine().connect() as conn:
        sha = _file_row(conn, c["file_id"])["sha256"]
    comp = {"part_no": c["part_no"], "name": c["name"].split(":")[0], "rev": c["rev"], "qty": qty, "stage": c["stage_no"], "sha": sha}
    with db.engine().connect() as conn:
        over = drawing_data.stored(conn, design_id).get(c["part_no"])
        hist = drawing_data.history(conn, design_id, c["part_no"])
    info = drawing_data.apply(c["info"], over, d["country"], d["as_of"])
    data = drawing_data.effective(info, over)
    meta = drawing_sheet.sheet_meta(d, hist, comp["name"], c["part_no"], c["rev"], "", sha, bool(info.get("is_example")))
    if fmt == "dxf":
        return tool_drawings.component_dxf(comp, info, shape, head, data, meta), f"{c['part_no']}_rev{c['rev']}.dxf"
    return tool_drawings.component_pdf(comp, info, shape, head, data, meta), f"{c['part_no']}_rev{c['rev']}.pdf"


def _data_version(conn, design_id: int, part_no: str):
    """Changes with every edit of the technical data of one part: the drawing cache key."""
    n = conn.execute(select(func.max(db.tool_drawing_edits.c.id)).where((db.tool_drawing_edits.c.design_id == design_id)
                     & (db.tool_drawing_edits.c.part_no.in_([part_no, "*"])))).scalar()
    return n or 0


# ------------------------------------------------------------ technical data of the drawings (review, change, download, upload)

def drawing_data_get(design_id: int) -> dict:
    """The data printed on each made part's drawing, the defaults it started from and the stored changes."""
    with db.engine().connect() as conn:
        d = _design_row(conn, design_id)
        comps = _comp_rows(conn, design_id)
        over = drawing_data.stored(conn, design_id)
    parts, seen = [], set()
    for c in comps:
        if c["kind"] != "made" or c["part_no"] in seen:
            continue
        seen.add(c["part_no"])
        o = over.get(c["part_no"])
        info = drawing_data.apply(c["info"], o, d["country"], d["as_of"])
        qty = sum(1 for x in comps if x["part_no"] == c["part_no"])
        parts.append({"part_no": c["part_no"], "name": c["name"].split(":")[0], "qty": qty, "rev": c["rev"], "component_id": c["id"],
                      "defaults": drawing_data.defaults(c["info"]), "effective": {k: v for k, v in drawing_data.effective(info, o).items() if k in drawing_data.FIELDS or k == "notes"},
                      "changed": bool(o and (o["fields"] or o["notes"])), "edited_by": (o or {}).get("edited_by"), "edited_at": (o or {}).get("edited_at"),
                      "stock_kg": info["stock"]["weight_kg"], "material_cost": info["material_cost"]})
    return {"design_id": design_id, "status": d["status"], "editable": d["status"] == "draft", "fields": drawing_data.LABELS, "parts": parts,
            "tolerance_tables": {k: drawing_data.tolerance_table(k, "linear") for k in drawing_data.CLASSES},
            "steels": sorted({r["code"] for r in mould_db.tables()["tool_materials"]})}


def drawing_data_save(design_id: int, items: list[dict]) -> dict:
    """items: [{part_no, fields: {...}, notes: [...]}]: replaces the stored data of those parts (an empty field goes back to the default)."""
    changed = []
    with db.engine().begin() as conn:
        d = _require_draft(conn, design_id)
        known = {c["part_no"] for c in _comp_rows(conn, design_id) if c["kind"] == "made"}
        for it in items:
            if it["part_no"] not in known:
                raise RegisterError(f"Part {it['part_no']} is not a made part of this design.")
            try:
                r = drawing_data.save(conn, design_id, it["part_no"], it.get("fields"), it.get("notes"), _user())
            except auth.AuthError as exc:
                raise RegisterError(f"{it['part_no']}: {exc}") from None
            if r:
                changed.append(r)
        if changed:
            register._log(conn, "part", d["part_id"], "tool_drawing_data", {"version": d["version"], "parts": [c["part_no"] for c in changed]})
    if changed:
        check_design(design_id)
    return {"changed": changed, "unchanged": len(items) - len(changed)}


def drawing_data_xlsx(design_id: int) -> tuple[bytes, str]:
    data = drawing_data_get(design_id)
    d = get_design(design_id)
    with db.engine().connect() as conn:
        hist = []
        for p in data["parts"]:
            for h in drawing_data.history(conn, design_id, p["part_no"]):
                hist.append(dict(h, part_no=p["part_no"]))
        for h in drawing_data.history(conn, design_id, None):
            hist.append(dict(h, part_no="(tool assembly)"))
        over = drawing_data.stored(conn, design_id)
    rows = [{"part_no": p["part_no"], "name": p["name"], "qty": p["qty"], "rev": p["rev"], "eff": p["effective"], "edited_by": p["edited_by"]} for p in data["parts"]]
    code = _safe(d["summary"].get("part_code", f"design{design_id}"))
    return drawing_data.workbook(design_id, rows, hist), f"{code}_v{d['version']}_rev{d['rev_label']}_drawing_data.xlsx"


def drawing_data_upload(design_id: int, content: bytes) -> dict:
    """Apply a changed drawing data sheet. Checked as a whole before anything is stored."""
    sheet = drawing_data.read_workbook(content)
    cur = {p["part_no"]: p for p in drawing_data_get(design_id)["parts"]}
    unknown = [k for k in sheet if k not in cur]
    if unknown:
        raise RegisterError("Part number(s) not in this design: " + ", ".join(unknown[:5]) + ". Download the sheet again from this design.")
    items = []
    for part_no, row in sheet.items():
        defaults = cur[part_no]["defaults"]
        # a cell equal to the default is not stored as an override
        def same(k, v):
            if k == "finish_ra":
                try:
                    return float(v) == float(defaults[k])
                except (TypeError, ValueError):
                    return False
            return str(v).strip() == str(defaults.get(k)).strip()
        fields = {k: v for k, v in row["fields"].items() if not same(k, v)}
        items.append({"part_no": part_no, "fields": fields, "notes": row["notes"]})
    out = drawing_data_save(design_id, items)
    out["parts_in_sheet"] = len(items)
    return out


def assembly_drawing(design_id: int) -> tuple[bytes, str]:
    from . import tool_drawings
    items = _live(design_id, ("made", "standard"))
    b = bill(design_id)
    numbers = {m["part_no"]: i for i, m in enumerate(b["made"], start=1)}
    parts = [{"shape": c["shape"], "group": c["group"], "color": c["color"] or (0.7, 0.7, 0.7), "item": numbers.get(c["part_no"]) if c["instance"] == 1 else None}
             for c in items]
    with db.engine().connect() as conn:
        d = _design_row(conn, design_id)
        head = _head(conn, d, _part_row(conn, d["part_id"]))
        hist = drawing_data.history(conn, design_id, None)
    meta = drawing_sheet.sheet_meta(d, hist, f"Tool assembly {head['part_code']}", f"{head['part_code']}-ASM", d["rev_label"])
    return tool_drawings.assembly_pdf(d, b, parts, head, meta), f"{head['part_code']}_assembly_{d['status']}_{d['rev_label']}.pdf"


# ------------------------------------------------------------ purchase-order pack

def po_file(design_id: int, kind: str) -> tuple[bytes, str]:
    """kind: xlsx, po-material, po-machining, po-bought-out, assembly-pdf, assembly-step."""
    from . import tool_po
    d = get_design(design_id)
    code = _safe(d["summary"].get("part_code", f"design{design_id}"))
    tag = f"{code}_v{d['version']}_rev{d['rev_label']}_{d['status']}"
    if kind == "xlsx":
        return tool_po.workbook(d), f"{tag}_purchase_pack.xlsx"
    if kind == "po-material":
        return tool_po.po_pdf(d, "material"), f"{tag}_PO_raw_material.pdf"
    if kind == "po-machining":
        return tool_po.po_pdf(d, "machining"), f"{tag}_PO_machining.pdf"
    if kind == "po-bought-out":
        return tool_po.po_pdf(d, "bought_out"), f"{tag}_PO_bought_out.pdf"
    if kind == "assembly-pdf":
        return assembly_drawing(design_id)
    if kind == "assembly-step":
        return assembly_step(design_id), f"{tag}_assembly.step"
    raise RegisterError("Unknown file.")


def supplier_package(design_id: int) -> tuple[bytes, str]:
    """ZIP for the supplier: STEP, PDF and DXF of every made part, the assembly, the Excel pack and the three order documents."""
    import zipfile
    d = get_design(design_id)
    code = _safe(d["summary"].get("part_code", f"design{design_id}"))
    buf = io.BytesIO()
    seen = set()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for c in d["components"]:
            if c["kind"] != "made" or c["part_no"] in seen:
                continue
            seen.add(c["part_no"])
            z.writestr(f"step/{c['part_no']}_rev{c['rev']}.step", component_step(design_id, c["id"])[0])
            z.writestr(f"drawings/{c['part_no']}_rev{c['rev']}.pdf", component_drawing(design_id, c["id"], "pdf")[0])
            z.writestr(f"dxf/{c['part_no']}_rev{c['rev']}.dxf", component_drawing(design_id, c["id"], "dxf")[0])
        z.writestr("drawing_data/drawing_data_sheet.xlsx", drawing_data_xlsx(design_id)[0])
        if d["kind"] != "mould":
            try:                                          # sheet metal: the strip layout and blank drawings of the part, with the same title block
                dp, rv, fp, sl, res = register._layout(d["part_id"])
                names = {"blank.dxf": "blank.dxf", "strip.dxf": "strip.dxf", "strip_layout.pdf": "strip_layout_A3.pdf", "tool_pack.pdf": "tool_pack.pdf"}
                for n_, b_ in design_pack._layout_files("x", dp, rv, fp, sl, res, "layout/").items():
                    z.writestr("layout/" + names[n_.split("/x_", 1)[1]], b_)
            except (ToolNotApplicable, RegisterError):
                pass
        for kind in ("xlsx", "po-material", "po-machining", "po-bought-out", "assembly-pdf", "assembly-step"):
            content, name = po_file(design_id, kind)
            z.writestr((f"assembly/{name}" if kind.startswith("assembly") else f"po/{name}"), content)
        z.writestr("README.txt", (
            f"Tool purchase pack: {d['summary'].get('part_code', '')}\nDesign version {d['version']}, drawing revision {d['rev_label']}, status {d['status']}.\n\n"
            "step/        one STEP file per made part (edit in any CAD; send the corrected file back for the next revision)\n"
            "drawings/    manufacturing drawing (PDF) per made part\n"
            "dxf/         the same views and features as DXF\n"
            "assembly/    assembly STEP and the assembly sheet with the bill of materials\n"
            "po/          Excel purchase pack and the three order documents (raw material, machining, bought-out)\n"
            "drawing_data/ the technical data of every part (material grade, hardness, tolerances, notes) as Excel\n"
            "layout/      sheet metal: blank and strip layout drawings (A3 PDF, DXF)\n\n"
            "The 3D (STEP) file governs where it differs from a drawing. Concept-level tool design generated by Design to Cost: check before cutting steel.\n"
            + ("\nDRAFT: this design is not released. Do not order from it.\n" if d["status"] != "released" else "")))
    return buf.getvalue(), f"{code}_tool_v{d['version']}_rev{d['rev_label']}_{d['status']}_supplier_package.zip"


# ------------------------------------------------------------ review

def review_decide(design_id: int, item_id: int, status: str, comment: str) -> dict:
    with db.engine().begin() as conn:
        d = _require_draft(conn, design_id)
        try:
            row = tool_review.decide(conn, design_id, item_id, status, comment, _user())
        except ValueError as exc:
            raise RegisterError(str(exc)) from None
        register._log(conn, "part", d["part_id"], "tool_review", {"design": d["version"], "item": item_id, "status": status})
    return row


def fea_file(design_id: int) -> tuple[bytes, str]:
    with db.engine().connect() as conn:
        d = _design_row(conn, design_id)
        fid = (d["summary"] or {}).get("fea_file_id")
        if not fid:
            raise RegisterError("No structural-check result file for this design (Gmsh and CalculiX are needed).")
        f = _file_row(conn, fid)
    return file_store.path_of(f).read_bytes(), f["original_name"]


# ------------------------------------------------------------ supplier quotes

QUOTE_COLUMNS = {"material": ("PO-1 Raw material", "Part no.", "Quoted rate per kg"),
                 "machining": ("PO-2 Machining", "Part no.", "Quoted price each"),
                 "bought_out": ("PO-3 Bought-out", "Description", "Quoted unit price")}


def quotes_import(design_id: int, content: bytes, supplier: str) -> dict:
    """Read the quoted prices a supplier typed into the purchase-pack workbook (the 'Quoted …' columns) and keep them with the design."""
    from openpyxl import load_workbook
    if not supplier.strip():
        raise RegisterError("Say which supplier quoted.")
    try:
        wb = load_workbook(io.BytesIO(content), data_only=True)
    except Exception:
        raise RegisterError("The file is not an Excel workbook (.xlsx).") from None
    d = get_design(design_id)
    parts = {m["part_no"]: m for m in d["bill"]["made"]}
    items = {l["item"] for l in d["bill"]["standard"]}
    found, skipped = [], []
    for kind, (sheet, key_head, val_head) in QUOTE_COLUMNS.items():
        if sheet not in wb.sheetnames:
            continue
        rows = list(wb[sheet].iter_rows(values_only=True))
        head_i = next((i for i, r in enumerate(rows) if r and key_head in [str(x) for x in r if x] and any(str(x or "").startswith(val_head) for x in r)), None)
        if head_i is None:
            continue
        head = [str(x or "") for x in rows[head_i]]
        kc, vc = head.index(key_head), next(i for i, h in enumerate(head) if h.startswith(val_head))
        for r in rows[head_i + 1:]:
            key, val = (r[kc] if kc < len(r) else None), (r[vc] if vc < len(r) else None)
            if key in (None, "") or val in (None, ""):
                continue
            try:
                val = float(val)
            except (TypeError, ValueError):
                skipped.append(f"{key}: '{val}' is not a number")
                continue
            if val <= 0 or (kind != "bought_out" and key not in parts) or (kind == "bought_out" and key not in items):
                skipped.append(f"{key}: not a part number / item of this design" if val > 0 else f"{key}: price must be more than 0")
                continue
            found.append((kind, str(key), val))
    if not found:
        raise RegisterError("No quoted prices found: type them into the 'Quoted …' columns of PO-1, PO-2 or PO-3 of the purchase pack and upload that file." +
                            (" " + "; ".join(skipped[:3]) if skipped else ""))
    with db.engine().begin() as conn:
        _require_draft_or_released(conn, design_id)
        conn.execute(db.tool_quotes.delete().where((db.tool_quotes.c.design_id == design_id) & (db.tool_quotes.c.supplier == supplier.strip())))
        for kind, key, val in found:
            steel = parts[key]["info"]["steel"] if kind == "material" else ""
            conn.execute(db.tool_quotes.insert().values(design_id=design_id, kind=kind, key=key, steel=steel, value=val,
                                                        currency=d["bill"]["currency"] if kind != "bought_out" else "INR", supplier=supplier.strip(),
                                                        created_by=_user()))
    return {"quotes": len(found), "skipped": skipped, "supplier": supplier.strip()}


def _require_draft_or_released(conn, design_id: int) -> dict:
    d = _design_row(conn, design_id)
    if d["status"] == "superseded":
        raise RegisterError("This design is superseded.")
    return d


def quotes_apply(design_id: int, approver: str = "") -> dict:
    """Make the quotes the approved values of the databases: quoted steel rates → approved tool-steel rates of the tool country (new
    version, not an example); quoted bought-out prices → Standard parts unit costs. Machining quotes only give the quote / estimate ratio."""
    d = get_design(design_id)
    quotes = d["quotes"]
    if not quotes:
        raise RegisterError("There are no quotes on this design.")
    month = rates.current_month()
    done = {"steel_rates": [], "standard_parts": [], "machining_ratio": None}
    by_steel = {}
    made = {m["part_no"]: m for m in d["bill"]["made"]}
    for q in quotes:
        if q["kind"] == "material" and q["key"] in made:
            kg = made[q["key"]]["info"]["stock"]["weight_kg"] * made[q["key"]]["qty"]
            a = by_steel.setdefault(q["steel"], [0.0, 0.0])
            a[0] += q["value"] * kg
            a[1] += kg
    for steel, (amt, kg) in by_steel.items():
        rate = round(amt / kg, 2)
        rates.create({"category": "tool_material", "item": steel, "country": d["country"], "value": rate, "unit": f"{d['bill']['currency']}/kg",
                      "currency": d["bill"]["currency"], "month": month, "source": f"Supplier quote ({', '.join(sorted({q['supplier'] for q in quotes}))}) "
                      f"for {d['summary'].get('part_code', '')} v{d['version']}", "comment": "", "is_example": False},
                     approve=True, origin="quote", approver=approver or _user())
        done["steel_rates"].append({"steel": steel, "rate": rate})
    db_ = mould_db.tables()
    changed = False
    for q in quotes:
        if q["kind"] == "bought_out":
            for r in db_["standard_parts"]:
                if r["item"] == q["key"]:
                    r["unit_cost"], r["is_example"], changed = q["value"], 0, True
                    done["standard_parts"].append({"item": q["key"], "unit_cost_inr": q["value"]})
    if changed:
        mould_db.write(db_)
    mq = [q for q in quotes if q["kind"] == "machining" and q["key"] in made]
    if mq:
        est = sum(made[q["key"]]["info"]["machining_cost"] * made[q["key"]]["qty"] for q in mq)
        done["machining_ratio"] = round(sum(q["value"] * made[q["key"]]["qty"] for q in mq) / est, 3) if est else None
    return done
