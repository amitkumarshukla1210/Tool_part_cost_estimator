"""Logo and company name ("Design to Cost" unless changed on Admin → Drawings) on every Excel and PDF the software hands out,
also inside ZIP packs (roadmap stage 35).

Applied centrally to the response of a download (middleware in main.py), so each exporter stays as it is:
    xlsx   a picture with the logo and the name to the right of the used cells of the first sheet (no cell is moved or
           written, so the layout and the values stay as the exporter made them) and the name in the print header of every sheet
    pdf    logo and name in the top margin of every page, only when the first page does not carry the name yet
    zip    the same for the xlsx and pdf files inside it
Files meant to be uploaded again are not stamped (SKIP): master-data, rates and setup downloads, the import template, the price master.
The logo and name are those of the drawings (drawing_style.py: Admin → Drawings).
"""

from __future__ import annotations

import io
import zipfile
from functools import lru_cache

SKIP = ("/download", "/template.xlsx", "/price-master/export.xlsx")
MEDIA = {"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "xlsx", "application/pdf": "pdf", "application/zip": "zip"}


def _brand() -> tuple:
    from . import drawing_style
    return str(drawing_style.logo_path()), drawing_style.settings()["company"]


@lru_cache(maxsize=4)
def _banner_png(logo: str, name: str) -> bytes:
    """Logo with the name next to it, as one picture for Excel."""
    from PIL import Image, ImageDraw, ImageFont
    h = 48
    lg = Image.open(logo).convert("RGBA")
    lg.thumbnail((h, h))
    font = None
    for f in ("DejaVuSans-Bold.ttf", "arialbd.ttf", "Arial Bold.ttf"):
        try:
            font = ImageFont.truetype(f, 22)
            break
        except OSError:
            continue
    font = font or ImageFont.load_default()
    w = lg.width + 12 + int(ImageDraw.Draw(Image.new("RGB", (1, 1))).textlength(name, font=font)) + 8
    im = Image.new("RGBA", (w, h), (255, 255, 255, 0))
    im.paste(lg, (0, (h - lg.height) // 2), lg)
    ImageDraw.Draw(im).text((lg.width + 12, h // 2), name, font=font, fill=(21, 34, 56, 255), anchor="lm")
    buf = io.BytesIO()
    im.save(buf, "PNG")
    return buf.getvalue()


def stamp_xlsx(data: bytes) -> bytes:
    from openpyxl import load_workbook
    from openpyxl.drawing.image import Image as XlImage
    from openpyxl.utils import get_column_letter
    logo, name = _brand()
    wb = load_workbook(io.BytesIO(data))
    if not wb.worksheets:
        return data
    first = wb.worksheets[0]
    if any(getattr(i, "_dtc_brand", False) for i in getattr(first, "_images", [])) or first.oddHeader.center.text == name:
        return data
    img = XlImage(io.BytesIO(_banner_png(logo, name)))
    img.height, img.width = 32, int(32 * img.width / img.height)
    first.add_image(img, f"{get_column_letter(first.max_column + 2)}1")
    for ws in wb.worksheets:
        ws.oddHeader.center.text = name
        ws.oddHeader.center.size = 10
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def stamp_pdf(data: bytes) -> bytes:
    from pypdf import PdfReader, PdfWriter
    from reportlab.lib.units import mm
    from reportlab.pdfgen import canvas
    logo, name = _brand()
    rd = PdfReader(io.BytesIO(data))
    if not rd.pages or name in (rd.pages[0].extract_text() or ""):
        return data
    wr = PdfWriter()
    for page in rd.pages:
        w, h = float(page.mediabox.width), float(page.mediabox.height)
        buf = io.BytesIO()
        c = canvas.Canvas(buf, pagesize=(w, h))
        size = 7 * mm
        x, y = 6 * mm, h - size - 2 * mm
        c.drawImage(logo, x, y, size, size, mask="auto", preserveAspectRatio=True)
        c.setFont("Helvetica-Bold", 9)
        c.setFillColorRGB(0.08, 0.13, 0.22)
        c.drawString(x + size + 2 * mm, y + 2.4 * mm, name)
        c.save()
        buf.seek(0)
        page.merge_page(PdfReader(buf).pages[0])
        wr.add_page(page)
    out = io.BytesIO()
    wr.write(out)
    return out.getvalue()


def stamp_zip(data: bytes) -> bytes:
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for info in src.infolist():
            z.writestr(info, stamp(src.read(info.filename), info.filename))
    return out.getvalue()


def stamp(data: bytes, name: str) -> bytes:
    """The file with the brand; unchanged when it cannot be read (a download must never fail because of the logo)."""
    ext = name.lower().rsplit(".", 1)[-1] if "." in name else ""
    fn = {"xlsx": stamp_xlsx, "pdf": stamp_pdf, "zip": stamp_zip}.get(ext)
    if not fn:
        return data
    try:
        return fn(data)
    except Exception:
        return data
