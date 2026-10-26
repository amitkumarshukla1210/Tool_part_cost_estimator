"""Price master (roadmap stage 29): month-wise prices of tool materials and part raw materials by country.

The prices are ordinary versioned rates of the rate library (rates.py); this module shows and edits them as an
item × month grid for one country:

    group           library category   items (master list)
    tool_material   tool_material      Mould & die database → Tool materials (type, grade, density, hardness)
    part_material   material           Master data → Materials (part type, type, grade, density)
    part_scrap      material_scrap     Master data → Materials (the scrap value of each grade)

Standard parts keep their catalogue price (Mould & die database → Standard parts); they are not in this grid.

A cell shows the approved value that costing uses in that month (the latest approved version on or before it) and
whether a version starts in that month. Typing a price saves a NEW version for that month: a draft for Rates → Rate
review, or approved at once when a user with the permission Approve rates ticks "Save & approve"; old versions are never changed, so stored costs only move when
you recalculate them. Nothing is filled in for months without a rate: the grid fills up as real months are entered.
New grades are added where they always are (Tool BOM → + New grade, Master data → Tool materials / Materials).
"""

from __future__ import annotations

import io
from contextlib import closing

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from . import costing_setup, master_data, mould_db, rates
from .table_set import EXAMPLE_FILL, HEADER_FILL

GROUPS = {
    "tool_material": {"title": "Tool materials (per kg)", "category": "tool_material", "per": "kg"},
    "part_material": {"title": "Part raw materials (per kg)", "category": "material", "per": "kg"},
    "part_scrap": {"title": "Part raw material scrap value (per kg)", "category": "material_scrap", "per": "kg"},
}
MONTHS = 24


class PriceError(ValueError):
    pass


def _group(group: str) -> dict:
    if group not in GROUPS:
        raise PriceError(f"Unknown price group '{group}' ({', '.join(GROUPS)}).")
    return GROUPS[group]


def add_months(month: str, n: int) -> str:
    y, m = map(int, month.split("-"))
    k = y * 12 + m - 1 + n
    return f"{k // 12:04d}-{k % 12 + 1:02d}"


def months(end: str | None = None, count: int = MONTHS) -> list[str]:
    end = end or rates.current_month()
    return [add_months(end, -k) for k in range(count - 1, -1, -1)]


def _countries() -> dict:
    return {c["code"]: c for c in costing_setup.tables()["countries"]}


def items(group: str) -> list[dict]:
    """The master list of the group: one row per library item, with what describes it."""
    _group(group)
    if group == "tool_material":
        return [{"item": r["code"], "type": r.get("material_type") or "", "grade": r["code"], "info": r["application"],
                 "density": r["density_g_cm3"], "hardness": r["hardness"], "is_example": r["is_example"]}
                for r in mould_db.tables()["tool_materials"]]
    return [{"item": r["material"], "type": r["material_type"], "grade": r["material"], "info": r["part_type"],
             "density": r["density_g_cm3"], "recommended": bool(r["recommended"]), "is_example": False}
            for r in master_data.tables()["Materials"]]


def grid(group: str, country: str, end: str | None = None, count: int = MONTHS) -> dict:
    """Items × months of one country: the approved value costing uses in each month, and any open drafts."""
    g = _group(group)
    country = (country or "").upper()
    countries = _countries()
    if country not in countries:
        raise PriceError(f"Country '{country}' is not in Master data → Countries.")
    end = end or rates.current_month()
    cols = months(add_months(end, 1), count + 1)  # one month ahead, to enter next month's prices
    with closing(rates._connect()) as con:
        rows = [rates._row(r) for r in con.execute(
            "SELECT * FROM rates WHERE category = ? AND country = ? AND month <= ? ORDER BY item, month, approved_date, id",
            (g["category"], country, cols[-1]))]
    by_item: dict[str, list] = {}
    for r in rows:
        by_item.setdefault(r["item"], []).append(r)
    out = []
    for it in items(group):
        versions = by_item.pop(it["item"], [])
        out.append(it | {"cells": _cells(versions, cols)})
    for item, versions in sorted(by_item.items()):  # rates of items no longer in the master list
        out.append({"item": item, "type": "", "grade": item, "info": "not in the master list", "is_example": False,
                    "cells": _cells(versions, cols), "orphan": True})
    return {"group": group, "title": g["title"], "category": g["category"], "country": country,
            "currency": countries[country]["currency"], "unit": f"{countries[country]['currency']}/{g['per']}",
            "months": cols, "current_month": rates.current_month(), "rows": out}


def _cells(versions: list[dict], cols: list[str]) -> dict:
    approved = [v for v in versions if v["status"] == "approved"]
    cells = {}
    for m in cols:
        best = None
        for v in approved:  # sorted by month, approval date, id: the last one on or before m wins
            if v["month"] <= m:
                best = v
        drafts = [v for v in versions if v["month"] == m and v["status"] in rates.OPEN]
        if best is None and not drafts:
            continue
        cells[m] = {"value": best["value"] if best else None, "id": best["id"] if best else None,
                    "is_example": bool(best and best["is_example"]), "starts": bool(best and best["month"] == m),
                    "source": best["source"] if best else "", "draft": drafts[-1]["value"] if drafts else None,
                    "draft_id": drafts[-1]["id"] if drafts else None}
    return cells


def set_price(group: str, item: str, country: str, month: str, value, approve: bool = False, comment: str = "",
              source: str = "") -> dict:
    """A new version of one price for that month (approved, or a draft for review)."""
    g = _group(group)
    country = (country or "").upper()
    countries = _countries()
    if country not in countries:
        raise PriceError(f"Country '{country}' is not in Master data → Countries.")
    if item not in {i["item"] for i in items(group)}:
        raise PriceError(f"'{item}' is not in the master list of {g['title']}. Add the grade first (Tool BOM → + New grade, or Master data).")
    cur = countries[country]["currency"]
    return rates.create({"category": g["category"], "item": item, "country": country, "value": value,
                         "unit": f"{cur}/{g['per']}", "currency": cur, "month": month, "comment": comment,
                         "source": source or "Price master (typed)", "is_example": False},
                        approve=approve, origin="price master")


# ---------- Excel ----------

def to_excel(group: str, country: str, end: str | None = None, count: int = MONTHS) -> bytes:
    """The grid as a sheet: item, type, grade, info, then one column per month (EXAMPLE values shaded)."""
    g = grid(group, country, end, count)
    wb = Workbook()
    ws = wb.active
    ws.title = "Prices"
    wb.properties.title = f"Design to Cost — Price master: {g['title']} · {g['country']}"
    wb.properties.creator = "Design to Cost"
    head = ["item", "type", "grade / make", "info"] + g["months"]
    ws.append(head)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
    for i, r in enumerate(g["rows"], start=2):
        ws.append([r["item"], r["type"], r["grade"], r["info"]] + [(r["cells"].get(m) or {}).get("value") for m in g["months"]])
        for j, m in enumerate(g["months"], start=5):
            if (r["cells"].get(m) or {}).get("is_example"):
                ws.cell(row=i, column=j).fill = EXAMPLE_FILL
    ws.append([])
    ws.append([f"Unit: {g['unit']}. Shaded = EXAMPLE value. Type a price under a month and upload the file on the Price "
               "master page: every changed cell becomes a new version for that month; the rest is left alone."])
    ws.freeze_panes = "E2"
    ws.column_dimensions["A"].width = 38
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def import_excel(group: str, country: str, content: bytes, approve: bool = False) -> dict:
    """Every cell that differs from the value costing uses in that month becomes a new version for that month."""
    current = grid(group, country, rates.current_month(), 1)  # only to check the country and group
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise PriceError(f"Could not read the uploaded file as Excel (.xlsx): {exc}") from exc
    ws = wb["Prices"] if "Prices" in wb.sheetnames else wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows or not rows[0] or str(rows[0][0] or "").strip() != "item":
        raise PriceError("The first column of the sheet must be 'item' (download the grid to get the layout).")
    month_cols = {j: str(h).strip() for j, h in enumerate(rows[0]) if h is not None and rates.MONTH.match(str(h).strip())}
    if not month_cols:
        raise PriceError("The sheet has no month columns (headers like 2026-10).")
    known = {i["item"] for i in items(group)}
    wanted = []
    for n, row in enumerate(rows[1:], start=2):
        item = str(row[0] or "").strip() if row else ""
        if not item or item.startswith("Unit:"):
            continue
        if item not in known:
            raise PriceError(f"Row {n}: '{item}' is not in the master list. Add the grade first (Tool BOM → + New grade, or Master data).")
        for j, m in month_cols.items():
            v = row[j] if j < len(row) else None
            if v in (None, ""):
                continue
            try:
                v = float(str(v).replace(",", ""))
            except ValueError:
                raise PriceError(f"Row {n}, {m}: '{v}' is not a number.") from None
            wanted.append((item, m, v))
    lo, hi = min(m for _, m, _ in wanted) if wanted else None, max(m for _, m, _ in wanted) if wanted else None
    made = 0
    if wanted:
        span = (int(hi[:4]) * 12 + int(hi[5:])) - (int(lo[:4]) * 12 + int(lo[5:])) + 1
        cells = {r["item"]: r["cells"] for r in grid(group, current["country"], add_months(hi, -1), span)["rows"]}
        for item, m, v in wanted:
            cell = cells.get(item, {}).get(m) or {}
            if cell.get("value") is not None and abs(float(cell["value"]) - v) < 1e-9:
                continue
            set_price(group, item, current["country"], m, v, approve, source="Price master (Excel upload)")
            made += 1
    return {"versions_created": made, "status": "approved" if approve else "draft"}
