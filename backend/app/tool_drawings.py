"""2D drawings and the release pack of a detailed tool design (roadmap stage 27).

    assembly drawing    iso view (closed), exploded iso view, section through the middle, bill of material with item
                        numbers, material, hardness, make / standard part, quantity; notes (shut height, press, fits)
    detail drawings     one sheet per made part: plates as true vector views (top view with every hole and opening,
                        front and side view), overall dimensions with tolerances, hole table from the datum corner
                        (X, Y, size, kind, fit, depth); other parts (punches, V-die, blade) as three projected views with
                        their overall dimensions; material, hardness, heat treatment, finish, quantity on the sheet
    every sheet         logo, company, drawing number, revision, general tolerance, projection, drawn / checked /
                        approved (drawing_style.py) - PDF (A3) and DXF
    release pack (ZIP)  3D: assembly STEP + a STEP per part; 2D: PDF and DXF per sheet and the full drawing set;
                        BOM (Excel); README
"""

from __future__ import annotations

import io
import zipfile
from datetime import date

import ezdxf

from . import drawing_style, render3d, tool_design


def _safe(s: str) -> str:
    s = s.replace("Ø", "D").replace("×", "x").replace("°", "deg")
    return "".join(ch if (ch.isascii() and ch.isalnum()) or ch in "-_." else "_" for ch in s).strip("_")


STD_SCALES = (10, 5, 4, 2.5, 2, 1.5, 1, 1 / 1.5, 1 / 2, 1 / 2.5, 1 / 4, 1 / 5, 1 / 10, 1 / 15, 1 / 20, 1 / 25, 1 / 50, 1 / 100)


def _std_scale(fit_pts_per_mm: float) -> float:
    """The largest standard drawing scale (ISO 5455) that fits, in points per model mm."""
    from reportlab.lib.units import mm
    r = fit_pts_per_mm / mm
    return next((s for s in STD_SCALES if s <= r + 1e-9), STD_SCALES[-1]) * mm


def drawing_no(part: dict, stage: int, comp_no: str) -> str:
    return f"{part['part_code']}-T{stage:02d}-{comp_no}"


# ---------- views ----------

def _render_png(design, size, az, el, explode=0.0, section=False, skip_ref=False) -> bytes:
    import cadquery as cq
    items = render3d.design_items(design, tool_design.COLORS, explode=explode, skip_ref=skip_ref)
    if section:   # cut away everything in front of the middle plane (y < 0)
        items = []
        for c in design["components"]:
            for (x, y, z) in c.locs:
                bb = c.shape.BoundingBox()
                if y + bb.ymax <= 0:
                    continue
                s = c.shape.moved(cq.Location(cq.Vector(x, y, z)))
                if y + bb.ymin < 0:
                    big = 1e4
                    s = s.cut(cq.Solid.makeBox(2 * big, big, 2 * big, cq.Vector(-big, -big, -big)))
                if not s.Solids():
                    continue
                v, t = render3d.mesh(s, tol=max(0.3, max(c.size) / 150))
                items.append((v, t, tool_design.COLORS.get(c.color, (0.6, 0.6, 0.6))))
    return render3d.png(render3d.render(items, size=size, az=az, el=el))


def _img(png: bytes, w, h):
    from reportlab.platypus import Image
    return Image(io.BytesIO(png), w, h)


def _dt(c, key: str) -> str:
    """Tolerance text of an overall dimension (stage 32), '' when the part has none attached."""
    dims = getattr(c, "dims", None)
    return dims["overall"][key]["text"] if dims else ""


def _dt_size(c, value: float) -> str:
    """Tolerance of a solid part's view dimension: the overall dimension (L, W, H) closest to the value."""
    dims = getattr(c, "dims", None)
    if not dims:
        return ""
    key = min(zip("LWH", c.size), key=lambda kv: abs(kv[1] - value))[0]
    return dims["overall"][key]["text"]


def _plate_views(c, theme, W, H):
    """Vector views of a plate: top view with holes and openings, front view below, side view at the right."""
    from reportlab.graphics.shapes import Circle, Drawing, Line, PolyLine, Rect, String
    from .drawings import _dim
    L, Wd, T = c.size
    gap = 22
    sc = _std_scale(min((W - 3 * gap - 20) / (L + T), (H - 3 * gap - 20) / (Wd + T)))
    d = Drawing(W, H)
    ox, oy = gap + 10, gap + T * sc + gap
    ink = theme["ink"]
    d.add(Rect(ox, oy, L * sc, Wd * sc, fillColor=None, strokeColor=ink, strokeWidth=0.8))
    for o in c.openings:
        d.add(PolyLine([v for x, y in o["pts"] for v in (ox + x * sc, oy + y * sc)], strokeColor=ink, strokeWidth=0.6))
    for i, h in enumerate(c.holes, start=1):
        d.add(Circle(ox + h["x"] * sc, oy + h["y"] * sc, max(h["d"] / 2 * sc, 0.6), fillColor=None, strokeColor=ink, strokeWidth=0.5))
        if h.get("cb_d"):
            d.add(Circle(ox + h["x"] * sc, oy + h["y"] * sc, h["cb_d"] / 2 * sc, fillColor=None, strokeColor=ink, strokeWidth=0.3,
                         strokeDashArray=None if h.get("top", True) else [1.5, 1]))
        d.add(String(ox + h["x"] * sc + h["d"] / 2 * sc + 1, oy + h["y"] * sc + 1, f"H{i}", fontSize=5.5, fillColor=theme["brand"]))
    d.add(String(ox - 2, oy - 8, "0,0", fontSize=6, fillColor=theme["brand"], textAnchor="end"))
    d.add(Circle(ox, oy, 1.6, fillColor=theme["brand"], strokeColor=theme["brand"]))
    _dim(d, ox, oy + Wd * sc + 8, ox + L * sc, oy + Wd * sc + 8, f"{L:.1f} {_dt(c, 'L')}".strip(), theme)
    _dim(d, ox - 10, oy, ox - 10, oy + Wd * sc, f"{Wd:.1f} {_dt(c, 'W')}".strip(), theme)
    d.add(String(ox, oy + Wd * sc + 20, "TOP VIEW (datum 0,0 = lower-left corner)", fontSize=7, fillColor=theme["grey"]))
    # front view
    fy = gap
    d.add(Rect(ox, fy, L * sc, T * sc, fillColor=None, strokeColor=ink, strokeWidth=0.8))
    for h in c.holes:
        x = ox + h["x"] * sc
        depth = (h["depth"] or T) * sc
        y0 = fy + (T * sc - depth if h.get("top", True) else 0)
        for dx in (-h["d"] / 2 * sc, h["d"] / 2 * sc):
            d.add(Line(x + dx, y0, x + dx, y0 + depth, strokeColor=theme["grey"], strokeWidth=0.3, strokeDashArray=[1.5, 1]))
    _dim(d, ox + L * sc + 8, fy, ox + L * sc + 8, fy + T * sc, f"{T:g} {_dt(c, 'H')}".strip(), theme)
    d.add(String(ox, fy - 9, "FRONT VIEW", fontSize=7, fillColor=theme["grey"]))
    # side view
    sx = ox + L * sc + gap + 10
    d.add(Rect(sx, oy, T * sc, Wd * sc, fillColor=None, strokeColor=ink, strokeWidth=0.8))
    d.add(String(sx, oy - 9, "SIDE VIEW", fontSize=7, fillColor=theme["grey"]))
    return d, sc


def _solid_views(c, theme, W, H):
    from reportlab.graphics.shapes import Drawing, PolyLine, String
    from .drawings import VIEWS, _dim, edge_polylines, project
    lines = edge_polylines(c.shape)
    vs = []
    for title, h, v in VIEWS:
        polys, (w, ht) = project(lines, h, v)
        vs.append({"title": title, "polys": polys, "w": max(w, 1e-6), "h": max(ht, 1e-6)})
    top, front, side = vs
    gap = 22
    sc = _std_scale(min((W - 3 * gap) / (front["w"] + side["w"]), (H - 3 * gap) / (front["h"] + top["h"])))
    d = Drawing(W, H)
    places = {"FRONT VIEW": (gap, gap), "TOP VIEW": (gap, gap + front["h"] * sc + gap),
              "SIDE VIEW": (gap + front["w"] * sc + gap, gap)}
    for v in vs:
        px, py = places[v["title"]]
        for poly in v["polys"]:
            d.add(PolyLine([q for x, y in poly for q in (px + x * sc, py + y * sc)], strokeColor=theme["ink"], strokeWidth=0.6))
        d.add(String(px, py - 10, v["title"], fontSize=7, fillColor=theme["grey"]))
    fx, fy = places["FRONT VIEW"]
    _dim(d, fx, fy - 4, fx + front["w"] * sc, fy - 4, f"{front['w']:.2f} {_dt_size(c, front['w'])}".strip(), theme, side=-1)
    _dim(d, fx - 6, fy, fx - 6, fy + front["h"] * sc, f"{front['h']:.2f} {_dt_size(c, front['h'])}".strip(), theme)
    tx, ty = places["TOP VIEW"]
    _dim(d, tx - 6, ty, tx - 6, ty + top["h"] * sc, f"{top['h']:.2f} {_dt_size(c, top['h'])}".strip(), theme)
    return d, sc


def _scale_text(sc_pts_per_mm: float) -> str:
    from reportlab.lib.units import mm
    r = sc_pts_per_mm / mm
    return f"1:{1 / r:g}" if r < 1 else f"{r:g}:1"


# ---------- PDF ----------

def _meta(part: dict, design: dict, c=None) -> tuple[dict | None, list | None]:
    """Release (design loop) and revision history of a drawing (stage 31); none for a part not in the register."""
    if not part.get("id"):
        return None, None
    from . import drawing_data
    return drawing_data.release(part["id"]), drawing_data.history(part["id"], design.get("key"), c.no if c is not None else None)


def _gen_tol(c, tol) -> str | None:
    """The general tolerance text of a detail drawing when its class was changed (else the Design rules text)."""
    tech = getattr(c, "tech", None)
    if not tech:
        return None
    from . import drawing_data
    return None if drawing_data.tolerance_class(tol.get("general_tolerance", "")) == tech["tolerance_class"] else f"ISO 2768-{tech['tolerance_class']}"


def _bom_rows(design):
    rows = [["Item", "Qty", "Description", "Material / grade", "Hardness", "Make / standard part", "Size (mm)"]]
    for c in design["components"]:
        if c.kind == "ref":
            continue
        L, W, H = c.size
        rows.append([c.no, str(c.qty), c.name[:44], (c.material or "-")[:30], (c.hardness or "-")[:22],
                     (c.std or c.make or "-")[:52], f"{L:.1f} × {W:.1f} × {H:.1f}"])
    return rows


def design_pdf(part: dict, design: dict, stage: int, rev: str = "", only: str | None = None) -> bytes:
    """Drawing set of one die / mould: assembly sheet + a detail sheet per made part (or only one sheet: 'ASM' / item no)."""
    from reportlab.lib.units import mm
    from reportlab.platypus import PageBreak, Paragraph, Spacer, Table, TableStyle
    from .drawings import _doc, _styles
    from .pdf_report import _table
    theme, ss, st = _styles()
    pt = design.get("part_type", "sheet_metal")
    tol = drawing_style.tolerances(pt)
    made = [c for c in design["components"] if c.kind == "made" and c.drawing != "none"]
    sheets = (1 + len(made))
    story = []
    buf = io.BytesIO()
    doc = _doc(buf, f"{part['part_code']} {design['name']} drawings")
    width = 396 * mm
    if only in (None, "ASM"):
        iso = _render_png(design, (1100, 800), -50, 28)
        exp = _render_png(design, (1100, 800), -55, 20, explode=0.9)
        sec = _render_png(design, (1100, 800), -90, 0, section=True)
        pics = Table([[_img(iso, 128 * mm, 93 * mm), _img(exp, 128 * mm, 93 * mm), _img(sec, 128 * mm, 93 * mm)],
                      ["ISO VIEW (closed, bottom of stroke)", "EXPLODED VIEW", "SECTION A-A (through the middle, from the front)"]],
                     colWidths=[132 * mm] * 3)
        pics.setStyle(TableStyle([("FONTSIZE", (0, 1), (-1, 1), 7), ("TEXTCOLOR", (0, 1), (-1, 1), theme["grey"]),
                                  ("BOX", (0, 0), (0, 0), 0.3, theme["border"]), ("BOX", (1, 0), (1, 0), 0.3, theme["border"]),
                                  ("BOX", (2, 0), (2, 0), 0.3, theme["border"])]))
        rows = [("Part", f"{part['part_code']} · {part['name']}"), ("Tool", design["name"]),
                ("Die set / base", design.get("die_set", "")), ("Shut height", f"{design.get('shut_height', 0):g} mm")]
        story += [drawing_style.pdf_header(st, "TOOL ASSEMBLY", f"{part['part_code']} · {design['name']}"), pics,
                  Spacer(1, 2 * mm), _left(_table(_bom_rows(design), theme, [12 * mm, 10 * mm, 80 * mm, 62 * mm, 44 * mm, 128 * mm, 60 * mm])),
                  Spacer(1, 2 * mm),
                  drawing_style.pdf_title_block(width, rows, pt, "as shown", f"1 / {sheets}", drawing_no(part, stage, "00"), rev,
                                                design.get("notes", []) + _fit_notes(tol), *_meta(part, design))]
    for k, c in enumerate(made, start=2):
        if only not in (None, c.no):
            continue
        if story:
            story.append(PageBreak())
        dw, dh = 300 * mm, 170 * mm
        views, sc = (_plate_views if c.drawing == "plate" else _solid_views)(c, theme, dw, dh)
        side = _detail_side(c, theme, tol)
        L, W, H = c.size
        tech = getattr(c, "tech", None)
        rows = [("Part (tool for)", f"{part['part_code']} · {design['name']}"), ("Description", c.name), ("Item / qty", f"{c.no} · {c.qty} off"),
                ("Material", c.material or "-"), ("Hardness", c.hardness or "-"),
                ("Size", f"{L:.2f} × {W:.2f} × {H:.2f} mm"), ("Make", c.make or "-"),
                ("Finish", f"Ra {tech['finish_ra']:g} µm unless stated" if tech else tol.get("finish_general", ""))]
        if tech:
            rows += [("Heat treatment", tech["heat_treatment"]), ("Surface treatment", tech["surface_treatment"])]
        story += [drawing_style.pdf_header(st, "DETAIL DRAWING", f"{c.no} · {c.name}"),
                  Table([[views, side]], colWidths=[dw + 4 * mm, 92 * mm], style=TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP")]))]
        if len(c.holes) > SIDE_HOLES:
            story += [Paragraph("Hole table (datum 0,0 = lower-left corner)", st["k"]),
                      _hole_table(c, theme, tol, [14 * mm, 24 * mm, 24 * mm, 16 * mm, 300 * mm])]
        rel, hist = _meta(part, design, c)
        story += [drawing_style.pdf_title_block(width, rows, pt, _scale_text(sc), f"{k} / {sheets}", drawing_no(part, stage, c.no), rev,
                                                c.notes, rel, hist, _gen_tol(c, tol))]
    doc.build(story)
    return buf.getvalue()


def _left(t):
    from reportlab.platypus import TableStyle
    t.setStyle(TableStyle([("ALIGN", (0, 0), (-1, -1), "LEFT"), ("FONTSIZE", (0, 0), (-1, -1), 7),
                           ("TOPPADDING", (0, 0), (-1, -1), 1.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 1.5)]))
    return t


def _fit_notes(tol):
    return [f"Fits: pillar {tol.get('fit_pillar', '')}, bush {tol.get('fit_bush', '')}, dowel {tol.get('fit_dowel', '')}"
            + (f", punch {tol['fit_punch']}" if tol.get("fit_punch") else "")]


SIDE_HOLES = 22      # more holes than this: the hole table goes full width under the views (and runs on to the next page)


def _widths7(widths):
    """The 5 column widths of the hole table spread over 7 columns (Ø tolerance and position added, stage 32)."""
    total = sum(widths)
    from reportlab.lib.units import mm
    if total < 120 * mm:                        # side panel
        return [7 * mm, 11 * mm, 11 * mm, 7 * mm, 17 * mm, 13 * mm, total - 66 * mm]
    return [14 * mm, 24 * mm, 24 * mm, 16 * mm, 44 * mm, 24 * mm, total - 146 * mm]


def _hole_table(c, theme, tol, widths):
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import Paragraph, Table, TableStyle
    small = ParagraphStyle("s", fontSize=6.5, leading=8)
    dims = getattr(c, "dims", None)
    rows = [["#", "X", "Y", "Ø", "Ø tol.", "Pos. ±", "Kind / fit"]] if dims else [["#", "X", "Y", "Ø", "Kind / fit"]]
    for i, h in enumerate(c.holes, start=1):
        kind = h["label"] or h["kind"]
        if h.get("depth"):
            kind += f", depth {h['depth']:g}" + ("" if h.get("top", True) else " from below")
        if h["kind"] == "fit" and not dims:
            kind += f", position {tol.get('position_tol', '')}"
        if dims:
            t = dims["holes"][f"H{i}"]
            rows.append([f"H{i}", f"{h['x']:.2f}", f"{h['y']:.2f}", f"{h['d']:g}", Paragraph(t["size"]["text"], small),
                         Paragraph(t["pos"]["text"].replace(" µm", ""), small), Paragraph(kind, small)])
        else:
            rows.append([f"H{i}", f"{h['x']:.2f}", f"{h['y']:.2f}", f"{h['d']:g}", Paragraph(kind, small)])
    if dims:
        widths = _widths7(widths)
    t = Table(rows, colWidths=widths, repeatRows=1)
    t.setStyle(TableStyle([("FONTSIZE", (0, 0), (-1, -1), 6.5), ("GRID", (0, 0), (-1, -1), 0.25, theme["border"]),
                           ("BACKGROUND", (0, 0), (-1, 0), theme["head"]), ("TEXTCOLOR", (0, 0), (-1, 0), theme["head_text"]),
                           ("TOPPADDING", (0, 0), (-1, -1), 1), ("BOTTOMPADDING", (0, 0), (-1, -1), 1)]))
    return t


def _detail_side(c, theme, tol):
    """Hole table (if short) and opening notes next to the views."""
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, Table, TableStyle
    small = ParagraphStyle("s", fontSize=6.5, leading=8)
    out = []
    if c.holes and len(c.holes) <= SIDE_HOLES:
        out.append([_hole_table(c, theme, tol, [8 * mm, 13 * mm, 13 * mm, 9 * mm, 47 * mm])])
    elif c.holes:
        out.append([Paragraph(f"<b>{len(c.holes)} holes</b>: hole table under the views.", small)])
    dims = getattr(c, "dims", None)
    for k, o in enumerate(c.openings, start=1):
        tol_txt = dims["openings"][f"O{k}"]["size"]["text"] if dims else o.get("tol")
        out.append([Paragraph(f"<b>{o['label']}</b>" + (f" · tolerance {tol_txt}" if tol_txt else "")
                              + " · profile in the DXF (wire cut / CNC)", small)])
    if dims:
        out.append([Paragraph(f"Tolerances in <b>{dims['unit']}</b>. Overall: L {dims['overall']['L']['text']} · W {dims['overall']['W']['text']} · "
                              f"H {dims['overall']['H']['text']}. Fits to ISO 286 (H7 = hole).", small)])
    if not out:
        out = [[Paragraph("Dimensions from the 3D model (STEP) of this part.", small)]]
    tech = getattr(c, "tech", None)
    if tech:
        out.append([_tech_table(tech, theme)])
    t = Table(out, colWidths=[92 * mm])
    t.setStyle(TableStyle([("LEFTPADDING", (0, 0), (-1, -1), 0), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    return t


def _tech_table(tech: dict, theme):
    """Technical data of a made part and the ISO 2768 tables of its class (stage 31)."""
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, Table, TableStyle
    small = ParagraphStyle("s", fontSize=6.5, leading=8)
    P = lambda x: Paragraph(str(x), small)
    rows = [["TECHNICAL DATA", ""], ["Material grade", P(tech["material"])], ["Hardness", P(tech["hardness"])],
            ["Heat treatment", P(tech["heat_treatment"])], ["Surface treatment", P(tech["surface_treatment"])],
            ["Surface roughness", P(f"Ra {tech['finish_ra']:g} µm unless stated")],
            ["General tolerance", P(f"ISO 2768-{tech['tolerance_class']} (tables below)")]]
    lin, ang = tech["tables"]["linear"], tech["tables"]["angular"]
    n = max(len(lin), len(ang))
    tab = [["Linear size (mm)", "±", "Angle, shorter leg (mm)", "±"]] + [
        [lin[i]["range"] if i < len(lin) else "", lin[i]["value"] if i < len(lin) else "",
         ang[i]["range"] if i < len(ang) else "", ang[i]["value"] if i < len(ang) else ""] for i in range(n)]
    t1 = Table(rows, colWidths=[28 * mm, 64 * mm])
    t2 = Table(tab, colWidths=[26 * mm, 16 * mm, 32 * mm, 18 * mm])
    for t in (t1, t2):
        t.setStyle(TableStyle([("FONTSIZE", (0, 0), (-1, -1), 6.5), ("GRID", (0, 0), (-1, -1), 0.25, theme["border"]),
                               ("BACKGROUND", (0, 0), (-1, 0), theme["head"]), ("TEXTCOLOR", (0, 0), (-1, 0), theme["head_text"]),
                               ("TOPPADDING", (0, 0), (-1, -1), 0.8), ("BOTTOMPADDING", (0, 0), (-1, -1), 0.8), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    out = Table([[t1], [t2]], colWidths=[92 * mm])
    out.setStyle(TableStyle([("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0)]))
    return out


# ---------- DXF ----------

def _dxf_doc():
    doc = ezdxf.new("R2010", setup=True)
    for name, col in (("PART", 7), ("HIDDEN", 8), ("DIM", 3), ("TEXT", 5), ("CENTER", 1)):
        doc.layers.add(name, color=col)
    return doc


def detail_dxf(part: dict, design: dict, stage: int, c, rev: str = "") -> bytes:
    """Detail drawing of one part at 1:1 in model space, with dimensions, hole table and the title block."""
    doc = _dxf_doc()
    msp = doc.modelspace()
    L, W, H = c.size
    tol = drawing_style.tolerances(design.get("part_type", "sheet_metal"))
    if c.drawing == "plate":
        msp.add_lwpolyline([(0, 0), (L, 0), (L, W), (0, W)], close=True, dxfattribs={"layer": "PART"})
        for o in c.openings:
            msp.add_lwpolyline([tuple(p) for p in o["pts"]], close=True, dxfattribs={"layer": "PART"})
        for i, h in enumerate(c.holes, start=1):
            msp.add_circle((h["x"], h["y"]), h["d"] / 2, dxfattribs={"layer": "PART"})
            if h.get("cb_d"):
                msp.add_circle((h["x"], h["y"]), h["cb_d"] / 2, dxfattribs={"layer": "PART" if h.get("top", True) else "HIDDEN"})
            msp.add_line((h["x"] - h["d"] * 0.7, h["y"]), (h["x"] + h["d"] * 0.7, h["y"]), dxfattribs={"layer": "CENTER"})
            msp.add_line((h["x"], h["y"] - h["d"] * 0.7), (h["x"], h["y"] + h["d"] * 0.7), dxfattribs={"layer": "CENTER"})
            msp.add_text(f"H{i}", dxfattribs={"layer": "TEXT", "height": 2.5}).set_placement((h["x"] + h["d"] / 2 + 1, h["y"] + 1))
        fy = -H - 25
        msp.add_lwpolyline([(0, fy), (L, fy), (L, fy + H), (0, fy + H)], close=True, dxfattribs={"layer": "PART"})
        u = lambda txt: txt.replace("µm", "um")
        msp.add_linear_dim(base=(0, W + 12), p1=(0, W), p2=(L, W), dimstyle="EZDXF", text=f"<> {u(_dt(c, 'L'))}".strip() if _dt(c, "L") else "<>",
                           dxfattribs={"layer": "DIM"}).render()
        msp.add_linear_dim(base=(-12, 0), p1=(0, 0), p2=(0, W), angle=90, dimstyle="EZDXF", text=f"<> {u(_dt(c, 'W'))}".strip() if _dt(c, "W") else "<>",
                           dxfattribs={"layer": "DIM"}).render()
        msp.add_linear_dim(base=(L + 12, fy), p1=(L, fy), p2=(L, fy + H), angle=90, dimstyle="EZDXF",
                           text=f"{H:g} {u(_dt(c, 'H')) if _dt(c, 'H') else tol.get('plate_thk_tol', '')}", dxfattribs={"layer": "DIM"}).render()
        # hole table
        tx, ty = L + 40, W
        msp.add_text("HOLE TABLE (datum 0,0 lower-left)", dxfattribs={"layer": "TEXT", "height": 3.5}).set_placement((tx, ty))
        dims = getattr(c, "dims", None)
        for i, h in enumerate(c.holes, start=1):
            tt = (f"  {dims['holes'][f'H{i}']['size']['text']}  pos {dims['holes'][f'H{i}']['pos']['text']}".replace("µm", "um") if dims else "")
            msp.add_text(f"H{i}  X {h['x']:.2f}  Y {h['y']:.2f}  Ø{h['d']:g}{tt}  {h['label'] or h['kind']}"
                         + (f" depth {h['depth']:g}" if h.get("depth") else ""),
                         dxfattribs={"layer": "TEXT", "height": 2.5}).set_placement((tx, ty - 6 * i))
        x1, y1 = tx + 160, W + 10
        y0 = fy - 10
    else:
        from .drawings import VIEWS, edge_polylines, project
        lines = edge_polylines(c.shape)
        x = 0.0
        for title, hh, vv in VIEWS:
            polys, (w, ht) = project(lines, hh, vv)
            for poly in polys:
                msp.add_lwpolyline([(x + a, b) for a, b in poly], dxfattribs={"layer": "PART"})
            msp.add_text(title, dxfattribs={"layer": "TEXT", "height": 3}).set_placement((x, -8))
            msp.add_linear_dim(base=(x, -16), p1=(x, 0), p2=(x + w, 0), dimstyle="EZDXF", dxfattribs={"layer": "DIM"}).render()
            x += w + 40
        x1, y1, y0 = x, max(c.size) + 10, -30
        for o in c.openings:
            msp.add_lwpolyline([(p[0], p[1] - max(c.size) - 60) for p in o["pts"]], close=True, dxfattribs={"layer": "PART"})
            dims = getattr(c, "dims", None)
            o_tol = dims["openings"][f"O{c.openings.index(o) + 1}"]["size"]["text"].replace("µm", "um") if dims else o.get("tol", "")
            msp.add_text(f"PROFILE: {o['label']} {o_tol}", dxfattribs={"layer": "TEXT", "height": 3}).set_placement(
                (0, -max(c.size) - 66))
            y0 = -max(c.size) - 80
    # frame + title block around the drawing
    frame_w, frame_h = max(x1, 420) + 30, y1 - y0 + 90
    off = (-15, y0 - 70)
    from ezdxf.math import Matrix44
    for e in list(msp):
        e.transform(Matrix44.translate(-off[0], -off[1], 0))
    tech = getattr(c, "tech", None)
    tech_rows = [("Heat treatment", tech["heat_treatment"]), ("Surface treatment", tech["surface_treatment"]),
                 ("Roughness", f"Ra {tech['finish_ra']:g} um unless stated")] if tech else []
    if tech:                                             # ISO 2768 tables of the class, left of the title block
        tx, ty = 0, 8
        msp.add_text(f"GENERAL TOLERANCES ISO 2768-{tech['tolerance_class']}", dxfattribs={"layer": "TEXT", "height": 3}).set_placement((tx, ty + 6 * 0 + 40))
        msp.add_text("Linear (mm): " + "  ".join(f"{r['range']}: {r['value']}" for r in tech["tables"]["linear"]),
                     dxfattribs={"layer": "TEXT", "height": 2.3}).set_placement((tx, ty + 34))
        msp.add_text("Angular (shorter leg mm): " + "  ".join(f"{r['range']}: {r['value']}" for r in tech["tables"]["angular"]),
                     dxfattribs={"layer": "TEXT", "height": 2.3}).set_placement((tx, ty + 29))
    rel, hist = _meta(part, design, c)
    drawing_style.dxf_frame(doc, msp, frame_w, frame_h, f"{c.no} {c.name}",
                            [("Part", part["part_code"]), ("Tool", design["name"]), ("Material", c.material or "-"),
                             ("Hardness", c.hardness or "-"), ("Qty", str(c.qty)), ("Size", f"{L:.2f} x {W:.2f} x {H:.2f}")] + tech_rows,
                            design.get("part_type", "sheet_metal"), "1:1", drawing_no(part, stage, c.no), rev, extra_notes=c.notes,
                            release=rel, history=hist, general_tolerance=_gen_tol(c, tol))
    return _bytes(doc)


def assembly_dxf(part: dict, design: dict, stage: int, rev: str = "") -> bytes:
    """Front and top views of the whole assembly (every part's edges, 1:1) with the BOM and the title block."""
    from .drawings import edge_polylines
    doc = _dxf_doc()
    msp = doc.modelspace()
    allx, ally, allz = [], [], []
    views = {"FRONT": (0, 2), "TOP": (0, 1)}
    polys = {k: [] for k in views}
    for c in design["components"]:
        lines = edge_polylines(c.shape)
        for (x, y, z) in c.locs:
            for line in lines:
                pts = [(p[0] + x, p[1] + y, p[2] + z) for p in line]
                for k, (a, b) in views.items():
                    polys[k].append([(p[a], p[b]) for p in pts])
                allx += [p[0] for p in pts]
                ally += [p[1] for p in pts]
                allz += [p[2] for p in pts]
    xmin, xmax, ymin, ymax, zmax = min(allx), max(allx), min(ally), max(ally), max(allz)
    dy_top = zmax + 60 - ymin
    for k, ps in polys.items():
        for pl in ps:
            q = [(a - xmin, b + (dy_top if k == "TOP" else 0)) for a, b in pl]
            if len(set(q)) > 1:
                msp.add_lwpolyline(q, dxfattribs={"layer": "PART"})
    msp.add_text("FRONT VIEW", dxfattribs={"layer": "TEXT", "height": 4}).set_placement((0, -12))
    msp.add_text("TOP VIEW", dxfattribs={"layer": "TEXT", "height": 4}).set_placement((0, dy_top + ymin - 12))
    bx = xmax - xmin + 40
    msp.add_text("BILL OF MATERIAL", dxfattribs={"layer": "TEXT", "height": 4}).set_placement((bx, dy_top + ymax))
    for i, r in enumerate(_bom_rows(design)[1:], start=1):
        msp.add_text(f"{r[0]}  {r[1]} x  {r[2]}  {r[3]}  {r[4]}  {r[5]}", dxfattribs={"layer": "TEXT", "height": 2.8}).set_placement(
            (bx, dy_top + ymax - 7 * i))
    w = bx + 260
    h = dy_top + ymax + 20
    from ezdxf.math import Matrix44
    for e in list(msp):
        e.transform(Matrix44.translate(15, 70, 0))
    drawing_style.dxf_frame(doc, msp, w + 30, h + 80, design["name"],
                            [("Part", f"{part['part_code']} {part['name']}"), ("Die set / base", design.get("die_set", "")),
                             ("Shut height", f"{design.get('shut_height', 0):g} mm")],
                            design.get("part_type", "sheet_metal"), "1:1", drawing_no(part, stage, "00"), rev,
                            extra_notes=design.get("notes", []), **dict(zip(("release", "history"), _meta(part, design))))
    return _bytes(doc)


def _bytes(doc) -> bytes:
    s = io.StringIO()
    doc.write(s)
    return s.getvalue().encode()


# ---------- BOM Excel and release pack ----------

def bom_xlsx(part: dict, design: dict, stage: int) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "Tool BOM"
    ws.append([f"{drawing_style.settings()['company']} - tool BOM {part['part_code']} {design['name']}", "", "", "", "",
               "", "", f"{date.today().isoformat()}"])
    ws.append([])
    ws.append(["Item", "Drawing no.", "Qty", "Description", "Material / grade", "Hardness", "Make / standard part", "L", "W", "H",
               "Notes", "Heat treatment", "Surface treatment", "Ra (µm)", "General tolerance"])
    for c in ws[3]:
        c.font = Font(bold=True)
    for c in design["components"]:
        if c.kind == "ref":
            continue
        L, W, H = c.size
        tech = getattr(c, "tech", None) or {}
        ws.append([c.no, drawing_no(part, stage, c.no) if c.kind == "made" else "standard", c.qty, c.name, c.material, c.hardness,
                   c.std or c.make, round(L, 2), round(W, 2), round(H, 2), "; ".join(c.notes), tech.get("heat_treatment", ""),
                   tech.get("surface_treatment", ""), tech.get("finish_ra", ""), f"ISO 2768-{tech['tolerance_class']}" if tech else ""])
    for col, w in zip("ABCDEFGHIJK", (6, 22, 6, 40, 28, 22, 50, 9, 9, 9, 80)):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def release_files(part: dict, design: dict, stage: int, rev: str = "") -> dict:
    """Every file of one tool for release: {path: bytes}."""
    stem = f"{part['part_code']}_T{stage:02d}_{_safe(design['key'])}"
    files = {}
    files[f"3D/{stem}_assembly.step"] = tool_design.step_bytes(tool_design.assembly(design, stem))
    for c in design["components"]:
        if c.kind == "ref":
            continue
        files[f"3D/parts/{c.no}_{_safe(c.name)}.step"] = tool_design.step_bytes(c.shape)
    files[f"2D/PDF/{stem}_drawing_set.pdf"] = design_pdf(part, design, stage, rev)
    files[f"2D/DXF/{stem}_00_assembly.dxf"] = assembly_dxf(part, design, stage, rev)
    for c in design["components"]:
        if c.kind == "made" and c.drawing != "none":
            files[f"2D/DXF/{stem}_{c.no}_{_safe(c.name)}.dxf"] = detail_dxf(part, design, stage, c, rev)
    files[f"{stem}_BOM.xlsx"] = bom_xlsx(part, design, stage)
    return files


def release_zip(part: dict, designs: list[tuple[int, dict]], rev: str = "", extra: dict | None = None) -> bytes:
    """ZIP with the release files of every tool given (one folder per tool), plus README."""
    buf = io.BytesIO()
    readme = [f"{drawing_style.settings()['company']} - tool design release for {part['part_code']} {part['name']}",
              f"Made {date.today().isoformat()}", "",
              "Per tool: 3D/ assembly STEP and one STEP per part; 2D/PDF the drawing set (assembly + every made part);",
              "2D/DXF the assembly and every made part at 1:1 with title block and logo; BOM (Excel).", "",
              "First-cut design made by the rules in Master data (Mould & die database -> Design rules): the tool designer",
              "checks and completes it before release for manufacture.", ""]
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for stage, design in designs:
            folder = f"T{stage:02d}_{_safe(design['name'])}"
            readme.append(f"{folder}: {design['name']}, {len([c for c in design['components'] if c.kind != 'ref'])} items, "
                          f"shut height {design.get('shut_height', 0):g} mm")
            for n in design.get("notes", []):
                readme.append(f"    - {n}")
            for path, data in release_files(part, design, stage, rev).items():
                z.writestr(f"{folder}/{path}", data)
        for path, data in (extra or {}).items():
            z.writestr(path, data)
        z.writestr("README.txt", "\n".join(readme))
    return buf.getvalue()
