"""Replace example values: a checklist of every value still marked EXAMPLE (roadmap stage 4b).

Two kinds of example values exist:

    rates   the approved version of a library rate (per category, item, country) that is still an EXAMPLE
    rows    master-data rows whose Example flag is still on (machines, BOM elements, tool processes, ...)

The checklist counts them by category / table and country, says how many stored costs (latest RM and tool cost
version of every part) use each rate, and lets you replace one:

    a rate   a new APPROVED version with your value and is_example off (old versions are kept, as always)
    a row    its value column changed (where a table has one obvious value) and its Example flag cleared,
             saved through the table's own checks
"""

from __future__ import annotations

import io
from collections import Counter

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import select

from . import costing_setup, country_rates, mould_db, rates, register_setup
from . import projects_db as db
from .table_set import EXAMPLE_FILL, HEADER_FILL

RATE_TITLES = {k: v[0] for k, v in country_rates.GROUPS.items()} | {
    "forex": "Forex (INR per unit)", "machine_price": "Machine prices", "consumables": "Machine consumables per hour",
    "labour_salary": "Labour salary per month", "labour": "Labour on-costs (employer %, supervision, productive time)",
    "power": "Electricity", "space": "Floor space rent", "finance": "Interest", "hours": "Working hours and shifts",
    "overhead": "Factory overhead %", "profit": "Profit %", "tool_commercial": "Toolmaker overhead, profit, packing, warranty %",
}
SETS = {
    "mould": ("Mould & die database", mould_db),
    "setup": ("Countries, machines & formulas", costing_setup),
    "register": ("Part register setup", register_setup),
}
# Tables with one obvious value you can type in on the checklist; other tables open in the Master data editor.
VALUE_COLUMN = {
    ("mould", "sizing_rules"): "value", ("mould", "time_parameters"): "value", ("mould", "mould_bases"): "cost",
    ("mould", "tool_materials"): "density_g_cm3", ("register", "insert_rules"): "value",
    ("register", "insert_materials"): "density_g_cm3",
}


class ExampleError(ValueError):
    pass


def _module(set_name: str):
    if set_name not in SETS:
        raise ExampleError(f"Unknown data set '{set_name}' ({', '.join(SETS)}).")
    return SETS[set_name][1]


def used_by() -> Counter:
    """Rate version id -> number of parts whose latest stored RM or tool cost uses it."""
    per_part: dict[int, set] = {}
    with db.engine().connect() as conn:
        for table, col in ((db.rm_costs, db.parts.c.rm_cost_id), (db.tool_costs, db.parts.c.tool_cost_id)):
            for part_id, used in conn.execute(select(db.parts.c.id, table.c.rates_used).join(db.parts, col == table.c.id)):
                per_part.setdefault(part_id, set()).update(r["id"] for r in used or [])
    count: Counter = Counter()
    for ids in per_part.values():
        count.update(ids)
    return count


def _row_label(t: dict, row: dict) -> str:
    return " / ".join(str(row[k]) for k in t["unique"])


def checklist(as_of: str | None = None) -> dict:
    as_of = as_of or rates.current_month()
    used = used_by()
    groups: dict[str, dict] = {}
    for r in rates.current(as_of):
        a = r["approved"]
        if not a:
            continue
        g = groups.setdefault(r["category"], {"category": r["category"], "title": RATE_TITLES.get(r["category"], r["category"]),
                                              "total": 0, "example": 0, "used_example": 0, "by_country": Counter(),
                                              "items": []})
        g["total"] += 1
        if a["is_example"]:
            g["example"] += 1
            g["by_country"][r["country"] or "global"] += 1
            n = used.get(a["id"], 0)
            g["used_example"] += n > 0
            g["items"].append({"id": a["id"], "item": a["item"], "country": a["country"], "value": a["value"],
                               "unit": a["unit"], "currency": a["currency"], "month": a["month"], "source": a["source"],
                               "used_by": n})
    for g in groups.values():
        g["items"].sort(key=lambda x: (-x["used_by"], x["item"], x["country"]))
        g["by_country"] = dict(sorted(g["by_country"].items()))
    rate_groups = sorted(groups.values(), key=lambda g: (-g["example"], g["title"]))

    tables = []
    for set_name, (set_title, module) in SETS.items():
        data = module.tables()
        for key, t in module.SET.tables.items():
            if not any(c["key"] == "is_example" for c in t["columns"]):
                continue
            rows = data[key]
            value_col = VALUE_COLUMN.get((set_name, key))
            col = next((c for c in t["columns"] if c["key"] == value_col), None)
            items = [{"key": {k: r[k] for k in t["unique"]}, "label": _row_label(t, r),
                      "value": r[value_col] if value_col else None,
                      "notes": r.get("notes") or r.get("description") or ""} for r in rows if r.get("is_example")]
            tables.append({"set": set_name, "set_title": set_title, "table": key, "title": t["title"],
                           "total": len(rows), "example": len(items),
                           "value_column": value_col, "value_label": col["label"] if col else "",
                           "items": items})
    tables.sort(key=lambda x: (-x["example"], x["set_title"], x["title"]))

    rate_total = sum(g["total"] for g in rate_groups)
    rate_left = sum(g["example"] for g in rate_groups)
    row_total = sum(t["total"] for t in tables)
    row_left = sum(t["example"] for t in tables)
    done = (rate_total - rate_left) + (row_total - row_left)
    return {"as_of": as_of, "approver_name": rates.settings()["approver_name"],
            "summary": {"rates_left": rate_left, "rates_total": rate_total, "rows_left": row_left, "rows_total": row_total,
                        "rates_used_in_costs": sum(g["used_example"] for g in rate_groups),
                        "replaced_pct": round(100 * done / (rate_total + row_total), 1) if rate_total + row_total else 100.0},
            "rates": rate_groups, "tables": tables}


def replace_rate(rate_id: int, value=None, source: str = "", comment: str = "") -> dict:
    """A new approved, non-example version of an example rate (the same value when `value` is None: 'confirmed')."""
    old = rates.get(rate_id)
    if old["status"] != "approved" or not old["is_example"]:
        raise ExampleError(f"Rate #{rate_id} is not an approved EXAMPLE value, so there is nothing to replace.")
    try:
        live = rates.lookup(old["category"], old["item"], old["country"], max(old["month"], rates.current_month()))
    except rates.RateMissing:
        live = old
    if live["id"] != rate_id:
        raise ExampleError(f"Rate #{rate_id} has already been replaced by version #{live['id']}.")
    month = max(old["month"], rates.current_month())
    new = rates.create({
        "category": old["category"], "item": old["item"], "country": old["country"],
        "value": old["value"] if value in (None, "") else value, "unit": old["unit"], "currency": old["currency"],
        "month": month, "is_example": False,
        "source": source.strip() or ("Confirmed on the Replace example values checklist" if value in (None, "")
                                     else "Replaced on the Replace example values checklist"),
        "comment": comment.strip() or f"Replaces EXAMPLE version #{rate_id} ({old['value']} {old['unit']}).",
    }, approve=True, origin="example checklist")
    return new


def replace_row(set_name: str, table: str, key: dict, value=None) -> dict:
    """Clear a master-data row's Example flag, optionally with a new value in its value column. Saved with the set's checks."""
    module = _module(set_name)
    t = module.SET.tables.get(table)
    if not t:
        raise ExampleError(f"'{table}' is not a table of {SETS[set_name][0]}.")
    data = module.tables()
    match = [r for r in data[table] if all(str(r[k]) == str(key.get(k)) for k in t["unique"])]
    if not match:
        raise ExampleError(f"{t['title']}: no row {' / '.join(str(key.get(k)) for k in t['unique'])}.")
    row = match[0]
    if value not in (None, ""):
        col = VALUE_COLUMN.get((set_name, table))
        if not col:
            raise ExampleError(f"{t['title']} has several values; edit this row in Master data.")
        row[col] = value
    row["is_example"] = False
    module.write(data)
    return {"table": t["title"], "row": _row_label(t, row)}


def to_excel(as_of: str | None = None) -> bytes:
    c = checklist(as_of)
    wb = Workbook()
    wb.properties.title = "Design to Cost — Replace example values"
    wb.properties.creator = "Design to Cost"

    def sheet(ws, title, head, rows):
        ws.title = title
        ws.append(head)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = HEADER_FILL
        for r in rows:
            ws.append(r)
            for cell in ws[ws.max_row]:
                cell.fill = EXAMPLE_FILL
        for i, h in enumerate(head, start=1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = min(
                max([len(str(h))] + [len(str(r[i - 1])) for r in rows]) + 2, 60)
        ws.freeze_panes = "A2"

    s = c["summary"]
    summary = wb.active
    summary.title = "Summary"
    summary.append(["Replace example values", f"as of {c['as_of']}"])
    summary["A1"].font = Font(bold=True, size=13)
    summary.append(["Example rates left", f"{s['rates_left']} of {s['rates_total']}"])
    summary.append(["Example master-data rows left", f"{s['rows_left']} of {s['rows_total']}"])
    summary.append(["Replaced", f"{s['replaced_pct']} %"])
    summary.append([])
    summary.append(["Category / table", "Example left", "Total"])
    for g in c["rates"]:
        summary.append([f"Rate · {g['title']}", g["example"], g["total"]])
    for t in c["tables"]:
        summary.append([f"{t['set_title']} · {t['title']}", t["example"], t["total"]])
    summary.column_dimensions["A"].width = 60
    summary.column_dimensions["B"].width = 16
    sheet(wb.create_sheet(), "Example rates",
          ["Rate id", "Category", "Item", "Country", "Value", "Unit", "Month", "Used by (parts)", "Source"],
          [[i["id"], g["category"], i["item"], i["country"], i["value"], i["unit"], i["month"], i["used_by"], i["source"]]
           for g in c["rates"] for i in g["items"]])
    sheet(wb.create_sheet(), "Example rows", ["Data set", "Table", "Row", "Value", "Notes"],
          [[t["set_title"], t["title"], i["label"], i["value"] if i["value"] is not None else "", i["notes"]]
           for t in c["tables"] for i in t["items"]])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
