"""A set of editable tables in the costing database.

Each table is edited as a whole (on the Master data page or through Excel):
a save validates every row first and then replaces the tables in one
transaction, so nothing is ever half-saved. Errors name the sheet and the
Excel row number.

Column kinds: text, number, bool, formula, plus any kind given a list of
choices (shown as a drop-down). Extra cross-table checks are passed in as
`check(clean_data)` and raise the set's error class.
"""

from __future__ import annotations

import io
from contextlib import closing
from typing import Callable, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from . import costing_db

# Design to Cost theme: table-head-dark and input-cell-bg
HEADER_FILL = PatternFill("solid", fgColor="0E1A2C")
EXAMPLE_FILL = PatternFill("solid", fgColor="FFF9DB")
SQL_TYPE = {"number": "REAL", "bool": "INTEGER"}


class TableSetError(ValueError):
    pass


def col(key, label, kind="text", choices=None, required=True):
    return {"key": key, "label": label, "kind": kind, "choices": list(choices) if choices else None, "required": required}


class TableSet:
    def __init__(self, title: str, tables: dict, error: type = TableSetError,
                 check: Optional[Callable[[dict], None]] = None, help_sheet: Optional[Callable[[], list]] = None):
        """tables: {table: {"title": sheet title, "columns": [col(...)], "unique": (keys...)}}"""
        self.title = title
        self.tables = tables
        self.error = error
        self.check = check
        self.help_sheet = help_sheet

    # ---------- storage ----------

    def _connect(self):
        con = costing_db.connect()
        for table, t in self.tables.items():
            defs = ", ".join(f"{c['key']} {SQL_TYPE.get(c['kind'], 'TEXT')}" for c in t["columns"])
            con.execute(f"CREATE TABLE IF NOT EXISTS {table} (id INTEGER PRIMARY KEY, {defs}, UNIQUE ({', '.join(t['unique'])}))")
            # Columns added in a newer version of the app are added to an existing database (data kept).
            have = {r["name"] for r in con.execute(f"PRAGMA table_info({table})")}
            for c in t["columns"]:
                if c["key"] not in have:
                    con.execute(f"ALTER TABLE {table} ADD COLUMN {c['key']} {SQL_TYPE.get(c['kind'], 'TEXT')}")
        return con

    def schema(self) -> dict:
        return {"tables": {k: {"title": t["title"], "columns": t["columns"]} for k, t in self.tables.items()}}

    def read(self) -> dict:
        with closing(self._connect()) as con:
            out = {}
            for table, t in self.tables.items():
                names = ", ".join(c["key"] for c in t["columns"])
                rows = con.execute(f"SELECT {names} FROM {table} ORDER BY id").fetchall()
                out[table] = [self._from_sql(dict(r), t["columns"]) for r in rows]
            return out

    @staticmethod
    def _from_sql(row: dict, cols) -> dict:
        for c in cols:
            v, kind = row[c["key"]], c["kind"]
            if kind == "bool":
                row[c["key"]] = bool(v)
            elif kind == "number" and v is not None and float(v).is_integer():
                row[c["key"]] = int(v)
            elif v is None:
                row[c["key"]] = None if kind == "number" else ""
        return row

    # ---------- validation ----------

    def _number(self, value, where, required):
        if isinstance(value, bool):
            raise self.error(f"{where}: '{value}' is not a number.")
        if value is None or str(value).strip() == "":
            if required:
                raise self.error(f"{where}: is empty.")
            return 0
        try:
            n = float(str(value).replace(",", "").strip())
        except ValueError:
            raise self.error(f"{where}: '{value}' is not a number.") from None
        if n < 0:
            raise self.error(f"{where}: must not be negative.")
        return int(n) if n.is_integer() else n

    @staticmethod
    def _bool(value) -> bool:
        if isinstance(value, str):
            return value.strip().lower() in ("1", "true", "yes", "y", "x", "example")
        return bool(value)

    def validate(self, raw: dict) -> dict:
        clean: dict = {}
        for table, t in self.tables.items():
            rows, seen = [], set()
            for i, row in enumerate(raw.get(table, []), start=2):  # row 1 is the Excel header
                where = f"{t['title']} row {i}"
                out = {}
                for c in t["columns"]:
                    v = row.get(c["key"])
                    if c["kind"] == "number":
                        out[c["key"]] = self._number(v, f"{where}, {c['label']}", c.get("required", True))
                    elif c["kind"] == "bool":
                        out[c["key"]] = self._bool(v)
                    else:
                        out[c["key"]] = "" if v is None else str(v).strip()
                        if c["choices"] and out[c["key"]] not in c["choices"]:
                            if not (out[c["key"]] == "" and "" in c["choices"]):
                                raise self.error(f"{where}: {c['key']} '{out[c['key']]}' must be one of "
                                                 f"{', '.join(x for x in c['choices'] if x)}.")
                key = tuple(out[k] for k in t["unique"])
                if not all(str(k) for k in key):
                    raise self.error(f"{where}: {', '.join(t['unique'])} must not be empty.")
                if key in seen:
                    raise self.error(f"{where}: {' / '.join(map(str, key))} is listed twice.")
                seen.add(key)
                rows.append(out)
            clean[table] = rows
        if self.check:
            self.check(clean)
        return clean

    def write(self, data: dict) -> None:
        clean = self.validate(data)
        with closing(self._connect()) as con, con:  # one transaction: all or nothing
            for table, t in self.tables.items():
                names = [c["key"] for c in t["columns"]]
                con.execute(f"DELETE FROM {table}")
                con.executemany(
                    f"INSERT INTO {table} ({', '.join(names)}) VALUES ({', '.join('?' * len(names))})",
                    [[int(r[c]) if isinstance(r[c], bool) else r[c] for c in names] for r in clean[table]],
                )

    # ---------- Excel ----------

    def to_excel(self) -> bytes:
        data = self.read()
        wb = Workbook()
        wb.remove(wb.active)
        wb.properties.title = f"Design to Cost — {self.title}"
        wb.properties.creator = "Design to Cost"
        for table, t in self.tables.items():
            ws = wb.create_sheet(t["title"])
            keys = [c["key"] for c in t["columns"]]
            ws.append(keys)
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = HEADER_FILL
            for r in data[table]:
                ws.append(["yes" if r[k] is True else "no" if r[k] is False else r[k] for k in keys])
                if r.get("is_example"):
                    for cell in ws[ws.max_row]:
                        cell.fill = EXAMPLE_FILL
            for i, k in enumerate(keys, start=1):
                width = max([len(k)] + [len(str(r[k])) for r in data[table]])
                ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 60)
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
        if self.help_sheet:
            ws = wb.create_sheet("Formula help")
            for i, line in enumerate(self.help_sheet(), start=1):
                ws.append(line)
                if i == 1:
                    ws["A1"].font = Font(bold=True, size=13)
                if line and line[0] == "variable":
                    for cell in ws[i]:
                        cell.font = Font(bold=True, color="FFFFFF")
                        cell.fill = HEADER_FILL
            ws.column_dimensions["A"].width = 26
            ws.column_dimensions["B"].width = 90
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def from_excel(self, content: bytes) -> None:
        try:
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:
            raise self.error(f"Could not read the uploaded file as Excel (.xlsx): {exc}") from exc
        data = {}
        for table, t in self.tables.items():
            if t["title"] not in wb.sheetnames:
                raise self.error(f"The Excel file has no '{t['title']}' sheet.")
            rows = list(wb[t["title"]].iter_rows(values_only=True))
            header = [str(h).strip() if h is not None else "" for h in rows[0]] if rows else []
            keys = [c["key"] for c in t["columns"]]
            missing = [k for k in keys if k not in header]
            if missing:
                raise self.error(f"Sheet '{t['title']}' is missing column(s): {', '.join(missing)}.")
            idx = {k: header.index(k) for k in keys}
            data[table] = [
                {k: row[idx[k]] if idx[k] < len(row) else None for k in keys}
                for row in rows[1:]
                if any(v not in (None, "") for v in row)
            ]
        wb.close()
        self.write(data)
