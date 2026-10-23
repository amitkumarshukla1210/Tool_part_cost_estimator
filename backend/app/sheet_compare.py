"""Progressive die or stage (multi-stage) dies? Side-by-side comparison of a sheet-metal part (roadmap stage 23).

For each die type a full production tool design is generated (the same one as the Tool design page: 3D, drawings, BOM, bought-out parts,
purchase-order pack) and these are compared:
    press        tonnage of every press, number of presses, die-set size
    raw material blank / strip layout: pitch, strip width, utilisation, gross weight per part, scrap
    process      strokes per minute, machine hour rate of the press(es) + operators × labour rate, per part
    tool         steel + machining + bought-out + assembly of the tool design, plus the toolroom mark-up
    per part     RM + process + rejection + tool amortisation over the programme
Rates come from the rate library and Machines; strokes per minute, operators, scrap credit and mark-up are Parameters (EXAMPLE until replaced).
The result says which die type is cheaper per part and from which volume. Indicative: your tool shop and press shop confirm.
"""

from __future__ import annotations

import io
import zipfile

from . import master_data, rate_engine, rates, register, register_setup, rm_cost, tool_components, tool_design_store as store

MODES = ("progressive", "multi_stage")
LABEL = {"progressive": "Progressive die", "multi_stage": "Stage dies (multi-stage)"}
EXAMPLE = "EXAMPLE - replace with your press shop / tool shop data"
PARAMS = {"spm_progressive": (60, "strokes/min", "Strokes per minute of the progressive die"),
          "spm_stage": (25, "parts/min", "Parts per minute of one stage die (feeding by hand or robot)"),
          "operators_stage_press": (1.0, "persons", "Operators per press in a stage line (progressive uses operators_per_machine)"),
          "scrap_credit_pct": (40, "%", "Scrap sold at this % of the material rate"),
          "tool_markup_pct": (25, "%", "Toolroom design, try-out, overhead and margin on the tool made-parts cost")}


class CompareError(ValueError):
    pass


def seed() -> int:
    md = master_data.tables()
    have = {(r["part_type"], r["parameter"]) for r in md["Parameters"]}
    cols = list(md["Parameters"][0].keys())
    new = []
    for k, (v, unit, desc) in PARAMS.items():
        if ("sheet_metal", k) not in have:
            new.append({c: "" for c in cols} | {"part_type": "sheet_metal", "parameter": k, "value": v, "unit": unit, "description": desc, "notes": EXAMPLE})
    if new:
        md["Parameters"] = md["Parameters"] + new
        master_data.write(md)
    return len(new)


def _latest(part_id: int, mode: str) -> dict | None:
    for d in store.list_designs(part_id):
        if d["status"] != "superseded" and (d["summary"] or {}).get("mode") == mode and d["kind"] == "sheet_metal":
            return d
    return None


def generate(part_id: int) -> dict:
    """Fresh production designs of both die types, then the comparison."""
    for m in MODES:
        store.generate(part_id, f"Comparison: {LABEL[m]}", m)
    return compare(part_id)


def _rate(category: str, item: str, country: str, month: str) -> float:
    try:
        return float(rates.lookup(category, item, country, month)["value"])
    except Exception:
        raise CompareError(f"No approved {category} rate for '{item}' in {country}: add it on the Rates page.") from None


def compare(part_id: int) -> dict:
    d = register.part_detail(part_id, with_rm=False)
    if register_setup.processes().get(d["process"], {}).get("part_type") != "sheet_metal":
        raise CompareError("The comparison is for sheet-metal parts.")
    from . import design_pack
    P = design_pack.sheet_prepare(part_id)           # features, strip layout and country of the part
    country, month = P["country"], d["costing_month"] or rates.current_month()
    par = master_data.parameters("sheet_metal")
    missing = [k for k in PARAMS if k not in par]
    if missing:
        raise CompareError(f"Master data → Parameters has no sheet_metal value for {', '.join(missing)}: run the seed (python -m backend.app.capacity_db).")
    density, _ = register._density(d["material"])
    t = P["t"]
    net_g = P["res"]["features"]["weight_g"]
    volume, _ = register._annual_volume(d)
    volume = volume or par["annual_volume"]
    years = d["programme_life_years"] or par["programme_years"]
    rate_kg = _rate("material", d["material"], country, month)
    labour = rate_engine.labour_rate(par["operator_grade"], country, month)["value"]
    oee = par["oee_pct"] / 100
    out = {"part_id": part_id, "part_code": d["part_code"], "material": d["material"], "thickness_mm": t, "country": country, "month": month,
           "currency": "INR", "annual_volume": volume, "programme_years": years, "net_weight_g": round(net_g, 1), "modes": {}, "example": True}
    for mode in MODES:
        dsg = _latest(part_id, mode)
        if not dsg:
            raise CompareError(f"No {LABEL[mode].lower()} design yet: press 'Generate comparison'.")
        full = store.get_design(dsg["id"])
        pre = P if P["res"]["mode"] == mode else design_pack.sheet_prepare(part_id, mode)
        design = pre["design"]
        presses = [s["press"] for s in design["stages"]]
        if not all(p.get("fits") for p in presses):
            raise CompareError(f"{LABEL[mode]}: " + "; ".join(p["reason"] for p in presses if not p.get("fits")))
        # raw material
        if mode == "progressive":
            v = pre["sl"]["best"]["values"]
            gross_g = v["pitch_mm"] * v["strip_width_mm"] * t * density / 1000
            rm = {"layout": f"strip {v['strip_width_mm']:.0f} mm wide, pitch {v['pitch_mm']:.1f} mm, feed {pre['sl']['best']['feed']}",
                  "pitch_mm": round(v["pitch_mm"], 2), "strip_width_mm": round(v["strip_width_mm"], 1), "utilisation_pct": round(v["utilisation_pct"], 1)}
        else:
            fp = pre["fp"]
            fit = rm_cost.grid_fit(par["sheet_length_mm"] - 2 * par["edge_mm"], par["sheet_width_mm"] - 2 * par["edge_mm"], fp["length_mm"], fp["width_mm"], par["gap_mm"])
            if not fit["per_sheet"]:
                raise CompareError("The blank does not fit on the sheet (Parameters → sheet_length_mm / sheet_width_mm).")
            gross_g = par["sheet_length_mm"] * par["sheet_width_mm"] * t * density / 1000 / fit["per_sheet"]
            rm = {"layout": f"blank {fp['length_mm']:.0f} × {fp['width_mm']:.0f} mm, {fit['along']} × {fit['across']} = {fit['per_sheet']} per sheet "
                            f"{par['sheet_length_mm']:g} × {par['sheet_width_mm']:g} mm", "blanks_per_sheet": fit["per_sheet"],
                  "utilisation_pct": round(net_g / gross_g * 100, 1)}
        scrap_g = max(gross_g - net_g, 0.0)
        rm_cost_pc = gross_g / 1000 * rate_kg - scrap_g / 1000 * rate_kg * par["scrap_credit_pct"] / 100
        rm |= {"gross_weight_g": round(gross_g, 1), "scrap_g": round(scrap_g, 1), "rm_rate_per_kg": rate_kg, "rm_cost_per_part": round(rm_cost_pc, 3)}
        # process: one press for the progressive die, one press per stage otherwise
        lines = []
        if mode == "progressive":
            p = presses[0]
            mhr = rate_engine.mhr(p["machine_id"], country, month)["value"]
            per_h = par["spm_progressive"] * 60 * oee
            lines.append({"press": p["name"], "tonnage_t": p["tonnage_t"], "mhr": round(mhr, 2), "operators": par["operators_per_machine"],
                          "parts_per_h": round(per_h), "cost_per_part": (mhr + par["operators_per_machine"] * labour) / per_h})
        else:
            per_h = par["spm_stage"] * 60 * oee
            for s in design["stages"]:
                p = s["press"]
                mhr = rate_engine.mhr(p["machine_id"], country, month)["value"]
                lines.append({"stage": s["n"], "name": s["name"], "press": p["name"], "tonnage_t": p["tonnage_t"], "mhr": round(mhr, 2),
                              "operators": par["operators_stage_press"], "parts_per_h": round(per_h),
                              "cost_per_part": (mhr + par["operators_stage_press"] * labour) / per_h})
        process_pc = sum(l["cost_per_part"] for l in lines)
        # tool
        from . import tool_po
        f = tool_po._facts(full)
        tool_base = f["total_inr"]
        tool_total = tool_base * (1 + par["tool_markup_pct"] / 100)
        made = full["bill"]["totals"]
        base_pc = rm_cost_pc + process_pc
        reject = base_pc * par["rejection_pct"] / 100
        amort = tool_total / max(volume * years, 1)
        out["modes"][mode] = {
            "label": LABEL[mode], "design_id": full["id"], "design_version": full["version"],
            "concept": {"stages": [{"n": s["n"], "name": s["name"], "force_t": round(s["force_t"], 1), "press": s["press"]["name"],
                                    "press_t": s["press"]["tonnage_t"], "die_set_mm": [round(s["values"]["die_set_l_mm"]), round(s["values"]["die_set_w_mm"])],
                                    "shut_height_mm": round(s["values"]["shut_height_mm"])} for s in design["stages"]],
                        "total_force_t": round(design["total_force_t"], 1), "presses": len(lines), "largest_press_t": max(p["tonnage_t"] for p in presses),
                        "tool_weight_kg": design["tool_weight_kg"]},
            "rm": rm, "process": {"lines": [{k: (round(v, 4) if isinstance(v, float) else v) for k, v in l.items()} for l in lines], "cost_per_part": round(process_pc, 3)},
            "tool": {"made_parts": made["made_parts"], "bought_out_items": len(full["bill"]["standard"]), "steel_kg": made["steel_kg"],
                     "material_inr": round(made["material_cost"]), "machining_inr": round(made["machining_cost"]), "bought_out_inr": round(made["standard_cost_inr"]),
                     "assembly_inr": f["assembly"], "assembly_hours": made["assembly_hours"], "base_inr": round(tool_base),
                     "markup_pct": par["tool_markup_pct"], "total_inr": round(tool_total)},
            "per_part": {"rm": round(rm_cost_pc, 3), "process": round(process_pc, 3), "rejection": round(reject, 3), "tool": round(amort, 3),
                         "total": round(base_pc + reject + amort, 3)},
            "checks": full["summary"].get("checks", {}).get("errors")}
    a, b = out["modes"]["progressive"], out["modes"]["multi_stage"]
    out["verdict"] = _verdict(a, b, volume, years)
    out["example"] = True            # strokes per minute, operators, mark-up and the rates are EXAMPLE until replaced
    return out


def _verdict(a: dict, b: dict, volume: float, years: float) -> dict:
    """Cheaper die type per part and the break-even volume (parts over the programme) where the dearer tool pays back."""
    var_a, var_b = a["per_part"]["rm"] + a["per_part"]["process"] + a["per_part"]["rejection"], b["per_part"]["rm"] + b["per_part"]["process"] + b["per_part"]["rejection"]
    cheaper = "progressive" if a["per_part"]["total"] <= b["per_part"]["total"] else "multi_stage"
    be, note = None, ""
    d_tool, d_var = a["tool"]["total_inr"] - b["tool"]["total_inr"], var_b - var_a
    if d_var > 0 and d_tool > 0:
        be = round(d_tool / d_var)
        note = f"The progressive tool costs INR {d_tool:,.0f} more but saves INR {d_var:.2f} per part: it pays back after {be:,} parts ({be / max(years, 1):,.0f} per year over {years:g} years)."
    elif d_var > 0 and d_tool <= 0:
        note = "The progressive die is cheaper in tool and per part."
    elif d_var <= 0 and d_tool < 0:
        note = "Stage dies run cheaper per part than the progressive die here, and the progressive tool is cheaper too."
    else:
        note = "Stage dies need less tool investment; the progressive die gives no saving per part at these inputs."
    return {"cheaper": cheaper, "label": LABEL[cheaper], "saving_per_part": round(abs(a["per_part"]["total"] - b["per_part"]["total"]), 3), "break_even_parts": be, "note": note,
            "programme_parts": round(volume * years)}


# ------------------------------------------------------------ documents

def workbook(c: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "Comparison"
    A, B = c["modes"]["progressive"], c["modes"]["multi_stage"]
    rows = [["Sheet-metal die comparison", c["part_code"]], ["Material", c["material"]], ["Thickness (mm)", c["thickness_mm"]], ["Country / month", f"{c['country']} {c['month']}"],
            ["Annual volume", c["annual_volume"]], ["Programme years", c["programme_years"]], ["", ""], ["", A["label"], B["label"]],
            ["Largest press (t)", A["concept"]["largest_press_t"], B["concept"]["largest_press_t"]],
            ["Presses", A["concept"]["presses"], B["concept"]["presses"]],
            ["Total force (t)", A["concept"]["total_force_t"], B["concept"]["total_force_t"]],
            ["Tool weight (kg)", A["concept"]["tool_weight_kg"], B["concept"]["tool_weight_kg"]],
            ["Raw material layout", A["rm"]["layout"], B["rm"]["layout"]],
            ["Net weight (g)", c["net_weight_g"], c["net_weight_g"]], ["Gross weight per part (g)", A["rm"]["gross_weight_g"], B["rm"]["gross_weight_g"]],
            ["Scrap per part (g)", A["rm"]["scrap_g"], B["rm"]["scrap_g"]], ["Material utilisation (%)", A["rm"]["utilisation_pct"], B["rm"]["utilisation_pct"]],
            ["RM cost per part (INR)", A["rm"]["rm_cost_per_part"], B["rm"]["rm_cost_per_part"]],
            ["Process cost per part (INR)", A["process"]["cost_per_part"], B["process"]["cost_per_part"]],
            ["Rejection per part (INR)", A["per_part"]["rejection"], B["per_part"]["rejection"]],
            ["Tool: made parts", A["tool"]["made_parts"], B["tool"]["made_parts"]], ["Tool: bought-out items", A["tool"]["bought_out_items"], B["tool"]["bought_out_items"]],
            ["Tool: steel (INR)", A["tool"]["material_inr"], B["tool"]["material_inr"]], ["Tool: machining (INR)", A["tool"]["machining_inr"], B["tool"]["machining_inr"]],
            ["Tool: bought-out (INR)", A["tool"]["bought_out_inr"], B["tool"]["bought_out_inr"]], ["Tool: assembly (INR)", A["tool"]["assembly_inr"], B["tool"]["assembly_inr"]],
            ["Tool total incl. mark-up (INR)", A["tool"]["total_inr"], B["tool"]["total_inr"]],
            ["Tool amortisation per part (INR)", A["per_part"]["tool"], B["per_part"]["tool"]], ["TOTAL per part (INR)", A["per_part"]["total"], B["per_part"]["total"]],
            ["", ""], ["Cheaper per part", c["verdict"]["label"]], ["Saving per part (INR)", c["verdict"]["saving_per_part"]], ["Break-even (parts)", c["verdict"]["break_even_parts"] or "-"],
            ["Note", c["verdict"]["note"]], ["", ""],
            ["EXAMPLE values", "Strokes per minute, operators, scrap credit, mark-up, rates and machine data are EXAMPLE rows until replaced in master data."]]
    for r in rows:
        ws.append(r)
    for r in (1, 8, 29):
        for cell in ws[r]:
            cell.font = Font(bold=True)
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 48
    ws.column_dimensions["C"].width = 48
    for mode, name in (("progressive", "Progressive stations"), ("multi_stage", "Stage dies")):
        s = wb.create_sheet(name)
        s.append(["Station / stage", "Operation", "Force (t)", "Press", "Press (t)", "Die set (mm)", "Shut height (mm)"])
        for x in c["modes"][mode]["concept"]["stages"]:
            s.append([x["n"], x["name"], x["force_t"], x["press"], x["press_t"], f"{x['die_set_mm'][0]} × {x['die_set_mm'][1]}", x["shut_height_mm"]])
        s.append([])
        s.append(["Process line", "", "", "Press", "MHR (INR/h)", "Operators", "Parts/h", "INR per part"])
        for l in c["modes"][mode]["process"]["lines"]:
            s.append([l.get("stage", ""), l.get("name", ""), "", l["press"], l["mhr"], l["operators"], l["parts_per_h"], l["cost_per_part"]])
        for col, w in zip("ABCDEFGH", (14, 24, 12, 24, 14, 16, 16, 14)):
            s.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def pdf(c: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    A, B = c["modes"]["progressive"], c["modes"]["multi_stage"]
    st = getSampleStyleSheet()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=24, rightMargin=24, topMargin=24, bottomMargin=24)
    rows = [["", A["label"], B["label"]],
            ["Largest press (t) / presses", f"{A['concept']['largest_press_t']:g} t / {A['concept']['presses']}", f"{B['concept']['largest_press_t']:g} t / {B['concept']['presses']}"],
            ["Total force (t)", A["concept"]["total_force_t"], B["concept"]["total_force_t"]],
            ["Raw material layout", Paragraph(A["rm"]["layout"], st["BodyText"]), Paragraph(B["rm"]["layout"], st["BodyText"])],
            ["Gross weight per part (g) / utilisation", f"{A['rm']['gross_weight_g']} g / {A['rm']['utilisation_pct']} %", f"{B['rm']['gross_weight_g']} g / {B['rm']['utilisation_pct']} %"],
            ["RM cost per part (INR)", A["rm"]["rm_cost_per_part"], B["rm"]["rm_cost_per_part"]],
            ["Process cost per part (INR)", A["process"]["cost_per_part"], B["process"]["cost_per_part"]],
            ["Tool cost (INR)", f"{A['tool']['total_inr']:,}", f"{B['tool']['total_inr']:,}"],
            ["  steel / machining / bought-out / assembly (INR)", f"{A['tool']['material_inr']:,} / {A['tool']['machining_inr']:,} / {A['tool']['bought_out_inr']:,} / {A['tool']['assembly_inr']:,}",
             f"{B['tool']['material_inr']:,} / {B['tool']['machining_inr']:,} / {B['tool']['bought_out_inr']:,} / {B['tool']['assembly_inr']:,}"],
            ["Tool amortisation per part (INR)", A["per_part"]["tool"], B["per_part"]["tool"]],
            ["TOTAL per part (INR)", A["per_part"]["total"], B["per_part"]["total"]]]
    t = Table(rows, colWidths=[250, 280, 280])
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.grey), ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey), ("FONTSIZE", (0, 0), (-1, -1), 8),
                           ("VALIGN", (0, 0), (-1, -1), "TOP"), ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold")]))
    story = [Paragraph(f"Progressive die or stage dies: {c['part_code']} · {c['material']} · {c['thickness_mm']:g} mm", st["Title"]),
             Paragraph(f"{c['country']} {c['month']} · {c['annual_volume']:,.0f} parts per year × {c['programme_years']:g} years", st["Normal"]), Spacer(1, 8), t, Spacer(1, 8),
             Paragraph(f"<b>{c['verdict']['label']}</b> is cheaper per part by INR {c['verdict']['saving_per_part']}. {c['verdict']['note']}", st["BodyText"]),
             Paragraph("EXAMPLE values: strokes per minute, operators, scrap credit, mark-up, rates and machine data stay EXAMPLE rows until you replace them in master data. "
                       "Indicative: your press shop and tool shop confirm.", st["Italic"])]
    for mode in ("progressive", "multi_stage"):
        m = c["modes"][mode]
        story += [Spacer(1, 10), Paragraph(m["label"], st["Heading3"])]
        r2 = [["Stage", "Operation", "Force (t)", "Press", "Die set (mm)", "Shut height (mm)"]] + [[x["n"], x["name"], x["force_t"], f"{x['press']} ({x['press_t']:g} t)",
              f"{x['die_set_mm'][0]} × {x['die_set_mm'][1]}", x["shut_height_mm"]] for x in m["concept"]["stages"]]
        t2 = Table(r2, colWidths=[40, 150, 70, 220, 110, 100])
        t2.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.grey), ("FONTSIZE", (0, 0), (-1, -1), 8), ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey)]))
        story.append(t2)
    doc.build(story)
    return buf.getvalue()


def package(part_id: int) -> tuple[bytes, str]:
    """Full package: comparison (xlsx, pdf), and for each die type the strip layout, supplier pack (STEP, drawings, DXF, BOM, POs, assembly sheet)."""
    from . import sheet_layout
    c = compare(part_id)
    code = store._safe(c["part_code"])
    d, rev, fp, sl, res = register._layout(part_id)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr(f"{code}_progressive_vs_stage_comparison.xlsx", workbook(c))
        z.writestr(f"{code}_progressive_vs_stage_comparison.pdf", pdf(c))
        z.writestr("layout/blank.dxf", sheet_layout.blank_dxf(fp, d["part_code"]))
        z.writestr("layout/strip.dxf", sheet_layout.strip_dxf(fp, sl, d["part_code"]))
        z.writestr("layout/strip_layout_and_tool_pack.pdf", sheet_layout.pack_pdf(fp, sl, res, d, rev["rev"]))
        for mode in MODES:
            data, _ = store.supplier_package(c["modes"][mode]["design_id"])
            with zipfile.ZipFile(io.BytesIO(data)) as inner:
                for n in inner.namelist():
                    z.writestr(f"{mode}/{n}", inner.read(n))
        z.writestr("README.txt", (f"Sheet-metal tool package: {c['part_code']}\n\n{c['verdict']['label']} is cheaper per part. {c['verdict']['note']}\n\n"
                                  "progressive/  production design of the progressive die: STEP and drawing of every made part, DXF, assembly sheet with BOM, bought-out list, POs\n"
                                  "multi_stage/  the same for the stage dies (every stage numbered S1-, S2- ...)\n"
                                  "layout/       blank and strip layout (DXF) and the strip / tool pack PDF\n"
                                  "The comparison workbook and PDF hold press tonnage, gross material, process cost, tool cost and the cost per part.\n"
                                  "EXAMPLE values stay flagged until replaced in master data. DRAFT designs are not for ordering.\n"))
    return buf.getvalue(), f"{code}_sheet_tool_package.zip"
