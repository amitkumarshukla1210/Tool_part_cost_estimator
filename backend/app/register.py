"""Part register: projects, variants, parts, revisions, BOM, history, search, dashboard.

Rules
- A part code is unique. A blank code is generated from the pattern in master
  data (Part register setup → Register settings → part_code_pattern).
- A new file for an existing part code is a NEW revision; old revisions, files
  and history are kept.
- Every change writes a history record: what changed (old → new), who, when.
- Process suggestion: material first (its family in Master data → Materials), then
  geometry rules, then moulded-in inserts; the reason is returned with every suggestion.
- Moulded-in inserts are parts of their own (usually bought-out) linked to the body
  part with a quantity. They are not BOM children: the body is one moulded part.
"""

from __future__ import annotations

import hashlib
import io
import json
import shutil
import uuid
import zipfile
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import and_, func, or_, select

from . import (cad_import, capacity, concept, costing_setup, file_store, formula, master_data, mould_db, rates, register_setup,
               rm_cost, tool_cost)
from . import projects_db as db
from .table_set import EXAMPLE_FILL, HEADER_FILL
from .trace import EXAMPLE_WARNING


class RegisterError(ValueError):
    pass


def user() -> str:
    from . import auth
    return auth.username() or rates.settings().get("approver_name") or "local user"


def _log(conn, entity, entity_id, action, detail=None):
    conn.execute(db.history.insert().values(entity=entity, entity_id=entity_id, action=action,
                                            detail=detail or {}, user=user()))


def _plain(row) -> dict:
    d = dict(row)
    for k, v in d.items():
        if isinstance(v, datetime):
            d[k] = v.isoformat(timespec="seconds")
    return d


def _num(v, where, kind=float, required=False):
    if v is None or str(v).strip() == "":
        if required:
            raise RegisterError(f"{where}: is empty.")
        return None
    try:
        n = kind(float(str(v).replace(",", "").strip()))
    except ValueError:
        raise RegisterError(f"{where}: '{v}' is not a number.") from None
    if n < 0:
        raise RegisterError(f"{where}: must not be negative.")
    return n


# ---------- projects & variants ----------

PROJECT_FIELDS = ("code", "name", "customer", "description", "capex_budget_inr", "sop_date")
VARIANT_FIELDS = ("code", "name", "target_bom_cost_inr", "annual_volume", "notes")


def _project_values(data, where="Project"):
    out = {}
    for k in PROJECT_FIELDS:
        if k in data:
            out[k] = _num(data[k], f"{where}: CAPEX budget") if k == "capex_budget_inr" else str(data[k] or "").strip()
    if out.get("sop_date"):
        try:
            datetime.strptime(out["sop_date"], "%Y-%m-%d")
        except ValueError:
            raise RegisterError(f"{where}: SOP date '{out['sop_date']}' must look like 2027-04-01.") from None
    if "code" in out:
        out["code"] = out["code"].upper()
        if not out["code"]:
            raise RegisterError(f"{where}: code must not be empty.")
    if "name" in out and not out["name"]:
        raise RegisterError(f"{where}: name must not be empty.")
    if out.get("price_month") and not (len(out["price_month"]) == 7 and out["price_month"][4] == "-"
                                       and out["price_month"].replace("-", "").isdigit()):
        raise RegisterError(f"{where}: price month '{out['price_month']}' must look like 2026-10.")
    return out


def _dimensions(v, where):
    """{"shape": code, <dimension>: mm, ...} checked against Part register setup → Insert shapes."""
    if v in (None, "", {}):
        return None
    if isinstance(v, str):
        try:
            v = json.loads(v)
        except ValueError:
            raise RegisterError(f"{where}: dimensions are not readable.") from None
    shapes = register_setup.insert_shapes()
    shape = str(v.get("shape") or "").strip()
    if not shape:
        return None
    if shape not in shapes:
        raise RegisterError(f"{where}: shape '{shape}' is not in Part register setup → Insert shapes "
                            f"({', '.join(shapes)}).")
    out = {"shape": shape}
    for d in shapes[shape]["dims"]:
        n = _num(v.get(d), f"{where}: {d}", required=True)
        if n <= 0:
            raise RegisterError(f"{where}: {d} must be more than 0.")
        out[d] = n
    return out


def create_project(data: dict) -> dict:
    vals = _project_values(data)
    for k in ("code", "name"):
        if k not in vals:
            raise RegisterError(f"Project: {k} must not be empty.")
    with db.engine().begin() as conn:
        if conn.execute(select(db.projects.c.id).where(db.projects.c.code == vals["code"])).first():
            raise RegisterError(f"Project code '{vals['code']}' already exists.")
        pid = conn.execute(db.projects.insert().values(**vals, created_by=user())).inserted_primary_key[0]
        _log(conn, "project", pid, "created", {"code": vals["code"]})
    return get_project(pid)


def update_project(pid: int, data: dict) -> dict:
    vals = _project_values(data)
    with db.engine().begin() as conn:
        old = conn.execute(select(db.projects).where(db.projects.c.id == pid)).mappings().first()
        if not old:
            raise RegisterError(f"Project {pid} does not exist.")
        changes = {k: {"old": old[k], "new": v} for k, v in vals.items() if old[k] != v}
        if changes:
            conn.execute(db.projects.update().where(db.projects.c.id == pid).values(**vals))
            _log(conn, "project", pid, "updated", changes)
    return get_project(pid)


def delete_project(pid: int, confirm_code: str = "", backup: bool = True) -> dict:
    """Delete a project with its variants, BOM, parts, revisions, cost versions and CAPEX lines.

    Safety: you type the project code to confirm; a part used by ANOTHER project (in its BOM or as a moulded-in
    insert) stops the delete; a backup ZIP is made first (dtc-before-delete-<code>-<time>.zip). CAD files stay in
    the storage folder (they are stored once by checksum and may be shared), and the history keeps a 'deleted' record.
    """
    with db.engine().connect() as conn:
        p = conn.execute(select(db.projects).where(db.projects.c.id == pid)).mappings().first()
    if not p:
        raise RegisterError(f"Project {pid} does not exist.")
    if str(confirm_code or "").strip().upper() != p["code"]:
        raise RegisterError(f"To delete project {p['code']}, type its code {p['code']} to confirm.")
    with db.engine().connect() as conn:
        part_ids = [i for (i,) in conn.execute(select(db.parts.c.id).where(db.parts.c.project_id == pid))]
        variant_ids = [i for (i,) in conn.execute(select(db.variants.c.id).where(db.variants.c.project_id == pid))]
        used = []
        if part_ids:
            bl = db.bom_lines
            for code, vcode, pcode in conn.execute(
                    select(db.parts.c.part_code, db.variants.c.code, db.projects.c.code)
                    .select_from(bl.join(db.variants, bl.c.variant_id == db.variants.c.id)
                                 .join(db.projects, db.variants.c.project_id == db.projects.c.id)
                                 .join(db.parts, bl.c.child_part_id == db.parts.c.id))
                    .where(bl.c.child_part_id.in_(part_ids) & ~bl.c.variant_id.in_(variant_ids or [-1]))):
                used.append(f"{code} in the BOM of {pcode} / {vcode}")
            body = db.parts.alias("body")
            for code, bcode in conn.execute(
                    select(db.parts.c.part_code, body.c.part_code)
                    .select_from(db.part_inserts.join(db.parts, db.part_inserts.c.insert_part_id == db.parts.c.id)
                                 .join(body, db.part_inserts.c.part_id == body.c.id))
                    .where(db.part_inserts.c.insert_part_id.in_(part_ids) & ~db.part_inserts.c.part_id.in_(part_ids))):
                used.append(f"{code} is moulded into {bcode}")
    if used:
        raise RegisterError(f"Project {p['code']} cannot be deleted: its parts are used elsewhere ("
                            + "; ".join(sorted(set(used))[:10]) + "). Remove those links first.")
    backup_file = ""
    if backup:
        from . import backup_tools
        backup_file = backup_tools.make_backup(label=f"dtc-before-delete-{p['code']}").name
    with db.engine().begin() as conn:
        if part_ids:
            ids = part_ids
            conn.execute(db.bom_lines.delete().where(db.bom_lines.c.child_part_id.in_(ids) | db.bom_lines.c.parent_part_id.in_(ids)))
            conn.execute(db.part_inserts.delete().where(db.part_inserts.c.part_id.in_(ids) | db.part_inserts.c.insert_part_id.in_(ids)))
            for t in (db.rm_costs, db.tool_costs, db.piece_costs):
                conn.execute(t.delete().where(t.c.part_id.in_(ids)))
            conn.execute(db.revisions.delete().where(db.revisions.c.part_id.in_(ids)))
            conn.execute(db.parts.delete().where(db.parts.c.id.in_(ids)))
        if variant_ids:
            conn.execute(db.bom_lines.delete().where(db.bom_lines.c.variant_id.in_(variant_ids)))
            conn.execute(db.variants.delete().where(db.variants.c.id.in_(variant_ids)))
        conn.execute(db.capex_items.delete().where(db.capex_items.c.project_id == pid))
        conn.execute(db.bom_extracts.update().where(db.bom_extracts.c.project_id == pid).values(project_id=None))
        conn.execute(db.projects.delete().where(db.projects.c.id == pid))
        _log(conn, "project", pid, "deleted", {"code": p["code"], "name": p["name"], "parts": len(part_ids),
                                               "variants": len(variant_ids), "backup": backup_file})
    return {"deleted": p["code"], "parts": len(part_ids), "variants": len(variant_ids), "backup": backup_file}


def remove_from_bom(part_id: int) -> dict:
    """Take the part out of the BOM of every variant of its project (the part stays in the register).
    BOM lines under it (its children, if it is an assembly) are removed too; the child parts stay."""
    with db.engine().begin() as conn:
        p = conn.execute(select(db.parts).where(db.parts.c.id == part_id)).mappings().first()
        if not p:
            raise RegisterError(f"Part {part_id} does not exist.")
        bl = db.bom_lines
        rows = conn.execute(select(bl.c.id, db.variants.c.code).join(db.variants, bl.c.variant_id == db.variants.c.id)
                            .where((bl.c.child_part_id == part_id) | (bl.c.parent_part_id == part_id))).all()
        if not rows:
            raise RegisterError(f"{p['part_code']} is not in the BOM of any variant.")
        conn.execute(bl.delete().where(bl.c.id.in_([r[0] for r in rows])))
        variants = sorted({r[1] for r in rows})
        _log(conn, "part", part_id, "removed_from_bom", {"variants": ", ".join(variants), "lines": len(rows)})
    return {"part_code": p["part_code"], "variants": variants, "lines": len(rows),
            "message": f"{p['part_code']} removed from the BOM of {', '.join(variants)} ({len(rows)} line(s)). The part stays in the register."}


def delete_part(part_id: int, confirm_code: str = "", backup: bool = True) -> dict:
    """Delete a part from the register: its revisions, RM / tool / part cost versions, suppliers, BOM lines and its
    tool CAPEX line. You type the part code to confirm; refused while it is moulded into another part; a backup
    ZIP is made first; CAD files stay in storage (shared by checksum); the history keeps a 'deleted' record."""
    with db.engine().connect() as conn:
        p = conn.execute(select(db.parts).where(db.parts.c.id == part_id)).mappings().first()
        if not p:
            raise RegisterError(f"Part {part_id} does not exist.")
        if str(confirm_code or "").strip().upper() != p["part_code"].upper():
            raise RegisterError(f"To delete part {p['part_code']}, type its code {p['part_code']} to confirm.")
        body = db.parts.alias("body")
        used = [c for (c,) in conn.execute(select(body.c.part_code).select_from(
            db.part_inserts.join(body, db.part_inserts.c.part_id == body.c.id))
            .where((db.part_inserts.c.insert_part_id == part_id) & (db.part_inserts.c.part_id != part_id)))]
    if used:
        raise RegisterError(f"{p['part_code']} is moulded into {', '.join(used)}. Unlink it there first (Moulded-in inserts).")
    backup_file = ""
    if backup:
        from . import backup_tools
        backup_file = backup_tools.make_backup(label=f"dtc-before-delete-{p['part_code']}").name
    with db.engine().begin() as conn:
        bl = db.bom_lines
        n_bom = conn.execute(bl.delete().where((bl.c.child_part_id == part_id) | (bl.c.parent_part_id == part_id))).rowcount
        conn.execute(db.part_inserts.delete().where((db.part_inserts.c.part_id == part_id) | (db.part_inserts.c.insert_part_id == part_id)))
        conn.execute(db.part_suppliers.delete().where(db.part_suppliers.c.part_id == part_id))
        for t in (db.rm_costs, db.tool_costs, db.piece_costs):
            conn.execute(t.delete().where(t.c.part_id == part_id))
        conn.execute(db.capex_items.delete().where((db.capex_items.c.project_id == p["project_id"])
                                                   & (db.capex_items.c.source_key == f"tool:{part_id}")))
        conn.execute(db.revisions.delete().where(db.revisions.c.part_id == part_id))
        conn.execute(db.parts.delete().where(db.parts.c.id == part_id))
        _log(conn, "part", part_id, "deleted", {"code": p["part_code"], "name": p["name"], "bom_lines": n_bom, "backup": backup_file})
    return {"deleted": p["part_code"], "bom_lines": n_bom, "backup": backup_file,
            "message": f"Part {p['part_code']} deleted" + (f" (removed from {n_bom} BOM line(s))" if n_bom else "")
                       + (f". Backup: {backup_file}" if backup_file else ".")}


def list_projects() -> list[dict]:
    from . import scope
    g = scope.groups()
    with db.engine().connect() as conn:
        cq = select(db.parts.c.project_id, func.count()).group_by(db.parts.c.project_id)
        pq = select(db.projects).order_by(db.projects.c.code)
        if g is not None:
            cq, pq = cq.where(scope.part_condition(g)), pq.where(scope.project_condition(g))
        counts = dict(conn.execute(cq).all())
        rows = conn.execute(pq).mappings().all()
        return [_plain(r) | {"part_count": counts.get(r["id"], 0)} for r in rows]


def get_project(pid: int) -> dict:
    with db.engine().connect() as conn:
        p = conn.execute(select(db.projects).where(db.projects.c.id == pid)).mappings().first()
        if not p:
            raise RegisterError(f"Project {pid} does not exist.")
        vs = conn.execute(select(db.variants).where(db.variants.c.project_id == pid).order_by(db.variants.c.code)).mappings().all()
        return _plain(p) | {"variants": [_plain(v) for v in vs]}


def _variant_values(data, where="Variant"):
    out = {}
    for k in VARIANT_FIELDS:
        if k in data:
            if k == "target_bom_cost_inr":
                out[k] = _num(data[k], f"{where}: target BOM cost")
            elif k == "annual_volume":
                out[k] = _num(data[k], f"{where}: annual volume", int)
            else:
                out[k] = str(data[k] or "").strip()
    if "code" in out:
        out["code"] = out["code"].upper()
        if not out["code"]:
            raise RegisterError(f"{where}: code must not be empty.")
    return out


def create_variant(pid: int, data: dict) -> dict:
    vals = _variant_values(data)
    if not vals.get("code") or not vals.get("name"):
        raise RegisterError("Variant: code and name must not be empty.")
    with db.engine().begin() as conn:
        if conn.execute(select(db.variants.c.id).where(and_(db.variants.c.project_id == pid,
                                                             db.variants.c.code == vals["code"]))).first():
            raise RegisterError(f"Variant '{vals['code']}' already exists in this project.")
        vid = conn.execute(db.variants.insert().values(project_id=pid, **vals)).inserted_primary_key[0]
        _log(conn, "variant", vid, "created", {"code": vals["code"], "project_id": pid})
    return get_project(pid)


def update_variant(vid: int, data: dict) -> dict:
    vals = _variant_values(data)
    with db.engine().begin() as conn:
        old = conn.execute(select(db.variants).where(db.variants.c.id == vid)).mappings().first()
        if not old:
            raise RegisterError(f"Variant {vid} does not exist.")
        changes = {k: {"old": old[k], "new": v} for k, v in vals.items() if old[k] != v}
        if changes:
            conn.execute(db.variants.update().where(db.variants.c.id == vid).values(**vals))
            _log(conn, "variant", vid, "updated", changes)
    return get_project(old["project_id"])


# ---------- parts & revisions ----------

PART_FIELDS = ("name", "material", "process", "sourcing", "status", "annual_volume", "programme_life_years",
               "target_piece_cost_inr", "notes", "unit_price", "price_currency", "price_source", "price_month",
               "price_is_example", "dimensions", "country", "costing_month", "rm_inputs", "tool_country", "tool_inputs", "piece_inputs", "sheet_inputs")


def next_part_code(conn, project_code: str) -> str:
    pattern = register_setup.settings()["part_code_pattern"]
    seq = conn.execute(select(func.count()).select_from(db.parts).join(db.projects)
                       .where(db.projects.c.code == project_code)).scalar() + 1
    while True:
        code = pattern.format(project=project_code, seq=seq)
        if not conn.execute(select(db.parts.c.id).where(db.parts.c.part_code == code)).first():
            return code
        seq += 1


def material_families(material: str) -> list[str]:
    """Part types (plastic, die_casting, foundry, sheet_metal) under which a material is listed in Master data → Materials."""
    return sorted({m["part_type"] for m in master_data.tables()["Materials"] if material and m["material"] == material})


def suggest_process(metrics: dict | None, material: str = "", insert_count: float = 0) -> dict:
    """Suggest a process: material first, then the geometry rules, then inserts. Returns the reason step by step."""
    family_label = {"plastic": "a plastic", "die_casting": "a die-casting alloy", "foundry": "a foundry casting material",
                    "sheet_metal": "a sheet-metal grade", "eps": "an EPS (thermocole) bead grade", "blow": "a PET bottle resin"}
    procs = register_setup.processes()
    label = lambda c: procs[c]["name"] if c in procs else c
    if metrics and metrics.get("assembly"):
        code = register_setup.assembly_process()
        return {"process": code, "scores": {}, "steps": ["This is an assembly (it has child parts)."],
                "reason": "This is an assembly (it has child parts)."}
    steps, scores = [], {}
    families = material_families(material)
    candidates = list(procs)
    if material and families:
        candidates = [c for c, p in procs.items() if p["part_type"] in families]
        fam = " and ".join(family_label.get(f, f) for f in families)
        steps.append(f"1 · Material: {material} is {fam} (Master data → Materials, part type "
                     f"{', '.join(families)}), so only "
                     f"{', '.join(label(c) for c in candidates)} are considered.")
    elif material and material in register_setup.insert_materials():
        code = register_setup.settings()["insert_default_process"]
        text = (f"1 · Material: {material} is an insert material (Part register setup → Insert materials), "
                f"so the part is taken as {label(code)}.")
        return {"process": code, "scores": {}, "steps": [text], "reason": text, "material_family": []}
    elif material:
        steps.append(f"1 · Material: {material} is not in Master data → Materials, so the geometry decides alone.")
    else:
        steps.append("1 · Material: not set yet, so this is a geometry-only suggestion. Set the material for a firmer one.")
    if not metrics:
        if material and families:
            code = next((c for c in candidates if bool(procs[c]["is_insert_moulding"]) == bool(insert_count)), candidates[0])
            steps.append(f"2 · Geometry: no 3D file yet, so the first matching process is taken: {label(code)}.")
            return {"process": code, "scores": {}, "steps": steps, "reason": " ".join(steps), "material_family": families}
        steps.append("2 · Geometry: no 3D file yet, so no suggestion.")
        return {"process": "", "scores": {}, "steps": steps, "reason": " ".join(steps), "material_family": families}

    env = {k: metrics.get(k, 0) for k in register_setup.SUGGEST_VARIABLES}
    env["insert_count"] = insert_count or 0
    for code, p in procs.items():
        try:
            scores[code] = formula.evaluate(p["suggest_rule"] or "0", env)
        except formula.FormulaError:
            scores[code] = 0
    steps.append(f"2 · Geometry: main wall {env['wall_mm']:g} mm over {env['sheet_share'] * 100:.0f} % of the surface, "
                 f"{env['bend_count']:g} bend{'s' if env['bend_count'] != 1 else ''}, fill {env['fill_ratio'] * 100:.0f} % "
                 f"of the bounding box, largest size {env['max_dim_mm']:g} mm.")
    if insert_count:
        steps.append(f"3 · Inserts: {insert_count:g} moulded-in insert{'s' if insert_count != 1 else ''} "
                     "(insert_count in the rules).")
    best = max(candidates, key=lambda c: scores[c]) if candidates else ""
    n = len(steps) + 1
    if best and scores[best] > 0:
        steps.append(f"{n} · {label(best)} scores highest ({scores[best]:g}) with its rule: {procs[best]['suggest_rule']}.")
    elif material and families:
        best = next((c for c in candidates if bool(procs[c]["is_insert_moulding"]) == bool(insert_count)), candidates[0])
        steps.append(f"{n} · No geometry rule of these processes fits, so the first one is taken: {label(best)}.")
    else:
        best = ""
        steps.append(f"{n} · No rule scores above 0. Please choose the process.")
    overall = max(scores, key=scores.get) if scores else ""
    if best and overall and overall != best and scores[overall] > scores.get(best, 0) and families:
        steps.append(f"Geometry alone would point to {label(overall)} ({scores[overall]:g}); the material decides.")
    return {"process": best, "scores": scores, "inputs": env, "candidates": candidates, "material_family": families,
            "steps": steps, "reason": " ".join(steps)}


def _part_values(data, where="Part"):
    out = {}
    procs, srcs = register_setup.processes(), register_setup.sourcing()
    statuses = {s["code"] for s in register_setup.statuses()}
    for k in PART_FIELDS:
        if k not in data:
            continue
        v = data[k]
        if k == "annual_volume":
            out[k] = _num(v, f"{where}: annual volume", int)
        elif k in ("programme_life_years", "target_piece_cost_inr", "unit_price"):
            out[k] = _num(v, f"{where}: {k.replace('_', ' ')}")
        elif k in ("country", "tool_country"):
            code = str(v or "").strip().upper()
            if code and code not in costing_setup.countries():
                raise RegisterError(f"{where}: country '{code}' is not in Master data → Countries.")
            out[k] = code
        elif k == "costing_month":
            m = str(v or "").strip()
            if m and not rates.MONTH.match(m):
                raise RegisterError(f"{where}: costing month '{m}' must look like 2026-10.")
            out[k] = m
        elif k == "rm_inputs":
            out[k] = _rm_inputs(v, where)
        elif k == "tool_inputs":
            out[k] = _tool_inputs(v, where)
        elif k == "piece_inputs":
            out[k] = _piece_inputs(v, where)
        elif k == "sheet_inputs":
            out[k] = _sheet_inputs(v, where)
        elif k == "price_is_example":
            out[k] = v is True or str(v).strip().lower() in ("1", "true", "yes", "on", "example")
        elif k == "dimensions":
            out[k] = _dimensions(v, where)
        else:
            out[k] = str(v or "").strip()
    if out.get("process") and out["process"] not in procs:
        raise RegisterError(f"{where}: process '{out['process']}' is not in Master data → Part register setup → "
                            f"Processes ({', '.join(procs)}).")
    if out.get("sourcing") and out["sourcing"] not in srcs:
        raise RegisterError(f"{where}: sourcing '{out['sourcing']}' must be one of {', '.join(srcs)}.")
    if out.get("status") and out["status"] not in statuses:
        raise RegisterError(f"{where}: status '{out['status']}' is not in Part statuses.")
    if "name" in out and not out["name"]:
        raise RegisterError(f"{where}: name must not be empty.")
    if out.get("price_month") and not (len(out["price_month"]) == 7 and out["price_month"][4] == "-"
                                       and out["price_month"].replace("-", "").isdigit()):
        raise RegisterError(f"{where}: price month '{out['price_month']}' must look like 2026-10.")
    return out


def _dimensions(v, where):
    """{"shape": code, <dimension>: mm, ...} checked against Part register setup → Insert shapes."""
    if v in (None, "", {}):
        return None
    if isinstance(v, str):
        try:
            v = json.loads(v)
        except ValueError:
            raise RegisterError(f"{where}: dimensions are not readable.") from None
    shapes = register_setup.insert_shapes()
    shape = str(v.get("shape") or "").strip()
    if not shape:
        return None
    if shape not in shapes:
        raise RegisterError(f"{where}: shape '{shape}' is not in Part register setup → Insert shapes "
                            f"({', '.join(shapes)}).")
    out = {"shape": shape}
    for d in shapes[shape]["dims"]:
        n = _num(v.get(d), f"{where}: {d}", required=True)
        if n <= 0:
            raise RegisterError(f"{where}: {d} must be more than 0.")
        out[d] = n
    return out


def _project_by_id(conn, pid):
    p = conn.execute(select(db.projects).where(db.projects.c.id == pid)).mappings().first()
    if not p:
        raise RegisterError(f"Project {pid} does not exist.")
    return p


def upsert_part(conn, project_id: int, code: str, values: dict, is_assembly=False) -> tuple[int, bool]:
    """Create the part, or return the existing one with that code. -> (part_id, created)."""
    project = _project_by_id(conn, project_id)
    code = (code or "").strip() or next_part_code(conn, project["code"])
    row = conn.execute(select(db.parts).where(db.parts.c.part_code == code)).mappings().first()
    if row:
        return row["id"], False
    vals = {"sourcing": next(iter(register_setup.sourcing())), "status": register_setup.initial_status()} | values
    if is_assembly:
        vals["process"] = vals.get("process") or register_setup.assembly_process()
    from . import scope
    scope.check_process(vals.get("process"))
    pid = conn.execute(db.parts.insert().values(part_code=code, project_id=project_id, is_assembly=is_assembly,
                                                created_by=user(), **vals)).inserted_primary_key[0]
    _log(conn, "part", pid, "created", {"part_code": code, **{k: v for k, v in vals.items() if v not in (None, "")}})
    return pid, True


def add_revision(conn, part_id: int, rev: str, *, step: bytes | None = None, file_name: str = "",
                 thumbnail: bytes | None = None, geometry: dict | None = None, change_note: str = "",
                 source: str = "") -> int:
    rev = (rev or "").strip()
    existing = [r for (r,) in conn.execute(select(db.revisions.c.rev).where(db.revisions.c.part_id == part_id))]
    if not rev:
        rev = _next_rev(existing)
    if rev in existing:
        code = conn.execute(select(db.parts.c.part_code).where(db.parts.c.id == part_id)).scalar()
        raise RegisterError(f"Part {code} already has revision '{rev}'. Use a new revision number.")
    file_row = file_store.put(step, file_name or "part.step", "step", conn) if step else None
    thumb_row = file_store.put(thumbnail, "thumbnail.png", "thumbnail", conn) if thumbnail else None
    rid = conn.execute(db.revisions.insert().values(
        part_id=part_id, rev=rev, file_id=file_row and file_row["id"], thumbnail_file_id=thumb_row and thumb_row["id"],
        change_note=change_note, geometry=geometry, source=source, created_by=user())).inserted_primary_key[0]
    conn.execute(db.parts.update().where(db.parts.c.id == part_id).values(current_revision_id=rid))
    _log(conn, "part", part_id, "revision_added", {"rev": rev, "file": file_name if step else "(file missing)",
                                                     "change_note": change_note, "source": source})
    if existing:  # design loop: a new revision can move the status on (Part statuses → after a new revision)
        old = conn.execute(select(db.parts.c.status).where(db.parts.c.id == part_id)).scalar()
        new = next((s["on_new_revision"] for s in register_setup.statuses() if s["code"] == old and s.get("on_new_revision")), "")
        if new and new != old:
            conn.execute(db.parts.update().where(db.parts.c.id == part_id).values(status=new))
            _log(conn, "part", part_id, "status_changed", {"old": old, "new": new, "comment": f"new revision {rev} uploaded"})
    return rid


def _next_rev(existing: list[str]) -> str:
    if not existing:
        return "A"
    last = sorted(existing)[-1]
    if last.isdigit():
        return str(int(last) + 1).zfill(len(last))
    if len(last) == 1 and last.isalpha() and last.upper() != "Z":
        return chr(ord(last.upper()) + 1)
    return f"{last}.1"


def _same(a, b) -> bool:
    """Equal, or both empty (None, "" and an unticked box are the same)."""
    empty = lambda x: x is None or x == "" or x is False or x == {}
    return a == b or (empty(a) and empty(b))


def _back_codes(code: str) -> list[str]:
    """Design loop: the earlier statuses (lower order) a part may step back to, nearest first, e.g. a step forgotten."""
    sts = register_setup.statuses()
    cur = next((s for s in sts if s["code"] == code), None)
    return [s["code"] for s in reversed(sts) if cur and s["sort_order"] < cur["sort_order"]]


def _is_step_back(old: str, new: str) -> bool:
    """A move to an earlier status that Part statuses does not allow as a next step (e.g. a forgotten step).
    An allowed move to an earlier status (Revised → Simulation) is a normal step."""
    st = {s["code"]: s for s in register_setup.statuses()}
    allowed = register_setup.next_codes(st[old]) if old in st else []
    return new in _back_codes(old) and new not in allowed


def _check_status_change(old: str, new: str, comment: str) -> None:
    """Design loop: the move must be allowed by Part statuses (next statuses), with a comment where required.
    A step back to an earlier status is always allowed and always needs a comment (why)."""
    st = {s["code"]: s for s in register_setup.statuses()}
    if new not in st:
        raise RegisterError(f"Status '{new}' is not in Part register setup → Part statuses.")
    if _is_step_back(old, new):
        if not str(comment or "").strip():
            raise RegisterError(f"Stepping back to '{st[new]['name']}' needs a comment (why).")
        return
    allowed = register_setup.next_codes(st[old]) if old in st else []
    if allowed and new not in allowed:
        names = ", ".join(st[c]["name"] for c in allowed if c in st)
        raise RegisterError(f"From '{st[old]['name']}' the part can only move to: {names}, or step back to an earlier "
                            "status with a comment (Part register setup → Part statuses).")
    if st[new].get("comment_required") and not str(comment or "").strip():
        raise RegisterError(f"Moving to '{st[new]['name']}' needs a comment (what has to change, or why).")


def set_status(part_id: int, status: str, comment: str = "") -> dict:
    return update_part(part_id, {"status": status, "status_comment": comment}, refresh=False)


def design_loop(part_id: int) -> dict:
    """Status flow of the part: current status, allowed next ones, and the timeline of status changes and revisions."""
    d = part_detail(part_id, with_rm=False)
    sts = register_setup.statuses()
    st = {s["code"]: s for s in sts}
    cur = st.get(d["status"], {})
    timeline = [{"when": h["at"], "who": h["user"], "action": h["action"], **(h["detail"] or {})}
                for h in reversed(d["history"]) if h["action"] in ("created", "status_changed", "revision_added")]
    days, last = {}, None
    for e in timeline + [{"when": datetime.now().isoformat(timespec="seconds"), "action": "now"}]:
        if last:
            dt = (datetime.fromisoformat(e["when"]) - datetime.fromisoformat(last[1])).total_seconds() / 86400
            days[last[0]] = days.get(last[0], 0) + dt
        if e["action"] == "status_changed":
            last = (e["new"], e["when"])
        elif e["action"] == "created":
            last = (e.get("status") or register_setup.initial_status(), e["when"])
    return {"part_id": part_id, "status": d["status"],
            "flow": [{"code": s["code"], "name": s["name"], "comment_required": bool(s.get("comment_required")),
                      "days": round(days.get(s["code"], 0), 2)} for s in sts],
            "next": register_setup.next_codes(cur) or [s["code"] for s in sts if s["code"] != d["status"]],
            "back": _back_codes(d["status"]),
            "timeline": timeline}


def update_part(part_id: int, data: dict, refresh: bool = True) -> dict:
    vals = _part_values(data)
    with db.engine().begin() as conn:
        old = conn.execute(select(db.parts).where(db.parts.c.id == part_id)).mappings().first()
        if not old:
            raise RegisterError(f"Part {part_id} does not exist.")
        changes = {k: {"old": old[k], "new": v} for k, v in vals.items() if not _same(old[k], v)}
        if "process" in changes:
            from . import scope
            scope.check_process(vals["process"])
        if "status" in changes:
            _check_status_change(old["status"], vals["status"], data.get("status_comment", ""))
        if changes:
            conn.execute(db.parts.update().where(db.parts.c.id == part_id).values(**vals))
            _log(conn, "part", part_id, "updated", changes)
        if "status" in changes:
            _log(conn, "part", part_id, "status_changed", {"old": old["status"], "new": vals["status"],
                                                           "comment": str(data.get("status_comment") or "").strip(),
                                                           "step_back": _is_step_back(old["status"], vals["status"])})
    if refresh and {"material", "process", "sourcing", "country", "costing_month", "rm_inputs", "tool_country",
                    "tool_inputs", "annual_volume", "programme_life_years"} & set(changes):
        refresh_costs([part_id], "part details changed")
    return part_detail(part_id)


def new_revision_from_upload(part_id: int, content: bytes, file_name: str, rev: str, change_note: str) -> dict:
    tmp = file_store.scratch_dir(uuid.uuid4().hex)
    try:
        path = tmp / "upload.step"
        path.write_bytes(content)
        tree = cad_import.analyse_file(str(path), file_name)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    root = tree["defs"][tree["root"]]
    geometry = root["metrics"]
    with db.engine().begin() as conn:
        add_revision(conn, part_id, rev, step=content, file_name=file_name, thumbnail=root.get("thumbnail"),
                     geometry=geometry, change_note=change_note, source=f"upload {file_name}")
    refresh_costs([part_id], "new revision")
    return part_detail(part_id)


# ---------- reading ----------

def _part_query():
    r = db.revisions.alias("cur")
    return (select(db.parts, db.projects.c.code.label("project_code"), r.c.rev, r.c.file_id, r.c.thumbnail_file_id,
                   r.c.geometry, r.c.created_at.label("revision_date"))
            .join(db.projects, db.projects.c.id == db.parts.c.project_id)
            .outerjoin(r, r.c.id == db.parts.c.current_revision_id)), r


def search(q="", project_id=None, process="", status="", sourcing="", rev="", file_missing=None,
           target_min=None, target_max=None, rm_min=None, rm_max=None, rm_missing=None, tool_min=None, tool_max=None,
           tool_missing=None, sort="part_code", limit=50, offset=0) -> dict:
    from . import scope
    query, r = _part_query()
    conds = []
    g = scope.groups()
    if g is not None:
        conds.append(scope.part_condition(g))
    if q:
        like = f"%{q.strip()}%"
        conds.append(or_(db.parts.c.part_code.ilike(like), db.parts.c.name.ilike(like), db.parts.c.material.ilike(like)))
    if project_id:
        conds.append(db.parts.c.project_id == project_id)
    for col, val in ((db.parts.c.process, process), (db.parts.c.status, status), (db.parts.c.sourcing, sourcing)):
        if val:
            conds.append(col == val)
    if rev:
        conds.append(r.c.rev == rev)
    if file_missing is True:
        conds.append(r.c.file_id.is_(None))
    if file_missing is False:
        conds.append(r.c.file_id.is_not(None))
    if target_min is not None:
        conds.append(db.parts.c.target_piece_cost_inr >= target_min)
    if target_max is not None:
        conds.append(db.parts.c.target_piece_cost_inr <= target_max)
    if rm_min is not None:
        conds.append(db.parts.c.rm_cost_inr >= rm_min)
    if rm_max is not None:
        conds.append(db.parts.c.rm_cost_inr <= rm_max)
    if rm_missing is True:
        conds.append(and_(db.parts.c.rm_cost_inr.is_(None), db.parts.c.is_assembly.is_(False)))
    if tool_min is not None:
        conds.append(db.parts.c.tool_cost_inr >= tool_min)
    if tool_max is not None:
        conds.append(db.parts.c.tool_cost_inr <= tool_max)
    if tool_missing is True:  # a part whose process has a tool cost (moulding, die casting) but none stored yet
        tool_procs = [c for c, p in register_setup.processes().items() if p["part_type"] in tool_cost.PART_TYPES]
        made = [c for c, x in register_setup.sourcing().items() if x["is_costed"]]
        conds.append(and_(db.parts.c.tool_cost_inr.is_(None), db.parts.c.is_assembly.is_(False),
                          db.parts.c.process.in_(tool_procs), db.parts.c.sourcing.in_(made)))
    if conds:
        query = query.where(and_(*conds))
    sort_cols = {"rm_cost": db.parts.c.rm_cost_inr.desc(), "tool_cost": db.parts.c.tool_cost_inr.desc(), "part_code": db.parts.c.part_code, "name": db.parts.c.name, "updated": db.parts.c.updated_at.desc(),
                 "status": db.parts.c.status, "process": db.parts.c.process}
    with db.engine().connect() as conn:
        total = conn.execute(select(func.count()).select_from(query.subquery())).scalar()
        rows = conn.execute(query.order_by(sort_cols.get(sort, db.parts.c.part_code))
                            .limit(min(int(limit), 500)).offset(int(offset))).mappings().all()
        example = _rm_example_parts(conn, [x["id"] for x in rows])
        tool_example = _tool_example_parts(conn, [x["id"] for x in rows])
    return {"total": total, "rows": [_plain(x) | {"rm_is_example": x["id"] in example, "tool_is_example": x["id"] in tool_example}
                                     for x in rows], "example_warning": EXAMPLE_WARNING}


def _rm_example_parts(conn, part_ids) -> set:
    """Parts whose stored (latest) RM cost used a rate that is still an EXAMPLE."""
    if not part_ids:
        return set()
    q = (select(db.parts.c.id, db.rm_costs.c.rates_used).join(db.rm_costs, db.rm_costs.c.id == db.parts.c.rm_cost_id)
         .where(db.parts.c.id.in_(list(part_ids))))
    return {r["id"] for r in conn.execute(q).mappings() if any(x.get("is_example") for x in r["rates_used"] or [])}


def _tool_example_parts(conn, part_ids) -> set:
    """Parts whose stored (latest) tool cost used an EXAMPLE rate or master-data row."""
    if not part_ids:
        return set()
    q = (select(db.parts.c.id, db.tool_costs.c.example).join(db.tool_costs, db.tool_costs.c.id == db.parts.c.tool_cost_id)
         .where(db.parts.c.id.in_(list(part_ids))))
    return {r["id"] for r in conn.execute(q).mappings() if (r["example"] or {}).get("any")}


def part_detail(part_id: int, with_rm: bool = True) -> dict:
    query, _ = _part_query()
    with db.engine().connect() as conn:
        p = conn.execute(query.where(db.parts.c.id == part_id)).mappings().first()
        if not p:
            raise RegisterError(f"Part {part_id} does not exist.")
        revs = conn.execute(select(db.revisions, db.files.c.original_name, db.files.c.size_bytes, db.files.c.sha256)
                            .outerjoin(db.files, db.files.c.id == db.revisions.c.file_id)
                            .where(db.revisions.c.part_id == part_id).order_by(db.revisions.c.id.desc())).mappings().all()
        hist = conn.execute(select(db.history).where(and_(db.history.c.entity == "part", db.history.c.entity_id == part_id))
                            .order_by(db.history.c.id.desc())).mappings().all()
        parent = db.parts.alias("parent")
        used = conn.execute(select(db.bom_lines.c.quantity, db.variants.c.code.label("variant"),
                                   parent.c.part_code.label("parent_code"), parent.c.name.label("parent_name"))
                            .join(db.variants, db.variants.c.id == db.bom_lines.c.variant_id)
                            .outerjoin(parent, parent.c.id == db.bom_lines.c.parent_part_id)
                            .where(db.bom_lines.c.child_part_id == part_id)).mappings().all()
        children = conn.execute(select(db.bom_lines.c.quantity, db.variants.c.code.label("variant"),
                                       db.parts.c.id, db.parts.c.part_code, db.parts.c.name)
                                .join(db.parts, db.parts.c.id == db.bom_lines.c.child_part_id)
                                .join(db.variants, db.variants.c.id == db.bom_lines.c.variant_id)
                                .where(db.bom_lines.c.parent_part_id == part_id)).mappings().all()
        inserts = _insert_rows(conn, part_id)
        body = db.parts.alias("body")
        moulded_into = conn.execute(select(db.part_inserts.c.quantity, body.c.id, body.c.part_code, body.c.name)
                                    .join(body, body.c.id == db.part_inserts.c.part_id)
                                    .where(db.part_inserts.c.insert_part_id == part_id)).mappings().all()
    insert_count = sum(i["quantity"] for i in inserts)
    suggestion = suggest_process(p["geometry"], p["material"], insert_count) if not p["is_assembly"] else None
    out = _plain(p) | {"revisions": [_plain(x) for x in revs], "history": [_plain(x) for x in hist],
                       "where_used": [_plain(x) for x in used], "children": [_plain(x) for x in children],
                       "inserts": inserts, "insert_count": insert_count,
                       "moulded_into": [_plain(x) for x in moulded_into], "suggestion": suggestion}
    if with_rm:
        out["rm"] = rm_summary(out)
        out["effective_country"] = (out["country"] or register_setup.settings()["default_country"]).upper()
        out["tool"] = tool_summary(out)
        out["effective_tool_country"] = (out["tool_country"] or out["effective_country"]).upper()
    return out


def revision_trace(part_id: int, rev_id: int | None = None) -> dict:
    """Traceable geometry and weight of a revision (see _geometry_trace)."""
    t, nodes, warnings, *_ = _geometry_trace(part_id, rev_id)
    out = t.to_dict(nodes[0])
    out["components"] = nodes
    out["warnings"] = warnings
    return out


def _geometry_trace(part_id: int, rev_id: int | None = None):
    """Traceable geometry and weight of a revision -> (Trace, node ids, warnings, part detail, revision).

    Body weight = (CAD volume − any modelled overlap with inserts) × material density.
    With moulded-in inserts: insert weight = Σ quantity × insert volume × insert density
    (volume from the insert's own 3D file, or from its shape formula and dimensions), and
    total weight = body weight + insert weight.
    """
    from .trace import Trace
    d = part_detail(part_id, with_rm=False)
    rev = next((r for r in d["revisions"] if rev_id in (None, r["id"])), None)
    if not rev or not rev["geometry"]:
        raise RegisterError("This revision has no 3D file yet, so there is no geometry to trace.")
    g = rev["geometry"]
    if g.get("assembly"):
        raise RegisterError("An assembly has no geometry of its own; open its parts to see their volume and weight.")
    cad = {"type": "cad", "file": rev["original_name"], "rev": rev["rev"], "sha256": rev["sha256"],
           "note": "Measured from the STEP file with OpenCascade (CadQuery)"}
    t = Trace()
    t.add_input("volume", "Volume", g["volume_cm3"], "cm³", cad)
    t.add_input("surface", "Surface area", g["surface_area_mm2"] / 100, "cm²", cad)
    t.add_input("proj", f"Projected area (along {g['pull_direction'].upper()})",
                g["projected_area_mm2"][g["pull_direction"]] / 100, "cm²", cad)
    nodes = ["volume", "surface", "proj"]
    warnings = []
    volume_node = "volume"
    if g.get("insert_overlap_mm3"):
        t.add_input("overlap", "Insert volume inside the body model", g["insert_overlap_mm3"] / 1000, "cm³",
                    {"type": "cad", "file": rev["original_name"], "rev": rev["rev"],
                     "note": "The body was modelled solid through its inserts; the shared volume is measured and removed"})
        t.compute("body_volume", "Body volume without inserts", "volume - overlap", "cm³")
        volume_node = "body_volume"
        nodes.insert(0, "body_volume")
    density, dsrc = _density(d["material"])
    if density:
        t.add_input("density", f"Density of {d['material']}", density, "g/cm³", dsrc)
        label = "Body weight (plastic / metal, without inserts)" if d["inserts"] else "Net weight"
        proc = register_setup.processes().get(d["process"]) or {}
        if proc.get("part_type") == "blow":
            # a bottle is a thin shell: the CAD volume is the outer envelope, so weight = outer surface × average wall × density
            typed = (d["rm_inputs"] or {}).get("wall_mm")
            wall = typed if typed else master_data.parameters("blow").get("blow_wall_mm")
            if not wall:
                raise RegisterError("Master data → Parameters has no 'blow_wall_mm' for blow, and no wall thickness is typed under RM inputs.")
            t.add_input("wall", "Average wall thickness", wall, "mm",
                        {"type": "user", "note": f"Typed under RM inputs for {d['part_code']}"} if typed else
                        {"type": "master", "sheet": "Parameters", "row": "blow / blow_wall_mm"})
            t.compute("weight", "Net bottle weight (shell)", "surface * wall / 10 * density", "g",
                      names={"surface": "surface", "wall": "wall", "density": "density"},
                      description="Outer surface area × average wall thickness × resin density (the CAD volume is not used for a hollow bottle)")
        else:
            t.compute("weight", label, "volume * density", "g", names={"volume": volume_node, "density": "density"},
                      description="Volume from the CAD file × material density from master data")
        nodes.insert(0, "weight")
    else:
        warnings.append("Set the part's material to calculate its weight (density from master data).")

    parts_ok = []
    for i, ins in enumerate(d["inserts"], start=1):
        x, code = ins["insert"], ins["insert"]["part_code"]
        vol = ins["volume"]
        if ins["missing"]:
            warnings.append(f"Insert {code}: {ins['missing']}, so its weight is not included.")
            continue
        if vol["method"] == "3d":
            t.add_input(f"i{i}_vol", f"Volume of insert {code}", vol["volume_cm3"], "cm³", vol["source"])
        else:
            names = {}
            for dim, val in vol["dims"].items():
                t.add_input(f"i{i}_{dim}", f"{code} {dim}", val, "mm",
                            {"type": "register", "note": f"Dimensions typed in the part register for {code}"})
                names[dim] = f"i{i}_{dim}"
            for k, v in register_setup.SHAPE_CONSTANTS.items():
                t.add_input(f"c_{k}", k, v, "", {"type": "base", "note": "Mathematical constant"})
                names[k] = f"c_{k}"
            t.compute(f"i{i}_mm3", f"Volume of insert {code} ({vol['shape']['name']})", vol["shape"]["volume_rule"],
                      "mm³", names=names, description="Shape formula from Part register setup → Insert shapes")
            t.compute(f"i{i}_vol", f"Volume of insert {code}", "v / 1000", "cm³", names={"v": f"i{i}_mm3"})
        t.add_input(f"i{i}_den", f"Density of {x['material']}", ins["density_g_cm3"], "g/cm³", ins["density_source"])
        t.add_input(f"i{i}_qty", f"Quantity of {code} in this part", ins["quantity"], "pcs",
                    {"type": "register", "note": "Insert link in the part register"
                     + (" (detected in the 3D assembly)" if ins["detected"] else "")})
        t.compute(f"i{i}", f"Insert {code} · {x['name']}", "qty * vol * den", "g",
                  names={"qty": f"i{i}_qty", "vol": f"i{i}_vol", "den": f"i{i}_den"})
        parts_ok.append(f"i{i}")
    if parts_ok:
        t.compute("insert_weight", "Insert weight", " + ".join(f"w{n}" for n in range(len(parts_ok))), "g",
                  names={f"w{n}": node for n, node in enumerate(parts_ok)},
                  description="Sum of quantity × volume × density of every insert")
        nodes.insert(0, "insert_weight")
        if density:
            t.compute("total_weight", "Total part weight", "body + inserts", "g",
                      names={"body": "weight", "inserts": "insert_weight"},
                      description="Body (plastic / metal) weight + insert weight")
            nodes.insert(0, "total_weight")
    return t, nodes, warnings, d, rev


# ---------- moulded-in inserts ----------

def _density(material: str):
    """(density g/cm³, source) from Insert materials first, then Master data → Materials; (None, None) if unknown."""
    im = register_setup.insert_materials().get(material)
    if im:
        return im["density_g_cm3"], {"type": "master", "sheet": "Part register setup → Insert materials", "row": material,
                                     "is_example": im["is_example"]}
    mat = next((m for m in master_data.tables()["Materials"] if material and m["material"] == material), None)
    if mat:
        return mat["density_g_cm3"], {"type": "master", "sheet": "Materials", "row": material, "part_type": mat["part_type"]}
    return None, None


def _current_geometry(conn, part_id):
    return conn.execute(select(db.revisions.c.geometry, db.revisions.c.rev, db.files.c.original_name, db.files.c.sha256)
                        .select_from(db.parts.join(db.revisions, db.revisions.c.id == db.parts.c.current_revision_id)
                                     .outerjoin(db.files, db.files.c.id == db.revisions.c.file_id))
                        .where(db.parts.c.id == part_id)).mappings().first()


def _insert_volume(conn, ins) -> dict:
    """Volume of one insert: from its own 3D file, else from its shape + dimensions."""
    cur = _current_geometry(conn, ins["id"])
    g = cur["geometry"] if cur else None
    if g and g.get("volume_cm3") is not None:
        return {"method": "3d", "volume_cm3": g["volume_cm3"],
                "source": {"type": "cad", "file": cur["original_name"], "rev": cur["rev"], "sha256": cur["sha256"],
                           "note": f"Insert {ins['part_code']}: measured from its STEP file"}}
    dims = ins["dimensions"] or {}
    shapes = register_setup.insert_shapes()
    if dims.get("shape") in shapes:
        sh = shapes[dims["shape"]]
        env = register_setup.SHAPE_CONSTANTS | {d: dims.get(d, 0) for d in sh["dims"]}
        return {"method": "dimensions", "shape": sh, "dims": {d: dims.get(d) for d in sh["dims"]},
                "volume_cm3": formula.evaluate(sh["volume_rule"], env) / 1000}
    return {"method": "none", "volume_cm3": None}


def _insert_rows(conn, part_id) -> list[dict]:
    ins = db.parts.alias("ins")
    rows = conn.execute(select(db.part_inserts.c.id.label("link_id"), db.part_inserts.c.quantity,
                               db.part_inserts.c.detected, db.part_inserts.c.notes, ins)
                        .join(ins, ins.c.id == db.part_inserts.c.insert_part_id)
                        .where(db.part_inserts.c.part_id == part_id).order_by(ins.c.part_code)).mappings().all()
    out = []
    for r in rows:
        r = dict(r)
        vol = _insert_volume(conn, r)
        density, dsrc = _density(r["material"])
        each = vol["volume_cm3"] * density if vol["volume_cm3"] is not None and density else None
        missing = ("no 3D file and no dimensions" if vol["volume_cm3"] is None else
                   "material has no density in master data" if not density else "")
        out.append({"link_id": r["link_id"], "quantity": r["quantity"], "detected": r["detected"], "notes": r["notes"],
                    "insert": {k: r[k] for k in ("id", "part_code", "name", "material", "sourcing", "process", "unit_price",
                                                 "price_currency", "price_source", "price_month", "price_is_example",
                                                 "dimensions", "current_revision_id")},
                    "volume": vol, "density_g_cm3": density, "density_source": dsrc,
                    "weight_each_g": each, "weight_g": each * r["quantity"] if each is not None else None,
                    "missing": missing})
    return out


def _set_insert_link(conn, part_id, insert_id, qty, detected=False, notes=""):
    if part_id == insert_id:
        raise RegisterError("A part cannot be its own insert.")
    body = conn.execute(select(db.parts).where(db.parts.c.id == part_id)).mappings().first()
    ins = conn.execute(select(db.parts).where(db.parts.c.id == insert_id)).mappings().first()
    if not body or not ins:
        raise RegisterError("Part not found.")
    if body["is_assembly"] or ins["is_assembly"]:
        raise RegisterError("Inserts link a moulded / cast part and a single insert part, not an assembly.")
    if conn.execute(select(db.part_inserts.c.id).where(and_(db.part_inserts.c.part_id == insert_id,
                                                            db.part_inserts.c.insert_part_id == part_id))).first():
        raise RegisterError(f"{body['part_code']} is already an insert of {ins['part_code']}; the link cannot go both ways.")
    if not qty or qty <= 0:
        raise RegisterError("Insert quantity must be more than 0.")
    old = conn.execute(select(db.part_inserts).where(and_(db.part_inserts.c.part_id == part_id,
                                                          db.part_inserts.c.insert_part_id == insert_id))).mappings().first()
    if old:
        if old["quantity"] != qty or (notes and old["notes"] != notes):
            conn.execute(db.part_inserts.update().where(db.part_inserts.c.id == old["id"])
                         .values(quantity=qty, notes=notes or old["notes"]))
            _log(conn, "part", part_id, "insert_updated", {"insert": ins["part_code"],
                                                           "quantity": {"old": old["quantity"], "new": qty}})
        return old["id"]
    lid = conn.execute(db.part_inserts.insert().values(part_id=part_id, insert_part_id=insert_id, quantity=qty,
                                                       detected=detected, notes=notes)).inserted_primary_key[0]
    _log(conn, "part", part_id, "insert_linked", {"insert": f"{ins['part_code']} · {ins['name']}", "quantity": qty,
                                                  "detected": "yes, in the 3D assembly" if detected else "no, typed in"})
    _log(conn, "part", insert_id, "moulded_into", {"part": f"{body['part_code']} · {body['name']}", "quantity": qty})
    return lid


def link_insert(part_id: int, data: dict) -> dict:
    """Link an insert to a moulded / cast part: an existing part (insert_part_code) or a new insert part."""
    qty = _num(data.get("quantity"), "Insert quantity", required=True)
    with db.engine().begin() as conn:
        body = conn.execute(select(db.parts).where(db.parts.c.id == part_id)).mappings().first()
        if not body:
            raise RegisterError(f"Part {part_id} does not exist.")
        code = str(data.get("insert_part_code") or "").strip()
        existing = conn.execute(select(db.parts.c.id).where(db.parts.c.part_code == code)).scalar() if code else None
        if existing:
            insert_id = existing
        else:
            if not str(data.get("name") or "").strip():
                raise RegisterError(f"Part code '{code}' is not in the register. Give a name to create the insert part."
                                    if code else "Give the insert's part code (existing part) or a name (new part).")
            s = register_setup.settings()
            vals = _part_values({"sourcing": s["insert_default_sourcing"], "process": s["insert_default_process"]}
                                | {k: data[k] for k in ("name", "material", "unit_price", "price_currency", "price_source",
                                                        "price_month", "price_is_example", "dimensions", "notes")
                                   if k in data}, "Insert")
            vals = {k: v for k, v in vals.items() if v not in (None, "")}
            insert_id, _ = upsert_part(conn, body["project_id"], code, vals)
        _set_insert_link(conn, part_id, insert_id, qty, notes=str(data.get("link_notes") or "").strip())
    refresh_costs([part_id], "insert linked")
    return part_detail(part_id)


def unlink_insert(part_id: int, link_id: int) -> dict:
    with db.engine().begin() as conn:
        row = conn.execute(select(db.part_inserts, db.parts.c.part_code).join(
            db.parts, db.parts.c.id == db.part_inserts.c.insert_part_id)
            .where(and_(db.part_inserts.c.id == link_id, db.part_inserts.c.part_id == part_id))).mappings().first()
        if not row:
            raise RegisterError("That insert link does not exist.")
        conn.execute(db.part_inserts.delete().where(db.part_inserts.c.id == link_id))
        _log(conn, "part", part_id, "insert_removed", {"insert": row["part_code"], "quantity": row["quantity"]})
        _log(conn, "part", row["insert_part_id"], "removed_from", {"part_id": part_id})
    refresh_costs([part_id], "insert removed")
    return part_detail(part_id)


# ---------- raw-material (RM) cost versions ----------
# The RM engine (rm_cost.py) works from the formulas in master data. Here its result is stored as a numbered,
# never-overwritten version of the part, with its full trace and the rate versions it used. A new version is made
# automatically when something the result depends on changes (new revision, material, process, country, month,
# typed inputs, inserts). Approving a newer rate does NOT change stored costs: the part shows "newer approved
# rates exist" and you recalculate on purpose.

class RMNotApplicable(Exception):
    """RM cannot be worked out for this part (the message says why and what to do)."""


RM_COLUMNS = ("id", "part_id", "version", "revision_id", "rev", "family", "process", "material", "country", "currency",
              "as_of", "inputs", "key", "net_weight_g", "gross_weight_g", "net_rm_cost", "net_rm_cost_inr",
              "rates_used", "reason", "created_by", "created_at")


def _rm_inputs(v, where):
    """{variable: number} typed over per part, checked against the RM formula inputs of every family."""
    if v in (None, "", {}):
        return None
    if isinstance(v, str):
        try:
            v = json.loads(v)
        except ValueError:
            raise RegisterError(f"{where}: RM inputs are not readable.") from None
    allowed = {i["variable"] for fam in rm_cost.FAMILIES for i in rm_cost.overridable(fam)}
    out = {}
    for k, val in v.items():
        if val in (None, ""):
            continue
        if k not in allowed:
            raise RegisterError(f"{where}: '{k}' is not an RM input you can type over ({', '.join(sorted(allowed))}).")
        n = _num(val, f"{where}: {k}", required=True)
        if n < 0:
            raise RegisterError(f"{where}: {k} must not be negative.")
        out[k] = n
    return out or None


def rm_eligibility(p: dict) -> str:
    """Why RM cannot be worked out for this part, or "" when it can."""
    if p["is_assembly"]:
        return "An assembly has no RM cost of its own; its parts do."
    src = register_setup.sourcing().get(p["sourcing"])
    if src and not src["is_costed"]:
        return f"{src['name']} parts have a bought-out price, not an RM cost."
    proc = register_setup.processes().get(p["process"])
    if not proc:
        return "Set the process first."
    if not proc["part_type"]:
        return f"RM for {proc['name']} parts (bar stock, bought-out) comes in a later stage."
    if not p["material"]:
        return "Set the material first."
    return ""


def _rm_family(p: dict) -> str:
    return register_setup.processes()[p["process"]]["part_type"]


def _formula_hash(family: str) -> str:
    setup = costing_setup.tables()
    scope = rm_cost.scope(family)
    rows = [r for k in ("cost_formulas", "formula_inputs") for r in setup[k] if r["scope"] == scope]
    return hashlib.sha1(json.dumps(rows, sort_keys=True, default=str).encode()).hexdigest()[:12]


def _rm_compute(part_id: int, country: str | None = None) -> dict:
    """Dry run: the traced RM result of the part's current revision. Raises RMNotApplicable / RegisterError / RMError."""
    from .geometry import blank_estimate
    d = part_detail(part_id, with_rm=False)
    why = rm_eligibility(d)
    if why:
        raise RMNotApplicable(why)
    family = _rm_family(d)
    try:
        t, nodes, warnings, d, rev = _geometry_trace(part_id)
    except RegisterError:
        raise RMNotApplicable("This part has no 3D file yet, so there is no volume to weigh. Upload a revision.") from None
    if "weight" not in t.nodes:
        raise RMNotApplicable(f"'{d['material']}' has no density in Master data → Materials. Add it there first.")
    g = rev["geometry"]
    cad = {"type": "cad", "file": rev["original_name"], "rev": rev["rev"], "sha256": rev["sha256"]}
    overrides = d["rm_inputs"] or {}
    part = {"net_weight_g": "weight", "density_g_cm3": "density"}
    if family == "sheet_metal":
        wall = g.get("wall_mm") or 0
        if wall <= 0 and "thickness_mm" not in overrides:
            raise RMNotApplicable("The sheet thickness could not be found in the 3D model. Type it under RM inputs.")
        blank = _developed_blank(rev, g, wall or overrides["thickness_mm"])
        part["thickness_mm"] = {"value": wall, "unit": "mm",
                                "source": cad | {"note": "Most common wall (sheet thickness) found in the 3D model"}}
        chosen = next((o for o in blank.get("options") or [] if o["name"] == blank.get("chosen")), None)
        how = blank["method"] + (f" ({chosen['yield_pct']:g} % material yield)" if chosen else "")
        part["blank_length_mm"] = {"value": blank["length"], "unit": "mm", "source": cad | {
            "note": f"Developed blank from the 3D model: {how}"}}
        part["blank_width_mm"] = {"value": blank["width"], "unit": "mm", "source": cad | {
            "note": f"Developed blank from the 3D model: {how}"}}
    country = (country or d["country"] or register_setup.settings()["default_country"]).upper()
    as_of = d["costing_month"] or rates.current_month()
    t.overrides = (d.get("trace_overrides") or {}).get("rm") or {}
    r = rm_cost.run(family, country=country, as_of=as_of, material=d["material"], part=part, given=overrides,
                    given_source={"type": "user", "note": f"Typed in the part register for {d['part_code']}"}, trace=t)
    r["blank"] = blank if family == "sheet_metal" else None
    r["key"] = {"blank": [blank["length"], blank["width"]] if family == "sheet_metal" else None,
                "overrides": t.overrides, "revision_id": rev["id"], "material": d["material"], "family": family, "country": country,
                "as_of": as_of, "inputs": overrides, "net_weight_g": round(t.value("weight"), 6),
                "rates": sorted(x["id"] for x in r["rates_used"]), "formulas": _formula_hash(family)}
    r["revision"] = {"id": rev["id"], "rev": rev["rev"]}
    r["process"] = d["process"]
    r["warnings"] = warnings
    return r


def _rm_row(row) -> dict:
    d = _plain({k: row[k] for k in RM_COLUMNS})
    d["is_example"] = any(x.get("is_example") for x in d["rates_used"] or [])
    return d


def rm_versions(part_id: int) -> list[dict]:
    with db.engine().connect() as conn:
        rows = conn.execute(select(db.rm_costs).where(db.rm_costs.c.part_id == part_id)
                            .order_by(db.rm_costs.c.version.desc())).mappings().all()
    return [_rm_row(r) for r in rows]


def _rate_changes(old_used: list, new_used: list) -> list[dict]:
    new = {(x["category"], x["item"], x["country"]): x for x in new_used}
    return [{"rate": f"{o['category']} / {o['item']}" + (f" · {o['country']}" if o["country"] else ""),
             "old_value": o["value"], "old_month": o["month"], "new_value": new[k]["value"], "new_month": new[k]["month"]}
            for o in old_used if (k := (o["category"], o["item"], o["country"])) in new and new[k]["id"] != o["id"]]


def _rates_used(r: dict) -> list[dict]:
    return [{"id": x["id"], "category": x["category"], "item": x["item"], "country": x["country"], "month": x["month"],
             "value": x["value"], "unit": x["unit"], "status": x["status"], "approved_by": x["approved_by"],
             "is_example": x["is_example"]} for x in r["rates_used"]]


def rm_summary(d: dict) -> dict:
    """What the part panel shows about RM: latest version, whether it is out of date, or why there is none."""
    versions = rm_versions(d["id"])
    latest = versions[0] if versions else None
    out = {"applicable": False, "problem": "", "outdated": "", "rate_changes": [], "latest": latest,
           "versions": len(versions), "overridable": [], "weights": None}
    why = rm_eligibility(d)
    if why:
        out["problem"] = why
        return out
    out["applicable"] = True
    out["overridable"] = rm_cost.overridable(_rm_family(d))
    try:
        r = _rm_compute(d["id"])
    except (RMNotApplicable, RegisterError, rm_cost.RMError) as exc:
        out["problem"] = str(exc)
        return out
    n = r["named"]
    out["weights"] = {k: n.get(k) for k in ("gross_g", "returns_g", "input_g", "material_yield_pct")}
    if r["missing"]:
        out["problem"] = " ".join(r["missing"])
    elif latest is None:
        out["outdated"] = "The RM cost has not been calculated yet."
    elif latest["key"] != r["key"]:
        strip = lambda k: {a: b for a, b in k.items() if a not in ("rates", "as_of")}
        changes = _rate_changes(latest["rates_used"] or [], _rates_used(r))
        if strip(latest["key"]) == strip(r["key"]) and (changes or latest["key"]["as_of"] != r["key"]["as_of"]):
            out["outdated"] = f"Newer approved rates exist than the ones used in version {latest['version']}."
            out["rate_changes"] = changes
        else:
            out["outdated"] = (f"Inputs, master data or formulas changed since version {latest['version']}.")
    return out


def compute_rm(part_id: int, reason: str = "calculated") -> dict:
    """Work out the RM cost and store it as a new version unless nothing changed."""
    try:
        r = _rm_compute(part_id)
    except RMNotApplicable as exc:
        raise RegisterError(str(exc)) from None
    except rm_cost.RMError as exc:
        raise RegisterError(str(exc)) from None
    if r["missing"]:
        raise RegisterError("RM cost not stored: " + " ".join(r["missing"]))
    n = r["named"]
    with db.engine().begin() as conn:
        last = conn.execute(select(db.rm_costs).where(db.rm_costs.c.part_id == part_id)
                            .order_by(db.rm_costs.c.version.desc()).limit(1)).mappings().first()
        if last and last["key"] == r["key"]:
            return {"stored": False, "version": last["version"], "message": f"Already up to date (version {last['version']})."}
        version = (last["version"] if last else 0) + 1
        rid = conn.execute(db.rm_costs.insert().values(
            part_id=part_id, version=version, revision_id=r["revision"]["id"], rev=r["revision"]["rev"],
            family=r["family"], process=r["process"], material=r["material"], country=r["country"],
            currency=r["currency"], as_of=r["as_of"], inputs=r["key"]["inputs"], key=r["key"],
            net_weight_g=r["key"]["net_weight_g"], gross_weight_g=n.get("gross_g"), net_rm_cost=r["value"],
            net_rm_cost_inr=r["value_inr"], rates_used=_rates_used(r), trace=r, reason=reason,
            created_by=user())).inserted_primary_key[0]
        conn.execute(db.parts.update().where(db.parts.c.id == part_id).values(rm_cost_id=rid, rm_cost_inr=r["value_inr"]))
        _log(conn, "part", part_id, "rm_cost_calculated", {
            "version": version, "net RM cost (INR)": round(r["value_inr"], 4), "material": r["material"],
            "country": r["country"], "costing month": r["as_of"], "reason": reason,
            "rate versions": ", ".join(f"#{x['id']}" for x in r["rates_used"])})
    return {"stored": True, "version": version, "message": f"RM cost version {version} stored."}


def refresh_rm(part_ids, reason: str) -> None:
    """Recalculate after a change. A part that cannot be costed yet is skipped; it never blocks the change itself."""
    for pid in dict.fromkeys(part_ids):
        try:
            compute_rm(pid, reason)
        except (RegisterError, rm_cost.RMError, RMNotApplicable):
            pass


def recalculate_project(project_id: int, reason: str = "recalculated with the latest approved rates") -> dict:
    with db.engine().connect() as conn:
        ids = [r for (r,) in conn.execute(select(db.parts.c.id).where(and_(db.parts.c.project_id == project_id,
                                                                         db.parts.c.is_assembly.is_(False))))]
    done = same = 0
    skipped = []
    for pid in ids:
        try:
            res = compute_rm(pid, reason)
        except RegisterError as exc:
            if rm_eligibility(part_detail(pid, with_rm=False)) == "":
                skipped.append({"part_id": pid, "why": str(exc)})
            continue
        done += res["stored"]
        same += not res["stored"]
    return {"new_versions": done, "unchanged": same, "problems": skipped}


def set_rm_inputs(part_id: int, rm_inputs: dict | None, country: str | None = None, costing_month: str | None = None,
                  reason: str = "RM inputs saved") -> dict:
    """Save the typed inputs / country / month of a part and recalculate (e.g. from the Estimate page)."""
    data = {"rm_inputs": rm_inputs or {}}
    if country is not None:
        data["country"] = country
    if costing_month is not None:
        data["costing_month"] = costing_month
    update_part(part_id, data, refresh=False)
    try:
        res = compute_rm(part_id, reason)
    except RegisterError as exc:
        res = {"stored": False, "version": None, "message": str(exc)}
    return res | {"part": part_detail(part_id)}


def rm_trace(part_id: int, version: int | None = None) -> dict:
    """The stored trace of a version, or (no version given and none stored) a live dry run."""
    if version is not None or rm_versions(part_id):
        with db.engine().connect() as conn:
            q = select(db.rm_costs).where(db.rm_costs.c.part_id == part_id)
            q = q.where(db.rm_costs.c.version == version) if version is not None else q.order_by(db.rm_costs.c.version.desc())
            row = conn.execute(q.limit(1)).mappings().first()
        if not row:
            raise RegisterError(f"Part {part_id} has no RM cost version {version}.")
        return row["trace"] | {"version": row["version"], "stored": True, "created_by": row["created_by"],
                               "created_at": row["created_at"].isoformat(timespec="seconds")}
    try:
        return _rm_compute(part_id) | {"version": None, "stored": False}
    except (RMNotApplicable, rm_cost.RMError) as exc:
        raise RegisterError(str(exc)) from None


def _node_source_text(node: dict) -> str:
    from .trace import source_text
    return source_text(node)


def export_rm_excel(part_id: int, version: int | None = None) -> bytes:
    d = part_detail(part_id, with_rm=False)
    tr = rm_trace(part_id, version)
    nodes = tr["nodes"]
    example = any(x.get("is_example") for x in tr["rates_used"])
    wb = Workbook()
    wb.properties.title = f"Design to Cost — RM cost {d['part_code']}"
    wb.properties.creator = "Design to Cost"

    def sheet(title, head, rows, first=False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = title
        if example:
            _excel_banner(ws, len(head))
        ws.append(head)
        top = ws.max_row
        for c in ws[top]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = HEADER_FILL
        for r in rows:
            ws.append(r)
            if len(r) and r[-1] == "EXAMPLE":
                for c in ws[ws.max_row]:
                    c.fill = EXAMPLE_FILL
        for i, h in enumerate(head, start=1):
            ws.column_dimensions[chr(64 + i)].width = min(max([len(str(h))] + [len(str(r[i - 1])) for r in rows]) + 2, 70)
        ws.freeze_panes = f"A{top + 1}"

    sheet("Summary", ["Item", "Value"], [
        ["Part code", d["part_code"]], ["Part name", d["name"]], ["Revision", tr.get("revision", {}).get("rev", "")],
        ["Material", tr["material"]], ["Process", tr.get("process", d["process"])], ["Producing country", tr["country"]],
        ["Costing month", tr["as_of"]], ["RM cost version", tr["version"] if tr["version"] else "not stored (live preview)"],
        [f"Net RM cost per part ({tr['currency']})", tr["value"]], ["Net RM cost per part (INR)", tr["value_inr"]],
        ["Rate versions used", ", ".join(f"#{x['id']}" for x in tr["rates_used"])],
        ["Includes EXAMPLE values", f"yes - {EXAMPLE_WARNING}. Replace them on the Rates page" if example else "no"],
        ["Calculated by / when", f"{tr.get('created_by', '')} {tr.get('created_at', '')}".strip()]], first=True)
    sheet("Steps", ["Step", "Formula", "Value", "Unit"],
          [[n["label"], n["formula"], n["value"], n["unit"]] for n in nodes.values() if n["kind"] == "calculated"])
    sheet("Inputs", ["Input", "Value", "Unit", "Comes from", ""],
          [[n["label"], n["value"], n["unit"], _node_source_text(n),
            "EXAMPLE" if (n.get("source") or {}).get("rate", {}).get("is_example") else ""]
           for n in nodes.values() if n["kind"] == "input"])
    sheet("Rates used", ["Version id", "Category", "Item", "Country", "Month", "Value", "Unit", "Status", "Approved by", ""],
          [[x["id"], x["category"], x["item"], x["country"], x["month"], x["value"], x["unit"], x["status"],
            x["approved_by"], "EXAMPLE" if x["is_example"] else ""] for x in tr["rates_used"]])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_rm_pdf(part_id: int, version: int | None = None) -> bytes:
    from .pdf_report import rm_cost_pdf
    return rm_cost_pdf(part_detail(part_id, with_rm=False), rm_trace(part_id, version))


def refresh_costs(part_ids, reason: str) -> None:
    """RM and tool cost after a change; a part that cannot be costed yet is skipped."""
    ids = list(dict.fromkeys(part_ids))
    refresh_rm(ids, reason)
    refresh_tool(ids, reason)


# ---------- tool cost versions (stage 4) ----------
# The tool cost engine (tool_cost.py) prices the mould / die of a part from the BOM rules, tool processes and
# formulas in master data with the approved rates of the tool country. Here its result is stored as a numbered,
# never-overwritten version, exactly like the RM cost: recalculated when the revision, process, material, tool
# country, month, tool features, volume or programme life change; newer approved rates never change a stored cost.

class ToolNotApplicable(Exception):
    """No tool cost for this part (the message says why and what to do)."""


TOOL_COLUMNS = ("id", "part_id", "version", "revision_id", "rev", "part_type", "country", "currency", "as_of", "inputs",
                "key", "cavities", "tools_count", "tool_cost", "tool_cost_inr", "tool_capex_inr", "rates_used", "example",
                "reason", "created_by", "created_at")
TOOL_INPUTS = {"cavities": "Cavities per tool", "hot_runner": "Hot runner", "slides_per_cavity": "Slides per cavity",
               "lifters_per_cavity": "Lifters per cavity", "inserts_per_cavity": "Small inserts per cavity",
               "warranty": "Warranty / spares included"}
TOOL_FLAGS = ("hot_runner", "warranty")


def _sheet_inputs(v, where):
    """Sheet-metal tool choices typed per part: {mode, operations: [codes], draw_depth_mm, bend_radius_mm}."""
    if v in (None, "", {}):
        return None
    out = {}
    if v.get("mode") not in (None, ""):
        if v["mode"] not in ("progressive", "multi_stage"):
            raise RegisterError(f"{where}: die type must be progressive or multi_stage.")
        out["mode"] = v["mode"]
    if v.get("operations") is not None:
        codes = {r["code"] for r in mould_db.tables()["sheet_operations"]}
        bad = [c for c in v["operations"] if c not in codes]
        if bad:
            raise RegisterError(f"{where}: operation '{bad[0]}' is not in Sheet operations ({', '.join(sorted(codes))}).")
        out["operations"] = list(v["operations"])
    for k in ("draw_depth_mm", "bend_radius_mm"):
        if v.get(k) not in (None, ""):
            n = _num(v[k], f"{where}: {k}", required=True)
            if n < 0:
                raise RegisterError(f"{where}: {k} must not be negative.")
            out[k] = n
    return out or None


def _developed_blank(rev: dict, g: dict, thickness: float, shape=None) -> dict:
    """The developed blank of a sheet-metal revision (best material yield, geometry.developed_blank), worked out
    once from the 3D file and kept in the revision's geometry; the old estimate if the file is not there."""
    from .geometry import blank_estimate, developed_blank, load_step
    key = f"v2:{thickness:g}"   # v2: main bend direction (stage 25 fix); older cached blanks are worked out again
    cached = (g.get("dev_blank") or {}).get(key)
    if cached:
        return cached
    try:
        if shape is None:
            with db.engine().connect() as conn:
                row = conn.execute(select(db.files).where(db.files.c.id == rev["file_id"])).mappings().first()
            path = file_store.path_of(dict(row)) if row else None
            if not path or not path.exists():
                raise FileNotFoundError
            shape = load_step(str(path))
        blank = developed_blank(shape, g["volume_mm3"], g["bbox_mm"], thickness)
    except Exception:  # no file: the estimate from volume and size
        b = blank_estimate(g["volume_mm3"], g["bbox_mm"], thickness)
        return b | {"method": "estimate from volume and size (3D file not found)", "options": [], "bends": g.get("bend_count") or 0}
    geo = dict(g)
    geo["dev_blank"] = dict(g.get("dev_blank") or {}) | {key: blank}
    with db.engine().begin() as conn:
        conn.execute(db.revisions.update().where(db.revisions.c.id == rev["id"]).values(geometry=geo))
    return blank


def _revision_shape(part_id: int):
    """(part detail, current revision, 3D shape) of a part; RegisterError if there is no 3D file."""
    from .geometry import load_step
    d = part_detail(part_id, with_rm=False)
    rev = d["revisions"][0] if d["revisions"] else None
    if not rev or not rev["file_id"]:
        raise RegisterError("This part has no 3D file yet. Upload a revision first.")
    with db.engine().connect() as conn:
        row = conn.execute(select(db.files).where(db.files.c.id == rev["file_id"])).mappings().first()
    path = file_store.path_of(dict(row))
    if not path.exists():
        raise RegisterError(f"The 3D file is missing from the storage folder ({row['stored_path']}).")
    return d, rev, load_step(str(path))


def revision_file_bytes(rev: dict) -> bytes:
    with db.engine().connect() as conn:
        row = conn.execute(select(db.files).where(db.files.c.id == rev["file_id"])).mappings().first()
    return file_store.path_of(dict(row)).read_bytes()


def part_drawing(part_id: int, kind: str = "pdf") -> bytes:
    """Part drawing of the current revision (roadmap stage 26): PDF (A3, three views, title block) or DXF."""
    from . import drawings
    d, rev, shape = _revision_shape(part_id)
    if kind == "dxf":
        return drawings.part_drawing_dxf(d, rev, shape)
    info = {}
    try:
        t = _geometry_trace(part_id)[0]
        if "weight" in t.nodes:
            info["weight"] = f"{t.value('weight'):,.1f} g"
    except Exception:
        pass
    proc = register_setup.processes().get(d["process"]) or {}
    if proc.get("part_type") == "sheet_metal":
        g = rev["geometry"]
        try:
            b = _developed_blank(rev, g, g.get("wall_mm") or 1.0, shape)
            info["blank"] = f"{b['length']:.1f} × {b['width']:.1f} mm ({b['method']})"
        except Exception:
            pass
    return drawings.part_drawing_pdf(d | {"process": proc.get("name") or d["process"]}, rev, shape, info)


def tool_drawing(part_id: int, kind: str = "pdf", die: str | None = None) -> bytes:
    """Concept drawing of the mould or of one die stage, from the stored (else live) tool cost."""
    from . import drawings
    d = part_detail(part_id, with_rm=False)
    tr = tool_trace(part_id)
    return (drawings.tool_drawing_dxf if kind == "dxf" else drawings.tool_drawing_pdf)(d, tr, die)


def _sheet_context(part_id: int):
    """(part detail, revision, 3D shape, blank, sheet-tools result) of a sheet-metal part's current revision."""
    from . import sheet_tools
    from .geometry import blank_estimate, load_step
    d = part_detail(part_id, with_rm=False)
    proc = register_setup.processes().get(d["process"])
    if d["is_assembly"] or not proc or proc["part_type"] != "sheet_metal":
        raise ToolNotApplicable("Sheet-metal tools are for parts whose process is sheet metal.")
    rev = d["revisions"][0] if d["revisions"] else None
    if not rev or not rev["file_id"]:
        raise ToolNotApplicable("This part has no 3D file yet. Upload a revision first.")
    with db.engine().connect() as conn:
        row = conn.execute(select(db.files).where(db.files.c.id == rev["file_id"])).mappings().first()
    path = file_store.path_of(dict(row))
    if not path.exists():
        raise ToolNotApplicable(f"The 3D file is missing from the storage folder ({row['stored_path']}).")
    density, _ = _density(d["material"])
    if not density:
        raise ToolNotApplicable("Set the part's material (with a density in Master data → Materials).")
    g, typed = rev["geometry"], d["sheet_inputs"] or {}
    volume, _ = _annual_volume(d)
    shape = load_step(str(path))
    blank = _developed_blank(rev, g, g.get("wall_mm") or 1.0, shape)
    feat = sheet_tools.features(shape, g, blank, typed, density, volume, d["programme_life_years"])
    try:
        by_process = {"sheet_metal_progressive": "progressive", "sheet_metal_multistage": "multi_stage"}.get(d["process"], "")
        res = sheet_tools.run(feat, d["material"], typed, by_process)
    except ValueError as exc:
        raise ToolNotApplicable(str(exc)) from None
    res = res | {"part_id": part_id, "rev": rev["rev"], "typed": typed, "process": d["process"], "blank": blank}
    return d, rev, shape, blank, res


def sheet_tools_result(part_id: int) -> dict:
    """Sheet-metal tools I for the part's current revision (roadmap stage 9)."""
    return _sheet_context(part_id)[4]


def _layout(part_id: int):
    from . import sheet_layout
    d, rev, shape, blank, res = _sheet_context(part_id)
    t = res["features"]["t"]
    fp = sheet_layout.flat_pattern(shape, rev["geometry"], blank, t)
    sl = sheet_layout.strip_layout(fp, t, len(res["operations"]))
    return d, rev, fp, sl, res


def sheet_layout_result(part_id: int) -> dict:
    """Flat pattern and progressive strip layout (roadmap stage 10)."""
    d, rev, fp, sl, res = _layout(part_id)
    return {"part_id": part_id, "rev": rev["rev"], "flat_pattern": fp, "strip": sl, "mode": res["mode"]}


def sheet_layout_file(part_id: int, kind: str) -> tuple[bytes, str]:
    from . import sheet_layout
    d, rev, fp, sl, res = _layout(part_id)
    if kind == "blank.dxf":
        return sheet_layout.blank_dxf(fp, d["part_code"]), "application/dxf"
    if kind == "strip.dxf":
        return sheet_layout.strip_dxf(fp, sl, d["part_code"]), "application/dxf"
    return sheet_layout.pack_pdf(fp, sl, res, d, rev["rev"]), "application/pdf"


def _piece_inputs(v, where):
    """Part-cost inputs typed per part: {cycle_s, cavities, secondary_ops_pct}."""
    from .piece_cost import PIECE_TYPED
    if v in (None, "", {}):
        return None
    out = {}
    for k, val in v.items():
        if val in (None, ""):
            continue
        if k not in PIECE_TYPED:
            raise RegisterError(f"{where}: '{k}' is not a part-cost input ({', '.join(PIECE_TYPED)}).")
        n = _num(val, f"{where}: {PIECE_TYPED[k].lower()}", required=True)
        if n < 0 or (k != "secondary_ops_pct" and n <= 0):
            raise RegisterError(f"{where}: {PIECE_TYPED[k].lower()} must be more than 0.")
        out[k] = n
    return out or None


def piece_versions(part_id: int) -> list[dict]:
    with db.engine().connect() as conn:
        rows = conn.execute(select(*[c for c in db.piece_costs.c if c.name != "trace"])
                            .where(db.piece_costs.c.part_id == part_id).order_by(db.piece_costs.c.version.desc())).mappings().all()
    return [_plain(r) for r in rows]


def piece_live(part_id: int) -> dict:
    from . import piece_cost
    try:
        return piece_cost.run(part_id)
    except piece_cost.PieceError as exc:
        raise RegisterError(str(exc)) from None


def piece_summary(part_id: int) -> dict:
    """Latest stored part cost, whether it is out of date, and the live result (or why there is none)."""
    versions = piece_versions(part_id)
    latest = versions[0] if versions else None
    out = {"latest": latest, "versions": versions, "problem": "", "outdated": "", "live": None}
    try:
        r = piece_live(part_id)
    except RegisterError as exc:
        out["problem"] = str(exc)
        return out
    out["live"] = r
    if latest is None:
        out["outdated"] = "The part cost has not been stored yet."
    elif latest["key"] != r["key"]:
        same = {k: v for k, v in latest["key"].items() if k not in ("rates", "as_of")} == {k: v for k, v in r["key"].items() if k not in ("rates", "as_of")}
        out["outdated"] = (f"Newer approved rates exist than the ones used in part cost version {latest['version']}." if same
                           else f"RM / tool cost, inputs or revision changed since part cost version {latest['version']}.")
    return out


def compute_piece(part_id: int, reason: str = "calculated") -> dict:
    """Store the part cost as a new version unless nothing changed. Stored versions are never changed."""
    r = piece_live(part_id)
    if r["missing"]:
        raise RegisterError("Part cost not stored: " + " ".join(r["missing"]))
    with db.engine().begin() as conn:
        last = conn.execute(select(db.piece_costs).where(db.piece_costs.c.part_id == part_id)
                            .order_by(db.piece_costs.c.version.desc()).limit(1)).mappings().first()
        if last and last["key"] == r["key"]:
            return {"stored": False, "version": last["version"], "message": f"Part cost already up to date (version {last['version']})."}
        version = (last["version"] if last else 0) + 1
        conn.execute(db.piece_costs.insert().values(
            part_id=part_id, version=version, revision_id=r["revision"]["id"], rev=r["revision"]["rev"], country=r["country"],
            currency=r["currency"], as_of=r["as_of"], key=r["key"], piece_cost=r["value"], piece_cost_inr=r["value_inr"],
            rates_used=_rates_used(r), trace=r, reason=reason, created_by=user()))
        conn.execute(db.parts.update().where(db.parts.c.id == part_id).values(piece_cost_inr=r["value_inr"]))
        _log(conn, "part", part_id, "piece_cost_calculated", {"version": version, "part cost (INR)": round(r["value_inr"], 4),
             "machine": r["machine"], "country": r["country"], "costing month": r["as_of"], "reason": reason,
             "rate versions": ", ".join(f"#{x['id']}" for x in r["rates_used"])})
    return {"stored": True, "version": version, "message": f"Part cost version {version} stored."}


def set_piece_inputs(part_id: int, inputs: dict | None, reason: str = "part-cost inputs saved") -> dict:
    update_part(part_id, {"piece_inputs": inputs or {}}, refresh=False)
    return compute_piece(part_id, reason)


def piece_trace(part_id: int, version: int | None = None) -> dict:
    """A stored part cost version's trace, else the live result."""
    if version is None and not piece_versions(part_id):
        return piece_live(part_id) | {"version": None, "stored": False}
    with db.engine().connect() as conn:
        q = select(db.piece_costs).where(db.piece_costs.c.part_id == part_id)
        q = q.where(db.piece_costs.c.version == version) if version is not None else q.order_by(db.piece_costs.c.version.desc())
        row = conn.execute(q.limit(1)).mappings().first()
    if not row:
        raise RegisterError(f"Part {part_id} has no part cost version {version}.")
    return row["trace"] | {"version": row["version"], "stored": True, "created_by": row["created_by"],
                           "created_at": row["created_at"].isoformat(timespec="seconds")}


def _trace_rows(tr: dict) -> tuple[list, list, list]:
    """(steps, inputs, rates) rows of a trace for Excel / PDF."""
    from .trace import source_text
    n = tr["nodes"]
    steps = [[n[s["node"]]["label"], n[s["node"]].get("formula", ""), n[s["node"]]["value"], n[s["node"]]["unit"]] for s in tr["steps"]]
    inputs = [[n[i["node"]]["label"], n[i["node"]]["value"], n[i["node"]]["unit"], source_text(n[i["node"]])] for i in tr["inputs"]]
    rts = [[x["id"], x["category"], x["item"], x["country"], x["month"], x["value"], x["unit"], x["status"],
            "EXAMPLE" if x.get("is_example") else ""] for x in tr["rates_used"]]
    return steps, inputs, rts


def piece_excel(part_id: int, version: int | None = None) -> bytes:
    """Part cost sheet: Summary, Steps, Inputs (with sources), Rates used. Full backup data."""
    tr, d = piece_trace(part_id, version), part_detail(part_id, with_rm=False)
    steps, inputs, rts = _trace_rows(tr)
    wb = Workbook()
    wb.properties.title = f"Design to Cost — Part cost {d['part_code']}"
    for i, (name, head, rows) in enumerate((
            ("Summary", ["Item", "Value"], [["Part", f"{d['part_code']} · {d['name']}"], ["Revision", tr["revision"]["rev"]],
                                             ["Country", tr["country"]], ["Costing month", tr["as_of"]], ["Machine", tr["machine"]],
                                             [f"Part cost ({tr['currency']})", tr["value"]], ["Part cost (INR)", tr["value_inr"]],
                                             ["Version", tr.get("version") or "live, not stored"],
                                             ["Example rates", "yes: not for negotiation" if tr["is_example"] else "no"]]),
            ("Steps", ["Step", "Formula", "Value", "Unit"], steps),
            ("Inputs", ["Input", "Value", "Unit", "Source"], inputs),
            ("Rates used", ["Rate #", "Category", "Item", "Country", "Month", "Value", "Unit", "Status", "Example"], rts))):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = name
        ws.append(head)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = HEADER_FILL
        for r in rows:
            ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def piece_pdf(part_id: int, version: int | None = None) -> bytes:
    from .pdf_report import piece_cost_pdf
    tr, d = piece_trace(part_id, version), part_detail(part_id, with_rm=False)
    return piece_cost_pdf(tr, d, *_trace_rows(tr))


def _tool_inputs(v, where):
    """Tool features typed per part: {cavities, hot_runner, slides_per_cavity, ..., warranty}."""
    if v in (None, "", {}):
        return None
    if isinstance(v, str):
        try:
            v = json.loads(v)
        except ValueError:
            raise RegisterError(f"{where}: tool inputs are not readable.") from None
    out = {}
    for k, val in v.items():
        if val in (None, "", {}):
            continue
        if k == "bom":
            bom = _bom_overrides(val, where)
            if bom:
                out["bom"] = bom
            continue
        if k not in TOOL_INPUTS:
            raise RegisterError(f"{where}: '{k}' is not a tool input ({', '.join(TOOL_INPUTS)}).")
        if k in TOOL_FLAGS:
            out[k] = val is True or str(val).strip().lower() in ("1", "true", "yes", "on")
            continue
        n = _num(val, f"{where}: {TOOL_INPUTS[k].lower()}", int, required=True)
        if n < (1 if k == "cavities" else 0) or n > 256:
            raise RegisterError(f"{where}: {TOOL_INPUTS[k].lower()} must be between {1 if k == 'cavities' else 0} and 256.")
        out[k] = n
    return out or None


BOM_FIELDS = ("grade", "make", "qty", "price_each", "std_part", "base")


def _bom_overrides(v, where) -> dict:
    """{line key: {grade, make, qty, price_each (INR), std_part}}: your changes to the tool BOM."""
    if not isinstance(v, dict):
        raise RegisterError(f"{where}: tool BOM changes are not readable.")
    out = {}
    for key, line in v.items():
        if not isinstance(line, dict):
            raise RegisterError(f"{where}: BOM line '{key}' is not readable.")
        clean = {}
        for f, val in line.items():
            if val in (None, ""):
                continue
            if f not in BOM_FIELDS:
                raise RegisterError(f"{where}: '{f}' cannot be changed on a BOM line ({', '.join(BOM_FIELDS)}).")
            if f in ("qty", "price_each"):
                n = _num(val, f"{where}: {key} {f.replace('_', ' ')}")
                clean[f] = n
            else:
                clean[f] = str(val).strip()
        if clean:
            out[str(key)] = clean
    return out


def tool_eligibility(p: dict) -> str:
    """Why there is no tool cost for this part, or "" when there is."""
    if p["is_assembly"]:
        return "An assembly has no tool of its own; its parts do."
    src = register_setup.sourcing().get(p["sourcing"])
    if src and not src["is_costed"]:
        return f"{src['name']} parts are bought ready-made, so there is no tool to cost here."
    proc = register_setup.processes().get(p["process"])
    if not proc:
        return "Set the process first."
    if proc["part_type"] not in tool_cost.PART_TYPES + ("sheet_metal",):
        return tool_cost.LATER.get(proc["part_type"], f"Tool cost for {proc['name']} parts comes in a later stage.")
    return ""


def _annual_volume(part: dict) -> tuple[float | None, str]:
    """The part's own annual volume, else Σ variant annual volume × quantity of the part in that variant."""
    if part["annual_volume"]:
        return part["annual_volume"], "annual volume of the part"
    total, used = 0.0, []
    for v in get_project(part["project_id"])["variants"]:
        if not v["annual_volume"]:
            continue
        qty = sum(b["total_quantity"] for b in variant_bom(v["id"]) if b["child_part_id"] == part["id"] and b["kind"] == "bom")
        if qty:
            total += v["annual_volume"] * qty
            used.append(f"{v['code']} {v['annual_volume']:,} × {qty:g}")
    if used:
        return total, "Σ variant annual volume × quantity (" + ", ".join(used) + ")"
    return None, "master data default (no annual volume on the part or its variants)"


def _tool_master_hash(part_type: str) -> str:
    """Fingerprint of every master-data row the tool cost depends on, to tell when it changed."""
    setup = costing_setup.tables()
    mdb = mould_db.tables()
    data = {
        "formulas": [r for k in ("cost_formulas", "formula_inputs") for r in setup[k]
                     if r["scope"] in costing_setup.TOOL_SCOPES + ("mhr", "labour")],
        "machines": setup["machines"], "grades": setup["labour_grades"],
        "mould": {k: [r for r in mdb[k] if r.get("part_type", part_type) == part_type]
                  for k in ("mould_bases", "bom_elements", "sizing_rules", "time_parameters", "tool_processes", "tool_materials",
                            "grade_rules", "standard_parts", "sheet_operations")},
        "parameters": master_data.parameters(part_type),
    }
    return hashlib.sha1(json.dumps(data, sort_keys=True, default=str).encode()).hexdigest()[:12]


def _tool_compute(part_id: int, country: str | None = None) -> dict:
    """Dry run: the traced tool cost of the part's current revision. Raises ToolNotApplicable / ToolCostError."""
    d = part_detail(part_id, with_rm=False)
    why = tool_eligibility(d)
    if why:
        raise ToolNotApplicable(why)
    part_type = register_setup.processes()[d["process"]]["part_type"]
    try:
        t, _, _, d, rev = _geometry_trace(part_id)
    except RegisterError:
        raise ToolNotApplicable("This part has no 3D file yet, so the mould cannot be sized. Upload a revision.") from None
    if "weight" not in t.nodes:
        raise ToolNotApplicable("Set the part's material (with a density in Master data → Materials) to cost its tool.")
    g = rev["geometry"]
    pull = g["pull_direction"]
    across = sorted((v for a, v in g["bbox_axes_mm"].items() if a != pull), reverse=True)
    cad = {"type": "cad", "file": rev["original_name"], "rev": rev["rev"], "sha256": rev["sha256"]}
    volume, volume_note = _annual_volume(d)
    cap_in = {"part_type": part_type, "footprint_length_mm": across[0], "footprint_width_mm": across[1]}
    if volume:
        cap_in["annual_volume"] = max(1, round(volume))
    if d["programme_life_years"]:
        cap_in["programme_years"] = max(1, round(d["programme_life_years"]))
    if part_type == "plastic" and g.get("wall_mm"):
        cap_in["cycle_time_s"] = capacity.plastic_cycle_estimate(g["wall_mm"])
    try:
        cap = capacity.calculate(capacity.CapacityInput(**cap_in))
    except ValueError as exc:
        raise ToolNotApplicable(f"Stage C (cavities and tools) cannot be worked out: {exc}") from None
    typed = d["tool_inputs"] or {}
    country = (country or d["tool_country"] or d["country"] or register_setup.settings()["default_country"]).upper()
    as_of = d["costing_month"] or rates.current_month()
    learned = _learned_for(part_id, part_type, d["material"])
    if part_type == "sheet_metal":
        return _die_compute(part_id, d, rev, t, g, cap, volume, volume_note, typed, country, as_of, learned)
    inp = tool_cost.ToolInput(
        annual_volume=cap["inputs"]["annual_volume"], programme_years=cap["inputs"]["programme_years"],
        bom=typed.get("bom") or {}, learned=learned, overrides=(d.get("trace_overrides") or {}).get("tool") or {},
        part_type=part_type, cavities=typed.get("cavities") or cap["cavities_per_tool"],
        part_length_mm=across[0], part_width_mm=across[1], part_height_mm=g["bbox_axes_mm"][pull],
        projected_area_cm2=g["projected_area_mm2"][pull] / 100, part_weight_g=t.value("weight"),
        part_volume_cm3=g["volume_cm3"], part_surface_cm2=g["surface_area_mm2"] / 100,
        hot_runner=bool(typed.get("hot_runner")) and part_type == "plastic",
        slides_per_cavity=typed.get("slides_per_cavity", 0),
        lifters_per_cavity=typed.get("lifters_per_cavity", 0) if part_type == "plastic" else 0,
        inserts_per_cavity=typed.get("inserts_per_cavity", 0), warranty=typed.get("warranty", True),
        tools_count=cap["tool_sets_over_programme"], country=country, as_of=as_of, material=d["material"])
    stage_c = (f"Stage C for {cap['inputs']['annual_volume']:,} parts/yr ({volume_note}), cycle "
               f"{cap['inputs']['cycle_time_s']:g} s, OEE {cap['inputs']['oee_pct']:g} %, "
               f"{cap['inputs']['programme_years']} programme years")
    you = {"type": "user", "note": f"Typed in the part register for {d['part_code']}"}
    sources = {k: cad | {"note": n} for k, n in (
        ("part_l", f"Bounding box across the parting plane (opening direction {pull.upper()})"),
        ("part_w", f"Bounding box across the parting plane (opening direction {pull.upper()})"),
        ("part_h", f"Bounding box along the opening direction {pull.upper()}"),
        ("projected_area_cm2", f"Shadow area along {pull.upper()}, through-holes removed"),
        ("part_volume_cm3", "Measured from the STEP file"), ("part_surface_cm2", "Measured from the STEP file"))}
    sources["part_weight_g"] = {"type": "register", "note": "Net (body) weight = CAD volume × material density (Materials)"}
    sources["cavities"] = you if typed.get("cavities") else {
        "type": "register", "note": f"{stage_c}: {cap['cavities_needed']:g} needed → {cap['cavities_per_tool']} per tool"}
    sources["tools_count"] = {"type": "register", "note": f"{stage_c}: max(tools for capacity {cap['tools_for_capacity']}, "
                                                          f"tools for tool life {cap['tools_for_life']})"}
    for k in ("hot_runner", "slides_per_cavity", "lifters_per_cavity", "inserts_per_cavity"):
        sources[k] = you if k in typed else {"type": "register", "note": "Default (none); set it under Tool features"}
    sources["warranty_on"] = you if "warranty" in typed else {"type": "register", "note": "Default: included"}
    r = tool_cost.run(inp, sources)
    r["key"] = {"revision_id": rev["id"], "part_type": part_type, "material": d["material"], "country": country,
                "as_of": as_of, "inputs": typed, "cavities": inp.cavities, "tools_count": inp.tools_count,
                "weight_g": round(t.value("weight"), 6), "rates": sorted(x["id"] for x in r["rates_used"]),
                "master": _tool_master_hash(part_type), "learned": learned, "overrides": inp.overrides}
    r["revision"] = {"id": rev["id"], "rev": rev["rev"]}
    r["process"] = d["process"]
    r["stage_c"] = {k: cap[k] for k in ("cavities_needed", "cavities_per_tool", "tools_for_capacity", "tools_for_life",
                                        "tool_sets_over_programme", "limited_by")} | {"note": stage_c}
    return r


def _die_compute(part_id, d, rev, t, g, cap, volume, volume_note, typed, country, as_of, learned) -> dict:
    """Sheet-metal dies, stage-wise (roadmap stage 23): the dies of the operations in order, each with its BOM and cost."""
    from . import die_cost
    try:
        sheet = sheet_tools_result(part_id)
    except (ToolNotApplicable, ValueError) as exc:
        raise ToolNotApplicable(f"Sheet-metal tools: {exc}") from None
    try:
        layout = sheet_layout_result(part_id) if sheet["mode"] == "progressive" else None
    except (ToolNotApplicable, ValueError):
        layout = None
    f = sheet["features"]
    inp = tool_cost.ToolInput(
        part_type="sheet_metal", cavities=1, part_length_mm=max(f["part_l"], 1), part_width_mm=max(f["part_w"], 1),
        part_height_mm=max(f["part_h"], 0.1), part_weight_g=t.value("weight"), part_volume_cm3=g["volume_cm3"],
        part_surface_cm2=g["surface_area_mm2"] / 100, warranty=typed.get("warranty", True),
        tools_count=cap["tool_sets_over_programme"], country=country, as_of=as_of, material=d["material"],
        annual_volume=cap["inputs"]["annual_volume"], programme_years=cap["inputs"]["programme_years"],
        bom=typed.get("bom") or {}, learned=learned, overrides=(d.get("trace_overrides") or {}).get("tool") or {})
    stage_c = (f"Stage C for {cap['inputs']['annual_volume']:,} parts/yr ({volume_note}), "
               f"{cap['inputs']['programme_years']} programme years")
    cad = {"type": "cad", "file": rev["original_name"], "rev": rev["rev"], "sha256": rev["sha256"],
           "note": "Sheet-metal tools (stage 9) from the 3D model"}
    sources = {k: cad for k in die_cost.FEATURES}
    sources["part_weight_g"] = {"type": "register", "note": "Net weight = CAD volume × material density (Materials)"}
    sources["tools_count"] = {"type": "register", "note": f"{stage_c}: max(tools for capacity {cap['tools_for_capacity']}, "
                                                          f"tools for tool life {cap['tools_for_life']})"}
    sources["annual_volume"] = sources["programme_years"] = {"type": "register", "note": stage_c}
    sources["warranty_on"] = {"type": "user", "note": "Tool features"} if "warranty" in typed else {"type": "register", "note": "Default: included"}
    try:
        r = die_cost.run(inp, sheet, layout, sources)
    except ValueError as exc:
        raise tool_cost.ToolCostError(str(exc)) from None
    r["key"] = {"revision_id": rev["id"], "part_type": "sheet_metal", "material": d["material"], "country": country,
                "as_of": as_of, "inputs": typed, "sheet_inputs": d["sheet_inputs"] or {}, "mode": sheet["mode"],
                "dies": [x["key"] for x in r["dies"]], "blank": [sheet["blank"]["length"], sheet["blank"]["width"]],
                "tools_count": inp.tools_count, "weight_g": round(t.value("weight"), 6),
                "rates": sorted(x["id"] for x in r["rates_used"]), "master": _tool_master_hash("sheet_metal"), "learned": learned,
                "overrides": inp.overrides}
    r["revision"] = {"id": rev["id"], "rev": rev["rev"]}
    r["process"] = d["process"]
    r["stage_c"] = {k: cap[k] for k in ("cavities_needed", "cavities_per_tool", "tools_for_capacity", "tools_for_life",
                                        "tool_sets_over_programme", "limited_by")} | {"note": stage_c}
    return r


# ---------- tool BOM: your changes, learning, new grades and standard parts (roadmap stage 23) ----------

def _learned_for(part_id: int, part_type: str, material: str) -> dict:
    """What you chose on earlier parts, per BOM element (latest; the same part material first).
    A part that already has a tool cost keeps the choices it was costed with, so its cost does not move."""
    with db.engine().connect() as conn:
        last = conn.execute(select(db.tool_costs.c.key).where(db.tool_costs.c.part_id == part_id)
                            .order_by(db.tool_costs.c.version.desc()).limit(1)).scalar()
        if last is not None:
            return (last or {}).get("learned") or {}
        rows = conn.execute(select(db.bom_learning).where((db.bom_learning.c.part_type == part_type)
                                                          & (db.bom_learning.c.part_id != part_id))
                            .order_by(db.bom_learning.c.at.desc(), db.bom_learning.c.id.desc())).mappings().all()
    out = {}
    for same_material in (True, False):
        for r in rows:
            if r["element"] in out or (same_material and (r["part_material"] or "") != (material or "")):
                continue
            n = sum(1 for x in rows if x["element"] == r["element"] and x["grade"] == r["grade"] and x["make"] == r["make"]
                    and x["std_part"] == r["std_part"])
            out[r["element"]] = {"grade": r["grade"], "make": r["make"], "std_part": r["std_part"],
                                 "note": f"{n}× · last on {r['part_code']}" + ("" if same_material else " (other material)")}
    return out


def save_bom_line(part_id: int, key: str, data: dict, as_default: bool = False) -> dict:
    """Change one tool BOM line of the part (grade, make, qty, price each, standard part; empty = back to suggested),
    remember the choice for the next parts, optionally make it the master-data default of the element, recalculate."""
    d = part_detail(part_id, with_rm=False)
    if tool_eligibility(d):
        raise RegisterError(tool_eligibility(d))
    part_type = register_setup.processes()[d["process"]]["part_type"]
    element = key.split("|", 1)[-1]
    typed = dict(d["tool_inputs"] or {})
    bom = dict(typed.get("bom") or {})
    line = _bom_overrides({key: data}, "BOM line").get(key, {})
    mdb = mould_db.tables()
    if line.get("grade") and line["grade"] not in {m["code"] for m in mdb["tool_materials"]}:
        raise RegisterError(f"Grade '{line['grade']}' is not in Tool materials. Add it first (+ New grade).")
    if line.get("std_part") and line["std_part"] not in {x["code"] for x in mdb["standard_parts"]}:
        raise RegisterError(f"Standard part '{line['std_part']}' is not in Standard parts. Add it first (+ New standard part).")
    if line:
        bom[key] = line
    else:
        bom.pop(key, None)
    typed["bom"] = bom
    update_part(part_id, {"tool_inputs": typed}, refresh=False)
    with db.engine().begin() as conn:
        if line and any(line.get(k) for k in ("grade", "make", "std_part")):
            conn.execute(db.bom_learning.insert().values(
                part_type=part_type, element=element, part_material=d["material"] or "", grade=line.get("grade", ""),
                make=line.get("make", ""), std_part=line.get("std_part", ""), part_id=part_id, part_code=d["part_code"],
                user=user()))
        _log(conn, "part", part_id, "tool_bom_changed", {"line": key, **({k: v for k, v in line.items()} or {"reset": "back to suggested"})})
    if as_default and line:
        set_bom_default(part_type, element, line.get("grade"), line.get("make"))
    try:
        res = compute_tool(part_id, f"tool BOM changed: {element}")
    except RegisterError as exc:
        res = {"stored": False, "version": None, "message": str(exc)}
    return res | {"part": part_detail(part_id)}


def set_bom_grade(part_id: int, keys: list[str], grade: str) -> dict:
    """One steel grade for several tool BOM lines at once (from the traceability window); other changes of the lines stay."""
    d = part_detail(part_id, with_rm=False)
    if tool_eligibility(d):
        raise RegisterError(tool_eligibility(d))
    if grade not in {m["code"] for m in mould_db.tables()["tool_materials"]}:
        raise RegisterError(f"Grade '{grade}' is not in Tool materials. Add it first (+ New grade).")
    keys = [str(k) for k in keys or [] if str(k).strip()]
    if not keys:
        raise RegisterError("Which BOM lines? (keys)")
    part_type = register_setup.processes()[d["process"]]["part_type"]
    typed = dict(d["tool_inputs"] or {})
    bom = {k: dict(v) for k, v in (typed.get("bom") or {}).items()}
    for k in keys:
        bom.setdefault(k, {})["grade"] = grade
    typed["bom"] = bom
    update_part(part_id, {"tool_inputs": typed}, refresh=False)
    with db.engine().begin() as conn:
        for k in keys:
            conn.execute(db.bom_learning.insert().values(
                part_type=part_type, element=k.split("|", 1)[-1], part_material=d["material"] or "", grade=grade,
                make=bom[k].get("make", ""), std_part=bom[k].get("std_part", ""), part_id=part_id, part_code=d["part_code"], user=user()))
        _log(conn, "part", part_id, "tool_bom_changed", {"lines": ", ".join(keys), "grade": grade})
    try:
        res = compute_tool(part_id, f"steel grade {grade} chosen for {len(keys)} line(s)")
    except RegisterError as exc:
        res = {"stored": False, "version": None, "message": str(exc)}
    return res


def grade_prices(country: str, as_of: str | None = None) -> dict:
    """Approved price per kg of every tool steel grade in a country (for the grade drop-down)."""
    out = {}
    for m in mould_db.tables()["tool_materials"]:
        try:
            r = rates.lookup("tool_material", m["code"], country, as_of or rates.current_month())
            out[m["code"]] = {"value": r["value"], "unit": r["unit"], "month": r["month"], "is_example": r["is_example"],
                              "hardness": m["hardness"], "make": m.get("make") or "", "application": m["application"]}
        except rates.RateMissing:
            out[m["code"]] = {"value": None, "unit": "", "missing": True, "hardness": m["hardness"], "make": m.get("make") or "",
                              "application": m["application"]}
    return out


def set_bom_default(part_type: str, element: str, grade: str | None, make: str | None) -> None:
    """Make your grade / make the master-data default of a BOM element (Mould & die database → BOM elements)."""
    mdb = mould_db.tables()
    row = next((r for r in mdb["bom_elements"] if r["part_type"] == part_type and r["element"] == element), None)
    if not row:
        return
    old = {"material_code": row["material_code"], "make": row.get("make") or ""}
    if grade and row["material_code"]:
        row["material_code"] = grade
    if make:
        row["make"] = make
    if old != {"material_code": row["material_code"], "make": row.get("make") or ""}:
        mould_db.write(mdb)
        with db.engine().begin() as conn:
            _log(conn, "master_data", 0, "saved", {"table": "mould-db/bom_elements", "element": f"{part_type} / {element}",
                                                    "default": {"old": old, "new": {"material_code": row["material_code"], "make": row.get("make")}}})


def add_tool_material(data: dict) -> dict:
    """A new steel / material grade for tool BOMs: Tool materials row + its price as a rate of the country
    (approved when you may approve rates, else a draft for the approver)."""
    from . import auth
    code = str(data.get("code") or "").strip()
    if not code:
        raise RegisterError("New grade: give the grade name, e.g. 'OHNS (O1)'.")
    density = _num(data.get("density_g_cm3"), "New grade: density", required=True)
    price = _num(data.get("price_per_kg"), "New grade: price per kg", required=True)
    if not density or not price:
        raise RegisterError("New grade: density and price per kg must be more than 0.")
    country = str(data.get("country") or register_setup.settings()["default_country"]).strip().upper()
    countries = {c["code"]: c for c in costing_setup.tables()["countries"]}
    if country not in countries:
        raise RegisterError(f"New grade: country '{country}' is not in Master data → Countries.")
    cur = countries[country]["currency"]
    mdb = mould_db.tables()
    if any(m["code"].lower() == code.lower() for m in mdb["tool_materials"]):
        raise RegisterError(f"Grade '{code}' already exists. Pick it from the list.")
    inr = price if cur == "INR" else None
    mdb["tool_materials"].append({"code": code, "application": str(data.get("application") or "").strip() or "die_insert",
                                  "density_g_cm3": density, "rate_per_kg": inr or 0, "hardness": str(data.get("hardness") or "").strip(),
                                  "make": str(data.get("make") or "").strip(), "is_example": False,
                                  "notes": str(data.get("notes") or "").strip() or f"Added from a tool BOM by {user()}"})
    mould_db.write(mdb)
    approve = auth.may("approve_rates") and (auth.server_mode() or bool(rates.settings()["approver_name"].strip()))
    rate = rates.create({"category": "tool_material", "item": code, "country": country, "value": price, "unit": f"{cur}/kg",
                         "currency": cur, "month": rates.current_month(), "source": str(data.get("source") or "your price"),
                         "comment": f"New grade added from a tool BOM by {user()}"}, approve=approve)
    with db.engine().begin() as conn:
        _log(conn, "master_data", 0, "saved", {"table": "mould-db/tool_materials", "added": code, "rate": rate["id"],
                                                "status": rate["status"]})
    return {"grade": code, "rate": rate, "approved": rate["status"] == "approved",
            "message": f"Grade {code} added with {price:g} {cur}/kg for {country}"
                       + ("" if rate["status"] == "approved" else " as a DRAFT rate: an approver approves it on Rates → Rate review.")}


def add_standard_part(data: dict) -> dict:
    """A new standard part (catalogue item) for tool BOMs: code, maker's part number, description, make, price."""
    code = str(data.get("code") or "").strip().upper()
    desc = str(data.get("description") or "").strip()
    make = str(data.get("make") or "").strip()
    pn = str(data.get("part_number") or "").strip()
    if not (code and desc and make and pn):
        raise RegisterError("New standard part: code, part number, description and make are all needed.")
    price = _num(data.get("price"), "New standard part: price", required=True)
    country = str(data.get("country") or "").strip().upper()
    countries = {c["code"]: c for c in costing_setup.tables()["countries"]}
    if country and country not in countries:
        raise RegisterError(f"New standard part: country '{country}' is not in Master data → Countries.")
    cur = str(data.get("currency") or (countries[country]["currency"] if country else "INR")).strip().upper()
    mdb = mould_db.tables()
    if any(x["code"] == code for x in mdb["standard_parts"]):
        raise RegisterError(f"Standard part '{code}' already exists. Pick it from the list.")
    row = {"code": code, "part_number": pn, "description": desc, "element": str(data.get("element") or "").strip(),
           "make": make, "country": country, "currency": cur, "price": price, "is_example": False,
           "notes": str(data.get("notes") or "").strip() or f"Added from a tool BOM by {user()}"}
    mdb["standard_parts"].append(row)
    mould_db.write(mdb)
    with db.engine().begin() as conn:
        _log(conn, "master_data", 0, "saved", {"table": "mould-db/standard_parts", "added": f"{code} {make} {pn}"})
    return {"standard_part": row, "message": f"Standard part {code} ({make} {pn}) added: {price:g} {cur}."}


def bom_choices() -> dict:
    """Drop-down lists for the tool BOM: grades (Tool materials) and standard parts."""
    mdb = mould_db.tables()
    return {"grades": [{"code": m["code"], "application": m["application"], "hardness": m["hardness"], "make": m.get("make") or "",
                        "is_example": m["is_example"]} for m in mdb["tool_materials"]],
            "standard_parts": [{k: x.get(k) for k in ("code", "part_number", "description", "element", "make", "country", "currency",
                                                      "price", "is_example")} for x in mdb["standard_parts"]],
            "makes": sorted({x for m in mdb["tool_materials"] for x in [m.get("make")] if x}
                            | {x for e in mdb["bom_elements"] for x in [e.get("make")] if x}
                            | {x["make"] for x in mdb["standard_parts"] if x.get("make")})}


def _concept_inputs(part_id: int):
    """(part detail, revision, part type, vertices, triangles, concept result, part facts) of the current revision."""
    from .geometry import load_step, tessellate
    d = part_detail(part_id, with_rm=False)
    why = tool_eligibility(d)
    if why:
        raise ToolNotApplicable(why)
    part_type = register_setup.processes()[d["process"]]["part_type"]
    rev = d["revisions"][0] if d["revisions"] else None
    if not rev or not rev["file_id"]:
        raise ToolNotApplicable("This part has no 3D file yet. Upload a revision first.")
    with db.engine().connect() as conn:
        row = conn.execute(select(db.files).where(db.files.c.id == rev["file_id"])).mappings().first()
    path = file_store.path_of(dict(row))
    if not path.exists():
        raise ToolNotApplicable(f"The 3D file is missing from the storage folder ({row['stored_path']}).")
    g = rev["geometry"]
    verts, tris = tessellate(load_step(str(path)))
    try:
        res = concept.analyze(verts, tris, part_type, g["projected_area_mm2"])
    except ValueError as exc:
        raise ToolNotApplicable(str(exc)) from None
    res.pop("_mask")
    axis = res["opening_axis"]
    across = sorted((v for a, v in g["bbox_axes_mm"].items() if a != axis), reverse=True)
    typed = d["tool_inputs"] or {}
    try:
        cavities = typed.get("cavities") or _tool_compute(part_id)["stage_c"]["cavities_per_tool"]
    except (RegisterError, ToolNotApplicable, tool_cost.ToolCostError):
        cavities = 1
    t = _geometry_trace(part_id)[0]
    facts = {"part_weight_g": t.value("weight") if "weight" in t.nodes else 0, "wall_mm": g.get("wall_mm") or 0,
             "part_l": across[0], "part_w": across[1], "part_h": g["bbox_axes_mm"][axis],
             "projected_area_cm2": g["projected_area_mm2"][axis] / 100, "volume_cm3": g["volume_cm3"],
             "cavities": cavities, "slides": typed.get("slides_per_cavity", res["slides"]),
             "lifters": typed.get("lifters_per_cavity", res["lifters"]), "undercut_area_cm2": res["undercut_area_cm2"],
             "inserts": d["insert_count"]}
    return d, rev, part_type, verts, tris, res, facts


def mould_concept(part_id: int) -> dict:
    """Concept v1 of the part's current revision: opening direction, undercuts -> slides / lifters, proposals."""
    d, rev, part_type, _, _, res, facts = _concept_inputs(part_id)
    typed = d["tool_inputs"] or {}
    return res | {"part_id": part_id, "rev": rev["rev"], "stored_pull_direction": rev["geometry"]["pull_direction"],
                  "facts": facts, "proposals": concept.proposals(part_type, facts),
                  "typed": {k: typed[k] for k in ("slides_per_cavity", "lifters_per_cavity") if k in typed}}


def _simulate(part_id: int):
    from . import simulation
    d, rev, part_type, verts, tris, _, facts = _concept_inputs(part_id)
    try:
        res = simulation.run(verts, tris, part_type, facts, rev["geometry"]["projected_area_mm2"])
    except ValueError as exc:
        raise ToolNotApplicable(str(exc)) from None
    return d, rev, res


def simulation_result(part_id: int) -> dict:
    """Simulation & analysis v1 of the part's current revision (roadmap stage 6)."""
    from . import simulation
    d, rev, res = _simulate(part_id)
    return simulation.public(res) | {"part_id": part_id, "rev": rev["rev"], "part_code": d["part_code"]}


def simulation_pdf(part_id: int) -> bytes:
    from . import simulation
    from .pdf_report import simulation_pdf as pdf
    d, rev, res = _simulate(part_id)
    return pdf(simulation.public(res), simulation.pictures(res), d, rev["rev"])


def _tool_row(row) -> dict:
    d = _plain({k: row[k] for k in TOOL_COLUMNS})
    d["is_example"] = bool((d["example"] or {}).get("any"))
    return d


def tool_versions(part_id: int) -> list[dict]:
    with db.engine().connect() as conn:
        rows = conn.execute(select(db.tool_costs).where(db.tool_costs.c.part_id == part_id)
                            .order_by(db.tool_costs.c.version.desc())).mappings().all()
    return [_tool_row(r) for r in rows]


def tool_summary(d: dict) -> dict:
    """What the part panel shows about the tool: latest version, whether it is out of date, or why there is none."""
    versions = tool_versions(d["id"])
    latest = versions[0] if versions else None
    out = {"applicable": False, "problem": "", "outdated": "", "rate_changes": [], "latest": latest,
           "versions": len(versions), "inputs": TOOL_INPUTS, "live": None}
    why = tool_eligibility(d)
    if why:
        out["problem"] = why
        return out
    out["applicable"] = True
    try:
        r = _tool_compute(d["id"])
    except (ToolNotApplicable, RegisterError, tool_cost.ToolCostError) as exc:
        out["problem"] = str(exc)
        return out
    out["live"] = {"cavities": r["cavities"], "tools_count": r["tools_count"], "stage_c": r["stage_c"],
                   "country": r["country"], "currency": r["currency"], "part_type": r["part_type"]}
    if r["missing"]:
        out["problem"] = " ".join(r["missing"])
    elif latest is None:
        out["outdated"] = "The tool cost has not been calculated yet."
    elif latest["key"] != r["key"]:
        strip = lambda k: {a: b for a, b in k.items() if a not in ("rates", "as_of")}
        changes = _rate_changes(latest["rates_used"] or [], _rates_used(r))
        if strip(latest["key"]) == strip(r["key"]) and (changes or latest["key"]["as_of"] != r["key"]["as_of"]):
            out["outdated"] = f"Newer approved rates exist than the ones used in tool cost version {latest['version']}."
            out["rate_changes"] = changes
        else:
            out["outdated"] = f"Inputs, volume, master data or formulas changed since tool cost version {latest['version']}."
    return out


def compute_tool(part_id: int, reason: str = "calculated") -> dict:
    """Work out the tool cost and store it as a new version unless nothing changed."""
    try:
        r = _tool_compute(part_id)
    except (ToolNotApplicable, tool_cost.ToolCostError) as exc:
        raise RegisterError(str(exc)) from None
    if r["missing"]:
        raise RegisterError("Tool cost not stored: " + " ".join(r["missing"]))
    with db.engine().begin() as conn:
        last = conn.execute(select(db.tool_costs).where(db.tool_costs.c.part_id == part_id)
                            .order_by(db.tool_costs.c.version.desc()).limit(1)).mappings().first()
        if last and last["key"] == r["key"]:
            return {"stored": False, "version": last["version"],
                    "message": f"Tool cost already up to date (version {last['version']})."}
        version = (last["version"] if last else 0) + 1
        tid = conn.execute(db.tool_costs.insert().values(
            part_id=part_id, version=version, revision_id=r["revision"]["id"], rev=r["revision"]["rev"],
            part_type=r["part_type"], country=r["country"], currency=r["currency"], as_of=r["as_of"],
            inputs=r["key"]["inputs"], key=r["key"], cavities=r["cavities"], tools_count=r["tools_count"],
            tool_cost=r["value"], tool_cost_inr=r["value_inr"], tool_capex_inr=r["capex_inr"],
            rates_used=_rates_used(r), example=r["example"], trace=r, reason=reason,
            created_by=user())).inserted_primary_key[0]
        conn.execute(db.parts.update().where(db.parts.c.id == part_id).values(
            tool_cost_id=tid, tool_cost_inr=r["value_inr"], tool_capex_inr=r["capex_inr"]))
        _log(conn, "part", part_id, "tool_cost_calculated", {
            "version": version, "tool cost per tool (INR)": round(r["value_inr"], 2),
            "tool CAPEX (INR)": round(r["capex_inr"], 2) if r["capex_inr"] is not None else None,
            "cavities": r["cavities"], "tools": r["tools_count"], "tool country": r["country"],
            "costing month": r["as_of"], "reason": reason,
            "rate versions": ", ".join(f"#{x['id']}" for x in r["rates_used"])})
    return {"stored": True, "version": version, "message": f"Tool cost version {version} stored."}


def refresh_tool(part_ids, reason: str) -> None:
    for pid in dict.fromkeys(part_ids):
        try:
            compute_tool(pid, reason)
        except (RegisterError, tool_cost.ToolCostError, ToolNotApplicable):
            pass


def recalculate_project_tools(project_id: int, reason: str = "recalculated with the latest approved rates") -> dict:
    with db.engine().connect() as conn:
        ids = [r for (r,) in conn.execute(select(db.parts.c.id).where(and_(db.parts.c.project_id == project_id,
                                                                         db.parts.c.is_assembly.is_(False))))]
    done = same = 0
    skipped = []
    for pid in ids:
        try:
            res = compute_tool(pid, reason)
        except RegisterError as exc:
            if tool_eligibility(part_detail(pid, with_rm=False)) == "":
                skipped.append({"part_id": pid, "why": str(exc)})
            continue
        done += res["stored"]
        same += not res["stored"]
    return {"new_versions": done, "unchanged": same, "problems": skipped}


def set_tool_inputs(part_id: int, tool_inputs: dict | None, tool_country: str | None = None,
                    reason: str = "tool inputs saved") -> dict:
    tool_inputs = dict(tool_inputs or {})
    if "bom" not in tool_inputs:  # the Tool features form does not carry the BOM changes: keep them
        old_bom = (part_detail(part_id, with_rm=False)["tool_inputs"] or {}).get("bom")
        if old_bom:
            tool_inputs["bom"] = old_bom
    data = {"tool_inputs": tool_inputs}
    if tool_country is not None:
        data["tool_country"] = tool_country
    update_part(part_id, data, refresh=False)
    try:
        res = compute_tool(part_id, reason)
    except RegisterError as exc:
        res = {"stored": False, "version": None, "message": str(exc)}
    return res | {"part": part_detail(part_id)}


def tool_trace(part_id: int, version: int | None = None) -> dict:
    """The stored trace of a tool cost version, or (no version given and none stored) a live dry run."""
    if version is not None or tool_versions(part_id):
        with db.engine().connect() as conn:
            q = select(db.tool_costs).where(db.tool_costs.c.part_id == part_id)
            q = q.where(db.tool_costs.c.version == version) if version is not None else q.order_by(db.tool_costs.c.version.desc())
            row = conn.execute(q.limit(1)).mappings().first()
        if not row:
            raise RegisterError(f"Part {part_id} has no tool cost version {version}.")
        return row["trace"] | {"version": row["version"], "stored": True, "created_by": row["created_by"],
                               "created_at": row["created_at"].isoformat(timespec="seconds")}
    try:
        return _tool_compute(part_id) | {"version": None, "stored": False}
    except (ToolNotApplicable, tool_cost.ToolCostError) as exc:
        raise RegisterError(str(exc)) from None


def export_tool_excel(part_id: int, version: int | None = None, die: str | None = None) -> bytes:
    d = part_detail(part_id, with_rm=False)
    return tool_cost_excel(tool_trace(part_id, version), d, die)


def export_tool_pdf(part_id: int, version: int | None = None, die: str | None = None) -> bytes:
    from .pdf_report import tool_cost_pdf
    return tool_cost_pdf(tool_trace(part_id, version), part_detail(part_id, with_rm=False), die)


def tool_cost_excel(tr: dict, part: dict | None = None, die: str | None = None) -> bytes:
    """Tool cost sheet: Summary, BOM & processes, Steps, Inputs, Rates used, Example values. Also used by the Estimate page."""
    from .trace import source_text
    nodes, cur = tr["nodes"], tr["currency"]
    ex = tr.get("example") or tool_cost.examples(tr)
    wb = Workbook()
    title = part["part_code"] if part else "Estimate"
    wb.properties.title = f"Design to Cost — Tool cost {title}"
    wb.properties.creator = "Design to Cost"

    def sheet(name, head, rows, first=False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = name
        if ex["any"]:
            _excel_banner(ws, len(head))
        ws.append(head)
        top = ws.max_row
        for c in ws[top]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = HEADER_FILL
        for r in rows:
            ws.append(r)
            if len(r) and r[-1] == "EXAMPLE":
                for c in ws[ws.max_row]:
                    c.fill = EXAMPLE_FILL
        for i, h in enumerate(head, start=1):
            ws.column_dimensions[ws.cell(top, i).column_letter].width = min(
                max([len(str(h))] + [len(str(r[i - 1])) for r in rows]) + 2, 70)
        ws.freeze_panes = f"A{top + 1}"

    v = lambda nid: nodes[nid]["value"] if nid and nid in nodes else None
    b = tr["base"]
    summary = []
    if part:
        summary += [["Part code", part["part_code"]], ["Part name", part["name"]],
                    ["Revision", tr.get("revision", {}).get("rev", "")], ["Process", tr.get("process", part["process"])]]
    summary += [["Part type", tr["part_type"]], ["Part material", tr.get("material") or "—"], ["Tool country", tr["country"]],
                ["Costing month", tr["as_of"]],
                ["Tool cost version", tr.get("version") or "not stored (live preview)"],
                ["Cavities per tool", tr["cavities"]], ["Mould base", f"{b['series']} {b['size_code']}"],
                ["Try-out machine", (tr["trial"]["machine"] or {}).get("machine_id", "— (missing)")],
                [f"Tool cost per tool ({cur})", tr["value"]], ["Tool cost per tool (INR)", tr["value_inr"]],
                ["Tools over the programme", tr["tools_count"]], [f"Tool CAPEX ({cur})", tr.get("capex")],
                ["Tool CAPEX (INR)", tr.get("capex_inr")],
                ["Rate versions used", ", ".join(f"#{x['id']}" for x in tr["rates_used"])],
                ["Includes EXAMPLE values", f"yes - {EXAMPLE_WARNING}. {ex['detail']}" if ex["any"] else "no"],
                ["Calculated by / when", f"{tr.get('created_by', '')} {tr.get('created_at', '')}".strip()]]
    dies = tr.get("dies") or []
    if die:
        dies = [x for x in dies if x["key"] == die]
        if not dies:
            raise ValueError(f"There is no die '{die}' in this tool cost.")
    sheet("Summary", ["Item", "Value"], summary, first=True)
    if dies:   # sheet-metal dies, stage-wise: summary + one sheet per die with its BOM
        sheet("Dies", ["Stage", "Die", "Operations", "Die set", "Plate material", "Die plate L (mm)", "W (mm)", "H (mm)", "Press",
                       "Force (t)", f"Tool cost ({cur})", "Tool cost INR", "CAPEX INR"],
              [[tr["dies"].index(x) + 1, x["name"], ", ".join(x["op_names"]), f"{x['base']['series']} {x['base']['size_code']}",
                x["base"].get("material") or "", x["die_plate"]["die_l"], x["die_plate"]["die_w"], x["die_plate"]["die_h"],
                (x["press"] or {}).get("machine_id") or "none", x["force_t"], x["value"], x["value_inr"], x.get("capex_inr")] for x in dies])
        for x in dies:
            n = tr["dies"].index(x) + 1
            sheet(f"Die {n} {x['key']}"[:31], ["Element", "Category", "Qty", "Grade", "Grade from", "Make", "Standard part", "Part number",
                                            "Weight each (kg)", f"Steel ({cur})", f"Bought-out ({cur})", "Process", "Hours",
                                            f"Process ({cur})", f"Line total ({cur})", ""],
                  [[l["element"], l["category"], l["quantity"], l.get("grade", ""), l.get("grade_reason", ""), l.get("make", ""),
                    l.get("std_part", ""), l.get("part_number", ""), l["weight_each_kg"], l["material_cost"], l["bought_out_cost"],
                    l["process"], l["hours"], l["process_cost"], l["total"], "EXAMPLE" if l["is_example"] else ""]
                   for l in tr["lines"] if l.get("die_key") == x["key"]])
        if die:
            buf = io.BytesIO()
            wb.save(buf)
            return buf.getvalue()
    sheet("Cost build-up", ["Step", "Formula", "Value", "Unit"],
          [[nodes[g["node"]]["label"], "Σ BOM lines of the group", v(g["node"]), cur] for g in tr["groups"]] +
          [[nodes[s["node"]]["label"], nodes[s["node"]]["formula"], v(s["node"]), nodes[s["node"]]["unit"]] for s in tr["steps"]])
    sheet("BOM & processes", ["Die", "Element", "Category", "Qty", "Grade", "Grade from", "Make", "Standard part", "Part number",
                              "Changed by you", "Steel weight each (kg)", f"Steel ({cur})",
                              f"Bought-out ({cur})", "Process", "Priced by", "Hours", "Amount", f"Process ({cur})",
                              f"Line total ({cur})", ""],
          [[l.get("die", ""), l["element"], l["category"], l["quantity"], l.get("grade", l["material"]), l.get("grade_reason", ""),
            l.get("make", ""), l.get("std_part", ""), l.get("part_number", ""), ", ".join(l.get("changed") or []) or
            (f"learned ({l['learned']})" if l.get("learned") else ""), l["weight_each_kg"], l["material_cost"],
            l["bought_out_cost"], l["process"], l["pricing"].replace("_", " "), l["hours"], l["amount"], l["process_cost"],
            l["total"], "EXAMPLE" if l["is_example"] else ""] for l in tr["lines"]])
    sheet("Steps", ["Step", "Formula", "Value", "Unit"],
          [[n["label"], n["formula"], n["value"], n["unit"]] for n in nodes.values() if n["kind"] == "calculated"])
    sheet("Inputs", ["Input", "Value", "Unit", "Comes from", ""],
          [[n["label"], n["value"], n["unit"], source_text(n),
            "EXAMPLE" if ((n.get("source") or {}).get("rate") or {}).get("is_example")
            or (n.get("source") or {}).get("is_example") else ""]
           for n in nodes.values() if n["kind"] == "input"])
    sheet("Rates used", ["Version id", "Category", "Item", "Country", "Month", "Value", "Unit", "Status", "Approved by", ""],
          [[x["id"], x["category"], x["item"], x["country"], x["month"], x["value"], x["unit"], x["status"],
            x["approved_by"], "EXAMPLE" if x["is_example"] else ""] for x in tr["rates_used"]])
    sheet("Example values", ["Still an EXAMPLE value", "Where to replace it", ""],
          [[f"Rate library #{i}", "Rates page → Rate library", "EXAMPLE"] for i in ex["rate_ids"]] +
          [[r, "Master data", "EXAMPLE"] for r in ex["rows"]])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------- BOM ----------

def set_bom_line(conn, variant_id: int, parent_id: int | None, child_id: int, qty: float) -> None:
    cond = and_(db.bom_lines.c.variant_id == variant_id, db.bom_lines.c.child_part_id == child_id,
                db.bom_lines.c.parent_part_id.is_(None) if parent_id is None else db.bom_lines.c.parent_part_id == parent_id)
    old = conn.execute(select(db.bom_lines).where(cond)).mappings().first()
    if old:
        if old["quantity"] != qty:
            conn.execute(db.bom_lines.update().where(db.bom_lines.c.id == old["id"]).values(quantity=qty))
            _log(conn, "bom", old["id"], "updated", {"quantity": {"old": old["quantity"], "new": qty}})
        return
    bid = conn.execute(db.bom_lines.insert().values(variant_id=variant_id, parent_part_id=parent_id,
                                                    child_part_id=child_id, quantity=qty)).inserted_primary_key[0]
    _log(conn, "bom", bid, "created", {"variant_id": variant_id, "parent": parent_id, "child": child_id, "quantity": qty})


def variant_bom(variant_id: int) -> list[dict]:
    with db.engine().connect() as conn:
        rows = conn.execute(select(db.bom_lines, db.parts.c.part_code, db.parts.c.name, db.parts.c.is_assembly,
                                   db.parts.c.sourcing, db.parts.c.status)
                            .join(db.parts, db.parts.c.id == db.bom_lines.c.child_part_id)
                            .where(db.bom_lines.c.variant_id == variant_id)).mappings().all()
        ins = db.parts.alias("ins")
        moulded: dict = {}
        for m in conn.execute(select(db.part_inserts.c.part_id, db.part_inserts.c.quantity, ins.c.id, ins.c.part_code,
                                     ins.c.name, ins.c.sourcing, ins.c.status)
                              .join(ins, ins.c.id == db.part_inserts.c.insert_part_id)).mappings():
            moulded.setdefault(m["part_id"], []).append(dict(m))
    by_parent: dict = {}
    for r in rows:
        by_parent.setdefault(r["parent_part_id"], []).append(dict(r))


    def tree(parent, depth, mult):
        out = []
        for r in sorted(by_parent.get(parent, []), key=lambda x: x["part_code"]):
            total = r["quantity"] * mult
            out.append(r | {"depth": depth, "total_quantity": total, "kind": "bom"})
            # moulded-in inserts: shown under their body for buying, but not BOM children
            for m in sorted(moulded.get(r["child_part_id"], []), key=lambda x: x["part_code"]):
                out.append({"child_part_id": m["id"], "parent_part_id": r["child_part_id"], "part_code": m["part_code"],
                            "name": m["name"], "is_assembly": False, "sourcing": m["sourcing"], "status": m["status"],
                            "quantity": m["quantity"], "depth": depth + 1, "total_quantity": m["quantity"] * total,
                            "kind": "insert"})
            out += tree(r["child_part_id"], depth + 1, total)
        return out

    return tree(None, 0, 1)


# ---------- assembly / single-part import (preview, then commit) ----------

def import_preview(content: bytes, file_name: str) -> dict:
    token = uuid.uuid4().hex
    folder = file_store.scratch_dir(token)
    src = folder / "source.step"
    src.write_bytes(content)
    try:
        tree = cad_import.analyse_file(str(src), file_name, insert_detection=register_setup.insert_detection())
    except cad_import.ImportError_ as exc:
        shutil.rmtree(folder, ignore_errors=True)
        raise RegisterError(str(exc)) from None
    defs = tree["defs"]
    settings = register_setup.settings()
    insert_count: dict[str, float] = {}
    for det in tree["inserts"]:
        det["body_name"], det["insert_name"] = defs[det["body"]]["name"], defs[det["insert"]]["name"]
        insert_count[det["body"]] = insert_count.get(det["body"], 0) + det["quantity"]
        defs[det["insert"]].setdefault("insert_of", []).append(det["body_name"])
        defs[det["insert"]]["default_sourcing"] = settings["insert_default_sourcing"]
    for det in tree["inserts"]:
        # an assembly that is only "body + its inserts" is really one moulded part: untick it by default
        a = defs[det["assembly"]]
        moulded = {det["body"]} | {x["insert"] for x in tree["inserts"] if x["assembly"] == det["assembly"]
                                   and x["body"] == det["body"]}
        if set(a["children"]) <= moulded and det["assembly"] != tree["root"]:
            a["only_moulded_part"] = det["body"]
    for key, d in tree["defs"].items():
        if d["is_assembly"]:
            # an assembly revision records its structure (child names and quantities)
            d["metrics"] = {"assembly": True, "children": sorted(
                [tree["defs"][c]["name"], q] for c, q in d["children"].items())}
        if d.get("step"):
            (folder / f"{key}.step").write_bytes(d.pop("step"))
        else:
            d.pop("step", None)
        if d.get("thumbnail"):
            (folder / f"{key}.png").write_bytes(d.pop("thumbnail"))
        d["insert_count"] = insert_count.get(key, 0)
        if d.get("insert_of"):
            code = settings["insert_default_process"]
            text = (f"Detected as a moulded-in insert of {', '.join(d['insert_of'])}, so it is taken as "
                    f"{register_setup.processes().get(code, {}).get('name', code)} (Register settings → "
                    "insert_default_process). Untick the insert on its body to treat it as a normal part.")
            d["suggestion"] = {"process": code, "steps": [text], "reason": text}
        else:
            d["suggestion"] = suggest_process(d["metrics"], "", d["insert_count"])
    manifest = {"token": token, "file_name": file_name, "root": tree["root"], "defs": tree["defs"],
                "inserts": tree["inserts"]}
    (folder / "manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
    return manifest


def import_thumbnail(token: str, key: str) -> Path:
    p = file_store.scratch_dir(token) / f"{Path(key).name}.png"
    if not p.exists():
        raise RegisterError("Thumbnail not found.")
    return p


def import_commit(token: str, project_id: int, variant_id: int | None, items: dict, rev: str = "",
                  change_note: str = "") -> dict:
    """items: {def_key: {part_code, name, material, process, sourcing, register(bool), inserts: [insert def keys]}}

    A detected insert that the engineer confirms on its body (items[body]["inserts"]) is
    linked to the body as a moulded-in insert instead of becoming a BOM child.
    """
    folder = file_store.scratch_dir(Path(token).name)
    mf = folder / "manifest.json"
    if not mf.exists():
        raise RegisterError("This import has expired. Please upload the file again.")
    manifest = json.loads(mf.read_text(encoding="utf-8"))
    defs = manifest["defs"]
    created, revised, ids = [], [], {}
    confirmed = {(det["assembly"], det["insert"]): det for det in manifest.get("inserts", [])
                 if det["insert"] in (items.get(det["body"], {}).get("inserts") or [])
                 and items.get(det["body"], {}).get("register", True)}
    for det in confirmed.values():
        if det["overlap_mm3"] and defs[det["body"]]["metrics"]:
            defs[det["body"]]["metrics"]["insert_overlap_mm3"] = det["overlap_mm3"]
    with db.engine().begin() as conn:
        if variant_id is not None:
            v = conn.execute(select(db.variants).where(db.variants.c.id == variant_id)).mappings().first()
            if not v or v["project_id"] != project_id:
                raise RegisterError("The chosen variant does not belong to the chosen project.")

        def make(key):
            if key in ids:
                return ids[key]
            d, choice = defs[key], items.get(key, {})
            if not choice.get("register", True):
                ids[key] = None
                return None
            defaults = {"sourcing": d["default_sourcing"]} if d.get("default_sourcing") else {}
            vals = _part_values(defaults | {k: choice[k] for k in ("material", "process", "sourcing") if choice.get(k)}
                                | {"name": choice.get("name") or d["name"]}, f"Part '{d['name']}'")
            pid, is_new = upsert_part(conn, project_id, choice.get("part_code", ""), vals, d["is_assembly"])
            is_root = key == manifest["root"]
            step_path = folder / ("source.step" if is_root else f"{key}.step")
            step = step_path.read_bytes() if step_path.exists() else None
            thumb_path = folder / f"{key}.png"
            if not is_new and _same_as_current(conn, pid, d["metrics"], step):
                unchanged.append(pid)  # same geometry / structure as the current revision: just link it
            else:
                add_revision(conn, pid, rev, step=step,
                             file_name=manifest["file_name"] if is_root else f"{d['name']}.step",
                             thumbnail=thumb_path.read_bytes() if thumb_path.exists() else None,
                             geometry=d["metrics"], change_note=change_note,
                             source=f"{manifest['file_name']} / {d['name']}")
                (created if is_new else revised).append(pid)
            ids[key] = pid
            return pid

        def link(key, parent_pid, qty):
            pid = make(key)
            if pid is not None and variant_id is not None:
                set_bom_line(conn, variant_id, parent_pid, pid, float(qty))
            # An unticked assembly passes its children (and its quantity) on to the next ticked level up.
            for child, child_qty in defs[key]["children"].items():
                if (key, child) in confirmed:
                    continue  # moulded into a sibling body: linked as an insert below, not a BOM child
                if pid is None:
                    link(child, parent_pid, child_qty * qty)
                else:
                    link(child, pid, child_qty)

        unchanged: list = []
        link(manifest["root"], None, 1)
        linked = 0
        for det in confirmed.values():
            body, ins = make(det["body"]), make(det["insert"])
            if body and ins:
                _set_insert_link(conn, body, ins, float(det["quantity"]), detected=True,
                                 notes=f"Detected in {manifest['file_name']}: {det['covered_share'] * 100:.0f} % of the "
                                       "insert surface is covered by the body")
                linked += 1
    shutil.rmtree(folder, ignore_errors=True)
    refresh_costs([i for i in ids.values() if i], "imported from a 3D file")
    return {"created": len(created), "new_revisions": len(revised), "unchanged": len(set(unchanged)),
            "inserts_linked": linked, "part_ids": [i for i in ids.values() if i]}


def _same_as_current(conn, part_id, metrics, step) -> bool:
    """Exported STEP files carry a timestamp, so compare geometry (or the file checksum when it is the original)."""
    cur = conn.execute(select(db.revisions.c.geometry, db.files.c.sha256).select_from(
        db.parts.join(db.revisions, db.revisions.c.id == db.parts.c.current_revision_id)
        .outerjoin(db.files, db.files.c.id == db.revisions.c.file_id)).where(db.parts.c.id == part_id)).first()
    if not cur:
        return False
    geometry, sha = cur
    if step is not None and sha == hashlib.sha256(step).hexdigest():
        return True
    if not geometry or not metrics:
        return False
    if metrics.get("assembly"):
        return geometry.get("children") == metrics["children"]
    return "volume_mm3" in geometry and cad_import.signature(geometry) == cad_import.signature(metrics)


# ---------- Excel bulk upload ----------

TEMPLATE_COLUMNS = [
    ("part_code", "Your part code; leave blank to generate one"),
    ("part_name", "Required"),
    ("rev", "Drawing revision, e.g. A or 01; blank = next revision"),
    ("file_name", "STEP file name, matched to the uploaded files / ZIP"),
    ("material", "Material name as in Master data → Materials"),
    ("process", "Process code (see Part register setup → Processes); blank = suggested from the 3D file"),
    ("sourcing", "make, bought_out, existing or not_required"),
    ("annual_volume", "Parts per year"),
    ("programme_life_years", "Years"),
    ("target_piece_cost_inr", "Optional target cost per piece (INR)"),
    ("variant", "Variant code; adds the part to that variant's BOM"),
    ("parent_part_code", "Assembly it belongs to (blank = top level of the variant)"),
    ("quantity", "Quantity in the parent (default 1)"),
    ("change_note", "What changed (for a new revision)"),
    ("insert_of_part_code", "Moulded-in insert: part code of the moulded / cast part this insert goes into"),
    ("insert_quantity", "How many of this insert go into that part (default 1)"),
    ("unit_price", "Bought-out price per piece (inserts and other BOP parts)"),
    ("price_currency", "Currency of unit_price, e.g. INR"),
    ("price_source", "Where the price comes from: supplier quote, catalogue, estimate"),
    ("price_month", "Month of the price, e.g. 2026-10"),
]


def excel_template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Parts"
    wb.properties.title = "Design to Cost — Part register bulk upload"
    wb.properties.creator = "Design to Cost"
    ws.append([c for c, _ in TEMPLATE_COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
    ws.append(["", "Example bracket", "A", "sheet_metal_bracket.step", "Mild steel CRCA / HR", "", "make", 200000, 5,
               "", "STD", "", 2, "first upload"])
    ws.append(["EX-HSG-01", "Example housing with inserts", "A", "housing.step", "PA66 GF30", "", "make", 200000, 5,
               "", "STD", "", 1, "first upload"])
    ws.append(["EX-INS-M3", "Example M3 brass threaded insert", "", "", "Brass CuZn39Pb3", "", "", "", "",
               "", "", "", "", "", "EX-HSG-01", 4, 2.5, "INR", "supplier quote", "2026-10"])
    for r in (2, 3, 4):
        for cell in ws[r]:
            cell.fill = EXAMPLE_FILL
    for i, (c, _) in enumerate(TEMPLATE_COLUMNS, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = max(14, len(c) + 4)
    help_ws = wb.create_sheet("Help")
    help_ws.append(["column", "meaning"])
    for cell in help_ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
    for c, h in TEMPLATE_COLUMNS:
        help_ws.append([c, h])
    help_ws.append([])
    help_ws.append(["note", "Rows 2-4 are examples (yellow); replace or delete them. Same part_code again = new revision."])
    help_ws.append(["inserts", "An insert row names its moulded part in insert_of_part_code. It gets sourcing / process "
                               "from Register settings when blank, and the moulded part's process is suggested as insert "
                               "moulding. Without a 3D file, add its shape and dimensions later on the Parts page."])
    help_ws.column_dimensions["A"].width = 24
    help_ws.column_dimensions["B"].width = 90
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _collect_files(uploads: list[tuple[str, bytes]]) -> dict[str, bytes]:
    """{lower-case file name: bytes} from STEP uploads and ZIPs."""
    out = {}
    for name, content in uploads:
        if name.lower().endswith(".zip"):
            try:
                with zipfile.ZipFile(io.BytesIO(content)) as z:
                    for info in z.infolist():
                        if not info.is_dir() and Path(info.filename).suffix.lower() in (".step", ".stp"):
                            out[Path(info.filename).name.lower()] = z.read(info)
            except zipfile.BadZipFile:
                raise RegisterError(f"{name} is not a valid ZIP file.") from None
        elif Path(name).suffix.lower() in (".step", ".stp"):
            out[Path(name).name.lower()] = content
    return out


def bulk_upload(project_id: int, excel: bytes, uploads: list[tuple[str, bytes]]) -> dict:
    try:
        wb = load_workbook(io.BytesIO(excel), read_only=True, data_only=True)
    except Exception as exc:
        raise RegisterError(f"Could not read the Excel file: {exc}") from None
    ws = wb["Parts"] if "Parts" in wb.sheetnames else wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        raise RegisterError("The Excel file is empty.")
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    if "part_name" not in header:
        raise RegisterError("The Parts sheet needs the template columns (download the template).")
    files = _collect_files(uploads)
    project = get_project(project_id)
    variants = {v["code"]: v["id"] for v in project["variants"]}

    # 1. check every row first; nothing is saved if any row is wrong
    records, errors = [], []
    codes_in_sheet = set()
    for i, raw in enumerate(rows[1:], start=2):
        if not any(v not in (None, "") for v in raw):
            continue
        r = {h: raw[j] if j < len(raw) else None for j, h in enumerate(header) if h}
        where = f"Parts row {i}"
        try:
            insert_of = str(r.get("insert_of_part_code") or "").strip()
            raw_vals = {"name": r.get("part_name"), "material": r.get("material"), "process": r.get("process"),
                        "sourcing": r.get("sourcing"), "annual_volume": r.get("annual_volume"),
                        "programme_life_years": r.get("programme_life_years"),
                        "target_piece_cost_inr": r.get("target_piece_cost_inr"), "unit_price": r.get("unit_price"),
                        "price_currency": r.get("price_currency"), "price_source": r.get("price_source"),
                        "price_month": r.get("price_month")}
            if insert_of:
                st = register_setup.settings()
                raw_vals["sourcing"] = raw_vals["sourcing"] or st["insert_default_sourcing"]
                raw_vals["process"] = raw_vals["process"] or st["insert_default_process"]
            vals = _part_values(raw_vals, where)
            vals = {k: v for k, v in vals.items() if v not in (None, "", False)}
            variant = str(r.get("variant") or "").strip().upper()
            if variant and variant not in variants:
                raise RegisterError(f"{where}: variant '{variant}' is not in project {project['code']}. Add it first.")
            qty = _num(r.get("quantity"), f"{where}: quantity") or 1
            fname = str(r.get("file_name") or "").strip()
            code = str(r.get("part_code") or "").strip()
            if code:
                if code in codes_in_sheet:
                    raise RegisterError(f"{where}: part_code '{code}' appears twice in the sheet.")
                codes_in_sheet.add(code)
            records.append({"row": i, "code": code, "values": vals, "rev": str(r.get("rev") or "").strip(),
                            "file_name": fname, "content": files.get(fname.lower()) if fname else None,
                            "variant": variant, "parent": str(r.get("parent_part_code") or "").strip(),
                            "quantity": qty, "change_note": str(r.get("change_note") or "").strip(),
                            "insert_of": insert_of,
                            "insert_quantity": _num(r.get("insert_quantity"), f"{where}: insert_quantity") or 1})
        except RegisterError as exc:
            errors.append(str(exc))
    for rec in records:
        if rec["parent"] and rec["parent"] not in codes_in_sheet:
            with db.engine().connect() as conn:
                if not conn.execute(select(db.parts.c.id).where(db.parts.c.part_code == rec["parent"])).first():
                    errors.append(f"Parts row {rec['row']}: parent_part_code '{rec['parent']}' is not in the sheet or the register.")
        if rec["parent"] and not rec["variant"]:
            errors.append(f"Parts row {rec['row']}: a parent_part_code needs a variant.")
        if rec["insert_of"]:
            if rec["insert_of"] == rec["code"]:
                errors.append(f"Parts row {rec['row']}: a part cannot be its own insert.")
            elif rec["insert_of"] not in codes_in_sheet:
                with db.engine().connect() as conn:
                    if not conn.execute(select(db.parts.c.id).where(db.parts.c.part_code == rec["insert_of"])).first():
                        errors.append(f"Parts row {rec['row']}: insert_of_part_code '{rec['insert_of']}' is not in the "
                                      "sheet or the register.")
    if errors:
        raise RegisterError("Nothing was saved. Please fix: " + " | ".join(errors[:20]) +
                            (f" | … and {len(errors) - 20} more" if len(errors) > 20 else ""))

    # 2. measure the files (outside the transaction)
    for rec in records:
        rec["geometry"] = rec["thumbnail"] = None
        if rec["content"]:
            tmp = file_store.scratch_dir(uuid.uuid4().hex)
            try:
                (tmp / "f.step").write_bytes(rec["content"])
                tree = cad_import.analyse_file(str(tmp / "f.step"), rec["file_name"])
                root = tree["defs"][tree["root"]]
                rec["geometry"], rec["thumbnail"] = root["metrics"], root.get("thumbnail")
            except cad_import.ImportError_ as exc:
                raise RegisterError(f"Nothing was saved. Parts row {rec['row']}: {rec['file_name']}: {exc}") from None
            finally:
                shutil.rmtree(tmp, ignore_errors=True)

    # 3. suggest a process where none is given: material first, then geometry and inserts
    inserts_into: dict[str, float] = {}
    for rec in records:
        if rec["insert_of"]:
            inserts_into[rec["insert_of"]] = inserts_into.get(rec["insert_of"], 0) + rec["insert_quantity"]
    for rec in records:
        rec["reason"], rec["steps"] = "", []
        if not rec["values"].get("process"):
            sug = suggest_process(rec["geometry"], rec["values"].get("material", ""), inserts_into.get(rec["code"], 0))
            if sug["process"]:
                rec["values"]["process"] = sug["process"]
            rec["reason"], rec["steps"] = sug["reason"], sug["steps"]

    # 4. save everything in one transaction
    report = []
    with db.engine().begin() as conn:
        ids = {}
        for rec in records:
            pid, is_new = upsert_part(conn, project_id, rec["code"], rec["values"])
            if not is_new and rec["values"]:
                old = conn.execute(select(db.parts).where(db.parts.c.id == pid)).mappings().one()
                changes = {k: {"old": old[k], "new": v} for k, v in rec["values"].items() if old[k] != v}
                if changes:
                    conn.execute(db.parts.update().where(db.parts.c.id == pid).values(**rec["values"]))
                    _log(conn, "part", pid, "updated", changes | {"source": f"Excel row {rec['row']}"})
            add_revision(conn, pid, rec["rev"], step=rec["content"], file_name=rec["file_name"],
                         thumbnail=rec["thumbnail"], geometry=rec["geometry"], change_note=rec["change_note"],
                         source=f"Excel row {rec['row']}")
            code = conn.execute(select(db.parts.c.part_code).where(db.parts.c.id == pid)).scalar()
            ids[code] = rec["pid"] = pid
            report.append({"row": rec["row"], "part_code": code, "result": "created" if is_new else "new revision",
                           "file": rec["file_name"] if rec["content"] else "File missing",
                           "process": rec["values"].get("process", ""), "process_reason": rec["reason"],
                           "process_steps": rec["steps"]})
        for rec in records:
            if rec["variant"]:
                parent = None
                if rec["parent"]:
                    parent = ids.get(rec["parent"]) or conn.execute(
                        select(db.parts.c.id).where(db.parts.c.part_code == rec["parent"])).scalar()
                set_bom_line(conn, variants[rec["variant"]], parent, rec["pid"], rec["quantity"])
        for rec in records:
            if rec["insert_of"]:
                body = ids.get(rec["insert_of"]) or conn.execute(
                    select(db.parts.c.id).where(db.parts.c.part_code == rec["insert_of"])).scalar()
                try:
                    _set_insert_link(conn, body, rec["pid"], rec["insert_quantity"], notes=f"Excel row {rec['row']}")
                except RegisterError as exc:
                    raise RegisterError(f"Nothing was saved. Parts row {rec['row']}: {exc}") from None
    refresh_costs([r["pid"] for r in records], "imported from Excel")
    return {"rows": report, "inserts_linked": sum(1 for r in records if r["insert_of"]), "created": sum(r["result"] == "created" for r in report),
            "new_revisions": sum(r["result"] == "new revision" for r in report),
            "file_missing": sum(r["file"] == "File missing" for r in report)}


# ---------- dashboard ----------

def _rm_rollup(bom: list[dict]) -> dict:
    """RM cost of a variant: total quantity × latest net RM cost of every costed part in its BOM (inserts excluded)."""
    lines = [b for b in bom if b["kind"] == "bom" and not b["is_assembly"]]
    ids = {b["child_part_id"] for b in lines}
    if not ids:
        return {"rm_cost_inr": None, "rm_parts_costed": 0, "rm_parts_missing": 0, "rm_is_example": False}
    costed = register_setup.sourcing()
    with db.engine().connect() as conn:
        rm = {r["id"]: r["rm_cost_inr"] for r in conn.execute(
            select(db.parts.c.id, db.parts.c.rm_cost_inr).where(db.parts.c.id.in_(ids))).mappings()}
        example = _rm_example_parts(conn, ids)
    total, have, missing = 0.0, 0, 0
    for b in lines:
        if rm.get(b["child_part_id"]) is not None:
            total += b["total_quantity"] * rm[b["child_part_id"]]
            have += 1
        elif costed.get(b["sourcing"], {}).get("is_costed"):
            missing += 1
    return {"rm_cost_inr": total if have else None, "rm_parts_costed": have, "rm_parts_missing": missing,
            "rm_is_example": any(b["child_part_id"] in example and rm.get(b["child_part_id"]) is not None for b in lines)}


def dashboard(project_id: int) -> dict:
    project = get_project(project_id)
    statuses = register_setup.statuses()
    with db.engine().connect() as conn:
        by = lambda col: dict(conn.execute(select(col, func.count()).where(db.parts.c.project_id == project_id)
                                           .group_by(col)).all())
        by_status, by_process, by_sourcing = by(db.parts.c.status), by(db.parts.c.process), by(db.parts.c.sourcing)
        r = db.revisions.alias("cur")
        missing = conn.execute(select(func.count()).select_from(db.parts.outerjoin(r, r.c.id == db.parts.c.current_revision_id))
                               .where(and_(db.parts.c.project_id == project_id, r.c.file_id.is_(None)))).scalar()
        total = sum(by_status.values())
        rm_costed = conn.execute(select(func.count()).where(and_(db.parts.c.project_id == project_id,
                                                                  db.parts.c.rm_cost_inr.is_not(None)))).scalar()
        targets = conn.execute(select(func.count(), func.sum(db.parts.c.target_piece_cost_inr))
                               .where(and_(db.parts.c.project_id == project_id,
                                           db.parts.c.target_piece_cost_inr.is_not(None)))).one()
        tool_rows = conn.execute(select(db.parts.c.id, db.parts.c.process, db.parts.c.sourcing, db.parts.c.tool_cost_inr,
                                        db.parts.c.tool_capex_inr)
                                 .where(and_(db.parts.c.project_id == project_id, db.parts.c.is_assembly.is_(False)))
                                 ).mappings().all()
        tool_example = _tool_example_parts(conn, [r["id"] for r in tool_rows if r["tool_capex_inr"] is not None])
    variants = []
    for v in project["variants"]:
        bom = variant_bom(v["id"])
        variants.append(v | {"bom_lines": len(bom), "unique_parts": len({b["child_part_id"] for b in bom}),
                             "bom_cost_inr": None} | _rm_rollup(bom))
    procs, srcs = register_setup.processes(), register_setup.sourcing()
    corrections = sum(by_status.get(s["code"], 0) for s in statuses if s.get("is_open_correction"))
    costed_tools = [r for r in tool_rows if r["tool_capex_inr"] is not None]
    tools_missing = sum(1 for r in tool_rows if r["tool_capex_inr"] is None
                        and procs.get(r["process"], {}).get("part_type") in tool_cost.PART_TYPES
                        and srcs.get(r["sourcing"], {}).get("is_costed"))
    tool_capex = sum(r["tool_capex_inr"] for r in costed_tools) if costed_tools else None
    budget = project["capex_budget_inr"]
    return {
        "project": project,
        "parts_total": total,
        "file_missing": missing,
        "status": [{"code": s["code"], "name": s["name"], "count": by_status.get(s["code"], 0)} for s in statuses],
        "process": [{"code": c, "name": procs[c]["name"] if c in procs else c or "—", "count": n} for c, n in by_process.items()],
        "sourcing": [{"code": c, "name": srcs[c]["name"] if c in srcs else c or "—", "count": n} for c, n in by_sourcing.items()],
        "variants": variants,
        "parts_with_target": targets[0], "target_piece_cost_sum_inr": targets[1],
        "capex": {"budget_inr": budget, "tool_cost_inr": tool_capex, "machine_cost_inr": None,
                  "tool_parts_costed": len(costed_tools), "tool_parts_missing": tools_missing,
                  "tool_is_example": bool(tool_example),
                  "tool_share_of_budget_pct": round(100 * tool_capex / budget, 1) if tool_capex and budget else None,
                  "tool_note": "Tool CAPEX = Σ over the project's parts of tool cost per tool × tools over the programme "
                               "(latest stored tool cost version; moulds and dies of plastic and die-cast parts)."},
        "open_corrections": corrections,
        "parts_rm_costed": rm_costed,
        "rm_is_example": any(v.get("rm_is_example") for v in variants),
        "tool_is_example": bool(tool_example),
        "example_warning": EXAMPLE_WARNING,
        "not_costed_note": "RM cost is the raw-material cost only (quantity × net RM cost per part, from the latest stored "
                           "version). Process cost and the full BOM cost come from the part-costing stage (roadmap "
                           "stage 8). Tool CAPEX covers the moulds and dies costed so far (stage 4).",
    }


# ---------- exports ----------

def _excel_banner(ws, width: int) -> None:
    """Red warning row at the top of a sheet whose costs use EXAMPLE rates."""
    ws.append([EXAMPLE_WARNING])
    ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=max(width, 2))
    cell = ws.cell(ws.max_row, 1)
    cell.font = Font(bold=True, color="FFFFFF", size=12)
    cell.fill = PatternFill("solid", fgColor="B91C1C")


def export_parts_excel(**filters) -> bytes:
    res = search(**filters, limit=500, offset=0)
    rows = res["rows"]
    offset = 500
    while offset < res["total"]:
        rows += search(**filters, limit=500, offset=offset)["rows"]
        offset += 500
    wb = Workbook()
    ws = wb.active
    ws.title = "Part register"
    wb.properties.title = "Design to Cost — Part register"
    wb.properties.creator = "Design to Cost"
    cols = ["part_code", "name", "project_code", "rev", "material", "process", "sourcing", "status", "annual_volume",
            "programme_life_years", "target_piece_cost_inr", "unit_price", "price_currency", "price_source", "price_month",
            "volume_cm3", "bbox_mm", "file", "country", "net_rm_cost_inr", "tool_country", "tool_cost_per_tool_inr",
            "tool_capex_inr", "updated_at"]
    if any(r["rm_is_example"] or r["tool_is_example"] for r in rows):
        _excel_banner(ws, len(cols))
    ws.append(cols)
    head = ws.max_row
    for cell in ws[head]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
    for r in rows:
        g = r.get("geometry") or {}
        b = g.get("bbox_mm") or {}
        ws.append([r["part_code"], r["name"], r["project_code"], r["rev"], r["material"], r["process"], r["sourcing"],
                   r["status"], r["annual_volume"], r["programme_life_years"], r["target_piece_cost_inr"],
                   r["unit_price"], r["price_currency"], r["price_source"], r["price_month"], g.get("volume_cm3"), f"{b.get('length')} x {b.get('width')} x {b.get('height')}" if b else "",
                   "yes" if r["file_id"] else "File missing", r["country"], r["rm_cost_inr"], r["tool_country"],
                   r["tool_cost_inr"], r["tool_capex_inr"], r["updated_at"]])
        if r["rm_is_example"] or r["tool_is_example"]:
            for cell in ws[ws.max_row]:
                cell.fill = EXAMPLE_FILL
    ws.freeze_panes = f"A{head + 1}"
    ws.auto_filter.ref = f"A{head}:{ws.cell(head, len(cols)).column_letter}{ws.max_row}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_dashboard_pdf(project_id: int) -> bytes:
    from .pdf_report import dashboard_pdf
    return dashboard_pdf(dashboard(project_id))


# ---------- detailed tool design (roadmap stage 27) ----------

def tool_designs(part_id: int, key: str | None = None) -> tuple[dict, dict, list]:
    """(part, revision, [(stage no, design)]) - every die stage of a sheet-metal part or the mould of a moulded part,
    designed in 3D from the current revision and the stored (else live) tool cost."""
    from . import die_cost, mould_design, tool_design
    d = part_detail(part_id, with_rm=False)
    tr = tool_trace(part_id)
    if tr.get("dies"):
        from . import sheet_layout
        d, rev, shape, blank, res = _sheet_context(part_id)
        fp = sheet_layout.flat_pattern(shape, rev["geometry"], blank, res["features"]["t"])
        sl = sheet_layout.strip_layout(fp, res["features"]["t"], len(res["operations"]))
        plan = {p["key"]: p for p in die_cost.plan(res, {"strip": sl})}
        out = []
        for k, die in enumerate(tr["dies"], start=1):
            if key and die["key"] != key:
                continue
            info = die | {f: plan.get(die["key"], {}).get(f) for f in ("progressive", "stations", "pitch", "strip_w")}
            lines = [x for x in tr["lines"] if x.get("die_key") == die["key"]]
            try:
                drawn = bool(fp.get("drawn"))
                if "draw" in die["ops"]:
                    des = tool_design.draw_die_design(d, shape, res["features"]["t"], info, lines)
                elif drawn and die["ops"] in (["trim"], ["pierce"], ["trim", "pierce"]) and "trim" in die["ops"]:
                    des = tool_design.draw_die_design(d, shape, res["features"]["t"], info, lines, trim=True)
                elif die["ops"] == ["bend"] and tool_design.u_shape(fp):
                    des = tool_design.u_die_design(d, fp, res["features"]["t"], info, lines)
                else:
                    des = tool_design.die_design(d, fp, res["features"]["t"], info, lines, sl["best"]["values"])
            except tool_design.DesignError as exc:
                raise RegisterError(f"{die['name']}: {exc}") from None
            out.append((k, des | {"part_type": "sheet_metal"}))
        if key and not out:
            raise RegisterError(f"There is no die '{key}' in this tool cost.")
        return d, rev, out
    part_type = (register_setup.processes().get(d["process"]) or {}).get("part_type")
    if part_type not in ("plastic", "die_casting", "eps", "blow", "foundry"):
        raise RegisterError(f"A detailed tool design is made for sheet-metal dies and for moulds (plastic, die casting, EPS, "
                            f"blow moulding); this part's process is {d['process']}.")
    d, rev, shape = _revision_shape(part_id)
    try:
        axis = _concept_inputs(part_id)[5]["opening_axis"]
    except Exception:
        axis = "z"
    try:
        des = (mould_design.foundry_design(d | {"wall_mm": (rev.get("geometry") or {}).get("wall_mm")}, shape, tr, axis)
               if part_type == "foundry" else mould_design.mould_design(d, shape, tr, part_type, axis))
    except tool_design.DesignError as exc:
        raise RegisterError(str(exc)) from None
    return d, rev, [(1, des)]


def tool_design_summary(part_id: int) -> dict:
    """What the part panel lists: each tool with its parts (no geometry)."""
    d, rev, designs = tool_designs(part_id)
    return {"part_id": part_id, "rev": rev["rev"], "tools": [
        {"stage": k, "key": des["key"], "name": des["name"], "notes": des["notes"], "shut_height": des["shut_height"],
         "die_set": des["die_set"], "rules_are_examples": des["rules_are_examples"],
         "components": [{"no": c.no, "name": c.name, "qty": c.qty, "kind": c.kind, "material": c.material, "hardness": c.hardness,
                         "std": c.std or c.make, "size": [round(v, 1) for v in c.size]} for c in des["components"] if c.kind != "ref"]}
        for k, des in designs]}


def tool_design_file(part_id: int, kind: str, key: str | None = None) -> tuple[bytes, str, str]:
    """(content, media type, file name) of a tool design download: zip (release pack), step (assembly), pdf (drawing
    set) or png (iso picture)."""
    from . import render3d, tool_design, tool_drawings
    d, rev, designs = tool_designs(part_id, key)
    code = d["part_code"]
    if kind == "zip":
        tag = f"_{tool_drawings._safe(key)}" if key else ""
        return tool_drawings.release_zip(d, designs, rev["rev"]), "application/zip", f"{code}_tool_design{tag}.zip"
    k, des = designs[0]
    stem = f"{code}_T{k:02d}_{tool_drawings._safe(des['key'])}"
    if kind == "step":
        return tool_design.step_bytes(tool_design.assembly(des, stem)), "application/step", f"{stem}_assembly.step"
    if kind == "pdf":
        return tool_drawings.design_pdf(d, des, k, rev["rev"]), "application/pdf", f"{stem}_drawing_set.pdf"
    if kind == "png":
        img = render3d.render(render3d.design_items(des, tool_design.COLORS), size=(900, 650))
        return render3d.png(img), "image/png", f"{stem}.png"
    raise RegisterError(f"Unknown tool design download '{kind}'.")


def forming_report(part_id: int, kind: str = "pdf") -> tuple[bytes, str, str]:
    """Sheet-metal forming simulation report of the current revision: (content, media type, file name); kind pdf or video."""
    from . import sim_report
    d, rev, fp, sl, res = _layout(part_id)
    res = res | {"blank": res.get("blank") or {}}
    code = d["part_code"]
    if kind == "pdf":
        return sim_report.pdf(res, fp, d, rev["rev"]), "application/pdf", f"{code}_rev{rev['rev']}_forming_simulation.pdf"
    data, ext = sim_report.video(res, fp, d)
    return data, "video/mp4" if ext == "mp4" else "image/gif", f"{code}_rev{rev['rev']}_forming_simulation.{ext}"


# ---------- 3D correction after the simulation (roadmap stage 27) ----------

def _correction(part_id: int, radius: float | None, hole_d: float | None, draft: float | None = None):
    """(part, revision, old shape, new shape, changes, defaults) of the current revision: sheet metal - bend radius and
    small holes; moulded parts (plastic, die casting, EPS, blow, foundry) - draft on the walls along the opening direction."""
    from . import part_fix
    d0 = part_detail(part_id, with_rm=False)
    pt = (register_setup.processes().get(d0["process"]) or {}).get("part_type")
    if pt != "sheet_metal":
        if pt not in ("plastic", "die_casting", "eps", "blow", "foundry"):
            raise RegisterError("The software corrects sheet-metal parts (bend radius, holes) and moulded / cast parts (draft).")
        sim = simulation_result(part_id)
        d, rev, shape = _revision_shape(part_id)
        min_draft = float(sim["results"].get("min_draft_deg") or 1) + 0.5     # margin over the minimum of the check
        defaults = {"draft_deg": min_draft, "axis": sim["opening_axis"], "low_draft_cm2": sim["maps"]["draft_low_cm2"],
                    "corrections": sim["corrections"], "kind": "mould"}
        if draft is None and sim["maps"]["draft_low_cm2"] <= 0:
            raise RegisterError("The simulation found nothing the software can correct in 3D (draft): correct walls, ribs and thick "
                                "sections in CAD.")
        try:
            new, changes = part_fix.add_draft(shape, sim["opening_axis"], draft if draft is not None else min_draft)
        except part_fix.FixError as exc:
            raise RegisterError(str(exc)) from None
        return d, rev, shape, new, changes, defaults
    d, rev, shape, blank, res = _sheet_context(part_id)
    t = res["features"]["t"]
    mv = res["material_values"]
    defaults = {"radius_mm": round(mv.get("min_bend_r_factor", 1) * t, 3), "hole_mm": round(mv.get("min_hole_factor", 1) * t, 3),
                "t": t, "checks": res.get("checks", []), "kind": "sheet"}
    new, changes = shape, []
    failed = {c["check"] for c in res.get("checks", []) if not c["passed"]}
    try:
        if radius is not None or "bend_radius" in failed:
            r = radius if radius is not None else defaults["radius_mm"]
            new, ch = part_fix.fix_bend_radius(new, t, r, only_below=r)
            changes += [{"what": "inside bend radius"} | c for c in ch]
        if hole_d is not None or ("hole_diameter" in failed and radius is None):
            h = hole_d if hole_d is not None else defaults["hole_mm"]
            new, ch = part_fix.fix_small_holes(new, t, h)
            changes += [{"what": "hole diameter"} | c for c in ch]
    except part_fix.FixError as exc:
        raise RegisterError(str(exc)) from None
    if not changes:
        raise RegisterError("The simulation found nothing the software can correct in 3D (bend radius, hole size): correct the model in CAD.")
    return d, rev, shape, new, changes, defaults


def correction_preview(part_id: int, radius: float | None = None, hole_d: float | None = None, draft: float | None = None) -> dict:
    d, rev, old, new, changes, defaults = _correction(part_id, radius, hole_d, draft)
    with db.engine().connect() as conn:
        existing = [r for (r,) in conn.execute(select(db.revisions.c.rev).where(db.revisions.c.part_id == part_id))]
    note = "; ".join(f"{c['what']} {c['from_mm']:g} → {c['to_mm']:g}" + ("°" if c["what"].startswith("draft") else " mm")
                     for c in changes[:3]) + (" …" if len(changes) > 3 else "")
    return {"part_id": part_id, "rev": rev["rev"], "next_rev": _next_rev(existing), "changes": changes, "defaults": defaults,
            "change_note": f"Corrected after simulation: {note}", "volume_before_cm3": round(old.Volume() / 1000, 3),
            "volume_after_cm3": round(new.Volume() / 1000, 3)}


def correction_file(part_id: int, kind: str, radius: float | None = None, hole_d: float | None = None,
                    draft: float | None = None) -> tuple[bytes, str, str]:
    """The corrected model as STEP, or a picture (before | after) as PNG."""
    from . import render3d, tool_design
    d, rev, old, new, changes, defaults = _correction(part_id, radius, hole_d, draft)
    if kind == "step":
        return tool_design.step_bytes(new), "application/step", f"{d['part_code']}_corrected_from_rev{rev['rev']}.step"
    import math
    from PIL import Image, ImageDraw
    from . import part_fix
    items_old = [(*render3d.mesh(old, 0.05), (0.6, 0.62, 0.66))]
    items_new = [(*render3d.mesh(new, 0.05), (0.35, 0.55, 0.8))]
    sph = render3d.sphere_of(items_old)
    img = Image.new("RGB", (1200, 840), "white")
    img.paste(render3d.render(items_old, size=(600, 420), sphere=sph, margin=0.05), (0, 0))
    img.paste(render3d.render(items_new, size=(600, 420), sphere=sph, margin=0.05), (600, 0))
    dr = ImageDraw.Draw(img)
    dr.text((10, 8), f"rev {rev['rev']} (now)", fill=(60, 60, 60))
    dr.text((610, 8), "corrected", fill=(60, 60, 60))
    if defaults.get("kind") == "mould":          # draft: the side elevation (square to the opening direction) shows the taper
        ax = defaults["axis"]
        az, el = {"z": (-90, 0), "x": (-90, 90), "y": (0, 0)}[ax]
        img.paste(render3d.render(items_old, size=(600, 420), az=az, el=el, sphere=sph, margin=0.05), (0, 420))
        img.paste(render3d.render(items_new, size=(600, 420), az=az, el=el, sphere=sph, margin=0.05), (600, 420))
        dr.text((10, 428), "side view: walls without draft", fill=(60, 60, 60))
        dr.text((610, 428), f"side view: {changes[0]['to_mm']:g}° draft", fill=(60, 60, 60))
        return render3d.png(img), "image/png", f"{d['part_code']}_correction.png"
    t = _sheet_context(part_id)[4]["features"]["t"]
    bs = part_fix.bends(old, t)
    if any(c["what"] == "inside bend radius" for c in changes) and bs:
        b = min(bs, key=lambda x: x["r"])        # zoom on the bend, looking along its axis (profile of the bend)
        a = b["axis"]
        centre = b["centre"] + a * (b["a0"] + b["a1"]) / 2
        az, el = math.degrees(math.atan2(a[1], a[0])), math.degrees(math.asin(max(-1, min(1, a[2]))))
        zoom = (centre, max(6 * t, 10.0))
        img.paste(render3d.render(items_old, size=(600, 420), az=az, el=el, sphere=zoom, margin=0.02), (0, 420))
        img.paste(render3d.render(items_new, size=(600, 420), az=az, el=el, sphere=zoom, margin=0.02), (600, 420))
        dr.text((10, 428), f"bend profile: inside radius {b['r']:g} mm", fill=(60, 60, 60))
        dr.text((610, 428), "bend profile: corrected", fill=(60, 60, 60))
    else:
        img = img.crop((0, 0, 1200, 420))
    return render3d.png(img), "image/png", f"{d['part_code']}_correction.png"


def correction_release(part_id: int, radius: float | None = None, hole_d: float | None = None, rev: str = "",
                       change_note: str = "", draft: float | None = None) -> dict:
    """Release the corrected model as a new revision (design loop → Revised, costs recalculated)."""
    from . import tool_design
    d, cur, old, new, changes, _ = _correction(part_id, radius, hole_d, draft)
    note = change_note or correction_preview(part_id, radius, hole_d, draft)["change_note"]
    data = tool_design.step_bytes(new)
    return new_revision_from_upload(part_id, data, f"{d['part_code']}_corrected.step", rev, note)


def simulation_video(part_id: int) -> tuple[bytes, str, str]:
    """Video of the simulation & analysis maps of a moulded part (plastic, die casting)."""
    from . import sim_report
    d, rev, res = _simulate(part_id)
    data, ext = sim_report.moulding_video(res, d)
    return data, "video/mp4" if ext == "mp4" else "image/gif", f"{d['part_code']}_rev{rev['rev']}_simulation.{ext}"
