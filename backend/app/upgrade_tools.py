"""Merge the master data of a new version into yours, so an update never loses your rates and edits.

    python -m backend.app.upgrade_tools merge --mine <your master_data folder> --new <new version's master_data folder>
    python -m backend.app.upgrade_tools fingerprints        (developers: rebuild shipped_rows.json from the git history)

Rules, table by table (costing_db.sqlite) and sheet by sheet (master_data.xlsx):
    a table, sheet or column only the new version has   added (with its rows)
    a row only the new version has (by its unique key)  added (new machines, formulas, example rates ...)
    a row you never changed                             takes the new version's values (fixes in example rows reach you)
    a row you changed                                   stays yours; listed in the report when the new version differs
    a row only you have                                 stays
    rates (the rate library)                            only the new version's (category, item, country, month) you do not
                                                        have are added; your versions are never changed
"A row you never changed" = the row is identical to a row some released version shipped (shipped_rows.json holds a
short fingerprint of every row ever shipped), so no record of your old version is needed.
The update kit (windows/dtc_manager.py) backs up before it calls this.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import shutil
import sqlite3
import subprocess
import sys
import tempfile
from contextlib import closing
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
FINGERPRINTS = Path(__file__).with_name("shipped_rows.json")
DB_FILE, XLSX_FILE = "costing_db.sqlite", "master_data.xlsx"
XLSX_KEYS = {"Parameters": ("part_type", "parameter"), "Materials": ("part_type", "material")}
RATE_KEY = ("category", "item", "country", "month")
SKIP_FINGERPRINT = {"rates"}                       # rates are merged by key only


# ---------- fingerprints ----------

def _norm(v):
    if isinstance(v, bool):
        return int(v)
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v


def row_hash(table: str, row: dict) -> str:
    items = sorted((k, _norm(v)) for k, v in row.items() if k != "id" and _norm(v) is not None)
    return hashlib.sha1(json.dumps([table, items], default=str).encode()).hexdigest()[:16]


def _db_rows(con: sqlite3.Connection, table: str, schema: str = "main") -> list[dict]:
    cur = con.execute(f'SELECT * FROM {schema}."{table}"')
    names = [d[0] for d in cur.description]
    return [dict(zip(names, r)) for r in cur.fetchall()]


def _tables(con: sqlite3.Connection, schema: str = "main") -> dict[str, str]:
    return {n: s for n, s in con.execute(f"SELECT name, sql FROM {schema}.sqlite_master WHERE type = 'table' AND name NOT LIKE 'sqlite_%'")}


def _xlsx_sheets(p: Path) -> dict[str, list[dict]]:
    from openpyxl import load_workbook
    from . import master_data
    wb = load_workbook(p, read_only=True, data_only=True)
    out = {}
    for sheet, columns in master_data.SHEETS.items():
        if sheet not in wb.sheetnames:
            continue
        rows = list(wb[sheet].iter_rows(values_only=True))
        header = [str(h).strip() if h is not None else "" for h in rows[0]] if rows else []
        out[sheet] = [{c: (row[header.index(c)] if c in header and header.index(c) < len(row) else None) for c in columns}
                      for row in rows[1:] if any(v not in (None, "") for v in row)]
    wb.close()
    return out


def fingerprints_of(folder: Path) -> set[str]:
    out: set[str] = set()
    db = folder / DB_FILE
    if db.exists():
        with closing(sqlite3.connect(db)) as con:
            for t in _tables(con):
                if t not in SKIP_FINGERPRINT:
                    out.update(row_hash(t, r) for r in _db_rows(con, t))
    x = folder / XLSX_FILE
    if x.exists():
        for sheet, rows in _xlsx_sheets(x).items():
            out.update(row_hash("xlsx:" + sheet, r) for r in rows)
    return out


def shipped() -> set[str]:
    return set(json.loads(FINGERPRINTS.read_text())["rows"]) if FINGERPRINTS.exists() else set()


def build_fingerprints(root: Path = ROOT) -> dict:
    """Every row of master_data/ in every commit of the git history, plus the working copy."""
    files = [f"master_data/{DB_FILE}", f"master_data/{XLSX_FILE}"]
    commits = subprocess.run(["git", "-C", str(root), "log", "--format=%H", "--", *files], capture_output=True, text=True, check=True).stdout.split()
    rows = fingerprints_of(root / "master_data")
    with tempfile.TemporaryDirectory() as tmp:
        for c in commits:
            d = Path(tmp) / c
            d.mkdir()
            for f in files:
                blob = subprocess.run(["git", "-C", str(root), "show", f"{c}:{f}"], capture_output=True)
                if blob.returncode == 0:
                    (d / Path(f).name).write_bytes(blob.stdout)
            try:
                rows |= fingerprints_of(d)
            except Exception:                       # an early commit with a different workbook layout
                pass
    data = {"note": "Fingerprints of every master-data row ever shipped; see upgrade_tools.py. Rebuild: python -m backend.app.upgrade_tools fingerprints",
            "commits": len(commits), "rows": sorted(rows)}
    FINGERPRINTS.write_text(json.dumps(data, indent=0) + "\n")
    return {"commits": len(commits), "rows": len(rows)}


# ---------- merge ----------

def _unique_key(con: sqlite3.Connection, table: str, schema: str) -> tuple[str, ...] | None:
    for idx in con.execute(f'PRAGMA {schema}.index_list("{table}")').fetchall():
        if idx[3] in ("u", "pk") and idx[2]:
            return tuple(r[2] for r in con.execute(f'PRAGMA {schema}.index_info("{idx[1]}")'))
    pk = [r[1] for r in con.execute(f'PRAGMA {schema}.table_info("{table}")') if r[5]]
    return tuple(pk) if pk and pk != ["id"] else None


def _same(a: dict, b: dict, cols) -> bool:
    return all(_norm(a.get(c)) == _norm(b.get(c)) for c in cols)


def merge_db(mine: Path, new: Path, known: set[str]) -> dict:
    rep = {"tables_added": [], "columns_added": [], "rows_added": {}, "rows_updated": {}, "kept_yours": {}, "rates_added": 0}
    with closing(sqlite3.connect(mine)) as con:
        con.execute("ATTACH DATABASE ? AS n", (str(new),))
        have = _tables(con)
        for t, sql in _tables(con, "n").items():
            if t not in have:
                con.execute(sql)
                con.execute(f'INSERT INTO main."{t}" SELECT * FROM n."{t}"')
                for (isql,) in con.execute("SELECT sql FROM n.sqlite_master WHERE type = 'index' AND tbl_name = ? AND sql IS NOT NULL", (t,)).fetchall():
                    con.execute(isql)
                rep["tables_added"].append(t)
                continue
            mine_cols = {r[1] for r in con.execute(f'PRAGMA main.table_info("{t}")')}
            new_info = con.execute(f'PRAGMA n.table_info("{t}")').fetchall()
            for r in new_info:
                if r[1] not in mine_cols:
                    con.execute(f'ALTER TABLE main."{t}" ADD COLUMN "{r[1]}" {r[2] or ""}')
                    rep["columns_added"].append(f"{t}.{r[1]}")
            cols = [r[1] for r in new_info if r[1] != "id"]
            col_sql = ", ".join(f'"{c}"' for c in cols)
            if t == "rates":
                cond = " AND ".join(f"m.{k} IS r.{k}" for k in RATE_KEY)
                cur = con.execute(f'INSERT INTO main.rates ({col_sql}) SELECT {", ".join("r." + c for c in cols)} FROM n.rates r '
                                  f'WHERE NOT EXISTS (SELECT 1 FROM main.rates m WHERE {cond})')
                rep["rates_added"] = cur.rowcount
                continue
            key = _unique_key(con, t, "n")
            if not key:
                continue
            mine_rows = {tuple(_norm(r[k]) for k in key): r for r in _db_rows(con, t)}
            for r in _db_rows(con, t, "n"):
                k = tuple(_norm(r[c]) for c in key)
                m = mine_rows.get(k)
                if m is None:
                    con.execute(f'INSERT INTO main."{t}" ({col_sql}) VALUES ({", ".join("?" * len(cols))})', [r[c] for c in cols])
                    rep["rows_added"][t] = rep["rows_added"].get(t, 0) + 1
                elif not _same(m, r, cols):
                    if row_hash(t, m) in known:
                        where = ("id = ?", [m["id"]]) if "id" in m else (" AND ".join(f'"{c}" IS ?' for c in key), [m[c] for c in key])
                        con.execute(f'UPDATE main."{t}" SET {", ".join(f"{chr(34)}{c}{chr(34)} = ?" for c in cols)} WHERE {where[0]}',
                                    [r[c] for c in cols] + where[1])
                        rep["rows_updated"][t] = rep["rows_updated"].get(t, 0) + 1
                    else:
                        rep["kept_yours"].setdefault(t, []).append(" / ".join(str(x) for x in k))
        con.commit()
    return rep


def merge_xlsx(mine: Path, new: Path, known: set[str]) -> dict:
    from . import master_data
    rep = {"rows_added": {}, "rows_updated": {}, "kept_yours": {}}
    a, b = _xlsx_sheets(mine), _xlsx_sheets(new)
    out = {}
    for sheet in master_data.SHEETS:
        key = XLSX_KEYS[sheet]
        rows = [dict(r) for r in a.get(sheet, [])]
        by_key = {tuple(_norm(r[c]) for c in key): r for r in rows}
        for r in b.get(sheet, []):
            k = tuple(_norm(r[c]) for c in key)
            m = by_key.get(k)
            if m is None:
                rows.append(dict(r))
                rep["rows_added"][sheet] = rep["rows_added"].get(sheet, 0) + 1
            elif not _same(m, r, r.keys()):
                if row_hash("xlsx:" + sheet, m) in known:
                    m.update(r)
                    rep["rows_updated"][sheet] = rep["rows_updated"].get(sheet, 0) + 1
                else:
                    rep["kept_yours"].setdefault(sheet, []).append(" / ".join(str(x) for x in k))
        out[sheet] = rows
    master_data.write(out, mine)
    return rep


def merge(mine: Path, new: Path, known: set[str] | None = None) -> dict:
    """Merge new/ into mine/ (both master_data folders). Works on copies; mine/ is replaced only when both files merged."""
    known = shipped() if known is None else known
    mine.mkdir(parents=True, exist_ok=True)
    report = {}
    with tempfile.TemporaryDirectory(dir=mine) as tmp:
        staged = {}
        for name, fn in ((DB_FILE, merge_db), (XLSX_FILE, merge_xlsx)):
            src, nf = mine / name, new / name
            if not nf.exists():
                continue
            work = Path(tmp) / name
            if src.exists():
                shutil.copy2(src, work)
                report[name] = fn(work, nf, known)
            else:
                shutil.copy2(nf, work)
                report[name] = {"copied": "you had none; the new version's file is used"}
            staged[name] = work
        for name, work in staged.items():
            shutil.copy2(work, mine / name)
    for f in new.iterdir():                        # any other master-data file you do not have yet
        if f.is_file() and f.name not in (DB_FILE, XLSX_FILE) and not (mine / f.name).exists():
            shutil.copy2(f, mine / f.name)
    return report


def report_text(rep: dict) -> str:
    lines = []
    for name, r in rep.items():
        lines.append(f"{name}:")
        if "copied" in r:
            lines.append(f"  {r['copied']}")
            continue
        for t in r.get("tables_added", []):
            lines.append(f"  new table added: {t}")
        for c in r.get("columns_added", []):
            lines.append(f"  new column added: {c}")
        for t, n in r.get("rows_added", {}).items():
            lines.append(f"  {t}: {n} new row(s) added")
        for t, n in r.get("rows_updated", {}).items():
            lines.append(f"  {t}: {n} row(s) you never changed updated to the new version")
        if r.get("rates_added"):
            lines.append(f"  rate library: {r['rates_added']} new rate row(s) added (your rates unchanged)")
        for t, keys in r.get("kept_yours", {}).items():
            lines.append(f"  {t}: {len(keys)} row(s) YOU changed were kept; the new version has other values for: {', '.join(keys[:20])}"
                         + (" ..." if len(keys) > 20 else ""))
        if len(lines) and lines[-1] == f"{name}:":
            lines.append("  nothing to merge")
    return "\n".join(lines)


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(prog="python -m backend.app.upgrade_tools")
    sub = ap.add_subparsers(dest="cmd", required=True)
    m = sub.add_parser("merge")
    m.add_argument("--mine", required=True)
    m.add_argument("--new", required=True)
    sub.add_parser("fingerprints")
    a = ap.parse_args(argv)
    if a.cmd == "fingerprints":
        print(json.dumps(build_fingerprints()))
    else:
        print(report_text(merge(Path(a.mine), Path(a.new))))
    return 0


if __name__ == "__main__":
    sys.exit(main())
