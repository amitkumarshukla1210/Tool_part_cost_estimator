"""One neutral report structure, three downloads: Excel, PDF and PowerPoint (roadmap stage 20).

report = {title, subtitle, example: bool, kpis: [(label, value)], sections: [{title, text: [str], table: {headers, rows},
          images: [(caption, png bytes)]}]}
"""

from __future__ import annotations

import io
import re
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
MAX_PPT_ROWS = 11


def _hex(name: str) -> str:
    css = (ROOT / "theme" / "dtc-theme.css").read_text(encoding="utf-8")
    return re.search(rf"--{name}:\s*(#[0-9a-fA-F]{{6}})", css).group(1).lstrip("#")


def _cell(v):
    if isinstance(v, float):
        return round(v, 3)
    return v


# ------------------------------------------------------------ Excel

def to_xlsx(report: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XlImage
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    head_fill = PatternFill("solid", fgColor=_hex("table-head-dark"))
    head_font = Font(bold=True, color=_hex("table-head-dark-text"))
    ws = wb.active
    ws.title = "Summary"
    r = 1
    if report.get("example"):
        ws.cell(r, 1, "Uses example rates – not for negotiation").font = Font(bold=True, color="FFFFFF")
        ws.cell(r, 1).fill = PatternFill("solid", fgColor="B91C1C")
        r += 2
    ws.cell(r, 1, report["title"]).font = Font(bold=True, size=14, color=_hex("brand-dark"))
    ws.cell(r + 1, 1, report["subtitle"]).font = Font(color=_hex("grey"))
    r += 3
    for label, value in report["kpis"]:
        ws.cell(r, 1, label).font = Font(bold=True)
        ws.cell(r, 2, _cell(value))
        r += 1
    r += 1
    for sec in report["sections"]:
        for line in sec.get("text", []):
            ws.cell(r, 1, f"{sec['title']}: {line}")
            ws.cell(r, 1).alignment = Alignment(wrap_text=False)
            r += 1
    ws.column_dimensions["A"].width = 48
    ws.column_dimensions["B"].width = 30
    used = {"Summary"}
    pics = wb.create_sheet("Pictures")
    prow = 1
    for sec in report["sections"]:
        tbl = sec.get("table")
        if tbl and tbl["rows"]:
            name = re.sub(r"[\[\]:*?/\\]", " ", sec["title"])[:28]
            base, k = name, 2
            while name in used:
                name, k = f"{base[:25]} {k}", k + 1
            used.add(name)
            sh = wb.create_sheet(name)
            sh.append(list(tbl["headers"]))
            for c in sh[1]:
                c.fill, c.font = head_fill, head_font
            for row in tbl["rows"]:
                sh.append([_cell(x) for x in row])
            for i, h in enumerate(tbl["headers"], start=1):
                width = max(len(str(h)), *(len(str(rw[i - 1])) for rw in tbl["rows"][:60]))
                sh.column_dimensions[sh.cell(1, i).column_letter].width = min(max(10, width + 2), 60)
            sh.freeze_panes = "A2"
        for cap, png in sec.get("images", []):
            pics.cell(prow, 1, f"{sec['title']}: {cap}").font = Font(bold=True)
            img = XlImage(io.BytesIO(png))
            img.width, img.height = img.width * 0.55, img.height * 0.55
            pics.add_image(img, f"A{prow + 1}")
            prow += int(img.height / 20) + 3
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------ PDF

def to_pdf(report: dict) -> bytes:
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Image, KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer
    from .pdf_report import _example_banner, _table, _theme
    theme = _theme()
    ss = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=ss["Title"], textColor=theme["ink"], fontSize=18, alignment=0, spaceAfter=2)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], textColor=theme["ink"], fontSize=12, spaceBefore=10, spaceAfter=4)
    note = ParagraphStyle("n", parent=ss["Normal"], textColor=theme["grey"], fontSize=8)
    body = ParagraphStyle("b", parent=ss["Normal"], fontSize=9, leading=12)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=14 * mm, rightMargin=14 * mm, topMargin=12 * mm,
                            bottomMargin=12 * mm, title=f"Design to Cost - {report['title']}", author="Design to Cost")
    story = [Paragraph("DESIGN TO COST · DESIGN PACK", ParagraphStyle("k", parent=ss["Normal"], textColor=theme["brand"], fontSize=8)),
             Paragraph(report["title"], h1), Paragraph(report["subtitle"], note)]
    if report.get("example"):
        story += [Spacer(1, 2 * mm), _example_banner(ss, "Rates and rules behind these numbers are still EXAMPLE values.")]
    story.append(Spacer(1, 3 * mm))
    story.append(_table([["Result", "Value"]] + [[k, str(_cell(v))] for k, v in report["kpis"]], theme, [90 * mm, 90 * mm]))
    for sec in report["sections"]:
        block = [Paragraph(sec["title"], h2)]
        for line in sec.get("text", []):
            block.append(Paragraph(line, body))
        story.append(KeepTogether(block))
        tbl = sec.get("table")
        if tbl and tbl["rows"]:
            n = len(tbl["headers"])
            width = 265 * mm / n
            story.append(_table([list(map(str, tbl["headers"]))] + [[str(_cell(x)) for x in r] for r in tbl["rows"]], theme, [width] * n))
        imgs = sec.get("images", [])
        row = []
        for cap, png in imgs:
            row.append([Image(io.BytesIO(png), width=84 * mm, height=63 * mm), Paragraph(cap, note)])
        if row:
            from reportlab.platypus import Table
            for i in range(0, len(row), 3):
                chunk = row[i:i + 3]
                story.append(Table([[c[0] for c in chunk], [c[1] for c in chunk]], colWidths=[88 * mm] * len(chunk)))
    doc.build(story)
    return buf.getvalue()


# ------------------------------------------------------------ PowerPoint

def to_pptx(report: dict) -> bytes:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt
    rgb = lambda n: RGBColor.from_string(_hex(n).upper())
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]

    def band(slide, text, sub=""):
        shp = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(0.9))
        shp.fill.solid(); shp.fill.fore_color.rgb = rgb("table-head-dark"); shp.line.fill.background()
        tf = shp.text_frame
        tf.text = text
        tf.paragraphs[0].font.size, tf.paragraphs[0].font.bold = Pt(26), True
        tf.paragraphs[0].font.color.rgb = rgb("white")
        if sub:
            p = tf.add_paragraph(); p.text = sub; p.font.size = Pt(11); p.font.color.rgb = rgb("blue-soft")
        if report.get("example"):
            w = slide.shapes.add_textbox(Inches(8.2), Inches(7.05), Inches(5), Inches(0.35))
            w.text_frame.text = "Uses example rates – not for negotiation"
            w.text_frame.paragraphs[0].font.size = Pt(10); w.text_frame.paragraphs[0].font.bold = True
            w.text_frame.paragraphs[0].font.color.rgb = rgb("danger")

    def table(slide, headers, rows, top, height):
        n = len(headers)
        shape = slide.shapes.add_table(len(rows) + 1, n, Inches(0.5), Inches(top), Inches(12.3), Inches(height))
        t = shape.table
        for j, h in enumerate(headers):
            c = t.cell(0, j); c.text = str(h)
            c.text_frame.paragraphs[0].font.size, c.text_frame.paragraphs[0].font.bold = Pt(11), True
            c.fill.solid(); c.fill.fore_color.rgb = rgb("table-head-dark")
            c.text_frame.paragraphs[0].font.color.rgb = rgb("table-head-dark-text")
        for i, r in enumerate(rows, start=1):
            for j, v in enumerate(r):
                c = t.cell(i, j); c.text = str(_cell(v))
                c.text_frame.paragraphs[0].font.size = Pt(10)
                c.fill.solid(); c.fill.fore_color.rgb = rgb("table-row-alt") if i % 2 == 0 else rgb("white")

    s = prs.slides.add_slide(blank)
    band(s, report["title"], report["subtitle"])
    table(s, ["Result", "Value"], [[k, v] for k, v in report["kpis"][:12]], 1.3, 0.4 * min(len(report["kpis"]), 12) + 0.4)
    for sec in report["sections"]:
        imgs, tbl = sec.get("images", []), sec.get("table")
        s = prs.slides.add_slide(blank)
        band(s, sec["title"])
        top = 1.1
        for line in sec.get("text", [])[:4]:
            tb = s.shapes.add_textbox(Inches(0.5), Inches(top), Inches(12.3), Inches(0.5))
            tb.text_frame.word_wrap = True
            tb.text_frame.text = re.sub(r"<[^>]+>", "", line)
            tb.text_frame.paragraphs[0].font.size = Pt(12)
            top += 0.45
        if imgs:
            x = 0.5
            for cap, png in imgs[:3]:
                w = 12.3 / min(len(imgs), 3) - 0.2
                s.shapes.add_picture(io.BytesIO(png), Inches(x), Inches(top + 0.1), width=Inches(w))
                tb = s.shapes.add_textbox(Inches(x), Inches(top + 0.1 + w * 0.75), Inches(w), Inches(0.3))
                tb.text_frame.text = cap
                tb.text_frame.paragraphs[0].font.size = Pt(9)
                x += w + 0.2
            top += 0.2 + (12.3 / min(len(imgs), 3)) * 0.75 + 0.4
        if tbl and tbl["rows"]:
            rows = tbl["rows"][:MAX_PPT_ROWS]
            table(s, tbl["headers"], rows, min(top, 6.0) if imgs else top, 0.34 * (len(rows) + 1))
            if len(tbl["rows"]) > MAX_PPT_ROWS:
                tb = s.shapes.add_textbox(Inches(0.5), Inches(7.0), Inches(7), Inches(0.3))
                tb.text_frame.text = f"{len(tbl['rows']) - MAX_PPT_ROWS} more rows are in the Excel download."
                tb.text_frame.paragraphs[0].font.size = Pt(9)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
