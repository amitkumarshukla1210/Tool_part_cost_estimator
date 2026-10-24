"""Stage 28: technical data on the tool drawings (review, change, download / upload) and the title block with logo, release data and revision history."""
import io
import zipfile

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader

from backend.app import drawing_data, mould_db
from backend.app.main import app
from tests.test_rm_cost import add_part, project

client = TestClient(app)


def mould(code="DD-1"):
    return add_part(project(code + "P"), "plastic_cover.step", "ABS", "injection_moulding", code)


def design(pid):
    r = client.post(f"/api/parts/{pid}/tool-designs", json={})
    assert r.status_code == 200, r.text
    return r.json()


def first_made(d):
    return next(c for c in d["components"] if c["kind"] == "made")


def pdf_text(content):
    return "\n".join(p.extract_text() for p in PdfReaderPages(content))


def PdfReaderPages(content):
    return PdfReader(io.BytesIO(content)).pages


def test_general_tolerance_tables_are_master_data_and_checked():
    rows = mould_db.tables()["general_tolerances"]
    assert {r["tol_class"] for r in rows} == {"f", "m", "c", "v"} and all(r["is_example"] for r in rows)
    assert drawing_data.tolerance_class("ISO 2768-c; fits H7", None) == "c" and drawing_data.tolerance_class("", None) == "m"
    assert drawing_data.fmt_angle(90) == "±1°30′" and drawing_data.fmt_angle(10) == "±0°10′"
    bad = mould_db.tables()
    bad["general_tolerances"].append({"tol_class": "x", "kind": "linear", "from_mm": 1, "to_mm": 2, "value": 0.1, "is_example": 1, "notes": ""})
    try:
        mould_db.write(bad)
        raise AssertionError("class x accepted")
    except mould_db.MouldDbError as exc:
        assert "class must be" in str(exc)


def test_drawing_data_defaults_and_change_follow_into_bill_drawing_and_history():
    d = design(mould())
    did, c = d["id"], first_made(d)
    data = client.get(f"/api/tool-designs/{did}/drawing-data").json()
    p = next(x for x in data["parts"] if x["part_no"] == c["part_no"])
    assert data["editable"] and p["effective"]["steel"] == p["defaults"]["steel"] and p["effective"]["tolerance_class"] in "fmcv" and not p["changed"]
    other = next(s for s in data["steels"] if s != p["defaults"]["steel"])
    r = client.put(f"/api/tool-designs/{did}/drawing-data", json={"items": [{"part_no": c["part_no"], "fields": {"steel": other, "hardness": "58 HRC", "tolerance_class": "f", "finish_ra": 0.8},
                                                                          "notes": ["Polish the cavity face to Ra 0.4", "Mark with the die number"]}]})
    assert r.status_code == 200 and len(r.json()["changed"]) == 1 and "material grade" in r.json()["changed"][0]["summary"]
    again = client.put(f"/api/tool-designs/{did}/drawing-data", json={"items": [{"part_no": c["part_no"], "fields": {"steel": other, "hardness": "58 HRC", "tolerance_class": "f", "finish_ra": 0.8},
                                                                               "notes": ["Polish the cavity face to Ra 0.4", "Mark with the die number"]}]})
    assert again.json()["changed"] == [] and again.json()["unchanged"] == 1                      # nothing new, no extra history
    g = client.get(f"/api/tool-designs/{did}").json()
    made = next(m for m in g["bill"]["made"] if m["part_no"] == c["part_no"])
    assert made["info"]["steel"] == other and made["info"]["hardness"] == "58 HRC" and made["info"]["finish_ra"] == 0.8 and made["info"]["tolerance_class"] == "f"
    comp_id = c["id"]
    pdf = client.get(f"/api/tool-designs/{did}/components/{comp_id}/drawing.pdf")
    assert pdf.status_code == 200
    text = pdf_text(pdf.content)
    for need in ("Design to Cost", other, "58 HRC", "ISO 2768-f", "Polish the cavity face to Ra 0.4", "NOT RELEASED", "REVISION HISTORY", "Third-angle", "TOLERANCES ISO 2768", "Technical data changed"):
        assert need.lower() in text.lower(), need
    dxf = client.get(f"/api/tool-designs/{did}/components/{comp_id}/drawing.dxf").content.decode()
    for need in (other, "58 HRC", "NOTE 1: Polish the cavity face", "NOT RELEASED", "REVISION HISTORY", "Design to Cost", "Third-angle projection"):
        assert need in dxf, need
    assert "SOLID" in dxf                                                                        # the logo is drawn as vectors
    # a bad grade or class is refused
    assert client.put(f"/api/tool-designs/{did}/drawing-data", json={"items": [{"part_no": c["part_no"], "fields": {"steel": "NOPE-STEEL"}}]}).status_code == 400
    assert client.put(f"/api/tool-designs/{did}/drawing-data", json={"items": [{"part_no": c["part_no"], "fields": {"tolerance_class": "z"}}]}).status_code == 400


def test_excel_download_upload_roundtrip_and_validation():
    d = design(mould("DD-2"))
    did, c = d["id"], first_made(d)
    x = client.get(f"/api/tool-designs/{did}/drawing-data.xlsx")
    assert x.status_code == 200
    wb = load_workbook(io.BytesIO(x.content))
    assert wb.sheetnames == ["Drawing data", "How to use", "Tolerance tables", "Revision history"]
    ws = wb["Drawing data"]
    top = next(r for r in range(1, 10) if ws.cell(r, 1).value == "Part no.")                      # logo rows sit above the header
    assert [c_.value for c_ in ws[top]][:6] == ["Part no.", "Name", "Qty per tool", "Revision", "Material grade", "Hardness"]
    assert "Design to Cost" in str([c_.value for r_ in ws.iter_rows(max_row=top) for c_ in r_])
    row = next(r for r in range(top + 1, ws.max_row + 1) if ws.cell(r, 1).value == c["part_no"])
    ws.cell(row, 6).value = "52 HRC"
    ws.cell(row, 12).value = "Deburr all edges\nStamp part number"
    for w_ in wb.worksheets:
        w_._images = []                                                                        # openpyxl cannot re-save the logo; Excel can
    buf = io.BytesIO(); wb.save(buf)
    up = client.post(f"/api/tool-designs/{did}/drawing-data/upload", files={"file": ("d.xlsx", buf.getvalue())})
    assert up.status_code == 200, up.text
    assert len(up.json()["changed"]) == 1 and up.json()["parts_in_sheet"] >= 1                   # only the edited part changed
    p = next(x for x in client.get(f"/api/tool-designs/{did}/drawing-data").json()["parts"] if x["part_no"] == c["part_no"])
    assert p["effective"]["hardness"] == "52 HRC" and p["effective"]["notes"] == ["Deburr all edges", "Stamp part number"]
    # unknown part and wrong file
    ws.cell(row, 1).value = "NO-SUCH-PART"
    buf2 = io.BytesIO(); wb.save(buf2)
    assert client.post(f"/api/tool-designs/{did}/drawing-data/upload", files={"file": ("d.xlsx", buf2.getvalue())}).status_code == 400
    assert client.post(f"/api/tool-designs/{did}/drawing-data/upload", files={"file": ("d.xlsx", b"not excel")}).status_code == 400
    # the download of a changed design carries the change in the revision history sheet
    wb2 = load_workbook(io.BytesIO(client.get(f"/api/tool-designs/{did}/drawing-data.xlsx").content))
    assert any("Technical data changed" in str(r[3]) for r in wb2["Revision history"].iter_rows(min_row=2, values_only=True))


def test_released_design_is_locked_shows_release_date_and_approver_and_carries_data_to_a_revision():
    d = design(mould("DD-3"))
    did, c = d["id"], first_made(d)
    client.put(f"/api/tool-designs/{did}/drawing-data", json={"items": [{"part_no": c["part_no"], "fields": {"hardness": "50 HRC"}, "notes": ["Keep burr-free"]}]})
    for it in client.get(f"/api/tool-designs/{did}").json()["review"]:
        client.post(f"/api/tool-designs/{did}/review/{it['id']}", json={"status": "ok", "comment": "checked"})
    rel = client.post(f"/api/tool-designs/{did}/release", json={"accept_warnings": True})
    assert rel.status_code == 200, rel.text
    released = rel.json()
    assert client.put(f"/api/tool-designs/{did}/drawing-data", json={"items": [{"part_no": c["part_no"], "notes": ["late change"]}]}).status_code == 400
    text = pdf_text(client.get(f"/api/tool-designs/{did}/components/{c['id']}/drawing.pdf").content)
    assert "NOT RELEASED" not in text and released["released_at"][:10] in text and "Approved by" in text
    asm = pdf_text(client.get(f"/api/tool-designs/{did}/files/assembly-pdf").content) if client.get(f"/api/tool-designs/{did}/files/assembly-pdf").status_code == 200 else ""
    if asm:
        assert "Design to Cost" in asm and released["released_at"][:10] in asm and "Revision history" in asm
    rev = client.post(f"/api/tool-designs/{did}/new-revision", json={"note": "next"}).json()
    p = next(x for x in client.get(f"/api/tool-designs/{rev['id']}/drawing-data").json()["parts"] if x["part_no"] == c["part_no"])
    assert p["effective"]["hardness"] == "50 HRC" and p["effective"]["notes"] == ["Keep burr-free"]       # the data is carried over
    t2 = pdf_text(client.get(f"/api/tool-designs/{rev['id']}/components/{rev["components"][0]["id"]}/drawing.pdf").content) if rev["components"][0]["kind"] == "made" else ""
    if t2:
        assert "NOT RELEASED" in t2                                                              # the new revision is a draft again


def test_strip_layout_drawings_carry_the_title_block():
    pid = add_part(project("SLP"), "sheet_metal_bracket.step", "Mild steel CRCA / HR", "sheet_metal_progressive", "SM-S")
    pdf = client.get(f"/api/parts/{pid}/sheet-layout/strip.pdf")
    assert pdf.status_code == 200
    t = pdf_text(pdf.content)
    for need in ("Design to Cost", "STRIP LAYOUT", "NOT RELEASED", "Approved by", "Revision history", "Third-angle"):
        assert need.lower() in t.lower(), need
    assert len(PdfReader(io.BytesIO(pdf.content)).pages) == 1
    for kind in ("strip.dxf", "blank.dxf"):
        dxf = client.get(f"/api/parts/{pid}/sheet-layout/{kind}").content.decode()
        assert "Design to Cost" in dxf and "REVISION HISTORY" in dxf and "NOT RELEASED" in dxf and "SOLID" in dxf
    pack = pdf_text(client.get(f"/api/parts/{pid}/sheet-layout/tool-pack.pdf").content)
    assert "Design to Cost" in pack and "Revision history" in pack
    z = zipfile.ZipFile(io.BytesIO(client.get(f"/api/parts/{pid}/design-pack.zip").content)) if client.get(f"/api/parts/{pid}/design-pack.zip").status_code == 200 else None
    if z:
        assert any(n.endswith("_strip_layout.pdf") for n in z.namelist())
