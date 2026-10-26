"""Stage 34: process parameters for the best part quality, from the simulation and the material window."""

import io

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.app import mould_db, process_params
from backend.app.main import app
from tests.test_drawing_data import bracket, pdf_text
from tests.test_rm_cost import add_part, project

client = TestClient(app)


def test_windows_and_rules_are_master_data_and_checked():
    t = mould_db.tables()
    assert len(t["process_windows"]) >= 22 and len(t["process_parameters"]) >= 60 and all(r["is_example"] for r in t["process_parameters"])
    assert process_params.window("ABS")["melt_c"] == 235
    bad = mould_db.tables()
    bad["process_parameters"].append({"part_type": "plastic", "step": 99, "name": "bad one", "formula": "1", "min_formula": "", "max_formula": "",
                                      "unit": "", "label": "x", "why": "", "kind": "parameter", "is_example": True, "notes": ""})
    try:
        mould_db.write(bad)
        raise AssertionError("bad name accepted")
    except mould_db.MouldDbError as exc:
        assert "one word" in str(exc)


def test_moulding_settings_follow_the_simulation_and_the_material():
    pid = add_part(project("PP1"), "plastic_cover.step", "ABS", "injection_moulding", "PP1-1")
    sim = client.get(f"/api/parts/{pid}/simulation").json()
    r = client.get(f"/api/parts/{pid}/process-parameters")
    assert r.status_code == 200, r.text
    p = {x["name"]: x for x in r.json()["parameters"]}
    assert 210 <= p["melt_temp_c"]["value"] <= 260 and p["melt_temp_c"]["min"] == 210        # ABS window, not the simulation's generic 230
    assert abs(p["cooling_time_s"]["value"] - sim["results"]["cooling_s"]) < 0.01 and abs(p["clamp_force_t"]["value"] - sim["results"]["clamp_t"]) < 0.01
    assert p["hold_pressure_bar"]["min"] < p["hold_pressure_bar"]["value"] < p["hold_pressure_bar"]["max"] and p["drying_temp_c"]["value"] == 80
    assert r.json()["is_example"] and "simulation" in r.json()["source"]
    wb = load_workbook(io.BytesIO(client.get(f"/api/parts/{pid}/process-parameters.xlsx").content))
    assert any("Melt temperature" in str(c.value) for row in wb.active.iter_rows() for c in row)
    assert "Melt temperature" in pdf_text(client.get(f"/api/parts/{pid}/process-parameters.pdf").content)


def test_sheet_metal_and_not_applicable():
    r = client.get(f"/api/parts/{bracket('PP2')}/process-parameters").json()
    p = {x["name"]: x for x in r["parameters"]}
    assert p["press_tonnage_t"]["value"] > 40 and p["die_clearance_mm"]["value"] > 0 and p["strokes_per_min"]["value"] > 0
    plain = add_part(project("PP3"), "plate_with_hole.step", "Mild steel CRCA / HR", "machined", "PP3-1")
    assert client.get(f"/api/parts/{plain}/process-parameters").status_code == 400


def test_every_moulded_part_type_gets_a_setting_sheet():
    cases = [("plastic_cover.step", "Aluminium ADC12 / A383", "die_casting", "fast_shot_speed_m_s"),
             ("plastic_cover.step", "EPS bead grade S (20 kg/m3)", "eps_moulding", "steam_pressure_bar"),
             ("pet_bottle.step", "PET bottle grade (IV 0.80)", "stretch_blow_moulding", "blow_pressure_bar"),
             ("plastic_cover.step", "Grey cast iron FG260", "foundry", "pour_temp_c")]
    for i, (f, mat, proc, key) in enumerate(cases):
        pid = add_part(project(f"PT{i}"), f, mat, proc, f"PT{i}-1")
        r = client.get(f"/api/parts/{pid}/process-parameters")
        assert r.status_code == 200, (proc, r.text)
        assert key in {x["name"] for x in r.json()["parameters"]}, proc
