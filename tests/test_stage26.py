"""Stage 26: users, head of department, admin; each user works only on the data of his / her tool groups; the admin moves responsibilities."""
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from backend.app.main import app

PW = "correct horse battery"
SAMPLES = Path(__file__).resolve().parents[1] / "samples"

PARTS = {"plastic": ("plastic_cover.step", "ABS", "injection_moulding"),
         "die_casting": ("plastic_cover.step", "Aluminium ADC12 / A383", "die_casting"),
         "sheet_metal": ("sheet_metal_bracket.step", "Mild steel CRCA / HR", "sheet_metal_multistage"),
         "eps": ("audit/eps_small.step", "EPS bead grade S (20 kg/m3)", "eps_moulding")}


@pytest.fixture
def server(monkeypatch):
    monkeypatch.setenv("DTC_MODE", "server")
    monkeypatch.setenv("DTC_ADMIN_PASSWORD", PW)
    with TestClient(app) as c:
        yield c


def login(c, user="admin", password=PW):
    r = c.post("/api/auth/login", json={"username": user, "password": password})
    assert r.status_code == 200, r.text
    return r


def logout(c):
    c.post("/api/auth/logout")


def admin_makes(c, username, role, groups=()):
    login(c)
    r = c.post("/api/users", json={"username": username, "password": PW, "role": role, "name": username.title(), "groups": list(groups)})
    assert r.status_code == 200, r.text
    logout(c)
    return r.json()


def add_part(c, project_id, kind, code):
    file, material, process = PARTS[kind]
    with open(SAMPLES / file, "rb") as fh:
        m = c.post("/api/import/step", files={"file": (Path(file).name, fh)})
    assert m.status_code == 200, m.text
    m = m.json()
    return c.post(f"/api/import/{m['token']}/commit", json={"project_id": project_id, "items": {m["root"]: {"material": material, "process": process, "part_code": code}}})


@pytest.fixture
def world(server):
    u1 = admin_makes(server, "user1", "user", ["sheet_metal", "eps"])
    u2 = admin_makes(server, "user2", "user", ["die_casting", "plastic"])
    admin_makes(server, "head1", "head")
    admin_makes(server, "head2", "head")
    ids = {}
    login(server, "user1")
    ids["p1"] = server.post("/api/projects", json={"code": "P1", "name": "User 1 project"}).json()["id"]
    ids["sheet"] = add_part(server, ids["p1"], "sheet_metal", "U1-SHEET").json()["part_ids"][0]
    ids["eps"] = add_part(server, ids["p1"], "eps", "U1-EPS").json()["part_ids"][0]
    logout(server)
    login(server, "user2")
    ids["p2"] = server.post("/api/projects", json={"code": "P2", "name": "User 2 project"}).json()["id"]
    ids["dc"] = add_part(server, ids["p2"], "die_casting", "U2-DC").json()["part_ids"][0]
    ids["pl"] = add_part(server, ids["p2"], "plastic", "U2-PL").json()["part_ids"][0]
    logout(server)
    return ids | {"u1": u1["id"], "u2": u2["id"]}


def codes(c):
    return sorted(r["part_code"] for r in c.get("/api/parts").json()["rows"])


def test_each_user_sees_and_works_only_on_the_data_of_his_groups(server, world):
    login(server, "user1")
    assert [p["code"] for p in server.get("/api/projects").json()] == ["P1"]
    assert codes(server) == ["U1-EPS", "U1-SHEET"]
    assert server.get(f"/api/parts/{world['sheet']}").status_code == 200
    for url in (f"/api/parts/{world['dc']}", f"/api/parts/{world['dc']}/rm-cost", f"/api/projects/{world['p2']}", f"/api/projects/{world['p2']}/dashboard"):
        assert server.get(url).status_code == 403, url
    assert server.put(f"/api/parts/{world['pl']}", json={"name": "hacked"}).status_code == 403
    assert server.delete(f"/api/parts/{world['pl']}").status_code == 403
    # end to end on his own data: edit, costs, design, delete
    assert server.put(f"/api/parts/{world['sheet']}", json={"name": "Bracket changed"}).status_code == 200
    assert server.get(f"/api/parts/{world['sheet']}").json()["name"] == "Bracket changed"
    assert server.post(f"/api/parts/{world['sheet']}/rm-cost", json={}).status_code == 200
    dash = server.get(f"/api/projects/{world['p1']}/dashboard").json()
    assert dash["parts_total"] == 2
    assert server.delete(f"/api/parts/{world['eps']}").status_code == 200
    assert codes(server) == ["U1-SHEET"]
    # he cannot start a part in a group that is not his
    r = add_part(server, world["p1"], "plastic", "U1-PL")
    assert r.status_code == 403 and "not responsible" in r.json()["detail"]
    # master data and rates are read-only for a user, user management is closed
    assert server.get("/api/master-data").status_code == 200
    assert server.put("/api/master-data", json={}).status_code == 403
    assert server.get("/api/users").status_code == 403 and server.get("/api/backups").status_code == 403
    logout(server)

    login(server, "user2")
    assert codes(server) == ["U2-DC", "U2-PL"] and [p["code"] for p in server.get("/api/projects").json()] == ["P2"]
    assert server.get(f"/api/parts/{world['sheet']}").status_code == 403
    assert server.get(f"/api/tool-designs/1").status_code in (403, 404, 400)


def test_a_project_that_mixes_groups_shows_each_user_only_his_part(server, world):
    login(server, "head1")
    r = add_part(server, world["p1"], "plastic", "HEAD-PL-IN-P1")                    # the head adds a plastic part to user 1's project
    assert r.status_code == 200
    logout(server)
    login(server, "user1")
    assert codes(server) == ["U1-EPS", "U1-SHEET"]                                   # user 1 does not see it
    assert server.get(f"/api/projects/{world['p1']}/dashboard").json()["parts_total"] == 2
    logout(server)
    login(server, "user2")                                                            # user 2 (plastic) sees it, and so the project P1 with that one part
    assert "HEAD-PL-IN-P1" in codes(server)
    assert sorted(p["code"] for p in server.get("/api/projects").json()) == ["P1", "P2"]
    assert next(p for p in server.get("/api/projects").json() if p["code"] == "P1")["part_count"] == 1
    assert server.delete(f"/api/projects/{world['p1']}").status_code == 403           # it also holds parts that are not his / hers


def test_head_of_department_sees_everything_and_can_edit_and_add(server, world):
    for head in ("head1", "head2"):
        login(server, head)
        assert sorted(p["code"] for p in server.get("/api/projects").json()) == ["P1", "P2"]
        assert codes(server) == ["U1-EPS", "U1-SHEET", "U2-DC", "U2-PL"]
        logout(server)
    login(server, "head1")
    assert server.put(f"/api/parts/{world['dc']}", json={"name": "Corrected by the head"}).status_code == 200
    assert add_part(server, world["p2"], "sheet_metal", "HEAD-SHEET").status_code == 200
    assert server.post(f"/api/parts/{world['dc']}/rm-cost", json={}).status_code == 200
    assert server.put("/api/master-data", json={}).status_code != 403                # may change shared data (the body is invalid: not 403)
    assert server.get("/api/users").status_code == 403                                # but does not manage users
    assert server.post("/api/auth/password", json={"old": PW, "new": "head one new password"}).status_code == 200
    logout(server)
    login(server, "head1", "head one new password")
    logout(server)
    login(server, "user2")
    assert server.get("/api/parts").json()["rows"] and next(r for r in server.get("/api/parts").json()["rows"] if r["part_code"] == "U2-DC")["name"] == "Corrected by the head"


def test_admin_moves_responsibilities_and_the_data_follows(server, world):
    login(server)
    assert server.get("/api/access-overview").status_code == 200
    ov = {g["group"]: g for g in server.get("/api/access-overview").json()["groups"]}
    assert sorted(ov["sheet_metal"]["seen_by"]) == ["admin", "head1", "head2", "user1"] and ov["die_casting"]["parts"] == 1
    # six months later: user 1 gets die casting and EPS, user 2 gets sheet metal and plastic
    assert server.put(f"/api/users/{world['u1']}", json={"groups": ["die_casting", "eps"]}).status_code == 200
    assert server.put(f"/api/users/{world['u2']}", json={"groups": ["sheet_metal", "plastic"]}).status_code == 200
    bad = server.put(f"/api/users/{world['u1']}", json={"groups": ["welding"]})
    assert bad.status_code == 400
    logout(server)
    login(server, "user1")
    assert codes(server) == ["U1-EPS", "U2-DC"]                                        # EPS (kept) and die casting (user 2's data)
    assert server.get(f"/api/parts/{world['dc']}").status_code == 200                  # the die-casting data user 2 made is now user 1's to work on
    assert server.get(f"/api/parts/{world['sheet']}").status_code == 403
    assert server.put(f"/api/parts/{world['dc']}", json={"name": "Now user 1"}).status_code == 200
    logout(server)
    login(server, "user2")
    assert codes(server) == ["U1-SHEET", "U2-PL"]
    assert server.get(f"/api/parts/{world['sheet']}").status_code == 200 and server.get(f"/api/parts/{world['dc']}").status_code == 403


def test_admin_creates_users_resets_passwords_and_deletes(server, world):
    login(server)
    u = server.post("/api/users", json={"username": "new.user", "password": PW, "role": "user", "groups": ["blow"]}).json()
    assert u["groups"] == ["blow"] and u["role"] == "user"
    h = server.post("/api/users", json={"username": "new.head", "password": PW, "role": "head"}).json()
    assert h["all_groups"] is True and len(h["groups"]) == 6
    assert server.put(f"/api/users/{u['id']}", json={"password": "reset by the admin"}).status_code == 200
    me = server.get("/api/auth/me").json()["user"]
    assert server.delete(f"/api/users/{me['id']}").status_code == 400                  # not oneself
    assert server.delete(f"/api/users/{u['id']}").status_code == 200
    logout(server)
    assert server.post("/api/auth/login", json={"username": "new.user", "password": "reset by the admin"}).status_code == 401   # deleted
    # a user changes his / her own password
    login(server, "user1")
    assert server.post("/api/auth/password", json={"old": "wrong", "new": "a brand new password"}).status_code == 400
    assert server.post("/api/auth/password", json={"old": PW, "new": "a brand new password"}).status_code == 200
    logout(server)
    login(server, "user1", "a brand new password")
    assert server.get("/account").status_code == 200


def test_legacy_roles_are_converted_and_pages_exist(server):
    from backend.app import auth, projects_db
    from sqlalchemy import select
    login(server)
    with projects_db.engine().begin() as c:
        c.execute(projects_db.users.insert().values(username="old.eng", name="", password_hash=auth.hash_password(PW), role="engineer", active=True))
        c.execute(projects_db.users.insert().values(username="old.app", name="", password_hash=auth.hash_password(PW), role="approver", active=True))
    assert auth.migrate_roles() == 2
    roles = {u["username"]: u["role"] for u in server.get("/api/users").json()}
    assert roles["old.eng"] == "user" and roles["old.app"] == "head"
    assert server.get("/users").status_code == 200 and server.get("/account").status_code == 200


def test_a_user_cannot_reach_other_peoples_data_through_import_or_code_reuse(server, world):
    login(server, "user2")
    r = add_part(server, world["p1"], "plastic", "SNEAKY")                    # P1 is user 1's: user 2 sees nothing in it
    assert r.status_code == 403
    mine = add_part(server, world["p2"], "plastic", "U1-SHEET")               # that code belongs to user 1's sheet-metal part
    assert mine.status_code == 403 and "already exists" in mine.json()["detail"]
    assert server.get("/api/parts/export.xlsx").status_code == 200
    from openpyxl import load_workbook
    import io
    wb = load_workbook(io.BytesIO(server.get("/api/parts/export.xlsx").content))
    text = " ".join(str(c.value) for row in wb.worksheets[0].iter_rows() for c in row if c.value)
    assert "U2-DC" in text and "U1-SHEET" not in text                           # the Excel export holds only his / her parts
