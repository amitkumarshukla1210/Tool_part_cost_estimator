import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.app import costing_setup, rate_engine, rates
from backend.app.main import app

client = TestClient(app)
MONTH = "2026-10"


def node(trace, name):
    return trace["nodes"][f"calc:{name}"]["value"]


def set_approver(name="Test Approver"):
    rates.save_settings({"approver_name": name})


# ---------- MHR and labour rate, checked by hand ----------

def test_mhr_india_vmc_by_hand():
    r = rate_engine.mhr("VMC-850", "IN", MONTH)
    hours = 300 * 2 * 8 * 0.80                                    # 3,840 h
    price, residual = 4_500_000, 450_000
    dep = (price - residual) / 10 / hours
    interest = (price + residual) / 2 * 0.10 / hours
    power = 15 * 0.60 * 8.5
    maint = price * 0.04 / hours
    space = 12 * 350 * 12 / hours
    subtotal = dep + interest + power + maint + space + 150
    assert node(r, "annual_hours") == pytest.approx(hours)
    assert node(r, "depreciation") == pytest.approx(dep)
    assert node(r, "interest") == pytest.approx(interest)
    assert node(r, "power") == pytest.approx(power)
    assert node(r, "floor_space") == pytest.approx(space)
    assert r["value"] == pytest.approx(subtotal * 1.15)
    assert r["value_inr"] == pytest.approx(r["value"])            # India: INR already
    assert r["currency"] == "INR"
    # every input names its source rate version
    assert len(r["rates_used"]) == 10
    assert all(x["status"] == "approved" for x in r["rates_used"])
    src = r["nodes"]["in:electricity_rate"]["source"]
    assert src["type"] == "rate" and src["rate"]["month"] == MONTH and src["rate"]["is_example"]
    assert r["nodes"]["in:power_kw"]["source"] == {"type": "machine", "sheet": "Machines", "machine_id": "VMC-850",
                                                   "field": "power_kw", "is_example": True,
                                                   "ref": {"store": "costing_setup", "table": "machines",
                                                           "match": {"machine_id": "VMC-850"}, "column": "power_kw"}}


def test_labour_rate_by_hand():
    r = rate_engine.labour_rate("toolmaker", "IN", MONTH)
    assert r["value"] == pytest.approx(45000 * 12 / (300 * 8) * 1.15 / 0.85)
    de = rate_engine.labour_rate("toolmaker", "DE", MONTH)       # Germany: +21 % employer cost
    assert de["value"] == pytest.approx(4400 * 12 * 1.21 / (230 * 8) * 1.15 / 0.85)
    assert de["currency"] == "EUR"
    assert de["value_inr"] == pytest.approx(de["value"] * 90.5)


def test_other_country_in_local_currency_and_inr():
    jp = rate_engine.mhr("VMC-850", "JP", MONTH)
    assert jp["currency"] == "JPY"
    assert jp["value_inr"] == pytest.approx(jp["value"] * 0.56)
    assert "fx:JPY" in jp["nodes"]


# ---------- approval and versions ----------

def test_only_approved_rates_are_used():
    before = rate_engine.mhr("VMC-850", "IN", MONTH)["value"]
    draft = rates.create({"category": "power", "item": "electricity_industrial", "country": "IN", "value": 10.5,
                          "unit": "INR/kWh", "currency": "INR", "month": MONTH})
    assert draft["status"] == "draft"
    assert rate_engine.mhr("VMC-850", "IN", MONTH)["value"] == pytest.approx(before)

    with pytest.raises(rates.RateError, match="name for approvals"):
        rates.set_status([draft["id"]], "approved")
    set_approver()
    approved = rates.set_status([draft["id"]], "approved", "checked tariff")[0]
    assert approved["approved_by"] == "Test Approver" and "checked tariff" in approved["comment"]
    after = rate_engine.mhr("VMC-850", "IN", MONTH)
    assert after["value"] == pytest.approx(before + 15 * 0.6 * (10.5 - 8.5) * 1.15)

    hist = rates.history("power", "electricity_industrial", "IN")
    assert [h["value"] for h in hist] == [8.5, 10.5]               # old value kept
    with pytest.raises(rates.RateError, match="cannot be changed"):
        rates.update(hist[0]["id"], {"value": 1})


def test_month_versions_keep_old_results_reproducible():
    set_approver()
    rates.create({"category": "finance", "item": "interest_pct", "country": "IN", "value": 12, "unit": "%",
                  "month": "2026-11"}, approve=True)
    october = rate_engine.mhr("VMC-850", "IN", "2026-10")
    november = rate_engine.mhr("VMC-850", "IN", "2026-11")
    assert october["nodes"]["in:interest_pct"]["value"] == 10
    assert november["nodes"]["in:interest_pct"]["value"] == 12
    with pytest.raises(rates.RateMissing, match="for 2026-09 or earlier"):
        rate_engine.mhr("VMC-850", "IN", "2026-09")


def test_review_shows_change_and_threshold():
    rates.create({"category": "forex", "item": "USD", "value": 88.0, "unit": "INR per USD", "currency": "INR",
                  "month": MONTH, "source_url": "https://www.ecb.europa.eu"})
    rates.create({"category": "forex", "item": "EUR", "value": 91.0, "unit": "INR per EUR", "currency": "INR",
                  "month": MONTH})
    items = {i["draft"]["item"]: i for i in rates.review()["items"]}
    assert items["USD"]["change_pct"] == pytest.approx((88 - 83.5) / 83.5 * 100, abs=0.01)
    assert items["USD"]["over_threshold"]                          # +5.4 % > 5 %
    assert not items["EUR"]["over_threshold"]                      # +0.6 %
    rates.save_settings({"review_threshold_pct": "10"})
    assert not rates.review()["items"][1]["over_threshold"]


def test_money_rate_in_other_currency_is_converted():
    set_approver()
    rates.create({"category": "machine_price", "item": "VMC-850", "country": "IN", "value": 60000,
                  "unit": "USD", "currency": "USD", "month": MONTH}, approve=True)
    r = rate_engine.mhr("VMC-850", "IN", MONTH)
    conv = r["nodes"]["in:machine_price:INR"]
    assert conv["value"] == pytest.approx(60000 * 83.5)
    assert {i["node"] for i in conv["inputs"]} == {"in:machine_price", "fx:USD", "fx:INR"}


# ---------- logic lives in data ----------

def test_formulas_are_editable_data():
    setup = costing_setup.tables()
    for f in setup["cost_formulas"]:
        if f["scope"] == "mhr" and f["name"] == "overhead":
            f["formula"] = "machine_cost * overhead_pct / 100 * 2"
    costing_setup.write(setup)
    r = rate_engine.mhr("VMC-850", "IN", MONTH)
    assert r["value"] == pytest.approx(node(r, "machine_cost") * 1.30)


@pytest.mark.parametrize("change, message", [
    (lambda s: s["cost_formulas"][0].update(formula="working_dayz * 2"), "unknown name 'working_dayz'"),
    (lambda s: s["cost_formulas"][1].update(formula="depreciation_per_year * 2"), "depends on itself"),
    (lambda s: s["cost_formulas"][0].update(is_result=True), "exactly one row"),
    (lambda s: s["formula_inputs"][10].update(item="colour"), "machine field must be one of"),
    (lambda s: s["countries"][0].update(currency="RUPEE"), "3-letter code"),
])
def test_bad_setup_is_rejected(change, message):
    setup = costing_setup.tables()
    change(setup)
    r = client.put("/api/costing-setup", json=setup)
    assert r.status_code == 400 and message in r.json()["detail"]


# ---------- API ----------

def test_api_overview_and_errors():
    r = client.get("/api/mhr", params={"country": "IN", "as_of": MONTH})
    body = r.json()
    assert r.status_code == 200 and len(body["machines"]) == 33 and len(body["labour"]) == 4
    assert all("error" not in m for m in body["machines"])
    # a month with no approved rates still renders, with the reason per row
    old = client.get("/api/mhr", params={"country": "IN", "as_of": "2020-01"}).json()
    assert "No approved rate" in old["machines"][0]["error"]
    assert client.get("/api/mhr/NOPE", params={"country": "IN"}).status_code == 400
    assert client.get("/api/mhr", params={"country": "XX"}).status_code == 400


def test_api_save_and_approve_and_edit_draft():
    client.put("/api/settings", json={"approver_name": "Amit"})
    new = client.post("/api/rates", json={"category": "space", "item": "rent_industrial", "country": "IN",
                                          "value": 400, "unit": "INR/m²/month", "currency": "INR",
                                          "month": MONTH, "approve": True}).json()
    assert new["status"] == "approved" and new["approved_by"] == "Amit"
    draft = client.post("/api/rates", json={"category": "space", "item": "rent_industrial", "country": "IN",
                                            "value": 420, "month": MONTH}).json()
    edited = client.put(f"/api/rates/{draft['id']}", json={**draft, "value": 410}).json()
    assert edited["value"] == 410 and edited["status"] == "draft"
    assert client.put(f"/api/rates/{new['id']}", json={**new, "value": 1}).status_code == 400
    assert client.post("/api/rates", json={"category": "x", "item": "y", "value": 1, "month": "Oct"}).status_code == 400
    assert client.post("/api/rates/status", json={"ids": [draft["id"]], "status": "rejected"}).json()[0]["status"] == "rejected"


def test_excel_export_and_import_as_drafts():
    xlsx = client.get("/api/rates/download")
    ws = load_workbook(io.BytesIO(xlsx.content))["Rates"]
    assert ws.max_row - 1 == len([1 for _ in rates.history("forex", "USD")]) + 591 + 748 + 36 + 97 + 126 + 256 + 9   # 9 = EN8 plate steel in every country (stage 25); 256 = EPS and stretch blow materials, machines, tool aluminium, steam and duty rates (stages 18, 19); 126 = new machine prices, consumables, wire conversion, bare-board rates (stage 15); 97 = freight, duty, import rates (stage 11); 748 = country material/tool-steel/process rates seeded in stage 3; 36 = toolmaker terms (stage 4)
    new = io.BytesIO()
    from openpyxl import Workbook
    wb = Workbook()
    wb.active.title = "Rates"
    wb.active.append(["category", "item", "country", "value", "unit", "currency", "month"])
    wb.active.append(["forex", "USD", "", 84.2, "INR per USD", "INR", "2026-11"])
    wb.save(new)
    r = client.post("/api/rates/upload", files={"file": ("rates.xlsx", new.getvalue())})
    assert r.json() == {"drafts_created": 1}
    assert rates.history("forex", "USD")[-1]["status"] == "draft"
    assert rate_engine.fx_to_inr  # still only approved values are used
    assert rates.lookup("forex", "USD", "", "2026-11")["value"] == 83.5
