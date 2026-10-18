"""Stage 2: RM cost on the register. Formulas, rates and defaults all come from master data; every result is traced."""
import io
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.app import costing_setup, register, rm_cost
from backend.app import projects_db as db
from backend.app.main import app

SAMPLES = Path(__file__).resolve().parents[1] / "samples"
client = TestClient(app)


def project(code="T"):
    return client.post("/api/projects", json={"code": code, "name": "Test"}).json()


def add_part(p, file="plastic_cover.step", material="ABS", process="injection_moulding", code="PC-1"):
    with open(SAMPLES / file, "rb") as fh:
        m = client.post("/api/import/step", files={"file": (file, fh)}).json()
    r = client.post(f"/api/import/{m['token']}/commit", json={
        "project_id": p["id"], "items": {m["root"]: {"material": material, "process": process, "part_code": code}}})
    assert r.status_code == 200, r.text
    return r.json()["part_ids"][0]


def approver():
    client.put("/api/settings", json={"approver_name": "Tester"})


def new_rate(item, value, category="material", month="2026-10", approve=True, country="IN"):
    r = client.post("/api/rates", json={"category": category, "item": item, "country": country, "value": value,
                                        "unit": "INR/kg", "currency": "INR", "month": month, "approve": approve})
    assert r.status_code == 200, r.text
    return r.json()


def detail(pid):
    return client.get(f"/api/parts/{pid}").json()


# ---------- calculation ----------

def test_plastic_rm_cost_from_the_rate_library_by_hand():
    pid = add_part(project())
    d = detail(pid)
    v = d["rm"]["latest"]
    net = d["revisions"][0]["geometry"]["volume_cm3"] * 1.04               # ABS density from Master data → Materials
    gross = net * 1.15                                                     # runner 15 % (Parameters)
    used = gross * 1.02                                                    # purge loss 2 %
    assert v["net_weight_g"] == pytest.approx(net, rel=1e-5)
    assert v["gross_weight_g"] == pytest.approx(gross, rel=1e-5)
    assert v["net_rm_cost"] == pytest.approx(used / 1000 * 180, rel=1e-5)  # ABS 180 INR/kg, scrap value 0
    assert v["net_rm_cost_inr"] == v["net_rm_cost"] and v["currency"] == "INR" and v["country"] == "IN"
    assert v["is_example"] and d["rm_cost_inr"] == v["net_rm_cost_inr"]
    assert {(r["category"], r["item"]) for r in v["rates_used"]} == {("material", "ABS"), ("material_scrap", "ABS")}


def test_sheet_metal_uses_the_nesting_and_traces_it():
    pid = add_part(project(), "sheet_metal_bracket.step", "Mild steel CRCA / HR", "sheet_metal_multistage", "SM-1")
    v = detail(pid)["rm"]["latest"]
    net = 11800 / 1000 * 7.85
    gross = 2500 * 1250 * 2 / 1000 * 7.85 / 450                            # 450 blanks per sheet (see test_gross_weight)
    assert v["gross_weight_g"] == pytest.approx(gross, rel=1e-4)
    cost = gross / 1000 * 65 - (gross - net) / 1000 * 30
    assert v["net_rm_cost"] == pytest.approx(cost, rel=1e-4)
    tr = client.get(f"/api/parts/{pid}/rm-cost/trace").json()
    nest = tr["nodes"]["in:blanks_per_sheet"]
    assert nest["value"] == 450 and nest["kind"] == "calculated" and "450" in nest["description"]
    assert tr["nodes"]["in:thickness_mm"]["source"]["type"] == "cad"       # thickness comes from the 3D file


def test_everything_is_traceable_and_every_input_has_a_source():
    pid = add_part(project())
    tr = client.get(f"/api/parts/{pid}/rm-cost/trace").json()
    root = tr["nodes"][tr["root"]]
    assert root["label"] == "Net RM cost per part" and root["formula"] == "rm_cost + scrap_credit"
    for n in tr["nodes"].values():
        if n["kind"] == "input":
            assert n["source"]["type"] in {"rate", "master", "cad", "base"}, n
    mat = tr["nodes"]["in:material_rate"]["source"]["rate"]
    assert mat["status"] == "approved" and mat["is_example"] and mat["month"] == "2026-10"
    assert tr["nodes"]["in:runner_pct"]["source"] == {"type": "master", "sheet": "Parameters", "row": "plastic / runner_pct"}
    assert tr["stored"] and tr["version"] == 1 and tr["rates_used"]


def test_typed_input_replaces_the_default_and_is_traced_as_yours():
    pid = add_part(project())
    base = detail(pid)["rm"]["latest"]["net_rm_cost"]
    r = client.post(f"/api/parts/{pid}/rm-cost", json={"rm_inputs": {"runner_pct": 0}, "reason": "hot runner"})
    assert r.status_code == 200 and r.json()["version"] == 2
    d = detail(pid)
    assert d["rm_inputs"] == {"runner_pct": 0}
    tr = client.get(f"/api/parts/{pid}/rm-cost/trace").json()
    assert tr["nodes"]["in:runner_pct"]["source"]["type"] == "user"
    assert d["rm"]["latest"]["net_rm_cost"] < base
    assert client.post(f"/api/parts/{pid}/rm-cost", json={"rm_inputs": {"not_an_input": 1}}).status_code == 400


# ---------- rates and versions ----------

def test_only_approved_rates_are_used_and_a_draft_changes_nothing():
    approver()
    pid = add_part(project())
    new_rate("ABS", 999, approve=False)
    assert detail(pid)["rm"]["outdated"] == ""
    assert client.post(f"/api/parts/{pid}/rm-cost", json={}).json()["stored"] is False


def test_new_approved_rate_does_not_change_stored_cost_until_you_recalculate():
    approver()
    pid = add_part(project())
    v1 = detail(pid)["rm"]["latest"]
    new_rate("ABS", 240)
    d = detail(pid)
    assert d["rm"]["latest"]["net_rm_cost"] == v1["net_rm_cost"]           # stored result is unchanged
    assert "Newer approved rates" in d["rm"]["outdated"]
    ch = d["rm"]["rate_changes"][0]
    assert (ch["old_value"], ch["new_value"], ch["new_month"]) == (180, 240, "2026-10")
    r = client.post(f"/api/parts/{pid}/rm-cost", json={}).json()
    assert r["stored"] and r["version"] == 2
    versions = client.get(f"/api/parts/{pid}/rm-cost").json()
    assert [x["version"] for x in versions] == [2, 1]
    assert versions[1]["net_rm_cost"] == v1["net_rm_cost"]                 # version 1 never overwritten
    assert versions[0]["net_rm_cost"] == pytest.approx(v1["net_rm_cost"] * 240 / 180, rel=1e-6)
    assert client.get(f"/api/parts/{pid}/rm-cost/trace", params={"version": 1}).json()["nodes"]["in:material_rate"]["value"] == 180
    assert detail(pid)["rm"]["outdated"] == ""


def test_costing_month_picks_the_rate_valid_then():
    approver()
    pid = add_part(project())
    new_rate("ABS", 240, month="2026-11")
    client.put(f"/api/parts/{pid}", json={"costing_month": "2026-10"})
    assert detail(pid)["rm"]["latest"]["version"] == 1                     # October still uses 180: nothing new
    client.put(f"/api/parts/{pid}", json={"costing_month": "2026-12"})
    assert detail(pid)["rm"]["latest"]["version"] == 2


def test_missing_rate_gives_a_clear_message_and_stores_nothing():
    pid = add_part(project(), material="PC (polycarbonate)")
    rates = client.get("/api/rates", params={"category": "material"}).json()["rows"]
    assert any(r["item"] == "PC (polycarbonate)" for r in rates)
    p2 = add_part(project("U"), material="PA6", code="PC-2")
    # a material without any rate: add it to Materials only
    data = client.get("/api/master-data").json()
    data["Materials"].append({"part_type": "plastic", "material": "PEEK", "density_g_cm3": 1.3, "notes": ""})
    assert client.put("/api/master-data", json=data).status_code == 200
    client.put(f"/api/parts/{p2}", json={"material": "PEEK"})
    d = detail(p2)
    assert "No approved rate 'material / PEEK'" in d["rm"]["problem"]
    assert d["rm"]["weights"]["gross_g"] > 0                               # weights still shown
    r = client.post(f"/api/parts/{p2}/rm-cost", json={})
    assert r.status_code == 400 and "PEEK" in r.json()["detail"]


def test_other_country_uses_its_own_rates_and_currency():
    pid = add_part(project())
    client.put(f"/api/parts/{pid}", json={"country": "CN"})
    assert "No approved rate" in detail(pid)["rm"]["problem"]              # no CN material rates yet (stage 3)
    approver()
    new_rate("ABS", 14, country="CN", month="2026-10")
    new_rate("ABS", 0, category="material_scrap", country="CN", month="2026-10")
    r = client.post(f"/api/parts/{pid}/rm-cost", json={}).json()
    v = detail(pid)["rm"]["latest"]
    assert v["country"] == "CN" and v["currency"] == "CNY"
    assert v["net_rm_cost_inr"] == pytest.approx(v["net_rm_cost"] * 11.6)  # CNY → INR with the approved forex rate


def test_formula_change_in_master_data_marks_the_cost_outdated_and_is_used():
    pid = add_part(project())
    setup = costing_setup.tables()
    for r in setup["cost_formulas"]:
        if r["scope"] == "rm_plastic" and r["name"] == "gross_g":
            r["formula"] = "net_weight_g * (1 + runner_pct / 100) * 2"
    costing_setup.write(setup)
    assert "formulas changed" in detail(pid)["rm"]["outdated"]
    client.post(f"/api/parts/{pid}/rm-cost", json={})
    assert detail(pid)["rm"]["latest"]["gross_weight_g"] == pytest.approx(detail(pid)["rm"]["latest"]["net_weight_g"] * 2.3)


# ---------- who gets an RM cost ----------

def test_bought_out_assembly_and_machined_parts_have_no_rm():
    pid = add_part(project())
    client.put(f"/api/parts/{pid}", json={"sourcing": "bought_out"})
    assert "bought-out price" in detail(pid)["rm"]["problem"]
    client.put(f"/api/parts/{pid}", json={"sourcing": "make", "process": "machined"})
    assert "later stage" in detail(pid)["rm"]["problem"]
    assert client.post(f"/api/parts/{pid}/rm-cost", json={}).status_code == 400


def test_insert_moulded_part_costs_the_body_material_only():
    pid = add_part(project())
    body = detail(pid)["rm"]["latest"]["net_weight_g"]
    client.post(f"/api/parts/{pid}/inserts", json={"name": "Brass insert", "material": "Brass", "quantity": 4,
                                                   "dimensions": {"shape": "bush", "od_mm": 6, "id_mm": 3, "length_mm": 8}})
    v = detail(pid)["rm"]["latest"]
    assert v["net_weight_g"] == pytest.approx(body)                       # insert weight is not in the RM weight


def test_new_revision_makes_a_new_version_automatically():
    pid = add_part(project())
    with open(SAMPLES / "plate_with_hole.step", "rb") as fh:
        client.post(f"/api/parts/{pid}/revisions", files={"file": ("x.step", fh)}, data={"rev": "B", "change_note": "bigger"})
    versions = client.get(f"/api/parts/{pid}/rm-cost").json()
    assert len(versions) == 2 and versions[0]["rev"] == "B"


# ---------- register, dashboard, exports ----------

def test_search_filter_dashboard_rollup_and_exports():
    approver()
    p = project()
    v = client.post(f"/api/projects/{p['id']}/variants", json={"code": "STD", "name": "Std", "target_bom_cost_inr": 1000}).json()
    a = add_part(p, code="PC-1")
    b = add_part(p, "sheet_metal_bracket.step", "Mild steel CRCA / HR", "sheet_metal_multistage", "SM-1")
    with db.engine().begin() as conn:
        register.set_bom_line(conn, v["id"] if "id" in v else v["variants"][0]["id"], None, a, 2)
        register.set_bom_line(conn, v["id"] if "id" in v else v["variants"][0]["id"], None, b, 3)
    rows = client.get("/api/parts", params={"project_id": p["id"], "sort": "rm_cost"}).json()["rows"]
    assert [r["part_code"] for r in rows][0] == "PC-1" and rows[0]["rm_cost_inr"] > rows[1]["rm_cost_inr"]
    cheap = client.get("/api/parts", params={"project_id": p["id"], "rm_max": 8}).json()["rows"]
    assert [r["part_code"] for r in cheap] == ["SM-1"]
    dash = client.get(f"/api/projects/{p['id']}/dashboard").json()
    var = dash["variants"][0]
    ra, rb = detail(a)["rm_cost_inr"], detail(b)["rm_cost_inr"]
    assert var["rm_cost_inr"] == pytest.approx(2 * ra + 3 * rb) and var["rm_parts_costed"] == 2
    assert client.get(f"/api/projects/{p['id']}/dashboard.pdf").content[:4] == b"%PDF"

    xl = client.get(f"/api/parts/{a}/rm-cost/export.xlsx")
    wb = load_workbook(io.BytesIO(xl.content))
    assert wb.sheetnames == ["Summary", "Steps", "Inputs", "Rates used"]
    steps = [r[0] for r in wb["Steps"].iter_rows(min_row=2, values_only=True)]
    assert "Net RM cost per part" in steps and "Gross weight (shot weight per part)" in steps
    inputs = {r[0]: r[3] for r in wb["Inputs"].iter_rows(min_row=2, values_only=True)}
    assert "Rate library #" in inputs["Material rate"] and "EXAMPLE" in inputs["Material rate"]
    assert client.get(f"/api/parts/{a}/rm-cost/export.pdf").content[:4] == b"%PDF"
    parts_xlsx = load_workbook(io.BytesIO(client.get("/api/parts/export.xlsx", params={"project_id": p["id"]}).content))
    assert "net_rm_cost_inr" in [c.value for c in parts_xlsx.active[1]]


def test_project_recalculation_uses_the_latest_rates():
    approver()
    p = project()
    a = add_part(p, code="PC-1")
    new_rate("ABS", 360)
    res = client.post(f"/api/projects/{p['id']}/rm-cost/recalculate").json()
    assert res["new_versions"] == 1 and res["problems"] == []
    assert detail(a)["rm"]["latest"]["version"] == 2


# ---------- the Estimate page API and master data ----------

def test_estimate_api_takes_the_rate_from_the_library_and_marks_typed_rates():
    body = {"part_type": "die_casting", "net_weight_g": 300, "density_g_cm3": 2.74, "material": "Aluminium ADC12 / A383"}
    r = client.post("/api/gross-weight", json=body).json()
    assert r["gross_weight_g"] == pytest.approx(500) and r["material_input_g"] == pytest.approx(515)
    assert r["rm_cost"] == pytest.approx(0.515 * 260) and r["scrap_credit"] == pytest.approx(0.2 * 180)
    assert r["nodes"]["in:material_rate"]["source"]["type"] == "rate"
    typed = client.post("/api/gross-weight", json=body | {"material_rate_per_kg": 300}).json()
    assert typed["nodes"]["in:material_rate"]["source"]["type"] == "user" and typed["rm_cost"] == pytest.approx(0.515 * 300)
    assert client.post("/api/gross-weight", json=body | {"material": ""}).json()["priced"] is False


def test_materials_no_longer_carry_rates_and_old_workbooks_still_upload():
    mats = client.get("/api/master-data").json()["Materials"]
    assert set(mats[0]) == {"part_type", "material", "density_g_cm3", "notes"}
    xlsx = client.get("/api/master-data/download").content
    wb = load_workbook(io.BytesIO(xlsx))
    ws = wb["Materials"]
    ws.insert_cols(4, 2)
    ws.cell(1, 4, "rate_per_kg")
    ws.cell(1, 5, "scrap_rate_per_kg")
    buf = io.BytesIO()
    wb.save(buf)
    up = client.post("/api/master-data/upload", files={"file": ("old.xlsx", buf.getvalue())})
    assert up.status_code == 200
    assert any("rate_per_kg" in n for n in up.json()["_notes"])


def test_setup_checks_for_rm_formulas():
    setup = costing_setup.tables()
    bad = [dict(r, scope="mhr") if r["variable"] == "net_weight_g" and r["scope"] == "rm_plastic" else r
           for r in setup["formula_inputs"]]
    with pytest.raises(costing_setup.SetupError, match="only for the RM scopes"):
        costing_setup.write(setup | {"formula_inputs": bad})
    bad = [r for r in setup["formula_inputs"] if not (r["scope"] == "rm_sheet_metal" and r["variable"] == "gap_mm")]
    with pytest.raises(costing_setup.SetupError, match="nesting also needs"):
        costing_setup.write(setup | {"formula_inputs": bad})
    assert rm_cost.overridable("plastic")[0]["variable"]
