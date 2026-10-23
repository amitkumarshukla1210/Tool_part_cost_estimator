import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Table

from backend.app import bom_extract
from backend.app.main import app

client = TestClient(app)

KBL = b"""<?xml version="1.0" encoding="UTF-8"?>
<KBL_container xmlns="http://www.prostep.org/KBL">
 <Harness_container>
  <Part id="p_h1"><Part_number>1-967640-1</Part_number><Company_name>TE</Company_name><Description>Housing 2-way</Description></Part>
  <Part id="p_t1"><Part_number>927777-3</Part_number><Description>Terminal</Description></Part>
  <General_wire id="p_w05"><Part_number>FLRY-B-0.5</Part_number><Company_name>Leoni</Company_name><Description>Wire FLRY-B 0.5</Description><Cross_section_area><Value_component>0.5</Value_component></Cross_section_area><Wire_type>FLRY-B</Wire_type><Cover_colour>RD</Cover_colour></General_wire>
  <Wire_occurrence id="w1" Part="p_w05"><Special_wire_id>W1</Special_wire_id><Length_information><Length_value><Value_component>350</Value_component></Length_value></Length_information></Wire_occurrence>
  <Wire_occurrence id="w2" Part="p_w05"><Special_wire_id>W2</Special_wire_id><Length_information><Length_value><Value_component>420</Value_component></Length_value></Length_information></Wire_occurrence>
  <Connector_occurrence id="c1" Part="p_h1"><Id>X1</Id></Connector_occurrence>
  <Connector_occurrence id="c2" Part="p_h1"><Id>X2</Id></Connector_occurrence>
  <Terminal_occurrence id="t1" Part="p_t1"/><Terminal_occurrence id="t2" Part="p_t1"/><Terminal_occurrence id="t3" Part="p_t1"/><Terminal_occurrence id="t4" Part="p_t1"/>
  <Wire_protection_occurrence id="s1"><Description>Corrugated tube</Description></Wire_protection_occurrence>
 </Harness_container>
</KBL_container>"""


def by_cat(lines):
    out = {}
    for l in lines:
        out.setdefault(l["category"], []).append(l)
    return out


def test_kbl_wires_are_grouped_with_length_and_connectors_counted():
    lines = by_cat(bom_extract.parse_xml(KBL))
    w = lines["wire"][0]
    assert w["qty"] == 2 and w["length_mm"] == 770 and w["cross_section_mm2"] == 0.5 and w["colour"] == "RD" and w["part_number"] == "FLRY-B-0.5"
    assert w["maker"] == "Leoni" and w["confidence"] == 1.0 and "W1" in w["ref"] and "W2" in w["ref"]
    assert lines["connector"][0]["qty"] == 2 and lines["connector"][0]["part_number"] == "1-967640-1" and lines["connector"][0]["maker"] == "TE"
    t = lines["terminal"][0]
    assert t["qty"] == 4 and t["confidence"] < 1.0 and "no manufacturer in the file" in t["flags"]        # TE terminal has no maker in the sample
    assert lines["sleeve"][0]["confidence"] < bom_extract.LOW and "no part number in the file" in lines["sleeve"][0]["flags"]


def test_vec_style_root_is_read_and_bad_xml_is_explained():
    vec = KBL.replace(b"KBL_container", b"VecContainer")
    assert bom_extract.parse_xml(vec)
    with pytest.raises(bom_extract.ExtractError, match="could not be read"):
        bom_extract.parse_xml(b"<a><b></a>")
    with pytest.raises(bom_extract.ExtractError, match="No wires"):
        bom_extract.parse_xml(b"<Root><Other/></Root>")


def make_pdf(rows, text_only=True):
    buf = io.BytesIO()
    t = Table(rows)
    t.setStyle([("GRID", (0, 0), (-1, -1), 0.5, (0, 0, 0))])
    SimpleDocTemplate(buf, pagesize=A4).build([t])
    return buf.getvalue()


def test_pdf_bom_table_with_low_confidence_cells():
    pdf = make_pdf([["Item", "Part number", "Description", "Qty", "Length (mm)", "Colour"],
                     ["1", "FLRY-B-0.5", "Wire FLRY 0.5 mm2", "4", "350", "RD"],
                     ["2", "1-967640-1", "Connector housing 2-way", "2", "", ""],
                     ["3", "", "Cable tie", "x", "", ""]])
    lines = bom_extract.parse_pdf(pdf)
    assert [l["category"] for l in lines] == ["wire", "connector", "clip"]
    assert lines[0]["qty"] == 4 and lines[0]["length_mm"] == 1400 and lines[0]["notes"] == "350 mm each" and lines[0]["colour"] == "RD" and lines[0]["confidence"] == 1.0
    assert lines[2]["confidence"] < bom_extract.LOW and "quantity not a number" in lines[2]["flags"]
    with pytest.raises(bom_extract.ExtractError, match="No BOM table"):
        bom_extract.parse_pdf(make_pdf([["a", "b"], ["c", "d"]]))
    with pytest.raises(bom_extract.ExtractError):
        bom_extract.parse_pdf(b"not a pdf")


KICAD = b"""Reference,Value,Footprint,Qty,MPN,Manufacturer
"C1, C2, C3",100nF,Capacitor_SMD:C_0603_1608Metric,3,GRM188R71H104KA93D,Murata
R1,10k,Resistor_SMD:R_0402_1005Metric,1,,
J1,Conn_2,Connector_PinHeader_2.54mm:PinHeader_1x02_P2.54mm_Vertical,1,61300211121,Wurth
U1,MCU,Package_QFP:LQFP-48,2,STM32F103C8T6,ST
"""
CPL = b"""Designator,Val,Package,Mid X,Mid Y,Rotation,Layer
C1,100nF,C_0603,10.0,5.0,0,top
C2,100nF,C_0603,40.5,5.0,0,top
C3,100nF,C_0603,10.0,30.0,0,top
R1,10k,R_0402,50.0,20.0,90,top
J1,Conn,PinHeader,5.0,2.0,0,top
U1,MCU,LQFP-48,25.0,15.0,0,top
"""


def test_pcb_bom_mount_counts_and_board_line():
    r = bom_extract.extract(KICAD, "bom.csv", typed={"layers": 4, "width_mm": 60, "height_mm": 40, "thickness_mm": 1.6})
    comp = {l["ref"]: l for l in r["lines"] if l["category"] == "component"}
    assert comp["C1, C2, C3"]["qty"] == 3 and comp["C1, C2, C3"]["mount"] == "SMT" and comp["R1"]["mount"] == "SMT" and comp["J1"]["mount"] == "THT"
    assert "quantity 2 differs from 1 designators" in comp["U1"]["flags"] and comp["U1"]["confidence"] < bom_extract.LOW
    assert comp["R1"]["confidence"] < 1 and "no manufacturer part number" in comp["R1"]["flags"]
    board = [l for l in r["lines"] if l["category"] == "pcb"][0]
    assert board["confidence"] == 1.0 and board["package"] == "4L" and "60.0 × 40.0 mm, 1.6 mm" in board["notes"]
    s = bom_extract.summary(r["lines"])
    assert s["by_mount"] == {"SMT": 6.0, "THT": 1.0} and s["low_confidence"] >= 1


def test_cpl_gives_extents_counts_and_cross_check():
    short = CPL.rsplit(b"U1", 1)[0]                                            # one placement missing
    r = bom_extract.extract(KICAD, "bom.csv", typed={}, cpl=(short, "cpl.csv"))
    board = [l for l in r["lines"] if l["category"] == "pcb"][0]
    assert "5 placements" in board["notes"] and "34.5" not in board["notes"]      # extents without U1: X 5 to 50, Y 2 to 30 = 45 × 28
    assert "45.0 × 28.0" in board["notes"] and board["confidence"] <= 0.5
    assert any("6 designators but the pick-and-place file 5 placements" in f for f in board["flags"])


def test_altium_and_excel_headers():
    wb = Workbook()
    ws = wb.active
    ws.append(["Designator", "Comment", "Footprint", "Quantity", "Manufacturer Part Number"])
    ws.append(["D1", "LED", "LED_0805", 1, "LTST-C170KRKT"])
    buf = io.BytesIO()
    wb.save(buf)
    r = bom_extract.extract(buf.getvalue(), "altium.xlsx", typed={"layers": 2})
    assert r["lines"][0]["part_number"] == "LTST-C170KRKT" and r["lines"][0]["mount"] == "SMT"
    with pytest.raises(bom_extract.ExtractError, match="No BOM header"):
        bom_extract.extract(b"a,b\n1,2\n", "x.csv")
    with pytest.raises(bom_extract.ExtractError, match="Unsupported"):
        bom_extract.extract(b"x", "drawing.step")


def test_api_review_edit_status_and_export():
    r = client.post("/api/bom-extracts", files={"file": ("harness.xml", KBL)})
    assert r.status_code == 200, r.text
    e = r.json()
    assert e["kind"] == "harness" and e["summary"]["low_confidence"] >= 1 and e["status"] == "draft"
    low = next(l for l in e["lines"] if l["category"] == "sleeve")
    e = client.put(f"/api/bom-extracts/{e['id']}/lines/{low['id']}", json={"part_number": "TUBE-12", "qty": 3}).json()
    fixed = next(l for l in e["lines"] if l["id"] == low["id"])
    assert fixed["confidence"] == 1.0 and fixed["edited"] and fixed["part_number"] == "TUBE-12" and fixed["qty"] == 3
    e = client.post(f"/api/bom-extracts/{e['id']}/lines", json={"category": "tape", "description": "Tesa 51036", "qty": 2}).json()
    assert any(l["category"] == "tape" for l in e["lines"])
    assert client.put(f"/api/bom-extracts/{e['id']}/lines/{low['id']}", json={"category": "nonsense"}).status_code == 400
    assert client.put(f"/api/bom-extracts/{e['id']}/lines/{low['id']}", json={"qty": -1}).status_code == 400
    e = client.post(f"/api/bom-extracts/{e['id']}/status", json={"status": "reviewed"}).json()
    assert e["status"] == "reviewed" and e["reviewed_by"]
    e = client.put(f"/api/bom-extracts/{e['id']}/lines/{low['id']}", json={"qty": 4}).json()
    assert e["status"] == "draft"                                                # an edit after review reopens it
    wb = load_workbook(io.BytesIO(client.get(f"/api/bom-extracts/{e['id']}/export.xlsx").content))
    assert wb.sheetnames == ["BOM", "Summary"] and wb["BOM"].max_row == len(e["lines"]) + 1 + 3          # + logo and name rows
    assert any(x["id"] == e["id"] for x in client.get("/api/bom-extracts").json())
    assert client.delete(f"/api/bom-extracts/{e['id']}/lines/{low['id']}").status_code == 200
    assert client.delete(f"/api/bom-extracts/{e['id']}").status_code == 200
    assert client.get(f"/api/bom-extracts/{e['id']}").status_code == 400
    bad = client.post("/api/bom-extracts", files={"file": ("x.step", b"x")})
    assert bad.status_code == 400 and "Unsupported" in bad.json()["detail"]


def test_pcb_upload_with_cpl_and_board_fields():
    r = client.post("/api/bom-extracts", files={"file": ("bom.csv", KICAD), "cpl": ("cpl.csv", CPL)}, data={"layers": "2", "thickness_mm": "1.6"})
    assert r.status_code == 200, r.text
    e = r.json()
    assert e["kind"] == "pcb" and e["format"] == "BOM + CPL" and e["board"]["layers"] == 2
