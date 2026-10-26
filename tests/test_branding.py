"""Stage 35: logo and company name on every Excel and PDF download (and inside ZIP packs); round-trip files stay plain."""

import io
import zipfile

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader
from reportlab.pdfgen import canvas

from backend.app import branding
from backend.app.main import app
from tests.test_rm_cost import add_part, project

client = TestClient(app)


def _plain_xlsx() -> bytes:
    from openpyxl import Workbook
    wb = Workbook()
    wb.active.append(["Part", "Cost"])
    wb.active.append(["A-1", 12.5])
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()


def test_excel_gets_the_banner_without_moving_any_cell():
    out = branding.stamp(_plain_xlsx(), "x.xlsx")
    ws = load_workbook(io.BytesIO(out)).active
    assert [c.value for c in ws[1]] == ["Part", "Cost"] and ws.cell(2, 2).value == 12.5     # values where the exporter put them
    assert len(ws._images) == 1 and ws.oddHeader.center.text == "Design to Cost"
    assert branding.stamp(out, "x.xlsx") == out                                              # once only


def test_pdf_without_the_name_is_stamped_once():
    buf = io.BytesIO(); c = canvas.Canvas(buf); c.drawString(100, 700, "plain report"); c.save()
    out = branding.stamp(buf.getvalue(), "x.pdf")
    assert "Design to Cost" in PdfReader(io.BytesIO(out)).pages[0].extract_text()
    assert branding.stamp(out, "x.pdf") == out


def test_downloads_are_stamped_and_round_trip_files_are_not():
    p = project("BR1")
    add_part(p, code="BR1-1")
    ws = load_workbook(io.BytesIO(client.get("/api/parts/export.xlsx").content)).worksheets[0]
    assert ws.oddHeader.center.text == "Design to Cost" and ws._images
    md = load_workbook(io.BytesIO(client.get("/api/master-data/download").content)).worksheets[0]
    assert md.oddHeader.center.text != "Design to Cost" and not md._images              # uploaded again: plain
    zbuf = io.BytesIO()
    with zipfile.ZipFile(zbuf, "w") as z:
        z.writestr("a.xlsx", _plain_xlsx())
    out = zipfile.ZipFile(io.BytesIO(branding.stamp(zbuf.getvalue(), "x.zip")))
    assert load_workbook(io.BytesIO(out.read("a.xlsx"))).active._images
