import io
import time
import zipfile
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import insert

from backend.app import projects_db as db
from backend.app import register, register_setup
from backend.app.main import app

SAMPLES = Path(__file__).resolve().parents[1] / "samples"
client = TestClient(app)


def project(code="SCT", **kw):
    p = client.post("/api/projects", json={"code": code, "name": "Scooter", "capex_budget_inr": 5_000_000, **kw}).json()
    return client.post(f"/api/projects/{p['id']}/variants",
                       json={"code": "STD", "name": "Standard", "target_bom_cost_inr": 45000, "annual_volume": 100000}).json()


def preview(name="demo_assembly.step"):
    with open(SAMPLES / name, "rb") as fh:
        r = client.post("/api/import/step", files={"file": (name, fh)})
    assert r.status_code == 200, r.text
    return r.json()


# ---------- projects ----------

def test_projects_and_variants():
    p = project()
    assert p["code"] == "SCT" and p["variants"][0]["code"] == "STD"
    assert client.post("/api/projects", json={"code": "sct", "name": "x"}).status_code == 400   # codes are upper-cased
    assert client.post(f"/api/projects/{p['id']}/variants", json={"code": "STD", "name": "dup"}).status_code == 400
    up = client.put(f"/api/projects/{p['id']}", json={"capex_budget_inr": 6_000_000}).json()
    assert up["capex_budget_inr"] == 6_000_000


# ---------- assembly split ----------

def test_assembly_preview_tree_and_lookalikes():
    m = preview()
    defs = m["defs"]
    by_name = {d["name"]: d for d in defs.values()}
    root = defs[m["root"]]
    assert root["is_assembly"] and root["name"] == "demo_assembly"
    assert by_name["bracket"]["total_quantity"] == 2          # 2 instances of one definition
    assert by_name["pin"]["total_quantity"] == 4              # "pin_1" tidied to "pin"
    assert by_name["pin"]["looks_like"] == ["locating_pin"]   # look-alike flagged, not merged
    assert by_name["base_plate"]["metrics"]["volume_cm3"] == pytest.approx(240)
    assert client.get(f"/api/import/{m['token']}/thumb/{m['root']}").status_code == 200


def test_assembly_commit_builds_parts_bom_and_files():
    p = project()
    m = preview()
    keys = {d["name"]: k for k, d in m["defs"].items()}
    items = {keys["pin"]: {"sourcing": "bought_out", "name": "Dowel pin 8x20"},
             keys["locating_pin"]: {"register": False},
             keys["base_plate"]: {"part_code": "BP-100", "material": "Mild steel CRCA / HR"}}
    r = client.post(f"/api/import/{m['token']}/commit",
                    json={"project_id": p["id"], "variant_id": p["variants"][0]["id"], "rev": "A", "items": items}).json()
    assert r["created"] == 5                                   # root, plate, group, bracket, pin
    parts = {x["name"]: x for x in client.get("/api/parts", params={"project_id": p["id"], "limit": 100}).json()["rows"]}
    assert parts["Base plate" if "Base plate" in parts else "base_plate"]["part_code"] == "BP-100"
    assert parts["Dowel pin 8x20"]["sourcing"] == "bought_out"
    assert parts["demo_assembly"]["is_assembly"] and parts["demo_assembly"]["process"] == register_setup.assembly_process()
    assert parts["bracket"]["part_code"].startswith("SCT-")   # generated from the pattern
    assert all(x["rev"] == "A" for x in parts.values())
    bom = {b["name"]: b for b in client.get(f"/api/variants/{p['variants'][0]['id']}/bom").json()}
    assert bom["bracket"]["quantity"] == 2 and bom["bracket"]["depth"] == 2
    assert bom["Dowel pin 8x20"]["total_quantity"] == 4
    assert "locating_pin" not in bom
    # a thumbnail and the part's own STEP file are stored and downloadable
    d = client.get(f"/api/parts/{parts['bracket']['id']}").json()
    assert client.get(f"/api/files/{d['thumbnail_file_id']}").headers["content-type"] == "image/png"
    step = client.get(f"/api/files/{d['file_id']}")
    assert step.status_code == 200 and b"ISO-10303-21" in step.content[:50]

    # importing the same file again links the unchanged parts and makes no new revisions
    m2 = preview()
    keys2 = {d["name"]: k for k, d in m2["defs"].items()}
    again = client.post(f"/api/import/{m2['token']}/commit", json={
        "project_id": p["id"], "variant_id": p["variants"][0]["id"],
        "items": {keys2["base_plate"]: {"part_code": "BP-100"}, keys2["locating_pin"]: {"register": False},
                  keys2["pin"]: {"part_code": parts["Dowel pin 8x20"]["part_code"]},
                  keys2["bracket"]: {"part_code": parts["bracket"]["part_code"]},
                  keys2["bracket_group"]: {"part_code": parts["bracket_group"]["part_code"]},
                  keys2["demo_assembly"]: {"part_code": parts["demo_assembly"]["part_code"]}}}).json()
    assert again["created"] == 0 and again["new_revisions"] == 0 and again["unchanged"] == 5


def test_single_part_and_new_revision_keeps_history():
    p = project()
    m = preview("sheet_metal_bracket.step")
    only = m["root"]
    assert m["defs"][only]["suggestion"]["process"].startswith("sheet_metal")    # rule-based proposal
    r = client.post(f"/api/import/{m['token']}/commit", json={
        "project_id": p["id"], "variant_id": p["variants"][0]["id"],
        "items": {only: {"part_code": "BRK-1", "material": "Mild steel CRCA / HR", "process": "sheet_metal_progressive"}}}).json()
    pid = r["part_ids"][0]
    with open(SAMPLES / "plate_with_hole.step", "rb") as fh:
        d = client.post(f"/api/parts/{pid}/revisions", files={"file": ("bracket_v2.step", fh)},
                        data={"rev": "B", "change_note": "thicker"}).json()
    assert [x["rev"] for x in d["revisions"]] == ["B", "A"]
    assert d["rev"] == "B" and d["revisions"][1]["geometry"]["wall_mm"] == pytest.approx(2.0)  # old rev kept
    with open(SAMPLES / "plate_with_hole.step", "rb") as fh:
        dup = client.post(f"/api/parts/{pid}/revisions", files={"file": ("x.step", fh)}, data={"rev": "B"})
    assert dup.status_code == 400 and "already has revision" in dup.json()["detail"]
    upd = client.put(f"/api/parts/{pid}", json={"status": "concept_tool", "target_piece_cost_inr": 12.5}).json()
    changes = [h for h in upd["history"] if h["action"] == "updated"][0]["detail"]
    assert changes["status"] == {"old": "uploaded", "new": "concept_tool"}
    assert [h["action"] for h in upd["history"]][-1] == "created"
    assert client.put(f"/api/parts/{pid}", json={"process": "welding"}).status_code == 400

    # traced geometry: weight = volume from the CAD file x density from master data
    t = client.get(f"/api/parts/{pid}/trace", params={"revision_id": d["revisions"][1]["id"]}).json()
    w = t["nodes"]["weight"]
    assert w["value"] == pytest.approx(11.8 * 7.85, rel=1e-6)
    assert t["nodes"]["volume"]["source"]["type"] == "cad" and t["nodes"]["density"]["source"]["sheet"] == "Materials"


# ---------- Excel bulk upload ----------

def _sheet(rows):
    wb = load_workbook(io.BytesIO(client.get("/api/import/template.xlsx").content))
    ws = wb["Parts"]
    ws.delete_rows(2, ws.max_row)  # the example rows
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_excel_bulk_with_zip_and_missing_file():
    p = project()
    cols = [c for c, _ in register.TEMPLATE_COLUMNS]
    row = lambda **kw: [kw.get(c, "") for c in cols]
    xlsx = _sheet([
        row(part_code="ASM-1", part_name="Frame assembly", rev="A", process="assembly", variant="STD"),
        row(part_code="P-1", part_name="Cover", rev="A", file_name="plastic_cover.step", material="ABS",
            annual_volume=100000, variant="STD", parent_part_code="ASM-1", quantity=1),
        row(part_name="Bracket", rev="01", file_name="sheet_metal_bracket.step", variant="STD", parent_part_code="ASM-1",
            quantity=2),
        row(part_code="P-3", part_name="Label", file_name="label.step", sourcing="bought_out", variant="STD",
            parent_part_code="ASM-1", quantity=1),
    ])
    z = io.BytesIO()
    with zipfile.ZipFile(z, "w") as zf:
        zf.write(SAMPLES / "plastic_cover.step", "cad/plastic_cover.step")
    files = [("files", ("parts.zip", z.getvalue())), ("files", ("sheet_metal_bracket.step", (SAMPLES / "sheet_metal_bracket.step").read_bytes()))]
    r = client.post("/api/import/excel", data={"project_id": p["id"]}, files=[("excel", ("parts.xlsx", xlsx))] + files)
    assert r.status_code == 200, r.text
    rep = r.json()
    assert rep["created"] == 4 and rep["file_missing"] == 2     # assembly row has no file, label.step not uploaded
    by_code = {x["part_code"]: x for x in rep["rows"]}
    assert by_code["P-3"]["file"] == "File missing"
    generated = [c for c in by_code if c.startswith("SCT-")]
    assert len(generated) == 1                                   # blank code generated
    bracket = client.get("/api/parts", params={"q": "Bracket"}).json()["rows"][0]
    assert bracket["process"].startswith("sheet_metal") and bracket["rev"] == "01"   # process suggested from the file
    bom = {b["part_code"]: b for b in client.get(f"/api/variants/{p['variants'][0]['id']}/bom").json()}
    assert bom[generated[0]]["quantity"] == 2 and bom["ASM-1"]["depth"] == 0
    assert client.get("/api/parts", params={"file_missing": "true", "project_id": p["id"]}).json()["total"] == 2

    # same code again = new revision; a bad row means nothing at all is saved
    again = _sheet([row(part_code="P-1", part_name="Cover", rev="B", change_note="new rib")])
    assert client.post("/api/import/excel", data={"project_id": p["id"]},
                       files=[("excel", ("x.xlsx", again))]).json()["new_revisions"] == 1
    bad = _sheet([row(part_code="P-9", part_name="New", rev="A"), row(part_code="P-10", part_name="Bad", process="welding")])
    r = client.post("/api/import/excel", data={"project_id": p["id"]}, files=[("excel", ("x.xlsx", bad))])
    assert r.status_code == 400 and "Parts row 3" in r.json()["detail"] and "Nothing was saved" in r.json()["detail"]
    assert client.get("/api/parts", params={"q": "P-9"}).json()["total"] == 0


# ---------- search, dashboard, exports ----------

def test_search_10000_parts_is_fast():
    p = project()
    with db.engine().begin() as conn:
        conn.execute(insert(db.parts), [
            {"part_code": f"SCT-{i:05d}", "name": f"Part {i} {'bracket' if i % 7 == 0 else 'cover'}",
             "project_id": p["id"], "process": "machined" if i % 3 else "injection_moulding", "status": "uploaded",
             "sourcing": "make", "target_piece_cost_inr": i % 100} for i in range(10_000)])
    t = time.perf_counter()
    r = client.get("/api/parts", params={"q": "bracket", "process": "injection_moulding", "target_min": 10,
                                         "target_max": 50, "limit": 50}).json()
    first = client.get("/api/parts", params={"project_id": p["id"], "limit": 50, "offset": 9950}).json()
    elapsed = time.perf_counter() - t
    expected = sum(1 for i in range(10_000) if i % 7 == 0 and i % 3 == 0 and 10 <= i % 100 <= 50)
    assert r["total"] == expected and len(first["rows"]) == 50
    assert elapsed < 2.0, f"search took {elapsed:.2f} s"


def test_dashboard_and_exports():
    p = project()
    m = preview()
    client.post(f"/api/import/{m['token']}/commit", json={"project_id": p["id"], "variant_id": p["variants"][0]["id"]})
    pid = client.get("/api/parts", params={"q": "bracket"}).json()["rows"][0]["id"]
    client.put(f"/api/parts/{pid}", json={"target_piece_cost_inr": 40})
    for st, c in (("concept_tool", ""), ("simulation", ""), ("corrections_required", "add draft")):   # design-loop order
        assert client.post(f"/api/parts/{pid}/status", json={"status": st, "comment": c}).status_code == 200
    d = client.get(f"/api/projects/{p['id']}/dashboard").json()
    assert d["parts_total"] == 6 and d["open_corrections"] == 1 and d["parts_with_target"] == 1
    assert {s["code"]: s["count"] for s in d["status"]}["uploaded"] == 5
    assert d["variants"][0]["unique_parts"] == 6 and d["variants"][0]["target_bom_cost_inr"] == 45000
    assert d["capex"]["budget_inr"] == 5_000_000 and d["capex"]["tool_cost_inr"] is None
    pdf = client.get(f"/api/projects/{p['id']}/dashboard.pdf")
    assert pdf.content[:4] == b"%PDF"
    xlsx = client.get("/api/parts/export.xlsx", params={"project_id": p["id"]})
    ws = load_workbook(io.BytesIO(xlsx.content))["Part register"]
    assert ws.max_row == 7 + 3                                  # + logo and name rows


# ---------- master data for the register ----------

@pytest.mark.parametrize("change, message", [
    (lambda s: s["register_settings"][0].update(value="{project}-{number}"), "not valid"),
    (lambda s: s["register_settings"][0].update(value="{project}"), "must contain {seq}"),
    (lambda s: s["part_statuses"][1].update(is_initial=True), "exactly one status"),
    (lambda s: s["processes"][0].update(suggest_rule="wall_mmm < 4"), "unknown name"),
])
def test_register_setup_is_checked(change, message):
    setup = client.get("/api/register-setup").json()
    change(setup)
    r = client.put("/api/register-setup", json=setup)
    assert r.status_code == 400 and message in r.json()["detail"]


def test_part_code_pattern_comes_from_master_data():
    setup = client.get("/api/register-setup").json()
    for s in setup["register_settings"]:
        if s["key"] == "part_code_pattern":
            s["value"] = "{project}/P{seq:03d}"
    client.put("/api/register-setup", json=setup)
    p = project()
    m = preview("plastic_cover.step")
    client.post(f"/api/import/{m['token']}/commit", json={"project_id": p["id"]})
    assert client.get("/api/parts").json()["rows"][0]["part_code"] == "SCT/P001"


def test_unticked_assembly_passes_children_up():
    p = project()
    m = preview()
    keys = {d["name"]: k for k, d in m["defs"].items()}
    client.post(f"/api/import/{m['token']}/commit", json={
        "project_id": p["id"], "variant_id": p["variants"][0]["id"],
        "items": {keys["bracket_group"]: {"register": False}}})
    bom = {b["name"]: b for b in client.get(f"/api/variants/{p['variants'][0]['id']}/bom").json()}
    assert "bracket_group" not in bom
    assert bom["bracket"]["depth"] == 1 and bom["bracket"]["quantity"] == 2   # now directly under the root
