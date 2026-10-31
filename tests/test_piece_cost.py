import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.app.main import app
from tests.test_rm_cost import add_part, project

client = TestClient(app)


def node(tr, name):
    return tr["nodes"][f"calc:{name}"]["value"]


def inp(tr, name):
    n = next(i["node"] for i in tr["inputs"] if i["variable"] == name)
    return tr["nodes"][n]["value"]


def costed_part(code="PC8"):
    pid = add_part(project(code))
    assert client.post(f"/api/parts/{pid}/tool-cost", json={}).status_code == 200
    return pid


def test_piece_cost_by_hand():
    pid = costed_part()
    s = client.get(f"/api/parts/{pid}/piece-cost").json()
    assert s["problem"] == "", s["problem"]
    tr = s["live"]
    v = {k: inp(tr, k) for k in ("rm_cost", "cycle_s", "cavities", "mhr", "labour_rate", "operators_per_machine",
                                 "rejection_pct", "part_overhead_pct", "profit_pct", "packing_pct", "tool_capex",
                                 "lifetime_volume", "total_weight_kg", "transport_rate")}
    sim = client.get(f"/api/parts/{pid}/simulation").json()["results"]["cycle_s"]
    assert v["cycle_s"] == pytest.approx(sim, abs=1e-3) and v["cycle_s"] != 25     # the calculated cycle, not the target of Parameters
    assert v["operators_per_machine"] == 0.5 and v["profit_pct"] == 10
    pph = 3600 / v["cycle_s"] * v["cavities"]
    process = (v["mhr"] + 0.5 * v["labour_rate"]) / pph
    rej = (v["rm_cost"] + process) * 2 / 98
    before = v["rm_cost"] + process + rej + process * 0.05
    total = before * 1.10 * 1.015 + v["tool_capex"] / v["lifetime_volume"] + v["total_weight_kg"] * v["transport_rate"]
    assert node(tr, "process_cost") == pytest.approx(process)
    assert tr["value"] == pytest.approx(total)
    assert tr["machine"].startswith("IMM") and tr["is_example"]
    # MHR is grafted with its whole trace: click through to the electricity rate
    assert any(k.startswith("mhr:in:electricity") for k in tr["nodes"])


def test_versions_typed_inputs_and_exports():
    pid = costed_part("PC9")
    r = client.post(f"/api/parts/{pid}/piece-cost", json={}).json()
    assert r["stored"] and r["version"] == 1
    assert client.post(f"/api/parts/{pid}/piece-cost", json={}).json()["stored"] is False   # nothing changed
    r = client.post(f"/api/parts/{pid}/piece-cost", json={"piece_inputs": {"cycle_s": 50}}).json()
    assert r["version"] == 2
    s = client.get(f"/api/parts/{pid}/piece-cost").json()
    assert [x["version"] for x in s["versions"]] == [2, 1] and s["outdated"] == ""
    v1 = client.get(f"/api/parts/{pid}/piece-cost/trace", params={"version": 1}).json()
    sim = client.get(f"/api/parts/{pid}/simulation").json()["results"]["cycle_s"]
    assert inp(s["live"], "cycle_s") == 50 and inp(v1, "cycle_s") == pytest.approx(sim, abs=1e-3) and s["live"]["value"] > v1["value"]
    wb = load_workbook(io.BytesIO(client.get(f"/api/parts/{pid}/piece-cost/export.xlsx").content))
    assert wb.sheetnames == ["Summary", "Steps", "Inputs", "Rates used"]
    assert client.get(f"/api/parts/{pid}/piece-cost/export.pdf").content[:4] == b"%PDF"
    assert client.put(f"/api/parts/{pid}", json={"piece_inputs": {"nope": 1}}).status_code == 400


def test_needs_tool_cost_first():
    pid = add_part(project("PC10"))
    from backend.app import projects_db as db
    with db.engine().begin() as conn:
        conn.execute(db.tool_costs.delete().where(db.tool_costs.c.part_id == pid))
    assert "tool cost" in client.get(f"/api/parts/{pid}/piece-cost").json()["problem"]


def test_inserts_add_cost_loading_and_rejection():
    pid = costed_part("PC11")
    r = client.post(f"/api/parts/{pid}/inserts", json={
        "insert_part_code": "NUT-M4X", "name": "M4 brass nut", "material": "Brass CuZn39Pb3", "quantity": 2,
        "dimensions": {"shape": "hex_nut", "af_mm": 7, "id_mm": 3.3, "l_mm": 3.2},
        "unit_price": 1.8, "price_currency": "INR", "price_source": "quote", "price_month": "2026-10"})
    assert r.status_code == 200, r.text
    client.post(f"/api/parts/{pid}/tool-cost", json={})
    tr = client.get(f"/api/parts/{pid}/piece-cost").json()["live"]
    assert node(tr, "insert_total") == pytest.approx(2 * 1.8 * 1.02)                 # + 2 % handling
    assert node(tr, "reject_pct") == pytest.approx(2 + 1.5)                         # + extra rejection
    assert node(tr, "load_s") == pytest.approx(2 * inp(tr, "cavities") * 3.0)       # manual loading per shot


def test_cycle_time_is_the_calculated_one_by_default_and_editable():
    pid = costed_part("PC12")
    sim = client.get(f"/api/parts/{pid}/simulation").json()["results"]["cycle_s"]
    tr = client.get(f"/api/parts/{pid}/piece-cost").json()["live"]
    n = tr["nodes"][next(i["node"] for i in tr["inputs"] if i["variable"] == "cycle_s")]
    assert n["value"] == pytest.approx(sim, abs=1e-3) and "simulation" in n["source"]["note"]
    # typed over on the part: the typed value wins, and going back to blank gives the calculated cycle again
    assert client.post(f"/api/parts/{pid}/piece-cost", json={"piece_inputs": {"cycle_s": 40}}).status_code == 200
    assert inp(client.get(f"/api/parts/{pid}/piece-cost").json()["live"], "cycle_s") == 40
    assert client.put(f"/api/parts/{pid}", json={"piece_inputs": {}}).status_code == 200
    assert inp(client.get(f"/api/parts/{pid}/piece-cost").json()["live"], "cycle_s") == pytest.approx(sim, abs=1e-3)


def test_sheet_metal_cycle_from_the_press_speed_rules():
    from backend.app import mould_db, piece_cost
    assert not piece_cost.seed()                                                  # rows are shipped; seeding twice adds nothing
    assert {"press_spm", "press_spm_manual", "press_spm_draw_factor"} <= set(mould_db.design_rules("sheet_metal"))
    pid = add_part(project("PC13"), "sheet_metal_bracket.step", "Mild steel CRCA / HR", "sheet_metal_multistage", "SM-1")
    res = client.get(f"/api/parts/{pid}/sheet-tools").json()
    ops = len(res["operations"])
    cyc, why = piece_cost.calculated_cycle(pid, "sheet_metal")
    spm = float(mould_db.design_rules("sheet_metal")["press_spm_manual" if res["mode"] != "progressive" else "press_spm"]["value"])
    strokes = 1 if res["mode"] == "progressive" else ops
    drawn = float(res["features"].get("draw_depth") or 0) > 0
    f = float(mould_db.design_rules("sheet_metal")["press_spm_draw_factor"]["value"]) if drawn else 1
    assert cyc == pytest.approx(strokes * 60 / (spm * f)) and "press_spm" in why
