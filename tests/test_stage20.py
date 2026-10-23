"""Stage 20: capacity 0-3000 t, 3D design correction, stage-wise sheet-metal tool design, design pack downloads."""
import io
import zipfile

import cadquery as cq
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.app import capacity_db, corrections, geometry, mould_db, software, tool_design
from backend.app.main import app
from tests.test_rm_cost import add_part, project

client = TestClient(app)


# ---------- capacity ----------

def test_machine_ladders_reach_3000_t_and_extend_is_idempotent():
    assert capacity_db.extend()["machines"] == [] and capacity_db.extend()["rates"] == 0
    lim = {r["process"]: r for r in capacity_db.process_limits()}
    for p in ("Injection moulding", "Die casting", "Press"):
        assert lim[p]["max_t"] == 3000
    assert not lim["EPS moulding"]["reaches_3000t"] if "reaches_3000t" in lim["EPS moulding"] else True


def test_coverage_every_step_has_machine_base_and_prices_in_all_countries():
    r = capacity_db.coverage(250)
    assert r["ok"], r["gaps"][:3]
    assert all(x["rate_countries"] == 9 for x in r["rows"])
    assert any(x["required_t"] == 3000 and x["machine"].endswith("3000") for x in r["rows"])


def test_every_material_at_every_tonnage_designs_a_tool_and_fits_the_tie_bars():
    r = capacity_db.validate_designs([100, 1000, 3000])
    assert r["ok"], r["problems"][:2]
    big = [x for x in r["rows"] if x["target_t"] == 3000]
    assert big and all(x["machine_t"] == 3000 or x["machine_t"] >= 2500 for x in big)


def test_sheet_stage_dies_pick_a_press_whose_bolster_takes_the_die_set():
    r = capacity_db.validate_sheet([100, 3000])
    assert r["ok"], r["problems"][:2]
    assert {x["press"] for x in r["rows"] if x["target_t"] == 3000} == {"PRESS-3000"}
    small = tool_design.press_fit(50, 3000, 2000)         # tonnage is enough, bed is not
    assert small["machine_id"] and small["bolster_mm"][0] >= 3000 or small["bolster_mm"][1] >= 3000
    assert not tool_design.press_fit(5000, 100, 100)["fits"]


# ---------- 3D correction ----------

def test_draft_sign_convention_and_correction_reduces_low_draft_area():
    s = geometry.load_step("samples/plastic_cover.step")
    r = corrections.apply_draft(s, "z", 1.0)
    assert r["faces_failed"] == [] and r["improved"]
    assert r["after"]["low_draft_cm2"] < 0.01 * r["before"]["low_draft_cm2"]
    assert r["shape"].isValid() and abs(r["volume_change_pct"]) < 3
    assert r["parting_z_mm"] == pytest.approx(15.0)       # open-top shell: parted at the rim
    plate = cq.Workplane("XY").box(100, 60, 20).faces(">Z").workplane().hole(20).val()
    p = corrections.apply_draft(plate, "z", 1.0)
    assert p["after"]["low_draft_cm2"] < p["before"]["low_draft_cm2"]


def test_draft_works_on_any_opening_axis():
    box = cq.Workplane("XY").box(40, 30, 60).val()
    r = corrections.apply_draft(box, "x", 1.0)
    assert r["axis"] == "x" and r["improved"]


# ---------- stage rules are data ----------

def test_stage_rules_are_master_data():
    db = mould_db.tables()
    assert any(r["name"] == "die_set_l_mm" for r in db["stage_rules"])
    next(r for r in db["stage_rules"] if r["name"] == "margin_mm")["formula"] = "200"
    mould_db.write(db)
    vals, _ = tool_design.rule_values("blank", {"t": 2, "shear_mpa": 300, "force_t": 10, "work_l": 100, "work_w": 50})
    assert vals["margin_mm"] == 200 and vals["die_set_l_mm"] == 500
    next(r for r in db["stage_rules"] if r["name"] == "margin_mm")["formula"] = "no_such_name + 1"
    with pytest.raises(mould_db.MouldDbError, match="Stage tool rules"):
        mould_db.write(db)


def test_deep_draw_is_split_into_stages_by_the_limiting_draw_ratio():
    from backend.app import sheet_tools
    cup = cq.Workplane("XY").circle(60).extrude(30).faces(">Z").shell(-2).val()
    m = geometry.analyze(cup, "sheet_metal", 7.85)
    g = {"volume_mm3": m.volume_mm3, "bbox_mm": m.bbox_mm, "wall_mm": 2.0, "bend_count": 0}
    f = sheet_tools.features(cup, g, geometry.blank_estimate(m.volume_mm3, m.bbox_mm, 2.0), {"draw_depth_mm": 28}, 7.85, 200000, 5)
    f["blank_dia"] = 400
    res = sheet_tools.run(f, "Mild steel CRCA / HR", {"mode": "multi_stage"})
    fp = {"length_mm": 400, "width_mm": 400, "outline": [[0, 0], [400, 0], [400, 400], [0, 400], [0, 0]], "holes": [], "area_mm2": 160000}
    d = tool_design.sheet_design(fp, {}, res, {"bbox_mm": m.bbox_mm})
    draws = [s for s in d["stages"] if s["code"] == "draw"]
    assert len(draws) >= 2 and draws[0]["workpiece"]["cup_dia_mm"] > draws[-1]["workpiece"]["cup_dia_mm"]
    assert draws[1]["force_t"] < draws[0]["force_t"]
    sol = tool_design.stage_solids(d, draws[0], fp, None)
    assert {x["group"] for x in sol} >= {"lower", "upper", "work"}


# ---------- design pack ----------

def test_software_status_lists_required_tools():
    rows = {s["id"]: s for s in client.get("/api/free-software").json()}
    assert rows["cadquery"]["installed"] and rows["pptx"]["installed"] and rows["gmsh"]["required"] is False
    assert software.ffmpeg()


def test_capacity_endpoints():
    r = client.get("/api/capacity/validation?step_t=1000").json()
    assert r["ok"] and r["totals"]["failed"] == 0
    assert client.get("/api/capacity/validation?step_t=5").status_code == 400
    assert client.get("/capacity").status_code == 200


def test_mould_design_pack_files_and_downloads():
    pid = add_part(project(), "plastic_cover.step", "ABS", "injection_moulding")
    r = client.get(f"/api/parts/{pid}/design-pack?video=none")
    assert r.status_code == 200, r.text
    out = r.json()
    s = out["summary"]
    assert s["kind"] == "mould" and s["draft_after_cm2"] < 0.01 * s["draft_before_cm2"] and s["faces_drafted"] > 0
    for kind in ("xlsx", "pdf", "pptx", "3d", "drawing", "drawing-dxf", "corrected-step", "tool-step", "tool-drawing"):
        assert out["kinds"][kind], kind
    wb = load_workbook(io.BytesIO(client.get(f"/api/parts/{pid}/design-pack/xlsx").content))
    assert "Summary" in wb.sheetnames and len(wb.sheetnames) >= 5
    from pptx import Presentation
    assert len(Presentation(io.BytesIO(client.get(f"/api/parts/{pid}/design-pack/pptx").content)).slides) >= 6
    assert client.get(f"/api/parts/{pid}/design-pack/pdf").content[:4] == b"%PDF"
    html = client.get(f"/api/parts/{pid}/design-pack/3d").text
    assert "three-src" in html and "Mould concept" in html
    step = client.get(f"/api/parts/{pid}/design-pack/corrected-step").content
    assert b"ISO-10303" in step[:200]
    z = zipfile.ZipFile(io.BytesIO(client.get(f"/api/parts/{pid}/design-pack.zip?video=none").content))
    assert any(n.endswith("_design_report.xlsx") for n in z.namelist()) and "README.txt" in z.namelist()
    assert client.get(f"/api/parts/{pid}/design-pack/nope").status_code == 404


def test_sheet_metal_design_pack_has_every_stage_and_a_video():
    pid = add_part(project("S"), "sheet_metal_bracket.step", "Mild steel CRCA / HR", "sheet_metal_multistage", "SM-1")
    out = client.get(f"/api/parts/{pid}/design-pack?video=short").json()
    s = out["summary"]
    assert s["kind"] == "sheet_metal" and s["stages"] >= 2 and s["mode"] == "multi_stage"
    names = [f["name"] for f in out["files"]]
    assert sum(n.endswith(".step") and "_stage" in n for n in names) == s["stages"]
    assert any(n.endswith("_simulation.mp4") for n in names) and any(n.endswith("_stage_tool_drawings.pdf") for n in names)
    mp4 = client.get(f"/api/parts/{pid}/design-pack/video?video=short").content
    assert mp4[4:8] == b"ftyp"
    assert client.get(f"/api/parts/{pid}/design-pack/blank-dxf?video=short").status_code == 200


def test_design_pack_refuses_unsupported_parts():
    pid = add_part(project("E"), "plate_with_hole.step", "Mild steel CRCA / HR", "machined", "MA-1")
    assert client.get(f"/api/parts/{pid}/design-pack").status_code == 400


# ---------- completion of the stage-20 limits ----------

def test_wall_thickness_ignores_wall_ends_so_the_cover_cycle_is_realistic():
    import numpy as np
    from backend.app import simulation
    s = geometry.load_step("samples/plastic_cover.step")
    v, t = geometry.tessellate(s)
    m = geometry.analyze(s, "plastic", 1.04)
    facts = {"part_weight_g": 50, "wall_mm": 0, "part_l": 120, "part_w": 80, "part_h": 30, "projected_area_cm2": 95.8,
             "volume_cm3": 46, "cavities": 1, "slides": 0, "lifters": 0, "undercut_area_cm2": 0, "inserts": 0}
    r = simulation.run(v, t, "plastic", facts, m.projected_area_mm2)
    assert r["wall_mm"] == pytest.approx(2.5, abs=0.05) and r["results"]["max_wall_mm"] == pytest.approx(2.5, abs=0.2)
    assert r["results"]["cycle_s"] < 30                          # was about 1800 s: the rim measured the wall LENGTH
    assert r["maps"]["sink_risk_cm2"] == 0                       # the rim is not a thick section


def test_forming_surfaces_come_from_the_part_and_the_mould_split_has_no_overlap():
    box = cq.Workplane("XY").box(60, 40, 30, centered=(True, True, False)).faces(">Z").shell(-2).val()
    cup = (cq.Workplane("XY").circle(30).extrude(20).faces(">Z").shell(-2)
           .union(cq.Workplane("XY").circle(45).circle(28).extrude(2).translate((0, 0, 18))).val())
    placed = tool_design.place_formed(cup, 0.0)
    assert tool_design.flange_level(placed) == pytest.approx(0.0) and placed.BoundingBox().zmin < -17     # cup points down
    fp = tool_design.forming_parts(placed, 0.0, 2.0, 200, 150, 80)
    assert fp["die_cut"].Volume() > 0 and fp["core"].Volume() > 1000                                       # punch core fills the cup
    tool = {"base": {"length_mm": 400, "width_mm": 400, "a_plate_mm": 80, "b_plate_mm": 80, "clamp_plates_mm": 25, "support_plate_mm": 40,
                     "spacer_block_mm": 80, "ejector_plates_mm": 45, "size_code": "T"},
            "layout": {"rows": 1, "cols": 1, "empty": 0, "each_l": 80, "each_w": 60, "pocket_l": 80, "pocket_w": 60}}
    m = tool_design.mould_3d(box, "z", 30.0, tool)
    a = next(c["shape"] for c in m["components"] if c["name"] == "A side insert")
    b = next(c["shape"] for c in m["components"] if c["name"] == "B side insert")
    assert a.Volume() > 0 and b.Volume() > 0
    # inserts are shown in their open positions; compare in the closed position (A back down by the opening gap)
    a_closed = a.translate(cq.Vector(0, 0, -m["open_gap_mm"]))
    assert a_closed.intersect(b).Volume() < 1.0


def test_sheet_model_is_corrected_bends_get_a_radius_and_small_holes_are_opened():
    s = geometry.load_step("samples/sheet_metal_bracket.step")
    r = corrections.correct_sheet(s, 2.0, 1.0, 1.0)
    assert r["changed"] and r["shape"].isValid()
    assert {c["kind"] for c in r["changes"]} == {"bend radius (inner)", "bend radius (outer)"}
    assert sum(1 for f in r["shape"].Faces() if f.geomType() == "CYLINDER") == 2
    assert not corrections.correct_sheet(r["shape"], 2.0, 1.0, 1.0)["changed"]       # already manufacturable
    plate = cq.Workplane("XY").box(80, 60, 2).faces(">Z").workplane().rarray(30, 20, 2, 2).hole(1).val()
    h = corrections.correct_sheet(plate, 2.0, 1.0, 1.0)
    assert h["changes"] == [{"kind": "hole opened to the minimum", "count": 4, "value_mm": 2.0}]
    assert h["shape"].Volume() < plate.Volume()


def test_springback_overbend_is_a_stage_rule():
    vals, _ = tool_design.rule_values("bend", {"t": 2, "shear_mpa": 300, "force_t": 5, "work_l": 80, "work_w": 50, "bend_radius": 2,
                                               "bend_angle_deg": 90, "springback_deg": 3})
    assert vals["bend_punch_angle_deg"] == 87


def test_fea_matches_the_hand_calculation():
    from backend.app import fea
    if not (fea.tools()["gmsh"] and fea.tools()["ccx"]):
        pytest.skip("Gmsh / CalculiX not installed")
    blk = cq.Workplane("XY").box(50, 50, 40, centered=(True, True, False)).val()
    r = fea.solve(blk, 100.0, "top", {})
    assert r["max_deflection_mm"] == pytest.approx(100 * 40 / 210000, rel=0.1)       # u = p·L / E for a fixed-end block
    assert r["ok"] and r["vtk"].startswith(b"# vtk")
    weak = fea.solve(blk, 1500.0, "top", {})
    assert not weak["ok"]


def test_design_pack_reports_the_fea_and_the_sheet_correction():
    pid = add_part(project("F"), "sheet_metal_bracket.step", "Mild steel CRCA / HR", "sheet_metal_multistage", "SM-2")
    s = client.get(f"/api/parts/{pid}/design-pack?video=none").json()["summary"]
    assert s["corrections"] and s["springback_overbend_deg"] is not None
    from backend.app import fea
    if fea.tools()["gmsh"] and fea.tools()["ccx"]:
        assert s["fea"]["available"] and s["fea"]["ok"]
