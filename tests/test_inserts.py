"""Process suggestion (material first, with reasons) and moulded-in inserts."""

import io
import math
import sqlite3
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select

from backend.app import cad_import, projects_db as db, register, register_setup
from backend.app.main import app

SAMPLES = Path(__file__).resolve().parents[1] / "samples"
client = TestClient(app)


def project(code="INS"):
    p = client.post("/api/projects", json={"code": code, "name": "Insert test"}).json()
    return client.post(f"/api/projects/{p['id']}/variants", json={"code": "STD", "name": "Standard"}).json()


def preview(name):
    with open(SAMPLES / name, "rb") as fh:
        r = client.post("/api/import/step", files={"file": (name, fh)})
    assert r.status_code == 200, r.text
    return r.json()


def metrics(name):
    t = cad_import.analyse_file(str(SAMPLES / name), name, with_files=False)
    return t["defs"][t["root"]]["metrics"]


def suggest(m, material="", inserts=0):
    r = client.post("/api/suggest-process", json={"metrics": m, "material": material, "insert_count": inserts})
    assert r.status_code == 200, r.text
    return r.json()


# ---------- material first, then geometry ----------

@pytest.mark.parametrize("material, family", [
    ("ABS", "injection_moulding"),
    ("Mild steel CRCA / HR", "sheet_metal"),
    ("Aluminium 5052", "sheet_metal"),
    ("Aluminium ADC12 / A383", "die_casting"),
    ("Grey cast iron FG260", "foundry"),
])
def test_material_decides_the_process_family(material, family):
    s = suggest(metrics("sheet_metal_bracket.step"), material)
    assert s["process"].startswith(family)
    assert material in s["steps"][0] and "Master data → Materials" in s["steps"][0]


def test_geometry_only_when_no_material_and_conflict_is_explained():
    bracket = metrics("sheet_metal_bracket.step")
    assert bracket["bend_count"] == 1 and bracket["sheet_share"] > 0.9
    s = suggest(bracket)
    assert s["process"] == "sheet_metal_progressive" and "geometry-only" in s["reason"]
    assert "1 bend" in s["reason"]
    # a plastic material on a sheet-like shape: the material wins, and the reason says so
    s = suggest(bracket, "ABS")
    assert s["process"] == "injection_moulding"
    assert "Geometry alone would point to" in s["reason"]


def test_inserts_turn_the_suggestion_into_insert_moulding():
    cover = metrics("plastic_cover.step")
    assert suggest(cover, "PA66 GF30", 4)["process"] == "insert_moulding"
    assert suggest(cover, "Aluminium ADC12 / A383", 2)["process"] == "insert_die_casting"
    assert suggest(cover, "PA66 GF30", 0)["process"] == "injection_moulding"
    assert suggest(None, "Brass CuZn39Pb3")["process"] == "bought_out"     # an insert material


# ---------- detection in a 3D assembly ----------

def test_housing_inserts_are_detected():
    m = preview("insert_moulded_housing.step")
    names = {k: d["name"] for k, d in m["defs"].items()}
    assert len(m["inserts"]) == 1
    det = m["inserts"][0]
    assert names[det["body"]] == "housing" and names[det["insert"]] == "threaded_insert"
    assert det["quantity"] == 4 and det["covered_share"] > 0.5
    body, ins = m["defs"][det["body"]], m["defs"][det["insert"]]
    assert body["insert_count"] == 4 and body["suggestion"]["process"] == "insert_moulding"
    assert ins["insert_of"] == ["housing"] and ins["suggestion"]["process"] == "bought_out"
    # "housing_with_inserts" holds only the body and its inserts: unticked by default in the tree
    assy = next(d for d in m["defs"].values() if d["name"] == "housing_with_inserts")
    assert assy["only_moulded_part"] == det["body"]
    # the cover sitting on the housing, and the demo assembly's pins on the plate, are not inserts
    assert preview("demo_assembly.step")["inserts"] == []


def _commit(m, confirm):
    p = project()
    det = m["inserts"][0]
    assy = next(k for k, d in m["defs"].items() if d["name"] == "housing_with_inserts")
    items = {det["body"]: {"material": "PA66 GF30", "process": "insert_moulding", "part_code": "HSG-1",
                           "inserts": [det["insert"]] if confirm else []},
             det["insert"]: {"material": "Brass CuZn39Pb3", "part_code": "INS-M3", "sourcing": "bought_out"},
             assy: {"register": False}}
    r = client.post(f"/api/import/{m['token']}/commit", json={"project_id": p["id"], "variant_id": p["variants"][0]["id"],
                                                               "items": items})
    assert r.status_code == 200, r.text
    return p, r.json()


def test_confirmed_inserts_are_linked_not_bom_children():
    m = preview("insert_moulded_housing.step")
    p, res = _commit(m, confirm=True)
    assert res["inserts_linked"] == 1
    bom = client.get(f"/api/variants/{p['variants'][0]['id']}/bom").json()
    rows = {b["part_code"]: b for b in bom}
    assert rows["HSG-1"]["kind"] == "bom" and rows["HSG-1"]["depth"] == 1        # body under the top assembly
    assert rows["INS-M3"]["kind"] == "insert" and rows["INS-M3"]["total_quantity"] == 4
    with db.engine().connect() as conn:
        assert conn.execute(select(db.bom_lines).join(db.parts, db.parts.c.id == db.bom_lines.c.child_part_id)
                            .where(db.parts.c.part_code == "INS-M3")).first() is None
    hsg = client.get("/api/parts", params={"q": "HSG-1"}).json()["rows"][0]
    d = client.get(f"/api/parts/{hsg['id']}").json()
    assert d["insert_count"] == 4 and d["inserts"][0]["detected"] and d["inserts"][0]["volume"]["method"] == "3d"
    assert d["suggestion"]["process"] == "insert_moulding"
    ins = client.get(f"/api/parts/{d['inserts'][0]['insert']['id']}").json()
    assert ins["moulded_into"][0]["part_code"] == "HSG-1" and ins["sourcing"] == "bought_out"

    # weights: body (PA66) and inserts (brass) separately, total = sum
    t = client.get(f"/api/parts/{hsg['id']}/trace").json()
    n = t["nodes"]
    insert_vol = math.pi * (6 ** 2 - 3.4 ** 2) / 4 * 8 / 1000
    assert n["insert_weight"]["value"] == pytest.approx(4 * insert_vol * 8.47, rel=1e-3)
    assert n["weight"]["label"].startswith("Body weight")
    assert n["total_weight"]["value"] == pytest.approx(n["weight"]["value"] + n["insert_weight"]["value"])
    assert t["root"] == "total_weight" and t["warnings"] == []


def test_unconfirmed_insert_stays_a_bom_child():
    m = preview("insert_moulded_housing.step")
    p, res = _commit(m, confirm=False)
    assert res["inserts_linked"] == 0
    rows = {b["part_code"]: b for b in client.get(f"/api/variants/{p['variants'][0]['id']}/bom").json()}
    assert rows["INS-M3"]["kind"] == "bom" and rows["INS-M3"]["quantity"] == 4


# ---------- linking inserts by hand ----------

def _part(code, **kw):
    p = client.get("/api/projects").json()[0]
    with db.engine().begin() as conn:
        pid, _ = register.upsert_part(conn, p["id"], code, {"name": code, **kw})
    return pid


def test_link_insert_with_dimensions_and_price():
    project()
    body = _part("BODY-1", material="PBT GF30")
    r = client.post(f"/api/parts/{body}/inserts", json={
        "insert_part_code": "NUT-M4", "name": "M4 brass nut", "material": "Brass CuZn39Pb3", "quantity": 2,
        "dimensions": {"shape": "hex_nut", "af_mm": 7, "id_mm": 3.3, "l_mm": 3.2},
        "unit_price": 1.8, "price_currency": "INR", "price_source": "supplier quote", "price_month": "2026-10"})
    assert r.status_code == 200, r.text
    d = r.json()
    link = d["inserts"][0]
    assert link["insert"]["part_code"] == "NUT-M4" and link["insert"]["unit_price"] == 1.8
    assert link["insert"]["sourcing"] == "bought_out" and link["volume"]["method"] == "dimensions"
    vol = (math.sqrt(3) / 2 * 7 ** 2 - math.pi * 3.3 ** 2 / 4) * 3.2 / 1000
    assert link["weight_g"] == pytest.approx(2 * vol * 8.47)
    assert d["suggestion"]["process"] == "insert_moulding"     # PBT + inserts, no 3D file yet
    assert any(h["action"] == "insert_linked" for h in d["history"])

    # same insert again updates the quantity; wrong links are refused
    assert client.post(f"/api/parts/{body}/inserts", json={"insert_part_code": "NUT-M4", "quantity": 3}).json()["inserts"][0]["quantity"] == 3
    nut = link["insert"]["id"]
    assert "both ways" in client.post(f"/api/parts/{nut}/inserts", json={"insert_part_code": "BODY-1", "quantity": 1}).json()["detail"]
    assert "own insert" in client.post(f"/api/parts/{body}/inserts", json={"insert_part_code": "BODY-1", "quantity": 1}).json()["detail"]
    assert "not in the register" in client.post(f"/api/parts/{body}/inserts", json={"insert_part_code": "NOPE", "quantity": 1}).json()["detail"]
    bad = client.post(f"/api/parts/{body}/inserts", json={"name": "x", "quantity": 1, "dimensions": {"shape": "bush", "od_mm": 6}})
    assert bad.status_code == 400 and "id_mm" in bad.json()["detail"]
    assert "2026-10" in client.post(f"/api/parts/{body}/inserts", json={"name": "y", "quantity": 1, "price_month": "Oct"}).json()["detail"]

    d = client.delete(f"/api/parts/{body}/inserts/{link['link_id']}").json()
    assert d["inserts"] == [] and d["history"][0]["action"] == "insert_removed"


def test_insert_without_size_is_reported_in_the_trace():
    project()
    p = client.get("/api/projects").json()[0]
    up = client.post(f"/api/parts/{_part('X')}/revisions", files={"file": ("c.step", (SAMPLES / "plastic_cover.step").read_bytes())})
    body = up.json()["id"]
    client.put(f"/api/parts/{body}", json={"material": "ABS"})
    client.post(f"/api/parts/{body}/inserts", json={"name": "Clip", "material": "Brass CuZn39Pb3", "quantity": 1})
    t = client.get(f"/api/parts/{body}/trace").json()
    assert "total_weight" not in t["nodes"] and "no 3D file and no dimensions" in t["warnings"][0]
    assert p["code"] == "INS"


# ---------- Excel ----------

def test_excel_insert_rows_link_and_suggest():
    p = project()
    wb = load_workbook(io.BytesIO(client.get("/api/import/template.xlsx").content))
    ws = wb["Parts"]
    cols = [c.value for c in ws[1]]
    ws.delete_rows(2, ws.max_row)
    row = lambda **kw: ws.append([kw.get(c, "") for c in cols])
    row(part_code="HSG-9", part_name="Housing", material="PA66 GF30", variant="STD")
    row(part_code="INS-9", part_name="Insert M3", material="Brass CuZn39Pb3", insert_of_part_code="HSG-9",
        insert_quantity=4, unit_price=2.5, price_currency="INR", price_source="quote", price_month="2026-10")
    buf = io.BytesIO()
    wb.save(buf)
    r = client.post("/api/import/excel", data={"project_id": p["id"]}, files=[("excel", ("p.xlsx", buf.getvalue()))])
    assert r.status_code == 200, r.text
    rep = r.json()
    rows = {x["part_code"]: x for x in rep["rows"]}
    assert rows["HSG-9"]["process"] == "insert_moulding" and "PA66 GF30" in rows["HSG-9"]["process_reason"]
    assert rows["INS-9"]["process"] == "bought_out" and rep["inserts_linked"] == 1
    hsg = client.get("/api/parts", params={"q": "HSG-9"}).json()["rows"][0]
    d = client.get(f"/api/parts/{hsg['id']}").json()
    assert d["inserts"][0]["quantity"] == 4 and d["inserts"][0]["insert"]["unit_price"] == 2.5

    # an insert row pointing at an unknown part: nothing is saved
    ws.delete_rows(2, ws.max_row)
    row(part_code="INS-X", part_name="Insert", insert_of_part_code="NOPE")
    buf = io.BytesIO()
    wb.save(buf)
    r = client.post("/api/import/excel", data={"project_id": p["id"]}, files=[("excel", ("p.xlsx", buf.getvalue()))])
    assert r.status_code == 400 and "insert_of_part_code 'NOPE'" in r.json()["detail"]


# ---------- master data checks and migration ----------

@pytest.mark.parametrize("change, message", [
    (lambda s: s["insert_shapes"][0].update(dimensions=""), "comma list"),
    (lambda s: s["insert_shapes"][0].update(volume_rule="d_mm * x_mm"), "unknown name"),
    (lambda s: s["insert_materials"][0].update(density_g_cm3=0), "more than 0"),
    (lambda s: next(r for r in s["register_settings"] if r["key"] == "insert_default_sourcing").update(value="nope"), "not a Sourcing"),
    (lambda s: next(r for r in s["register_settings"] if r["key"] == "insert_detect_min_share").update(value="x"), "number"),
])
def test_insert_setup_is_checked(change, message):
    data = register_setup.tables()
    change(data)
    r = client.put("/api/register-setup", json=data)
    assert r.status_code == 400 and message in r.json()["detail"]


def test_old_database_gets_the_new_columns(tmp_path, monkeypatch):
    path = tmp_path / "old.sqlite"
    with sqlite3.connect(path) as con:
        con.execute("CREATE TABLE projects (id INTEGER PRIMARY KEY, code TEXT NOT NULL UNIQUE, name TEXT NOT NULL)")
        con.execute("CREATE TABLE parts (id INTEGER PRIMARY KEY, part_code TEXT NOT NULL UNIQUE, name TEXT NOT NULL, project_id INTEGER)")
        con.execute("INSERT INTO projects (code, name) VALUES ('OLD', 'Old project')")
        con.execute("INSERT INTO parts (part_code, name, project_id) VALUES ('P1', 'Kept', 1)")
    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{path}")
    with db.engine().connect() as conn:
        row = conn.execute(select(db.parts.c.name, db.parts.c.unit_price, db.parts.c.dimensions)).one()
    assert row.name == "Kept" and row.unit_price is None
