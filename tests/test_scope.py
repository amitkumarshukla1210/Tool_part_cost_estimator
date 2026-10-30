"""Stage 30: head of department, data access by tool group (built on the logins and roles of stages 17 and 22)."""
from backend.app import auth
from tests.test_server_mode import PW, login, server  # noqa: F401  (fixture)

SAMPLES = __import__("pathlib").Path(__file__).resolve().parents[1] / "samples"


def _part(c, project_id, file, material, process, code):
    with open(SAMPLES / file, "rb") as fh:
        m = c.post("/api/import/step", files={"file": (file, fh)}).json()
    return c.post(f"/api/import/{m['token']}/commit", json={"project_id": project_id, "items": {m["root"]: {"material": material, "process": process, "part_code": code}}})


def _user(c, username, role, groups):
    assert login(c).status_code == 200
    r = c.post("/api/users", json={"username": username, "password": PW, "role": role, "name": username.title(), "groups": groups})
    assert r.status_code == 200, r.text
    c.post("/api/auth/logout")
    return r.json()


def test_head_role_sits_between_approver_and_admin():
    assert auth.ROLES == ("viewer", "engineer", "approver", "head", "admin")
    head = {"role": "head"}
    assert auth.has_role("approver", head) and not auth.has_role("admin", head)


def test_users_see_only_their_tool_groups_and_the_admin_can_move_them(server):
    login(server)
    pl = server.post("/api/projects", json={"code": "PL", "name": "Plastic project"}).json()
    sm = server.post("/api/projects", json={"code": "SM", "name": "Sheet project"}).json()
    assert _part(server, pl["id"], "plastic_cover.step", "ABS", "injection_moulding", "PL-1").status_code == 200
    assert _part(server, sm["id"], "sheet_metal_bracket.step", "Mild steel CRCA / HR", "sheet_metal_multistage", "SM-1").status_code == 200
    server.post("/api/auth/logout")
    u1 = _user(server, "user1", "engineer", ["sheet_metal", "eps"])
    _user(server, "boss", "head", [])
    assert u1["groups"] == ["sheet_metal", "eps"]
    assert server.post("/api/users", json={"username": "x", "password": PW, "groups": ["wood"]}).status_code in (400, 401, 403)

    login(server, "user1")
    codes = {p["code"] for p in server.get("/api/projects").json()}
    assert codes == {"SM"}
    parts = server.get("/api/parts").json()["rows"]
    assert [p["part_code"] for p in parts] == ["SM-1"]
    pl_part = server.get("/api/parts", params={"q": "PL-1"}).json()["rows"]
    assert pl_part == []
    # opening a part or project of another group is refused
    login(server)
    pid_pl = server.get("/api/parts", params={"q": "PL-1"}).json()["rows"][0]["id"]
    server.post("/api/auth/logout")
    login(server, "user1")
    r = server.get(f"/api/parts/{pid_pl}")
    assert r.status_code == 403 and "tool group" in r.json()["detail"]
    assert server.get(f"/api/projects/{pl['id']}").status_code == 403
    # a new part only in my groups; my own new project stays visible
    mine = server.post("/api/projects", json={"code": "U1", "name": "Mine"}).json()
    assert "U1" in {p["code"] for p in server.get("/api/projects").json()}
    r = _part(server, mine["id"], "plastic_cover.step", "ABS", "injection_moulding", "U1-PL")
    assert r.status_code == 403 and "Plastic moulding" in r.json()["detail"]
    server.post("/api/auth/logout")

    login(server, "boss")                                   # head of department sees everything
    assert {"PL", "SM"} <= {p["code"] for p in server.get("/api/projects").json()}
    assert server.get(f"/api/parts/{pid_pl}").status_code == 200
    assert server.get("/api/users").status_code == 403      # but does not manage users
    server.post("/api/auth/logout")

    login(server)                                           # six months later: user1 moves to plastic
    r = server.put(f"/api/users/{u1['id']}", json={"groups": ["plastic", "die_casting"]})
    assert r.status_code == 200 and r.json()["groups"] == ["plastic", "die_casting"]
    server.post("/api/auth/logout")
    login(server, "user1")
    assert server.get(f"/api/parts/{pid_pl}").status_code == 200
    assert {p["part_code"] for p in server.get("/api/parts").json()["rows"]} == {"PL-1"}


def test_user_without_groups_sees_everything(server):
    login(server)
    p = server.post("/api/projects", json={"code": "A1", "name": "a"}).json()
    _part(server, p["id"], "plastic_cover.step", "ABS", "injection_moulding", "A1-1")
    server.post("/api/auth/logout")
    _user(server, "old", "engineer", [])
    login(server, "old")
    assert {"A1"} <= {x["code"] for x in server.get("/api/projects").json()}
    assert server.get("/api/parts").json()["total"] == 1


def test_files_and_variants_of_other_groups_are_refused(server):
    login(server)
    pl = server.post("/api/projects", json={"code": "FP", "name": "Plastic files"}).json()
    assert _part(server, pl["id"], "plastic_cover.step", "ABS", "injection_moulding", "FP-1").status_code == 200
    part = server.get("/api/parts", params={"q": "FP-1"}).json()["rows"][0]
    detail = server.get(f"/api/parts/{part['id']}").json()
    file_id = detail["revisions"][0]["file_id"]
    vid = next(v["id"] for v in server.post(f"/api/projects/{pl['id']}/variants", json={"code": "V1", "name": "v"}).json()["variants"] if v["code"] == "V1")
    server.post("/api/auth/logout")
    _user(server, "sm1", "engineer", ["sheet_metal"])
    login(server, "sm1")
    r = server.get(f"/api/files/{file_id}")
    assert r.status_code == 403 and "tool group" in r.json()["detail"]
    assert server.get(f"/api/variants/{vid}/bom").status_code == 403
    server.post("/api/auth/logout")
    login(server)
    assert server.get(f"/api/files/{file_id}").status_code == 200


def _part_in_variant(c, project_id, variant_id, file, material, process, code):
    with open(SAMPLES / file, "rb") as fh:
        m = c.post("/api/import/step", files={"file": (file, fh)}).json()
    r = c.post(f"/api/import/{m['token']}/commit", json={"project_id": project_id, "variant_id": variant_id,
                                                          "items": {m["root"]: {"material": material, "process": process, "part_code": code}}})
    assert r.status_code == 200, r.text
    pid = r.json()["part_ids"][0]
    for u in ("rm-cost", "tool-cost", "piece-cost"):
        c.post(f"/api/parts/{pid}/{u}", json={})
    return pid


def test_project_dashboard_and_leader_view_count_only_my_tool_groups(server):
    """Stage 39: a limited user's project totals (dashboard, leader view, exports) cover their tool groups only."""
    login(server)
    p = server.post("/api/projects", json={"code": "MIX", "name": "Mixed project"}).json()
    vid = next(v["id"] for v in server.post(f"/api/projects/{p['id']}/variants", json={"code": "V1", "name": "Base", "annual_volume": 10000,
                                                                                       "target_bom_cost_inr": 100}).json()["variants"])
    pl = _part_in_variant(server, p["id"], vid, "plastic_cover.step", "ABS", "injection_moulding", "MIX-PL")
    sm = _part_in_variant(server, p["id"], vid, "example_bracket_revA.step", "Galvanised steel GI", "sheet_metal_multistage", "MIX-SM")
    full = server.get(f"/api/projects/{p['id']}/dashboard").json()
    lead_full = server.get(f"/api/projects/{p['id']}/leader").json()
    assert full["parts_total"] == 2 and not full["scope"]["limited"]
    server.post("/api/auth/logout")
    _user(server, "smu", "engineer", ["sheet_metal"])
    login(server, "smu")
    d = server.get(f"/api/projects/{p['id']}/dashboard").json()
    sm_capex = server.get(f"/api/parts/{sm}").json()["tool_capex_inr"]
    assert d["parts_total"] == 1 and d["scope"]["limited"] and d["scope"]["hidden_parts"] == 1 and "not counted" in d["scope"]["note"]
    assert abs(d["capex"]["tool_cost_inr"] - sm_capex) < 0.01 < full["capex"]["tool_cost_inr"] - sm_capex
    lead = server.get(f"/api/projects/{p['id']}/leader").json()
    lines = lead["variants"][0]["lines"]
    assert [l["part_code"] for l in lines] == ["MIX-SM"] and len(lead_full["variants"][0]["lines"]) == 2
    assert lead["capex"]["tool_inr"] < lead_full["capex"]["tool_inr"]
    assert all(m["machine_id"].startswith("PRESS") for m in lead["machines"]["rows"])           # no moulding machine of the plastic part
    from openpyxl import load_workbook
    import io
    wb = load_workbook(io.BytesIO(server.get(f"/api/projects/{p['id']}/leader.xlsx").content))
    assert any("not counted" in str(c.value) for row in wb.worksheets[0].iter_rows() for c in row)
    assert server.get(f"/api/projects/{p['id']}/dashboard.pdf").status_code == 200
    server.post("/api/auth/logout")
    _user(server, "hod", "head", [])
    login(server, "hod")
    assert server.get(f"/api/projects/{p['id']}/dashboard").json()["parts_total"] == 2


def test_machined_parts_are_the_tool_group_machining(server):
    """Machined parts (stage 37) are a tool group of their own: only users with Machining ticked see them."""
    login(server)
    p = server.post("/api/projects", json={"code": "MCH", "name": "Machined"}).json()
    assert _part(server, p["id"], "plate_with_hole.step", "Mild steel CRCA / HR", "machined", "MCH-1").status_code == 200
    server.post("/api/auth/logout")
    _user(server, "mach", "engineer", ["machined"])
    _user(server, "sheet2", "engineer", ["sheet_metal"])
    login(server, "mach")
    assert [r["part_code"] for r in server.get("/api/parts").json()["rows"]] == ["MCH-1"]
    server.post("/api/auth/logout")
    login(server, "sheet2")
    assert server.get("/api/parts").json()["rows"] == []
