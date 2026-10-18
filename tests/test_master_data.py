import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.app import master_data
from backend.app.capacity import CapacityInput, calculate
from backend.app.main import app

client = TestClient(app)


def set_param(data, part_type, name, value):
    for row in data["Parameters"]:
        if row["part_type"] == part_type and row["parameter"] == name:
            row["value"] = value
            return
    raise KeyError(name)


def test_every_part_type_has_its_defaults():
    for pt in master_data.PART_TYPES:
        assert client.get("/api/gross-weight/defaults").json()[pt]["currency"] == "INR"
        assert client.get("/api/capacity/defaults").json()["by_part_type"][pt]["oee_pct"] > 0
    tables = client.get("/api/master-data").json()
    assert {m["part_type"] for m in tables["Materials"]} == set(master_data.PART_TYPES)


def test_saved_change_is_used_by_calculations(master_data_copy):
    data = client.get("/api/master-data").json()
    set_param(data, "die_casting", "max_cavities", 2)
    set_param(data, "die_casting", "standard_cavities", "1; 2 ,4")
    assert client.put("/api/master-data", json=data).status_code == 200

    r = calculate(CapacityInput(part_type="die_casting", footprint_length_mm=100, footprint_width_mm=60,
                                annual_volume=1_500_000))
    assert r["cavities_per_tool"] == 2
    # Written back to Excel with spaces so Excel keeps it as text.
    ws = load_workbook(master_data_copy)["Parameters"]
    values = [row[2] for row in ws.iter_rows(values_only=True) if row[0] == "die_casting" and row[1] == "standard_cavities"]
    assert values == ["1, 2, 4"]


def test_bad_value_is_rejected_and_file_unchanged(master_data_copy):
    before = master_data_copy.read_bytes()
    data = client.get("/api/master-data").json()
    set_param(data, "plastic", "oee_pct", "abc")
    r = client.put("/api/master-data", json=data)
    assert r.status_code == 400
    assert "oee_pct" in r.json()["detail"] and "not a number" in r.json()["detail"]
    assert master_data_copy.read_bytes() == before

    data = client.get("/api/master-data").json()
    data["Materials"][0]["density_g_cm3"] = 0
    assert "density" in client.put("/api/master-data", json=data).json()["detail"]

    data = client.get("/api/master-data").json()
    data["Materials"][0]["part_type"] = "forging"
    assert "part_type" in client.put("/api/master-data", json=data).json()["detail"]


def test_missing_parameter_gives_clear_message():
    data = client.get("/api/master-data").json()
    data["Parameters"] = [r for r in data["Parameters"] if not (r["part_type"] == "foundry" and r["parameter"] == "yield_pct")]
    client.put("/api/master-data", json=data)
    r = client.post("/api/gross-weight", json={"part_type": "foundry", "net_weight_g": 650, "density_g_cm3": 7.2})
    assert r.status_code == 422
    assert "yield_pct" in r.text and "Parameters" in r.text


def test_download_and_upload_round_trip():
    xlsx = client.get("/api/master-data/download")
    assert xlsx.status_code == 200
    wb = load_workbook(io.BytesIO(xlsx.content))
    ws = wb["Parameters"]
    for row in ws.iter_rows(min_row=2):
        if row[0].value == "plastic" and row[1].value == "runner_pct":
            row[2].value = 20
    buf = io.BytesIO()
    wb.save(buf)

    up = client.post("/api/master-data/upload", files={"file": ("master_data.xlsx", buf.getvalue())})
    assert up.status_code == 200
    assert client.get("/api/gross-weight/defaults").json()["plastic"]["runner_pct"] == 20


def test_upload_rejects_bad_files():
    assert client.post("/api/master-data/upload", files={"file": ("m.csv", b"a,b")}).status_code == 400
    r = client.post("/api/master-data/upload", files={"file": ("m.xlsx", b"not excel")})
    assert r.status_code == 400 and "Excel" in r.json()["detail"]
