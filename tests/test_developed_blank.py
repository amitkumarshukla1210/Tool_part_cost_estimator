import cadquery as cq
import pytest
from fastapi.testclient import TestClient

from backend.app.cad_import import measure
from backend.app.geometry import developed_blank
from backend.app.main import app
from tests.test_rm_cost import project

client = TestClient(app)


def l_bracket(t=2.0, strip=20.0, long=300.0, short=100.0, holes=True):
    a = cq.Workplane("XY").box(long, strip, t).translate((long / 2, strip / 2, t / 2))
    b = cq.Workplane("XY").box(t, strip, short).translate((t / 2, strip / 2, short / 2))
    s = a.union(b)
    if holes:
        s = s.cut(cq.Workplane("XY").pushPoints([(100, strip / 2), (200, strip / 2)]).circle(3).extrude(10).translate((0, 0, -5)))
    return s.val()


def blank_of(shape, t=2.0):
    m = measure(shape)
    return developed_blank(shape, m["volume_mm3"], m["bbox_mm"], t)


def test_l_bracket_is_a_developed_straight_strip_with_a_bending_stage():
    r = blank_of(l_bracket())
    assert r["chosen"].startswith("Developed straight strip") and r["width"] == pytest.approx(20, abs=0.1)
    assert r["length"] == pytest.approx(398, abs=1)                         # 300 + 100 - 2 (sharp corner, mid-line)
    yields = {o["name"]: o["yield_pct"] for o in r["options"]}
    assert yields["Rectangle keeping the formed width"] < 30 and r["options"][-1]["yield_pct"] > 99


def test_u_channel_tray_and_flat_part():
    u = cq.Workplane("XZ").polyline([(0, 0), (100, 0), (100, 40), (98, 40), (98, 2), (2, 2), (2, 40), (0, 40)]).close().extrude(-200).val()
    r = blank_of(u)
    assert (r["length"], r["width"]) == (pytest.approx(176, abs=1), pytest.approx(200, abs=0.1))
    tray = blank_of(cq.Workplane("XY").box(100, 80, 30).faces(">Z").shell(-2).val())
    assert tray["chosen"].startswith("Tray") and (tray["length"], tray["width"]) == (156, 136)
    flat = blank_of(cq.Workplane("XY").box(100, 50, 2).val())
    assert flat["chosen"] == "Flat part" and (flat["length"], flat["width"]) == (100, 50)


def test_register_gross_weight_uses_the_developed_blank(tmp_path):
    f = tmp_path / "l_bracket.step"
    cq.exporters.export(cq.Workplane().add(l_bracket()), str(f))
    p = project("LB")
    with open(f, "rb") as fh:
        m = client.post("/api/import/step", files={"file": ("l_bracket.step", fh)}).json()
    r = client.post(f"/api/import/{m['token']}/commit", json={"project_id": p["id"], "items": {
        m["root"]: {"material": "Mild steel CRCA / HR", "process": "sheet_metal_multistage", "part_code": "LB-1"}}})
    pid = r.json()["part_ids"][0]
    rm = client.get(f"/api/parts/{pid}/rm-cost/trace").json()
    assert rm["nodes"]["in:blank_width_mm"]["value"] == pytest.approx(20, abs=0.1)
    assert "Developed straight strip" in rm["nodes"]["in:blank_width_mm"]["source"]["note"]
    gross, net = rm["nodes"]["calc:gross_g"]["value"], rm["nodes"]["in:net_weight_g"]["value"] if "in:net_weight_g" in rm["nodes"] else None
    assert gross < 2 * 15920 / 1000 * 7.85                                   # not the 300 × 100 rectangle any more
    sheet = client.get(f"/api/parts/{pid}/sheet-tools").json()
    assert any(o["code"] == "bend" for o in sheet["operations"]) and sheet["blank"]["chosen"].startswith("Developed")


def test_l_bracket_with_a_small_tab_is_still_a_strip_not_a_tray():
    t = 1.0
    a = cq.Workplane("XY").box(570, 40, t).translate((285, 20, t / 2))
    b = cq.Workplane("XY").box(t, 40, 120).translate((t / 2, 20, 60))
    tab = cq.Workplane("XY").box(10, t, 8).translate((300, 40 + t / 2, 4 + t))
    r = blank_of(a.union(b).union(tab).val(), t)
    assert r["chosen"].startswith("Developed straight strip") and r["width"] == pytest.approx(41, abs=0.5)


def test_stage_wise_tool_exports_and_country_selection():
    from tests.test_rm_cost import add_part
    pid = add_part(project("SW"), "sheet_metal_bracket.step", "Mild steel CRCA / HR", "sheet_metal_multistage", "SW-1")
    tr = client.get(f"/api/parts/{pid}/tool-cost/trace").json()
    key = tr["dies"][0]["key"]
    pdf = client.get(f"/api/parts/{pid}/tool-cost/export.pdf")
    one = client.get(f"/api/parts/{pid}/tool-cost/export.pdf", params={"die": key})
    assert pdf.status_code == one.status_code == 200 and pdf.content[:4] == b"%PDF" and len(one.content) < len(pdf.content)
    from io import BytesIO
    from openpyxl import load_workbook
    wb = load_workbook(BytesIO(client.get(f"/api/parts/{pid}/tool-cost/export.xlsx").content))
    assert "Dies" in wb.sheetnames and any(n.startswith("Die 1") for n in wb.sheetnames) and any(n.startswith("Die 2") for n in wb.sheetnames)
    assert client.get(f"/api/parts/{pid}/tool-cost/export.pdf", params={"die": "nope"}).status_code == 404
    res = client.get(f"/api/parts/{pid}/country-comparison", params={"countries": "IN,CN"}).json()
    assert [r["country"] for r in res["rows"]] == ["IN", "CN"]
