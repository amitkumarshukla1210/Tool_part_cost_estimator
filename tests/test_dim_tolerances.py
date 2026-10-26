"""Stage 32: a tolerance on every dimension of the tool detail drawings, in mm or µm."""

import io
import re
import zipfile

import ezdxf
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.app import dim_tolerances as dt
from backend.app import mould_db
from backend.app.main import app
from tests.test_drawing_data import bracket, die_plate, pdf_text

client = TestClient(app)


def test_fit_limits_and_texts():
    r = dt.resolve
    assert dt.fmt(r("H7", 8, "m")) == "H7 +0.015/0" and dt.fmt(r("H7", 12, "m")) == "H7 +0.018/0"
    assert dt.fmt(r("g6", 20, "m"), "µm") == "g6 -7/-20 µm" and dt.fmt(r("m6", 12, "m")) == "m6 +0.018/+0.007"
    assert dt.fmt(r("general", 250, "m")) == "±0.5" and dt.fmt(r("class", 10, "m", "Ø10 m6 / H7 dowel")) == "H7 +0.015/0"
    assert dt.fmt(r("+0.01/0", 30, "m"), "µm") == "+10/0 µm" and dt.check_spec("0.05") == "±0.05"
    rows = mould_db.tables()["fit_tolerances"]
    assert len(rows) >= 117 and all(r["is_example"] for r in rows)
    assert {r["parameter"] for r in mould_db.tables()["design_rules"]} >= {"thread_class", "hole_clear_tol"}
    bad = mould_db.tables()
    bad["fit_tolerances"].append({"fit": "Z9", "from_mm": 0, "to_mm": 3, "upper_um": "-5", "lower_um": "+5", "is_example": True, "notes": ""})
    try:
        mould_db.write(bad)
        raise AssertionError("upper below lower accepted")
    except mould_db.MouldDbError as exc:
        assert "upper must not be below lower" in str(exc)


def test_every_dimension_has_a_tolerance_from_the_design_rules_and_can_be_changed():
    pid = bracket("DT1")
    _, tool, dp = die_plate(pid)
    dims = {x["key"]: x for x in dp["dimensions"]}
    assert dims["H"]["tolerance"] == "±0.02"                                 # plate_thk_tol of the Design rules
    assert dims["L"]["tolerance"].startswith("±") and dims["W"]["tolerance"].startswith("±")
    fitted = next(k for k, x in dims.items() if re.fullmatch(r"H\d+", k) and x["tolerance"].startswith("H7"))
    assert dims[fitted + "@pos"]["tolerance"] == "±0.01"                     # position_tol
    tap = next(k for k, x in dims.items() if re.fullmatch(r"H\d+", k) and x["tolerance"] == "6H")
    assert tap and any(k.startswith("O") for k in dims)                      # thread class and the die opening
    r = client.put(f"/api/parts/{pid}/drawing-data", json={"items": [{"tool": tool["key"], "item": dp["item"],
        "fields": {"tolerance_unit": "µm", "features": {"L": "±0.05", fitted: "H6"}}}]})
    assert r.status_code == 200, r.text
    assert "tolerance L" in r.json()["changed"][0]["summary"]
    _, _, dp2 = die_plate(pid)
    d2 = {x["key"]: x for x in dp2["dimensions"]}
    assert d2["L"]["tolerance"] == "±50 µm" and d2["L"]["changed"] and d2[fitted]["tolerance"].startswith("H6 +") and d2["H"]["tolerance"] == "±20 µm"
    text = pdf_text(client.get(f"/api/parts/{pid}/tool-design.pdf", params={"die": tool["key"]}).content)
    assert "±50 µm" in text and "H6 +" in text and "Tolerances in µm" in text and "Ø tol." in text
    z = zipfile.ZipFile(io.BytesIO(client.get(f"/api/parts/{pid}/tool-design.zip", params={"die": tool["key"]}).content))
    dxf = next(n for n in z.namelist() if n.endswith("_Die_plate.dxf"))
    doc = ezdxf.read(io.StringIO(z.read(dxf).decode()))
    texts = " ".join(t.dxf.text for t in doc.modelspace().query("TEXT"))
    assert "H6 +" in texts and "6H" in texts
    for bad in ({"features": {fitted: "X99"}}, {"features": {"Q1": "H7"}}, {"tolerance_unit": "inch"}):
        assert client.put(f"/api/parts/{pid}/drawing-data", json={"items": [{"tool": tool["key"], "item": dp["item"], "fields": bad}]}).status_code == 400


def test_dimension_sheet_round_trip():
    pid = bracket("DT2")
    _, tool, dp = die_plate(pid)
    wb = load_workbook(io.BytesIO(client.get(f"/api/parts/{pid}/drawing-data.xlsx").content))
    ws = wb["Dimension tolerances"]
    top = next(r for r in range(1, 10) if ws.cell(r, 1).value == "Tool" and ws.cell(r, 3).value == "Dimension")
    row = next(r for r in range(top + 1, ws.max_row + 1) if ws.cell(r, 1).value == tool["key"] and str(ws.cell(r, 2).value) == dp["item"]
               and ws.cell(r, 3).value == "W")
    ws.cell(row, 6).value = "+0.1/-0.05"
    for w in wb.worksheets:
        w._images = []
    buf = io.BytesIO(); wb.save(buf)
    up = client.post(f"/api/parts/{pid}/drawing-data/upload", files={"file": ("d.xlsx", buf.getvalue())})
    assert up.status_code == 200, up.text
    _, _, dp2 = die_plate(pid)
    assert {x["key"]: x for x in dp2["dimensions"]}["W"]["tolerance"] == "+0.1/-0.05"
