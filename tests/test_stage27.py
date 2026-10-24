"""Stage 27: the machine the tool runs on (typed parameters, fit, recommendation), process parameters, design skills."""
import io
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from backend.app import machine_spec
from backend.app.main import app
from tests.test_rm_cost import add_part, project

client = TestClient(app)
SAMPLES = Path(__file__).resolve().parents[1] / "samples"


def mould(code="MC-1", file="plastic_cover.step", material="ABS", process="injection_moulding"):
    return add_part(project(code + "P"), file, material, process, code)


def design(pid):
    r = client.post(f"/api/parts/{pid}/tool-designs", json={})
    assert r.status_code == 200, r.text
    return r.json()


def test_machine_form_fields_and_catalogue_per_family():
    pid = mould()
    m = client.get(f"/api/parts/{pid}/machine").json()
    assert m["family"] == "moulding" and m["spec"] is None
    assert {"tonnage_t", "tie_bar_h_mm", "min_mould_h_mm", "opening_stroke_mm"} <= {f["key"] for f in m["fields"]}
    assert m["machines"] and all(x["tonnage_t"] for x in m["machines"])
    sheet = add_part(project("SMP"), "sheet_metal_bracket.step", "Mild steel CRCA / HR", "sheet_metal_multistage", "SM-M")
    assert client.get(f"/api/parts/{sheet}/machine").json()["family"] == "press"
    plain = add_part(project("MCP"), "plate_with_hole.step", "Mild steel CRCA / HR", "machined", "MCH-1")
    assert client.get(f"/api/parts/{plain}/machine").status_code == 400                  # a machined part has no tool machine


def test_typed_machine_is_saved_checked_and_cleaned():
    pid = mould("MC-2")
    r = client.put(f"/api/parts/{pid}/machine", json={"spec": {"tonnage_t": "350", "tie_bar_h_mm": 520, "tie_bar_v_mm": 520, "min_mould_h_mm": 250, "max_mould_h_mm": 600, "junk": 5}})
    assert r.status_code == 200 and r.json()["spec"] == {"tonnage_t": 350.0, "tie_bar_h_mm": 520.0, "tie_bar_v_mm": 520.0, "min_mould_h_mm": 250.0, "max_mould_h_mm": 600.0}
    assert client.put(f"/api/parts/{pid}/machine", json={"spec": {"tonnage_t": "abc"}}).status_code == 400
    assert client.put(f"/api/parts/{pid}/machine", json={"spec": {"min_mould_h_mm": 700, "max_mould_h_mm": 600}}).status_code == 400
    assert client.put(f"/api/parts/{pid}/machine", json={"spec": {"machine_id": "NO-SUCH"}}).status_code == 400
    assert client.put(f"/api/parts/{pid}/machine", json={"spec": None}).json()["spec"] is None


def test_without_a_typed_machine_the_software_recommends_one():
    d = design(mould("MC-3"))
    m = d["summary"]["machine"]
    assert m["source"] == "recommended" and m["fit"] == [] and m["fit_ok"] == -1
    rec = m["recommended"]
    assert rec["rec_tonnage_t"] >= d["summary"]["facts"]["clamp_t"] and rec["rec_tie_bar_mm"] > 0 and rec["rec_platen_l_mm"] >= rec["rec_platen_w_mm"]
    assert rec["rec_min_mould_h_mm"] <= rec["rec_max_mould_h_mm"] and rec["rec_opening_stroke_mm"] > 0 and m["nearest_machine"]["tonnage_t"] >= rec["rec_tonnage_t"]
    assert "No machine was typed" in m["note"]


def test_typed_machine_shapes_the_mould_and_the_fit_is_checked():
    pid = mould("MC-4")
    base = design(pid)["summary"]["machine"]["recommended"]
    # a machine with a smaller shut height than the mould needs is refused by the fit; a machine with a higher minimum gets a riser plate
    need = base["rec_min_mould_h_mm"]
    client.put(f"/api/parts/{pid}/machine", json={"spec": {"tonnage_t": 800, "tie_bar_h_mm": 800, "tie_bar_v_mm": 800, "platen_l_mm": 1100, "platen_w_mm": 1100,
                                                          "min_mould_h_mm": need + 60, "max_mould_h_mm": need + 400, "opening_stroke_mm": 900, "locating_ring_dia_mm": 125}})
    d = design(pid)
    m = d["summary"]["machine"]
    assert m["source"] == "typed" and m["fit_ok"] == 1, m["fit"]
    names = [c["name"] for c in d["components"]]
    assert "Height riser plate" in names and d["summary"]["facts"]["machine_fit"] == 1
    ring = next(c for c in d["components"] if c["name"] == "Locating ring")
    assert abs(ring["info"]["bbox_mm"][0] - 125) < 0.5                                     # the locating ring has the machine's diameter
    # a machine that is too small
    client.put(f"/api/parts/{pid}/machine", json={"spec": {"tonnage_t": 50, "tie_bar_h_mm": 200, "tie_bar_v_mm": 200}})
    r = client.post(f"/api/parts/{pid}/tool-designs", json={})
    if r.status_code == 200:                                                                # the base cannot pass 200 mm tie bars: the design says so
        m = r.json()["summary"]["machine"]
        assert m["fit_ok"] == 0 and any(not x["ok"] for x in m["fit"])
    else:
        assert "tie bars" in r.text or "typed machine" in r.text


def test_sheet_metal_press_recommended_typed_and_riser():
    pid = add_part(project("SMX"), "sheet_metal_bracket.step", "Mild steel CRCA / HR", "sheet_metal_multistage", "SM-X")
    d = design(pid)
    rec = d["summary"]["machine"]["recommended"]
    assert d["summary"]["machine"]["source"] == "recommended" and rec["rec_tonnage_t"] >= 10 and rec["rec_bolster_l_mm"] >= rec["rec_bolster_w_mm"] > 0
    assert rec["rec_shut_height_min_mm"] < rec["rec_shut_height_max_mm"]
    # your press: big enough, but its smallest shut height is above the die: parallels are added
    client.put(f"/api/parts/{pid}/machine", json={"spec": {"tonnage_t": 400, "bolster_l_mm": 1500, "bolster_w_mm": 900, "shut_height_min_mm": rec["rec_shut_height_max_mm"] + 100,
                                                          "shut_height_max_mm": rec["rec_shut_height_max_mm"] + 300, "stroke_mm": 400}})
    d2 = design(pid)
    m = d2["summary"]["machine"]
    assert m["source"] == "typed" and m["fit_ok"] == 1, m["fit"]
    assert any("riser" in c["name"].lower() for c in d2["components"]) and all(s == "Your press 400 t" for s in d2["summary"]["presses"])
    # a press that is too small for the stage force
    client.put(f"/api/parts/{pid}/machine", json={"spec": {"tonnage_t": 5}})
    d3 = design(pid)
    assert d3["summary"]["machine"]["fit_ok"] == 0 and d3["summary"]["facts"]["machine_fit"] == 0
    item = next(i for i in d3["review"] if "press you typed" in i["text"])
    assert item["status"] == "issue" and item["blocking"]                                  # a press that does not suit the die blocks the release


def test_foundry_flask_recommended_and_checked():
    pid = add_part(project("FDX"), "plastic_cover.step", "Grey cast iron FG260", "foundry", "FD-X")
    d = design(pid)
    rec = d["summary"]["machine"]["recommended"]
    assert rec["rec_flask_l_mm"] >= rec["rec_flask_w_mm"] > 0 and rec["rec_cope_h_mm"] > 0
    client.put(f"/api/parts/{pid}/machine", json={"spec": {"flask_l_mm": 300, "flask_w_mm": 200}})
    d2 = design(pid)
    assert d2["summary"]["machine"]["fit_ok"] == 0


# ---------- process parameters ----------

def test_process_parameters_follow_the_simulation_and_the_material():
    pid = mould("PP-1")
    r = client.get(f"/api/parts/{pid}/process-parameters")
    assert r.status_code == 200, r.text
    res = r.json()
    p = {x["name"]: x for x in res["parameters"]}
    assert {"melt_temp_c", "mould_temp_c", "fill_time_s", "hold_pressure_bar", "hold_time_s", "cooling_time_s", "clamp_force_t", "cycle_time_s"} <= set(p)
    abs_win = res["window"]
    assert abs_win["melt_min_c"] <= p["melt_temp_c"]["value"] <= abs_win["melt_max_c"] and p["melt_temp_c"]["min"] == abs_win["melt_min_c"]
    assert p["hold_pressure_bar"]["value"] < p["injection_pressure_bar"]["value"] and p["hold_time_s"]["value"] < p["cooling_time_s"]["value"]
    # the material changes the settings: PC is processed much hotter than ABS
    pc = client.get(f"/api/parts/{mould('PC-1', material='PC (polycarbonate)')}/process-parameters").json()
    assert {x["name"]: x for x in pc["parameters"]}["melt_temp_c"]["value"] > p["melt_temp_c"]["value"] + 30
    assert res["is_example"] and "simulation" in res["source"]


def test_process_parameters_for_every_family_and_downloads():
    cases = [("die", "plastic_cover.step", "Aluminium ADC12 / A383", "die_casting", {"metal_temp_c", "die_temp_c", "fast_shot_speed_m_s", "intensification_pressure_bar"}),
             ("sheet", "sheet_metal_bracket.step", "Mild steel CRCA / HR", "sheet_metal_multistage", {"press_tonnage_t", "strokes_per_min", "die_clearance_mm", "springback_comp_deg"}),
             ("fdry", "plastic_cover.step", "Grey cast iron FG260", "foundry", {"pour_temp_c", "pour_time_s", "solidification_min"}),
             ("eps", "audit/eps_small.step", "EPS bead grade S (20 kg/m3)", "eps_moulding", {"steam_pressure_bar", "steam_time_s", "cooling_time_s"}),
             ("blow", "audit/blow_small.step", "PET bottle grade (IV 0.80)", "stretch_blow_moulding", {"preform_temp_c", "blow_pressure_bar", "mould_cooling_s"})]
    for code, file, material, process, expect in cases:
        pid = add_part(project(code.upper() + "PP"), file, material, process, code.upper() + "-1")
        r = client.get(f"/api/parts/{pid}/process-parameters")
        assert r.status_code == 200, (code, r.text)
        assert expect <= {x["name"] for x in r.json()["parameters"]}, code
    pdf = client.get(f"/api/parts/{pid}/process-parameters.pdf")
    assert pdf.status_code == 200 and pdf.content[:4] == b"%PDF"
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(client.get(f"/api/parts/{pid}/process-parameters.xlsx").content))
    assert any("Process" in n for n in wb.sheetnames)
    assert client.get(f"/api/parts/{pid}/process-parameters.docx").status_code == 404


def test_quality_risks_come_from_the_simulation_maps():
    # a thick-walled plastic part: sink risk is reported with the remedy, and hold pressure / time go up
    pid = mould("PP-RISK", file="plate_with_hole.step")
    res = client.get(f"/api/parts/{pid}/process-parameters").json()
    assert isinstance(res["risks"], list)
    p = {x["name"]: x for x in res["parameters"]}
    assert p["hold_pressure_bar"]["value"] >= 0.55 * p["injection_pressure_bar"]["value"] - 1e-6


def test_material_windows_and_rules_are_master_data():
    from backend.app import mould_db
    d = mould_db.tables()
    assert any(w["material"] == "ABS" for w in d["process_windows"]) and any(r["name"] == "melt_temp_c" and r["part_type"] == "plastic" for r in d["process_parameters"])
    next(r for r in d["process_parameters"] if r["name"] == "melt_temp_c")["formula"] = "no_such_name + 1"
    with pytest.raises(mould_db.MouldDbError, match="melt_temp_c"):
        mould_db.write(d)


# ---------- design skills ----------

def _plate_z(d, name):
    return next(c for c in d["components"] if c["name"] == name)["info"]["bbox_mm"][2]


def test_a_corrected_plate_becomes_a_skill_and_the_next_design_uses_it():
    pid = mould("SK-1")
    d1 = design(pid)
    comp = next(c for c in d1["components"] if c["name"] == "A plate (cavity plate)")
    z1 = comp["info"]["bbox_mm"][2]
    import cadquery as cq
    raw = client.get(f"/api/tool-designs/{d1['id']}/components/{comp['id']}/step").content
    p = Path(__file__).parent / "_tmp_plate.step"
    p.write_bytes(raw)
    old = cq.importers.importStep(str(p)).val()
    bb = old.BoundingBox()
    thick = cq.Workplane("XY").box(bb.xlen, bb.ylen, bb.zlen * 1.25, centered=False).translate((bb.xmin, bb.ymin, bb.zmin)).val()
    cq.exporters.export(cq.Workplane("XY").add(thick), str(p))
    r = client.post(f"/api/tool-designs/{d1['id']}/components/{comp['id']}/upload", files={"file": ("a.step", p.read_bytes())}, data={"note": "thicker A plate"})
    p.unlink()
    assert r.status_code == 200, r.text
    learned = r.json()["learned"]
    assert learned and learned[0]["target"] == "base.a_plate_mm" and 1.2 < learned[0]["value"] < 1.3 and "thickness" in learned[0]["text"]
    skills = client.get("/api/skills").json()
    assert any(s["target"] == "base.a_plate_mm" and s["mode"] == "factor" and s["evidence"] == 1 for s in skills)
    # the next design of the same kind of tool is made with the learned thickness
    d2 = design(pid)
    z2 = _plate_z(d2, "A plate (cavity plate)")
    assert z2 >= 1.15 * z1 and any(a["target"] == "base.a_plate_mm" for a in d2["summary"]["skills"]["applied"]) and "Your skills used" in d2["summary"]["notes"]
    # a design can be made without skills; a skill can be switched off or deleted
    r = client.post(f"/api/parts/{pid}/tool-designs", json={"use_skills": False}).json()
    assert abs(_plate_z(r, "A plate (cavity plate)") - z1) < 6 and r["summary"]["skills"]["applied"] == []
    sid = next(s["id"] for s in skills if s["target"] == "base.a_plate_mm")
    assert client.put(f"/api/skills/{sid}", json={"active": False}).json()["active"] is False
    assert abs(_plate_z(design(pid), "A plate (cavity plate)") - z1) < 6
    assert client.put(f"/api/skills/{sid}", json={"value": 5}).status_code == 400
    assert client.delete(f"/api/skills/{sid}").status_code == 200 and not client.get("/api/skills").json()


def test_cooling_corrections_are_learned_as_values_and_repeat_corrections_average():
    import cadquery as cq
    from backend.app import design_skills
    old = cq.Workplane("XY").box(200, 100, 60).faces(">X").workplane().pushPoints([(-20, 0), (20, 0)]).hole(8).val()
    new = cq.Workplane("XY").box(200, 100, 60).faces(">X").workplane().pushPoints([(-30, 0), (-10, 0), (10, 0), (30, 0)]).hole(10).val()
    pid = mould("SK-2")
    d = design(pid)
    comp = {"name": "B side insert"}
    got = design_skills.learn({"id": d["id"], "kind": "mould", "part_type": "plastic"}, comp, old, new, 300.0)
    assert {g["target"] for g in got} == {"rule.cool_dia_mm", "rule.cool_channels"}
    vals = {g["target"]: g["value"] for g in got}
    assert vals["rule.cool_dia_mm"] == 10 and vals["rule.cool_channels"] == 4
    again = design_skills.learn({"id": d["id"], "kind": "mould", "part_type": "plastic"}, comp, old, old, 300.0)          # no change: nothing learned
    assert again == []
    design_skills.learn({"id": d["id"], "kind": "mould", "part_type": "plastic"}, {"name": "A side insert"},
                        old, cq.Workplane("XY").box(200, 100, 60).faces(">X").workplane().pushPoints([(-30, 0), (0, 0), (30, 0)]).hole(12).val(), 300.0)
    s = {x["target"]: x for x in client.get("/api/skills").json()}
    assert s["rule.cool_dia_mm"]["value"] == 11.0 and s["rule.cool_dia_mm"]["evidence"] == 2          # (10 + 12) / 2
    # the next design cools with the learned channels
    d2 = design(pid)
    rules = d2["summary"].get("skills", {}).get("applied", [])
    assert any(a["target"] == "rule.cool_dia_mm" and a["to"] == 11.0 for a in rules)


def test_skill_knobs_are_master_data():
    from backend.app import mould_db
    d = mould_db.tables()
    assert any(k["component"] == "A plate" and k["target"] == "base.a_plate_mm" for k in d["skill_knobs"])
    d["skill_knobs"][0]["measure"] = "colour"
    with pytest.raises(mould_db.MouldDbError, match="measure"):
        mould_db.write(d)
