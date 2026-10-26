"""Stage 31: technical data, release (design loop) and revision history on the tool drawings; review and Excel round trip."""

import io
import zipfile

import ezdxf
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader

from backend.app import drawing_data, mould_db, register_setup
from backend.app.main import app
from tests.test_rm_cost import add_part, project

client = TestClient(app)


def bracket(code):
    return add_part(project(code), "example_bracket_revA.step", "Galvanised steel GI", "sheet_metal_multistage", f"{code}-1")


def pdf_text(content: bytes) -> str:
    return "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(content)).pages)


def die_plate(pid):
    data = client.get(f"/api/parts/{pid}/drawing-data").json()
    tool = data["tools"][0]
    return data, tool, next(i for i in tool["items"] if i["name"] == "Die plate")


def test_general_tolerances_are_master_data():
    rows = mould_db.tables()["general_tolerances"]
    assert len(rows) >= 50 and {r["tol_class"] for r in rows} == {"f", "m", "c", "v"} and all(r["is_example"] for r in rows)
    assert drawing_data.tolerance_class("ISO 2768-mK") == "m" and drawing_data.fmt_angle(90) == "±1°30′"
    assert drawing_data.tables_text("m")["linear"][3] == {"range": "30 – 120", "value": "±0.3"}


def test_technical_data_defaults_change_and_drawings():
    pid = bracket("DD1")
    data, tool, dp = die_plate(pid)
    assert data["editable"] and data["release"]["released"] is False
    assert dp["effective"]["material"] == "D2 (1.2379)" and dp["effective"]["tolerance_class"] == "m" and dp["effective"]["finish_ra"] == 1.6
    assert dp["defaults"]["heat_treatment"] != "none (as supplied)"                  # D2 is hardened
    r = client.put(f"/api/parts/{pid}/drawing-data", json={"items": [{"tool": tool["key"], "item": dp["item"],
        "fields": {"surface_treatment": "TiN coat cutting edges", "tolerance_class": "f", "finish_ra": 0.8}, "notes": ["Stress relieve before wire cut", "Mark die no."]}]})
    assert r.status_code == 200 and len(r.json()["changed"]) == 1
    again = client.put(f"/api/parts/{pid}/drawing-data", json={"items": [{"tool": tool["key"], "item": dp["item"],
        "fields": {"surface_treatment": "TiN coat cutting edges", "tolerance_class": "f", "finish_ra": 0.8}, "notes": ["Stress relieve before wire cut", "Mark die no."]}]})
    assert again.json()["changed"] == []
    assert client.put(f"/api/parts/{pid}/drawing-data", json={"items": [{"tool": tool["key"], "item": dp["item"], "fields": {"material": "X"}}]}).status_code == 400
    assert client.put(f"/api/parts/{pid}/drawing-data", json={"items": [{"tool": tool["key"], "item": "99", "fields": {}}]}).status_code == 400
    text = pdf_text(client.get(f"/api/parts/{pid}/tool-design.pdf", params={"die": tool["key"]}).content)
    for need in ("TECHNICAL DATA", "TiN coat cutting edges", "ISO 2768-f", "Ra 0.8", "Stress relieve before wire cut", "NOT RELEASED",
                 "Revision history", "Technical data changed"):
        assert need in text, need
    z = zipfile.ZipFile(io.BytesIO(client.get(f"/api/parts/{pid}/tool-design.zip", params={"die": tool["key"]}).content))
    dxf = next(n for n in z.namelist() if n.endswith("_Die_plate.dxf"))
    texts = " ".join(t.dxf.text for t in ezdxf.read(io.StringIO(z.read(dxf).decode())).modelspace().query("TEXT"))
    assert "TiN coat cutting edges" in texts and "REVISION HISTORY" in texts and "NOT RELEASED" in texts and "ISO 2768-f" in texts
    bom = next(n for n in z.namelist() if n.endswith("_BOM.xlsx"))
    ws = load_workbook(io.BytesIO(z.read(bom))).active
    assert "Surface treatment" in [c.value for c in ws[3]]


def test_release_from_the_design_loop_locks_the_data():
    pid = bracket("DD2")
    _, tool, dp = die_plate(pid)
    order = [x["code"] for x in register_setup.statuses()]
    for _ in range(10):
        loop = client.get(f"/api/parts/{pid}/design-loop").json()
        if loop["status"] == order[-1]:
            break
        nxt = max((c for c in loop["next"] if c != "corrections_required"), key=order.index)
        r = client.put(f"/api/parts/{pid}", json={"status": nxt, "status_comment": "ok"})
        assert r.status_code == 200, r.text
    data, _, _ = die_plate(pid)
    assert data["release"]["released"] and not data["editable"]
    assert client.put(f"/api/parts/{pid}/drawing-data", json={"items": [{"tool": tool["key"], "item": dp["item"], "notes": ["late"]}]}).status_code == 400
    text = pdf_text(client.get(f"/api/parts/{pid}/tool-design.pdf", params={"die": tool["key"]}).content)
    assert "NOT RELEASED" not in text and data["release"]["date"] in text and "Released" in text


def test_excel_round_trip():
    pid = bracket("DD3")
    _, tool, dp = die_plate(pid)
    wb = load_workbook(io.BytesIO(client.get(f"/api/parts/{pid}/drawing-data.xlsx").content))
    assert wb.sheetnames == ["Drawing data", "Dimension tolerances", "How to use", "Tolerance tables", "Revision history"]
    ws = wb["Drawing data"]
    top = next(r for r in range(1, 10) if ws.cell(r, 1).value == "Tool")
    row = next(r for r in range(top + 1, ws.max_row + 1) if ws.cell(r, 1).value == tool["key"] and str(ws.cell(r, 2).value) == dp["item"])
    ws.cell(row, 8).value = "Nitride 0.1 mm"
    ws.cell(row, 11).value = "Deburr\nEngrave part no."
    for w in wb.worksheets:
        w._images = []
    buf = io.BytesIO(); wb.save(buf)
    up = client.post(f"/api/parts/{pid}/drawing-data/upload", files={"file": ("d.xlsx", buf.getvalue())})
    assert up.status_code == 200, up.text
    assert len(up.json()["changed"]) == 1
    _, _, dp2 = die_plate(pid)
    assert dp2["effective"]["surface_treatment"] == "Nitride 0.1 mm" and dp2["changed"]
    assert client.post(f"/api/parts/{pid}/drawing-data/upload", files={"file": ("d.xlsx", b"nope")}).status_code == 400


def test_strip_layout_and_part_drawings_carry_release_and_history():
    pid = bracket("DD4")
    for kind in ("blank.dxf", "strip.dxf"):
        r = client.get(f"/api/parts/{pid}/sheet-layout/{kind}")
        if r.status_code != 200:
            continue
        texts = " ".join(t.dxf.text for t in ezdxf.read(io.StringIO(r.content.decode())).modelspace().query("TEXT"))
        assert "REVISION HISTORY" in texts and "NOT RELEASED" in texts
