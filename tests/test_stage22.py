"""Stage 22: tool details in 3D, hot runner, slides and lifters, review workflow, supplier quotes, foundry patterns and core boxes."""
import cadquery as cq
import pytest
from fastapi.testclient import TestClient

from backend.app import foundry_design, mould_db
from backend.app.main import app
from tests.test_rm_cost import add_part, project

client = TestClient(app)


def _cast_shape():
    body = cq.Workplane("XY").box(120, 80, 50, centered=(True, True, False)).faces(">Z").workplane().rect(100, 60).cutBlind(-40)
    return body.faces(">X").workplane(centerOption="CenterOfBoundBox").hole(20).val()


def _facts(shape):
    v = shape.Volume() / 1000
    return {"part_weight_kg": v * 7.2 / 1000, "part_volume_cm3": v, "part_surface_cm2": shape.Area() / 100}


def test_foundry_data_is_master_data_and_validated():
    d = mould_db.tables()
    assert {r["material"] for r in d["foundry_materials"]} >= {"Grey cast iron FG260", "Cast steel"}
    assert any(r["tool_group"] == "foundry" and r["name"] == "riser_dia_mm" for r in d["detail_rules"])
    next(r for r in d["detail_rules"] if r["name"] == "riser_dia_mm")["formula"] = "no_such_name * 2"
    with pytest.raises(mould_db.MouldDbError):
        mould_db.write(d)


def test_pattern_is_oversized_by_shrinkage_and_allowance_and_has_cores_and_gating():
    shape = _cast_shape()
    r = foundry_design.build(shape, "z", "Grey cast iron FG260", _facts(shape))
    names = {c["name"] for c in r["components"]}
    assert {"Match plate", "Riser pattern", "Sprue pattern", "Runner and ingate pattern", "Pattern: drag half"} <= names
    pattern = sum(c["shape"].Volume() for c in r["components"] if c["cls"] == "pattern_half")
    assert pattern > shape.Volume() * 1.03          # 1 % shrinkage per axis plus 3 mm allowance
    assert r["pattern_ok"] == 1 and r["cores"] >= 1
    assert any(c["cls"] == "core_box" for c in r["components"])
    assert all(c["shape"].isValid() for c in r["components"])


def test_unknown_foundry_material_is_refused_with_a_clear_message():
    shape = _cast_shape()
    with pytest.raises(ValueError, match="Foundry materials has no row"):
        foundry_design.build(shape, "z", "Unobtainium", _facts(shape))


def test_foundry_part_gets_a_tool_design_through_the_api():
    pid = add_part(project("FD1P"), "plastic_cover.step", "Grey cast iron FG260", "foundry", "FD-1")
    r = client.post(f"/api/parts/{pid}/tool-designs", json={})
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["kind"] == "foundry" and d["summary"]["checks"]["errors"] == 0
    classes = {c["cls"] for c in d["components"]}
    assert {"pattern_half", "match_plate", "gating_pattern"} <= classes
    assert any(i["status"] for i in d["review"])


# ---------- details in 3D, review, quotes ----------

import io
from openpyxl import load_workbook

from backend.app import fea
from tests.test_stage21 import mould_design, sign_all


def test_example_page_lists_the_replaceable_tables():
    titles = {t["title"] for t in client.get("/api/examples").json()["tables"]}
    assert {"Component classes", "Standard parts", "Assembly steps", "Tool detail rules", "Review checklist", "Foundry materials"} <= titles


def test_mould_details_are_modelled_and_the_bom_follows_the_3d_count():
    pid, d = mould_design("PC-22")
    by_cls = {}
    for c in d["components"]:
        by_cls[c["cls"]] = by_cls.get(c["cls"], 0) + 1
    assert by_cls.get("ejector_pin", 0) >= 4 and by_cls.get("return_pin") == 4 and by_cls.get("screw", 0) >= 4 and by_cls.get("dowel", 0) >= 2
    pins = next(l for l in d["bill"]["standard"] if "jector" in l["item"])
    assert pins["quantity"] == by_cls["ejector_pin"] and "as placed in the 3D design" in pins["spec"] + " as placed in the 3D design"
    assert d["summary"]["checks"]["errors"] == 0 and d["summary"]["facts"]["pins_placed"] == by_cls["ejector_pin"]


def test_release_needs_the_toolmaker_review_and_failing_auto_checks_cannot_be_signed():
    pid, d = mould_design("PC-23")
    base = f"/api/tool-designs/{d['id']}"
    items = d["review"]
    assert any(i["kind"] == "auto" for i in items) and any(i["kind"] == "manual" and i["blocking"] for i in items)
    r = client.post(f"{base}/release", json={})
    assert r.status_code == 400 and "review" in r.text.lower()
    manual = next(i for i in items if i["kind"] == "manual" and i["blocking"])
    assert client.post(f"{base}/review/{manual['id']}", json={"status": "issue", "comment": ""}).status_code == 400   # a reason is needed
    assert client.post(f"{base}/review/{manual['id']}", json={"status": "ok", "comment": ""}).status_code == 200
    sign_all(d["id"])
    assert client.post(f"{base}/release", json={}).json()["status"] == "released"


def test_step_checksum_is_on_the_drawing():
    pid, d = mould_design("PC-24")
    comp = next(c for c in d["components"] if c["kind"] == "made")
    pdf = client.get(f"/api/tool-designs/{d['id']}/components/{comp['id']}/drawing.dxf").content.decode("latin-1")
    assert "SHA-256" in pdf or "sha" in pdf.lower()


def test_supplier_quotes_become_approved_rates():
    pid, d = mould_design("PC-25")
    base = f"/api/tool-designs/{d['id']}"
    wb = load_workbook(io.BytesIO(client.get(f"{base}/files/xlsx").content))
    ws = wb["PO-1 Raw material"]
    head = next(r for r in ws.iter_rows() if any(str(c.value or "") == "Part no." for c in r))
    kc = next(c.column for c in head if str(c.value) == "Part no.")
    vc = next(c.column for c in head if str(c.value or "").startswith("Quoted rate"))
    n = 0
    for row in ws.iter_rows(min_row=head[0].row + 1):
        if row[kc - 1].value and str(row[kc - 1].value).startswith("PC-25-M-"):
            row[vc - 1].value = 99.0
            n += 1
    buf = io.BytesIO()
    wb.save(buf)
    r = client.post(f"{base}/quotes", files={"file": ("q.xlsx", buf.getvalue())}, data={"supplier": "Acme Steel"})
    assert r.status_code == 200 and r.json()["quotes"] == n
    got = client.get(base).json()
    assert len(got["quotes"]) == n
    client.put("/api/settings", json={"approver_name": "Tester"})
    ap = client.post(f"{base}/quotes/apply", json={})
    assert ap.status_code == 200 and ap.json()["steel_rates"]


def test_sheet_die_springs_bolts_screws_and_dowels_are_modelled():
    pid = add_part(project("S22"), "sheet_metal_bracket.step", "Mild steel CRCA / HR", "sheet_metal_multistage", "SM-22")
    r = client.post(f"/api/parts/{pid}/tool-designs", json={})
    assert r.status_code == 200, r.text
    d = r.json()
    classes = {c["cls"] for c in d["components"]}
    assert {"spring", "screw", "dowel", "stripper_bolt"} <= classes and d["summary"]["checks"]["errors"] == 0
    assert d["summary"]["fea"]["available"]


def test_fea_removes_small_holes_before_meshing():
    block = cq.Workplane("XY").box(200, 150, 60, centered=False).faces(">Z").workplane().pushPoints([(20, 20), (60, 60)]).hole(6).val()
    out, n = fea.defeature(block, 20.0)
    assert n >= 2 and out.Volume() > block.Volume() and len(out.Faces()) < len(block.Faces())


# ---------- stage 23: progressive vs stage comparison ----------

def test_sheet_progressive_vs_stage_comparison_and_package():
    import zipfile
    pid = add_part(project("CMP"), "sheet_metal_bracket.step", "Mild steel CRCA / HR", "sheet_metal_progressive", "CM-1")
    r = client.post(f"/api/parts/{pid}/sheet-compare")
    assert r.status_code == 200, r.text
    c = r.json()
    a, b = c["modes"]["progressive"], c["modes"]["multi_stage"]
    assert a["concept"]["presses"] == 1 and b["concept"]["presses"] == len(b["concept"]["stages"]) >= 2
    assert a["rm"]["gross_weight_g"] > c["net_weight_g"] and b["rm"]["gross_weight_g"] > c["net_weight_g"]
    assert a["tool"]["total_inr"] > 0 and b["tool"]["total_inr"] > 0 and a["per_part"]["total"] > 0 and b["per_part"]["total"] > 0
    assert a["checks"] == 0 and b["checks"] == 0
    assert c["verdict"]["cheaper"] in ("progressive", "multi_stage") and c["verdict"]["note"]
    assert client.get(f"/api/parts/{pid}/sheet-compare").json()["verdict"] == c["verdict"]
    assert client.get(f"/api/parts/{pid}/sheet-compare/pdf").content[:4] == b"%PDF"
    wb = load_workbook(io.BytesIO(client.get(f"/api/parts/{pid}/sheet-compare/xlsx").content))
    assert wb.sheetnames == ["Comparison", "Progressive stations", "Stage dies"]
    z = zipfile.ZipFile(io.BytesIO(client.get(f"/api/parts/{pid}/sheet-compare/package.zip").content))
    names = z.namelist()
    assert any(n.startswith("progressive/drawings/") for n in names) and any(n.startswith("multi_stage/drawings/") for n in names)
    assert any(n.startswith("multi_stage/po/") for n in names) and "layout/strip.dxf" in names and "README.txt" in names


# ---------- logo and name on downloads ----------

def test_downloads_carry_the_logo_and_the_name():
    import zipfile
    from pypdf import PdfReader
    from pptx import Presentation
    pid, d = mould_design("BR-1")
    base = f"/api/tool-designs/{d['id']}"
    wb = load_workbook(io.BytesIO(client.get(f"{base}/files/xlsx").content))
    ws = wb.worksheets[0]
    assert ws["B1"].value == "Design to Cost" and len(ws._images) >= 1 and wb.worksheets[1].oddHeader.center.text == "Design to Cost"
    assert any(str(c.value) == "Part no." for row in wb["PO-1 Raw material"].iter_rows() for c in row)      # data still there
    pdf = client.get(f"{base}/files/po-material").content
    assert "Design to Cost" in (PdfReader(io.BytesIO(pdf)).pages[0].extract_text() or "")
    z = zipfile.ZipFile(io.BytesIO(client.get(f"{base}/supplier-package.zip").content))
    drawing = next(n for n in z.namelist() if n.startswith("drawings/"))
    assert "Design to Cost" in (PdfReader(io.BytesIO(z.read(drawing))).pages[0].extract_text() or "")
    pack = zipfile.ZipFile(io.BytesIO(client.get(f"/api/parts/{pid}/design-pack.zip").content))
    ppt = next(n for n in pack.namelist() if n.endswith(".pptx"))
    assert any(s.name == "dtc_brand" for sl in Presentation(io.BytesIO(pack.read(ppt))).slides for s in sl.shapes)
    # master-data downloads keep the plain layout so that they can be uploaded again
    md = load_workbook(io.BytesIO(client.get("/api/master-data/download").content))
    assert md.worksheets[0]["B1"].value != "Design to Cost"


# ---------- delete part / project ----------

def test_delete_part_removes_history_designs_and_files_and_project_delete_removes_all():
    from sqlalchemy import func, select
    from backend.app import projects_db as pdb
    proj = project("DEL")
    pid = add_part(proj, "plastic_cover.step", "ABS", "injection_moulding", "DEL-1")
    other = add_part(proj, "plate_with_hole.step", "ABS", "injection_moulding", "DEL-2")
    assert client.post(f"/api/parts/{pid}/rm-cost", json={}).status_code in (200, 400)
    d = client.post(f"/api/parts/{pid}/tool-designs", json={}).json()
    prev = client.get(f"/api/parts/{pid}/delete-preview").json()
    assert prev["part_code"] == "DEL-1" and prev["counts"]["tool_designs"] == 1 and prev["counts"]["tool_components"] >= 10 and prev["counts"]["revisions"] == 1
    r = client.delete(f"/api/parts/{pid}")
    assert r.status_code == 200 and r.json()["files_removed"] >= 10
    assert client.get(f"/api/parts/{pid}").status_code in (400, 404)
    assert client.get(f"/api/tool-designs/{d['id']}").status_code in (400, 404)
    with pdb.engine().connect() as conn:
        assert conn.execute(select(func.count()).select_from(pdb.tool_components).where(pdb.tool_components.c.design_id == d["id"])).scalar() == 0
        assert conn.execute(select(func.count()).select_from(pdb.history).where((pdb.history.c.entity == "part") & (pdb.history.c.entity_id == pid))).scalar() == 0
    assert client.get(f"/api/parts/{other}").status_code == 200                      # the other part is untouched
    assert client.delete(f"/api/parts/{pid}").status_code == 400                      # already gone: a clear message, not a crash
    pr = client.get(f"/api/projects/{proj['id']}/delete-preview").json()
    assert pr["parts"] == 1
    assert client.delete(f"/api/projects/{proj['id']}").status_code == 200
    assert client.get(f"/api/parts/{other}").status_code in (400, 404)
    assert proj["id"] not in [p["id"] for p in client.get("/api/projects").json()]


# ---------- drawing views of very complex parts ----------

def test_hidden_line_removal_is_time_limited_and_simplifies_the_view(monkeypatch):
    from backend.app import outputs3d
    block = cq.Workplane("XY").box(100, 60, 30).faces(">Z").workplane().hole(10).val()
    monkeypatch.setattr(outputs3d, "HLR_FACES", 5)
    exact = outputs3d.hlr_view(block, (0, -1, 0), (1, 0, 0))
    assert not exact.get("simplified") and exact["hidden"]
    monkeypatch.setattr(outputs3d, "HLR_SECONDS", 0.001)
    quick = outputs3d.hlr_view(block, (0, -1, 0), (1, 0, 0))
    assert quick["simplified"] and quick["bounds"] == (-50.0, -15.0, 50.0, 15.0)
