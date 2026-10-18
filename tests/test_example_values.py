"""Stage 4b: Replace example values checklist (counts by category, replace / keep rates and rows)."""
import io
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.app import mould_db, rates, tool_cost
from backend.app.main import app

SAMPLES = Path(__file__).resolve().parents[1] / "samples"
MONTH = "2026-10"
client = TestClient(app)


def checklist():
    r = client.get("/api/examples")
    assert r.status_code == 200
    return r.json()


def group(c, category):
    return next(g for g in c["rates"] if g["category"] == category)


def table(c, set_name, name):
    return next(t for t in c["tables"] if t["set"] == set_name and t["table"] == name)


@pytest.fixture
def approver():
    client.put("/api/settings", json={"approver_name": "Tester"})


def test_counts_by_category_and_country():
    c = checklist()
    s = c["summary"]
    every = [r for r in rates.current(MONTH) if r["approved"]]
    assert s["rates_total"] == len(every)
    assert s["rates_left"] == sum(r["approved"]["is_example"] for r in every)
    tc = group(c, "tool_commercial")
    assert tc["example"] == tc["total"] == 36 and tc["by_country"]["IN"] == 4
    assert group(c, "forex")["by_country"] == {"global": 8}
    tp = table(c, "mould", "time_parameters")
    assert tp["example"] == tp["total"] == len(mould_db.tables()["time_parameters"]) and tp["value_column"] == "value"
    assert table(c, "setup", "machines")["value_column"] is None
    assert s["rows_left"] == sum(t["example"] for t in c["tables"])


def test_replace_a_rate_makes_a_new_approved_real_version(approver):
    old = rates.lookup("tool_material", "H13 (1.2344)", "IN", MONTH)
    before = group(checklist(), "tool_material")["example"]
    r = client.post("/api/examples/rate", json={"rate_id": old["id"], "value": 690})
    assert r.status_code == 200, r.text
    new = rates.lookup("tool_material", "H13 (1.2344)", "IN", MONTH)
    assert new["id"] == r.json()["id"] and new["value"] == 690 and not new["is_example"] and new["status"] == "approved"
    assert rates.get(old["id"])["value"] == old["value"]                 # old version kept
    assert group(checklist(), "tool_material")["example"] == before - 1
    again = client.post("/api/examples/rate", json={"rate_id": old["id"], "value": 700})
    assert again.status_code == 400 and "already been replaced" in again.json()["detail"]
    assert client.post("/api/examples/rate", json={"rate_id": new["id"]}).status_code == 400   # not an example


def test_keep_confirms_the_value(approver):
    old = rates.lookup("power", "electricity_industrial", "IN", MONTH)
    client.post("/api/examples/rate", json={"rate_id": old["id"]})
    new = rates.lookup("power", "electricity_industrial", "IN", MONTH)
    assert new["value"] == old["value"] and not new["is_example"] and "Confirmed" in new["source"]


def test_replacing_needs_an_approver_name():
    old = rates.lookup("power", "electricity_industrial", "IN", MONTH)
    r = client.post("/api/examples/rate", json={"rate_id": old["id"], "value": 9})
    assert r.status_code == 400 and "name for approvals" in r.json()["detail"]


def test_replace_and_keep_master_data_rows():
    key = {"part_type": "plastic", "parameter": "cnc_finish_cm2_per_h"}
    r = client.post("/api/examples/row", json={"set": "mould", "table": "time_parameters", "key": key, "value": 55})
    assert r.status_code == 200, r.text
    row = mould_db.time_parameters("plastic")["cnc_finish_cm2_per_h"]
    assert row["value"] == 55 and not row["is_example"]
    res = tool_cost.run(tool_cost.ToolInput(part_type="plastic", cavities=2, part_length_mm=120, part_width_mm=80,
                                            part_height_mm=30, projected_area_cm2=95.78, part_weight_g=48.66,
                                            part_surface_cm2=420, material="ABS", as_of=MONTH))
    assert "Process time parameters · plastic / cnc_finish_cm2_per_h" not in res["example"]["rows"]

    keep = client.post("/api/examples/row", json={"set": "setup", "table": "machines", "key": {"machine_id": "VMC-850"}})
    assert keep.status_code == 200
    assert table(checklist(), "setup", "machines")["example"] == 32
    bad = client.post("/api/examples/row", json={"set": "setup", "table": "machines", "key": {"machine_id": "VMC-850"},
                                                 "value": 5})
    assert bad.status_code == 400 and "edit this row in Master data" in bad.json()["detail"]
    missing = client.post("/api/examples/row", json={"set": "mould", "table": "sizing_rules", "key": {"part_type": "x", "parameter": "y"}})
    assert missing.status_code == 400 and "no row" in missing.json()["detail"]


def test_used_by_counts_stored_costs():
    p = client.post("/api/projects", json={"code": "EX", "name": "Ex"}).json()
    with open(SAMPLES / "plastic_cover.step", "rb") as fh:
        m = client.post("/api/import/step", files={"file": ("plastic_cover.step", fh)}).json()
    client.post(f"/api/import/{m['token']}/commit", json={
        "project_id": p["id"], "items": {m["root"]: {"material": "ABS", "process": "injection_moulding", "part_code": "EX-1"}}})
    c = checklist()
    abs_in = next(i for i in group(c, "material")["items"] if i["item"] == "ABS" and i["country"] == "IN")
    assert abs_in["used_by"] == 1 and group(c, "material")["items"][0]["used_by"] >= 1   # used ones come first
    assert c["summary"]["rates_used_in_costs"] > 0


def test_excel_export_and_pages():
    wb = load_workbook(io.BytesIO(client.get("/api/examples/export.xlsx").content))
    assert wb.sheetnames == ["Summary", "Example rates", "Example rows"]
    assert wb["Example rates"].max_row - 1 == checklist()["summary"]["rates_left"]
    assert client.get("/master-data/examples").status_code == 200
    assert "/master-data/examples" in client.get("/master-data").text
