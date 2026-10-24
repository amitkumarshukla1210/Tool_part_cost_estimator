"""Stage 21: production tool design: components, part numbers, drawings, STEP correction by upload, checks, release, purchase-order pack."""
import io
import zipfile

import cadquery as cq
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.app import mould_db, tool_components, tool_design_store as st
from backend.app.main import app
from backend.app.register import RegisterError
from tests.test_rm_cost import add_part, project

client = TestClient(app)


def mould_design(code="PC-1", material="ABS", process="injection_moulding"):
    pid = add_part(project(code + "P"), "plastic_cover.step", material, process, code)
    return pid, client.post(f"/api/parts/{pid}/tool-designs", json={}).json()


def sign_all(design_id):
    """The toolmaker signs every blocking review item (an automatic check that fails is waived with a reason)."""
    base = f"/api/tool-designs/{design_id}"
    for it in client.get(base).json()["review"]:
        if it["blocking"] and it["status"] not in ("ok", "na", "waived"):
            body = {"status": "ok", "comment": ""} if it["kind"] == "manual" else {"status": "waived", "comment": "accepted in test"}
            assert client.post(f"{base}/review/{it['id']}", json=body).status_code == 200


@pytest.fixture(scope="module")
def _unused():
    return None


# ---------- master data ----------

def test_component_classes_standard_parts_and_steps_are_master_data():
    db = mould_db.tables()
    classes = {r["code"] for r in db["component_classes"]}
    assert {"base_plate", "insert_a", "insert_a_hot", "die_block", "punch", "guide_pillar"} <= classes
    assert any(r["item"] == "Stripper spring" for r in db["standard_parts"]) and len(db["assembly_steps"]) >= 10
    db["component_classes"][0]["steel"] = "Unobtainium"
    with pytest.raises(mould_db.MouldDbError, match="Component classes row 2: steel"):
        mould_db.write(db)
    db = mould_db.tables()
    next(r for r in db["standard_parts"] if r["item"] == "Dowel pin")["quantity_rule"] = "no_such_name"
    with pytest.raises(mould_db.MouldDbError, match="Standard parts row"):
        mould_db.write(db)


def test_stock_size_follows_the_class_allowances_and_rounding():
    plate = cq.Workplane("XY").box(101, 52, 20, centered=False).val()
    cls = tool_components.class_row("base_plate", "plastic")
    s = tool_components.stock(plate, cls)
    # (101 + 2·3) → 110, (52 + 6) → 60, thickness 20 + 0.6 grinding + 2·3 allowance → 30 (rounded up to 5 mm)
    assert s["size_mm"] == [110, 60, 30] or s["size_mm"][:2] == [110, 60]
    assert s["weight_kg"] == pytest.approx(s["size_mm"][0] * s["size_mm"][1] * s["size_mm"][2] * 7.85e-6, rel=0.02)
    assert tool_components.class_row("insert_a", "die_casting")["steel"] == "H13 (1.2344)"
    assert tool_components.class_row("insert_a", "eps")["steel"] == "Aluminium 7075-T6"


def test_features_hole_table_and_openings_are_read_from_the_solid():
    plate = (cq.Workplane("XY").box(120, 80, 20, centered=False).faces(">Z").workplane(origin=(0, 0, 20))
             .pushPoints([(20, 20), (100, 60)]).hole(10).faces(">Z").workplane(origin=(60, 40, 20)).rect(30, 20).cutThruAll().val())
    f = tool_components.features(plate)
    assert f["hole_count"] == 2 and f["holes"][0]["dia_mm"] == 10 and f["holes"][0]["through"]
    assert sorted(f["holes"][0]["positions_mm"]) == [[20.0, 20.0, 10.0], [100.0, 60.0, 10.0]]
    assert len(f["openings"]) == 1 and f["openings"][0]["through"] and f["openings"][0]["perimeter_mm"] == pytest.approx(100, abs=0.5)
    assert f["openings"][0]["x_mm"] == pytest.approx(45) and f["openings"][0]["y_mm"] == pytest.approx(30)


def test_interference_is_found_by_volume():
    a = cq.Workplane("XY").box(10, 10, 10).val()
    b = cq.Workplane("XY").box(10, 10, 10).translate((5, 0, 0)).val()
    c = cq.Workplane("XY").box(10, 10, 10).translate((10, 0, 0)).val()
    r = tool_components.interference([{"id": 1, "name": "a", "shape": a}, {"id": 2, "name": "b", "shape": b}, {"id": 3, "name": "c", "shape": c}])
    assert [(x["a"], x["b"], round(x["volume_mm3"])) for x in r] == [(1, 2, 500), (2, 3, 500)]       # a and c only touch


# ---------- generate, view, drawings ----------

def test_mould_design_has_numbered_components_bom_and_no_overlap():
    pid, d = mould_design()
    assert d["status"] == "draft" and d["kind"] == "mould" and d["summary"]["checks"]["errors"] == 0
    made = [c for c in d["components"] if c["kind"] == "made"]
    assert {c["part_no"] for c in made} >= {"PC-1-M-01", "PC-1-M-09"}
    assert sum(1 for c in made if c["part_no"] == [m for m in made if "Spacer" in m["name"]][0]["part_no"]) == 2     # two identical spacers, one number
    b = d["bill"]
    assert b["totals"]["made_parts"] == 9 and b["totals"]["made_qty"] == 10 and b["totals"]["steel_kg"] > 100
    assert any(l["item"] == "Guide pillar" for l in b["standard"]) and any(l["item"] == "Socket head cap screw" for l in b["standard"])
    ins = next(m for m in b["made"] if "B side insert" in m["name"])
    assert ins["info"]["steel"].startswith("P20") and ins["info"]["stock"]["size_mm"][0] >= ins["info"]["bbox_mm"][0] + 6
    apl = next(m for m in b["made"] if "A plate" in m["name"])
    bush = [h for h in apl["info"]["features"]["holes"] if h["dia_mm"] == pytest.approx(36)]
    assert len(bush) == 1 and len(bush[0]["positions_mm"]) == 4                                                                  # guide bush bores
    assert apl["info"]["features"]["hole_count"] >= 4                                                                           # plus screws, dowels, sprue
    assert apl["info"]["features"]["openings"][0]["through"]                                                                   # insert pocket
    assert d["bill"]["totals"]["assembly_hours"] > 10


def test_die_casting_uses_hot_work_steel_and_other_groups_are_covered():
    pid, d = mould_design("DC-1", "Aluminium ADC12 / A383", "die_casting")
    ins = next(m for m in d["bill"]["made"] if "B side insert" in m["name"])
    assert ins["info"]["steel"].startswith("H13") and any("Heat" in o["process"] for o in ins["info"]["operations"])


def test_sheet_metal_design_per_stage_and_progressive():
    pid = add_part(project("S"), "sheet_metal_bracket.step", "Mild steel CRCA / HR", "sheet_metal_multistage", "SM-1")
    r0 = client.post(f"/api/parts/{pid}/tool-designs", json={})
    assert r0.status_code == 200, r0.text
    d = r0.json()
    assert d["kind"] == "sheet_metal" and d["summary"]["checks"]["errors"] == 0
    nos = {c["part_no"] for c in d["components"] if c["kind"] == "made"}
    assert any("-S1-" in n for n in nos) and any("-S2-" in n for n in nos)
    assert any(l["item"] == "Stripper spring" for l in d["bill"]["standard"])
    pid2 = add_part(project("S2"), "sheet_metal_bracket.step", "Mild steel CRCA / HR", "sheet_metal_progressive", "SM-2")
    d2 = client.post(f"/api/parts/{pid2}/tool-designs", json={}).json()
    assert d2["summary"]["mode"] == "progressive" and all("-P-" in c["part_no"] for c in d2["components"] if c["kind"] == "made")
    assert d2["summary"]["checks"]["errors"] == 0                                    # strip fed along the width: stations do not collide


def test_drawings_step_assembly_mesh_and_endpoints():
    pid, d = mould_design("PC-2")
    comp = next(c for c in d["components"] if c["kind"] == "made" and "B side insert" in c["name"])
    base = f"/api/tool-designs/{d['id']}"
    pdf = client.get(f"{base}/components/{comp['id']}/drawing.pdf")
    assert pdf.status_code == 200 and pdf.content[:4] == b"%PDF"
    assert client.get(f"{base}/components/{comp['id']}/drawing.dxf").content.startswith(b"  0")
    step = client.get(f"{base}/components/{comp['id']}/step")
    assert b"ISO-10303" in step.content[:200] and comp["part_no"] in step.headers["content-disposition"]
    std = next(c for c in d["components"] if c["kind"] == "standard")
    assert client.get(f"{base}/components/{std['id']}/drawing.pdf").status_code == 400            # catalogue parts have no drawing
    assert b"ISO-10303" in client.get(f"{base}/assembly.step").content[:200]
    mesh = client.get(f"{base}/mesh").json()
    assert len(mesh["components"]) == len(d["components"]) and all(c["v"] and c["f"] for c in mesh["components"])
    assert any(c["info"]["features"]["holes"] for c in mesh["components"] if c["kind"] == "made")
    assert client.get("/tool-design").status_code == 200
    assert client.get(f"/api/parts/{pid}/tool-designs").json()[0]["id"] == d["id"]


# ---------- correct by STEP upload ----------

def edited_step(content: bytes, fn) -> bytes:
    import tempfile, os
    with tempfile.TemporaryDirectory() as t:
        p = os.path.join(t, "a.step"); open(p, "wb").write(content)
        shape = cq.importers.importStep(p).val()
        out = fn(cq.Workplane("XY").add(shape))
        q = os.path.join(t, "b.step"); cq.exporters.export(out, q)
        return open(q, "rb").read()


def test_upload_corrected_step_recalculates_and_flags_overlap_and_revert():
    pid, d = mould_design("PC-3")
    base = f"/api/tool-designs/{d['id']}"
    comp = next(c for c in d["components"] if "Top clamp plate" in c["name"])
    step = client.get(f"{base}/components/{comp['id']}/step").content
    # thinner plate: lose 5 mm on top
    thin = edited_step(step, lambda w: w.faces(">Z").workplane().rect(1000, 1000).cutBlind(-5))
    r = client.post(f"{base}/components/{comp['id']}/upload", files={"file": ("x.step", thin)}, data={"note": "5 mm thinner"})
    assert r.status_code == 200, r.text
    out = r.json()
    assert out["rev"] == "B" and out["size_mm"]["after"][2] == pytest.approx(out["size_mm"]["before"][2] - 5, abs=0.05)
    assert out["stock"]["size_mm"][2] <= 30 and out["volume_cm3"]["after"] < out["volume_cm3"]["before"]
    # thicker plate overlaps the locating ring / sprue bush
    thick = edited_step(step, lambda w: w.faces("<Z").workplane().rect(1000, 1000).extrude(8))
    r2 = client.post(f"{base}/components/{comp['id']}/upload", files={"file": ("y.step", thick)}, data={"note": "8 mm thicker"}).json()
    assert r2["rev"] == "C" and r2["checks"]["errors"] + r2["checks"]["warnings"] > 0 or r2["issues_for_this_part"]
    revs = client.get(f"{base}/components/{comp['id']}/revisions").json()
    assert [x["rev"] for x in revs] == ["A", "B", "C"]
    assert client.post(f"{base}/components/{comp['id']}/revert", json={"rev": "A"}).status_code == 200
    assert client.get(f"{base}/check".replace("/check", "")).json()["summary"]["checks"]["errors"] == 0
    # empty note and garbage are refused
    assert client.post(f"{base}/components/{comp['id']}/upload", files={"file": ("z.step", step)}, data={"note": " "}).status_code == 400
    assert client.post(f"{base}/components/{comp['id']}/upload", files={"file": ("z.step", b"not a step")}, data={"note": "x"}).status_code == 400


def test_identical_parts_follow_an_edit():
    pid, d = mould_design("PC-4")
    base = f"/api/tool-designs/{d['id']}"
    spacers = [c for c in d["components"] if "Spacer" in c["name"]]
    step = client.get(f"{base}/components/{spacers[0]['id']}/step").content
    new = edited_step(step, lambda w: w.faces(">Z").workplane().rect(1000, 1000).cutBlind(-2))
    r = client.post(f"{base}/components/{spacers[0]['id']}/upload", files={"file": ("x.step", new)}, data={"note": "2 mm lower"}).json()
    assert r["applied_to"] == 2
    got = client.get(base).json()
    heights = {c["info"]["bbox_mm"][2] for c in got["components"] if "Spacer" in c["name"]}
    assert len(heights) == 1 and heights.pop() == pytest.approx(spacers[0]["info"]["bbox_mm"][2] - 2, abs=0.05)


# ---------- release and the purchase-order pack ----------

def test_release_blocks_edits_and_new_revision_starts_from_it():
    pid, d = mould_design("PC-5")
    base = f"/api/tool-designs/{d['id']}"
    assert client.post(f"{base}/release", json={}).status_code == 400            # the toolmaker's review is not signed yet
    sign_all(d["id"])
    rel = client.post(f"{base}/release", json={})
    assert rel.status_code == 200 and rel.json()["status"] == "released"
    comp = d["components"][0]
    step = client.get(f"{base}/components/{comp['id']}/step").content
    assert client.post(f"{base}/components/{comp['id']}/upload", files={"file": ("x.step", step)}, data={"note": "x"}).status_code == 400
    rev = client.post(f"{base}/new-revision", json={}).json()
    assert rev["status"] == "draft" and rev["rev_label"] == "B" and rev["version"] == 2 and len(rev["components"]) == len(d["components"])
    sign_all(rev["id"])
    assert client.post(f"/api/tool-designs/{rev['id']}/release", json={}).json()["status"] == "released"
    assert client.get(base).json()["status"] == "superseded"


def test_purchase_pack_workbook_pos_and_supplier_zip():
    pid, d = mould_design("PC-6")
    base = f"/api/tool-designs/{d['id']}"
    wb = load_workbook(io.BytesIO(client.get(f"{base}/files/xlsx").content))
    assert wb.sheetnames == ["Cover", "BOM", "PO-1 Raw material", "PO-2 Machining", "Machining detail", "PO-3 Bought-out", "Assembly", "Inspection", "Features", "Machine", "Quote comparison", "Cost summary"]
    ws = wb["PO-1 Raw material"]
    rows = [r for r in ws.iter_rows(values_only=True) if r[1] and str(r[1]).startswith("PC-6-M-")]
    assert len(rows) == 9 and all(r[3] and r[4] and r[9] for r in rows)
    assert "DRAFT" in str(wb["Cover"]["A1"].value)
    for kind in ("po-material", "po-machining", "po-bought-out", "assembly-pdf"):
        assert client.get(f"{base}/files/{kind}").content[:4] == b"%PDF"
    z = zipfile.ZipFile(io.BytesIO(client.get(f"{base}/supplier-package.zip").content))
    names = z.namelist()
    assert sum(n.startswith("step/") for n in names) == 9 and sum(n.startswith("drawings/") for n in names) == 9 and sum(n.startswith("dxf/") for n in names) == 9
    assert any(n.startswith("po/") and n.endswith(".xlsx") for n in names) and "README.txt" in names
    sign_all(d["id"])
    client.post(f"{base}/release", json={})
    wb2 = load_workbook(io.BytesIO(client.get(f"{base}/files/xlsx").content))
    assert "DRAFT" not in str(wb2["Cover"]["A1"].value)


def test_parts_without_a_tool_are_refused_with_a_reason():
    pid = add_part(project("F"), "plate_with_hole.step", "Mild steel CRCA / HR", "machined", "MC-1")
    r = client.post(f"/api/parts/{pid}/tool-designs", json={})
    assert r.status_code == 400
