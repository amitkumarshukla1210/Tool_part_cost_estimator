"""Stage 29: price master (month-wise prices by country), material types, recommended part grade, new EXAMPLE grades, V0."""
import io
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.app import master_data, mould_db, price_master, rates, tool_cost
from backend.app.main import app

SAMPLES = Path(__file__).resolve().parents[1] / "samples"
ROOT = Path(__file__).resolve().parents[1]
MONTH = "2026-10"
client = TestClient(app)


def approver():
    client.put("/api/settings", json={"approver_name": "Tester"})


def row(group, country, item, **kw):
    g = client.get("/api/price-master/grid", params={"group": group, "country": country, **kw}).json()
    return g, next(r for r in g["rows"] if r["item"] == item)


# ---------- master data ----------

def test_types_on_every_row_and_one_recommended_grade_per_type():
    steels = {r["code"]: r for r in mould_db.tables()["tool_materials"]}
    assert all(r["material_type"] for r in steels.values())
    assert steels["BeCu"]["material_type"] == "Copper alloy" and steels["D2 (1.2379)"]["material_type"] == "Cold-work tool steel"
    mats = master_data.tables()["Materials"]
    assert all(r["material_type"] for r in mats)
    per_type = {}
    for r in mats:
        per_type.setdefault((r["part_type"], r["material_type"]), []).append(r["recommended"])
    assert all(v.count("yes") == 1 for v in per_type.values())


def test_new_grades_are_example_rows_with_example_rates_in_every_country():
    steels = {r["code"]: r for r in mould_db.tables()["tool_materials"]}
    for code in ("Beryllium copper C17200", "Beryllium copper C17510", "NAK80", "1.2316 ESR (stainless)"):
        assert steels[code]["is_example"]
        for c in ("IN", "CN", "DE", "JP", "KR", "TH", "TW", "US", "VN"):
            assert rates.lookup("tool_material", code, c, MONTH)["is_example"]
    for grade in ("PP copolymer (impact)", "CRCA D513 EDD (deep drawing)", "Aluminium AlSi10Mg (EN AC-43400)"):
        for cat in ("material", "material_scrap"):
            assert rates.lookup(cat, grade, "DE", MONTH)["is_example"]


def test_new_grade_in_a_tool_shows_the_example_warning():
    """Main's own grade change on a BOM line (stage 23) picks the new grade; its EXAMPLE rate raises the warning."""
    base = dict(part_type="plastic", part_length_mm=120, part_width_mm=80, part_height_mm=30, projected_area_cm2=95.78,
                part_weight_g=48.66, part_volume_cm3=46.3, part_surface_cm2=420, material="ABS", as_of=MONTH, cavities=2)
    r = tool_cost.run(tool_cost.ToolInput(**base, bom={"Cavity insert (steel)": {"grade": "Beryllium copper C17200"}}))
    line = next(l for l in r["lines"] if l["element"] == "Cavity insert (steel)")
    assert line["grade"] == "Beryllium copper C17200" and line["material_cost"] > 0
    becu = next(x for x in r["rates_used"] if x["item"] == "Beryllium copper C17200")
    assert becu["id"] in r["example"]["rate_ids"] and r["example"]["warning"]


def test_old_workbook_without_the_new_columns_still_reads(tmp_path, monkeypatch):
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)
    src = master_data.tables()
    for sheet, cols in (("Parameters", master_data.SHEETS["Parameters"]), ("Materials", ["part_type", "material", "density_g_cm3", "notes"])):
        ws = wb.create_sheet(sheet)
        ws.append(cols)
        for r in src[sheet]:
            ws.append([", ".join(map(str, r[c])) if isinstance(r[c], list) else r[c] for c in cols])
    p = tmp_path / "old.xlsx"
    wb.save(p)
    monkeypatch.setenv("MASTER_DATA_PATH", str(p))
    mats = master_data.tables()["Materials"]
    assert mats and all(r["material_type"] == "" and r["recommended"] == "" for r in mats)


def test_two_recommended_grades_of_one_type_are_refused():
    md = master_data.tables()
    rows = [dict(r) for r in md["Materials"]]
    rows.append(dict(next(r for r in rows if r["material"] == "ABS"), material="ABS 2"))
    with pytest.raises(master_data.MasterDataError, match="already has a recommended grade"):
        master_data.write({"Parameters": md["Parameters"], "Materials": rows})


def test_part_panel_recommends_the_grade_of_the_material_type():
    p = client.post("/api/projects", json={"code": "G", "name": "Grades"}).json()
    with open(SAMPLES / "plastic_cover.step", "rb") as fh:
        m = client.post("/api/import/step", files={"file": ("plastic_cover.step", fh)}).json()
    pid = client.post(f"/api/import/{m['token']}/commit", json={"project_id": p["id"], "items": {
        m["root"]: {"material": "PP copolymer (impact)", "process": "injection_moulding", "part_code": "G-1"}}}).json()["part_ids"][0]
    rec = client.get(f"/api/parts/{pid}").json()["material_recommendation"]
    assert rec["material"] == "PP (polypropylene)" and not rec["is_recommended"]
    assert {"PP copolymer (impact)", "PP talc filled 20 %"} <= set(rec["grades"])
    client.put(f"/api/parts/{pid}", json={"material": "PP (polypropylene)"})
    assert client.get(f"/api/parts/{pid}").json()["material_recommendation"]["is_recommended"]


# ---------- price master ----------

def test_grid_starts_from_the_rates_in_main_and_shows_no_made_up_history():
    g, r = row("tool_material", "IN", "P20 (1.2311)")
    assert g["months"][-1] == "2026-11" and len(g["months"]) == 25
    assert r["cells"][MONTH]["starts"] and r["cells"]["2026-11"]["value"] == r["cells"][MONTH]["value"]
    assert not r["cells"]["2026-11"]["starts"]  # carried over, not a version of its own
    assert "2026-09" not in r["cells"]           # no rate before 2026-10: the cell stays empty
    assert set(client.get("/api/price-master/meta").json()["groups"]) == {"tool_material", "part_material", "part_scrap"}
    assert client.get("/api/price-master/grid", params={"group": "standard_part", "country": "IN"}).status_code == 400


def test_typed_price_is_a_new_version_from_that_month_on():
    approver()
    r = client.post("/api/price-master/price", json={"group": "tool_material", "item": "NAK80", "country": "IN",
                                                     "month": "2026-11", "value": 700})
    assert r.status_code == 200 and r.json()["status"] == "approved" and not r.json()["is_example"]
    _, cells = row("tool_material", "IN", "NAK80")
    assert cells["cells"]["2026-11"]["value"] == 700 and cells["cells"]["2026-11"]["starts"]
    assert rates.lookup("tool_material", "NAK80", "IN", MONTH)["value"] == 650  # the earlier month is unchanged
    draft = client.post("/api/price-master/price", json={"group": "part_material", "item": "ABS", "country": "IN",
                                                         "month": "2026-11", "value": 999, "approve": False}).json()
    assert draft["status"] == "draft" and rates.lookup("material", "ABS", "IN", "2026-11")["value"] != 999
    bad = client.post("/api/price-master/price", json={"group": "tool_material", "item": "Nope", "country": "IN",
                                                       "month": "2026-11", "value": 1})
    assert bad.status_code == 400 and "master list" in bad.json()["detail"]


def test_excel_round_trip_only_saves_changed_cells():
    approver()
    x = client.get("/api/price-master/export.xlsx", params={"group": "part_scrap", "country": "DE"})
    up = lambda content: client.post("/api/price-master/upload", data={"group": "part_scrap", "country": "DE", "approve": "true"},
                                     files={"file": ("p.xlsx", content)}).json()
    assert up(x.content) == {"versions_created": 0, "status": "approved"}
    wb = load_workbook(io.BytesIO(x.content))
    ws = wb.active
    head = [c.value for c in ws[1]]
    ws.cell(row=2, column=head.index("2026-11") + 1).value = 9.5
    buf = io.BytesIO()
    wb.save(buf)
    assert up(buf.getvalue())["versions_created"] == 1
    assert rates.lookup("material_scrap", ws.cell(row=2, column=1).value, "DE", "2026-11")["value"] == 9.5


def test_rate_library_download_is_quick_and_marks_examples():
    import time
    t = time.time()
    x = client.get("/api/rates/download")
    assert x.status_code == 200 and time.time() - t < 30
    ws = load_workbook(io.BytesIO(x.content), read_only=True)["Rates"]
    head = next(ws.iter_rows(max_row=1, values_only=True))
    assert "is_example" in head


# ---------- V0 ----------

def test_release_name_v0_in_the_api_and_admin():
    assert (ROOT / "VERSION").read_text().strip() == "V0"
    assert client.get("/api/auth/me").json()["version"] == "V0"
    assert app.version == "V0"
    assert client.get("/api/admin/about").json()["release"] == "V0"
