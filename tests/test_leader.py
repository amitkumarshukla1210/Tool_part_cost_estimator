import io
import math

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.app import leader
from backend.app.main import app
from tests.test_rm_cost import add_part

client = TestClient(app)


def setup_project(code="LV"):
    p = client.post("/api/projects", json={"code": code, "name": "Leader test", "capex_budget_inr": 5_000_000, "sop_date": "2027-04-01"}).json()
    return p


def costed(p, code="LP-1"):
    pid = add_part(p, code=code)
    assert client.post(f"/api/parts/{pid}/tool-cost", json={}).status_code == 200
    assert client.post(f"/api/parts/{pid}/piece-cost", json={}).status_code == 200
    return pid


def test_bom_cost_vs_target_by_hand():
    p = setup_project()
    pid = costed(p)
    vid = client.post(f"/api/projects/{p['id']}/variants", json={"code": "STD", "name": "Std", "target_bom_cost_inr": 10, "annual_volume": 100000}).json()["variants"][0]["id"]
    from backend.app import projects_db as db, register
    with db.engine().begin() as conn:
        conn.execute(db.bom_lines.insert().values(variant_id=vid, parent_part_id=None, child_part_id=pid, quantity=3))
    v = client.get(f"/api/projects/{p['id']}/leader").json()
    piece = register.piece_summary(pid)["latest"]["piece_cost_inr"]
    x = v["variants"][0]
    assert x["bom_cost_inr"] == pytest.approx(3 * piece) and x["delta_inr"] == pytest.approx(3 * piece - 10)
    assert x["state"] == ("over target" if 3 * piece > 10 else "within target") and x["parts_missing"] == 0 and x["is_example"]


def test_machine_need_capex_and_total_vs_budget():
    p = setup_project("LV2")
    pid = costed(p)
    from backend.app import register
    v = client.get(f"/api/projects/{p['id']}/leader").json()
    tr = register.piece_trace(pid)
    m = v["machines"]["rows"][0]
    vol = 100000                                                              # Parameters default annual volume
    hours = vol / tr["nodes"]["calc:parts_per_h"]["value"] / 0.75             # OEE 75 % (plastic)
    assert m["machine_id"] == tr["machine"] and m["hours_needed"] == pytest.approx(hours)
    assert m["machines"] == math.ceil(hours / (300 * 2 * 8)) and m["capex_inr"] == pytest.approx(m["machines"] * m["price_inr"])
    cap = v["capex"]
    assert cap["tool_inr"] > 0 and cap["total_inr"] == pytest.approx(cap["tool_inr"] + cap["machine_inr"])
    assert cap["over_budget"] == (cap["total_inr"] > 5_000_000)


def test_plan_build_milestones_timeline_and_payments():
    p = setup_project("LV3")
    pid = costed(p)
    r = client.post(f"/api/projects/{p['id']}/capex/build").json()
    assert r["created"] == 2 and client.post(f"/api/projects/{p['id']}/capex/build").json() == {"created": 0, "refreshed": 0}
    plan = client.get(f"/api/projects/{p['id']}/leader").json()["plan"]
    tool = next(i for i in plan["items"] if i["kind"] == "tool")
    assert tool["need_by_effective"] == "2027-04-01" and [m["pct"] for m in tool["milestones"]] == [30, 30, 30, 10]
    assert [m["due"] for m in tool["milestones"]] == ["2026-10-15", "2026-12-10", "2027-02-18", "2027-04-29"]       # 24, 16, 6 weeks before; 4 after
    assert sum(m["amount_inr"] for m in tool["milestones"]) == pytest.approx(tool["amount_inr"])
    client.put(f"/api/capex/{tool['id']}", json={"supplier": "ABC Tools", "country": "CN", "payments": {"Advance with PO": "2026-10-20"}})
    plan = client.get(f"/api/projects/{p['id']}/leader").json()["plan"]
    assert plan["paid_inr"] == pytest.approx(0.3 * tool["amount_inr"]) and ("ABC Tools", tool["amount_inr"]) in [tuple(x) for x in plan["by_supplier"]]
    assert sum(a for _, a in plan["by_month"]) == pytest.approx(plan["total_inr"])
    assert client.put(f"/api/capex/{tool['id']}", json={"payments": {"nope": "2026-10-20"}}).status_code == 400
    assert client.post(f"/api/projects/{p['id']}/capex", json={"kind": "other", "description": "Gauges", "amount_inr": 50000, "need_by": "2027-01-15"}).status_code == 200
    assert client.get(f"/api/projects/{p['id']}/leader").json()["capex"]["other_inr"] == 50000
    assert client.delete(f"/api/capex/{tool['id']}").status_code == 200


def test_open_corrections_exports_and_recalculate():
    p = setup_project("LV4")
    pid = costed(p)
    for st, c in (("concept_tool", ""), ("simulation", ""), ("corrections_required", "add draft")):
        client.post(f"/api/parts/{pid}/status", json={"status": st, "comment": c})
    v = client.get(f"/api/projects/{p['id']}/leader").json()
    assert [c["part_code"] for c in v["corrections"]] == ["LP-1"]
    wb = load_workbook(io.BytesIO(client.get(f"/api/projects/{p['id']}/leader.xlsx").content))
    assert "CAPEX plan" in wb.sheetnames and "Payments" in wb.sheetnames and wb["Open corrections"].max_row == 2
    assert client.get(f"/api/projects/{p['id']}/leader.pdf").content[:4] == b"%PDF"
    assert client.post(f"/api/projects/{p['id']}/piece-cost/recalculate").json()["unchanged"] == 1


def test_setup_checks():
    assert client.post("/api/projects", json={"code": "BAD", "name": "x", "sop_date": "01/04/2027"}).status_code == 400
    s = client.get("/api/register-setup").json()
    s["capex_milestones"][0]["pct"] = 50
    r = client.put("/api/register-setup", json=s)
    assert r.status_code == 400 and "100 %" in r.text
